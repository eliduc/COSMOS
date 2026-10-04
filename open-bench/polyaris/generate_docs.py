# -*- coding: utf-8 -*-
r"""Генератор документов «Полярис»: docx + xlsx + jpg из M0.

Форма — по скелету документов одной из линий (решение владельца 2026-08-20):
docx — альбомный A4, интро-таблица и таблицы правил по разделам с
одной вложенной таблицей; xlsx — журнал замечаний на два листа с
закреплением области E2 и колонками журнального типа; jpg — «памятка
планировщику» (изображение страницы). Содержательных токенов исходной
линии нет — гейт линтует выход стоп-словарём.

Дефекты вносятся по реестру defects.py (дозировка §5 постановки).

После сборки документов генератор прогоняет ШТАТНУЮ заморозку CDPM,
сопоставляет клаузы формулировкам и вписывает трассировку (sources)
в M0 — итоговая модель обязана проходить validate_model начисто
(кроме jpg: его клаузы появятся после приёмника изображений ПИ-1;
до тех пор правила памятки трассируются на docx-дубли).

Запуск: ..\..\.venv\Scripts\python.exe generate_docs.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).parent
CREW = HERE.parent.parent
sys.path.insert(0, str(CREW))
sys.path.insert(0, str(HERE))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import defects as defects_mod                     # noqa: E402
import formulations as forms                      # noqa: E402
from m0 import build as build_m0                  # noqa: E402

DOCS = HERE / "docs"
ETALON = HERE / "etalon"
# Имена входных файлов — константы: производные наборы (Лайт, мини,
# нано) подменяют их вместе с DOCS/ETALON (директива владельца
# 2026-08-21: PMI1/PME1/PMC1 и PNI1/PNE1/PNC1).
DOCX_NAME = "PI1-Pravila-Polyaris.docx"
XLSX_NAME = "PE1-Zhurnal-Polyaris.xlsx"
JPG_NAME = "PR1-Pamyatka-Polyaris.jpg"

# Доли баз для вложенной таблицы M-051 (из кредитных ограничений
# GERAD inst1: 16.4 / 64.3 / 19.3).
BASE_SHARES = [("Северогорск (SVG)", "16%"),
               ("Полярный (PLY)", "64%"),
               ("Гранитный (GRT)", "19%")]

# Шумовые строки журнала: замечания без правил (реализм Е1).
XLSX_NOISE = [
    "Дубли рейсов в выгрузке расписания за 3-е число.",
    "Не подтянулись фамилии в отчёте по цепочкам.",
    "Просьба добавить фильтр по базе в отчёт по резерву.",
    "Расхождение нумерации рейсов между выгрузками.",
    "Медленно строится отчёт по месяцу.",
    "Опечатка в названии аэропорта Пункт-07 в справочнике.",
    "Не выгружается список отпусков за прошлый период.",
    "Продублировать журнал в формате PDF.",
    "Уточнить порядок сортировки в отчёте по экипажам.",
    "В выгрузке смен пустая колонка примечаний.",
    "Сценарий расчёта завис на шаге проверки допусков.",
    "Разъяснить обозначения статусов в отчёте.",
    "Добавить экспорт журнала за произвольный период.",
    "Отчёт по налёту не учитывает последний день месяца.",
    "Пропали цвета статусов после обновления.",
]

HEADERS = ["ID", "Замечание", "Описание", "Приоритет", "Статус",
           "Категория", "Название задачи на расчет в Borealis.Crew",
           "Название сценария в Borealis.Crew",
           "Дата добавления в журнал", "Тип расчета",
           "Ссылка на задачу в PLS", "Ответственный",
           "Инициатор (автор)", "Доп.комментарий", "", ""]


def _apply_replace(text: str, doc: str, rid: str | None) -> str:
    """Дефекты вида replace для данного документа/правила."""
    for did, _k, d_doc, target, how, payload, _e, _r in \
            defects_mod.DEFECTS:
        if d_doc != doc or how != "replace" or target != rid:
            continue
        old, new = payload
        assert old in text, f"{did}: подстрока {old!r} не найдена"
        text = text.replace(old, new)
    return text


def _extras(doc: str) -> list[str]:
    return [payload for _d, _k, d_doc, _t, how, payload, _e, _r in
            defects_mod.DEFECTS if d_doc == doc and how == "extra"]


def build_docx(path: Path) -> dict[str, str]:
    """PI1-Pravila-Polyaris.docx по форме И1. Возвращает rid → сниппет
    формулировки (для сопоставления клауз)."""
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm, Pt

    d = docx.Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    sec.left_margin = sec.right_margin = Mm(20)

    # Интро-таблица 5×2 — как в И1.
    intro = d.add_table(rows=5, cols=2)
    intro.style = "Table Grid"
    for row, (k, v) in zip(intro.rows, [
            ("Документ", "Правила оптимизации планирования экипажей"),
            ("Авиакомпания", "«Полярис» (PLR)"),
            ("Версия", "1.0"),
            ("Дата", "01.08.2026"),
            ("Подразделение", "Отдел планирования экипажей")]):
        row.cells[0].text = k
        row.cells[1].text = v

    d.add_heading("Правила оптимизации планирования экипажей "
                  "авиакомпании «Полярис»", level=2)

    snippets: dict[str, str] = {}
    by_section: dict[int, list[str]] = {1: [], 2: [], 3: [], 4: []}
    for rid, (sect, *_rest) in forms.DOCX.items():
        by_section[sect].append(rid)

    # Дефекты-extra docx раскладываются по разделам: G-01→2, A-03→2,
    # G-02→4, D-01→3 (см. defects.py; порядок фиксированный).
    extra_by_section = {2: [], 3: [], 4: []}
    for did, _k, doc, _t, how, payload, _e, _r in defects_mod.DEFECTS:
        if doc != "docx" or how != "extra":
            continue
        sect = {"A-03": 2, "G-01": 2, "D-01": 3, "G-02": 4}.get(did, 2)
        extra_by_section[sect].append(payload)

    for sect in (1, 2, 3, 4):
        d.add_heading(f"Раздел {sect}. {forms.SECTIONS[sect]}",
                      level=2)
        rows = by_section[sect]
        extras = extra_by_section.get(sect, [])
        t = d.add_table(rows=1 + len(rows) + len(extras), cols=6)
        t.style = "Table Grid"
        head = t.rows[0].cells
        for i, h in enumerate(["№", "Наименование", "Формулировка",
                               "Характер", "Значение", "Примечание"]):
            head[i].text = h
        for n, rid in enumerate(rows, start=1):
            _s, name, text, kind, value, note = forms.DOCX[rid]
            text = _apply_replace(text, "docx", rid)
            cells = t.rows[n].cells
            cells[0].text = f"{sect}.{n}"
            cells[1].text = name
            cells[2].text = text
            cells[3].text = kind
            cells[4].text = value
            cells[5].text = note
            snippets[rid] = text[:60]
            if rid == "M-051":
                # Вложенная таблица долей баз — как в И1 (форма).
                nested = cells[5].add_table(rows=1 + len(BASE_SHARES),
                                            cols=2)
                nested.rows[0].cells[0].text = "База"
                nested.rows[0].cells[1].text = "Доля кредита"
                for rr, (b, share) in zip(nested.rows[1:],
                                          BASE_SHARES):
                    rr.cells[0].text = b
                    rr.cells[1].text = share
        for j, payload in enumerate(extras):
            cells = t.rows[1 + len(rows) + j].cells
            cells[0].text = f"{sect}.{len(rows) + 1 + j}"
            cells[1].text = "Дополнительно"
            cells[2].text = payload
            cells[3].text = "—"
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return snippets


def build_xlsx(path: Path) -> dict[str, str]:
    """PE1-Zhurnal-Polyaris.xlsx по форме Е1 (2 листа, закрепление E2)."""
    import openpyxl
    from datetime import date

    wb = openpyxl.Workbook()
    snippets: dict[str, str] = {}

    def fill(ws, title, rows):
        ws.title = title
        ws.append(HEADERS)
        ws.freeze_panes = "E2"
        for i, (note, descr) in enumerate(rows, start=1):
            ws.append([f"PLR-{i:03d}", note, descr,
                       ["Низкий", "Средний", "Высокий"][i % 3],
                       ["Открыто", "В работе", "Закрыто"][i % 3],
                       "Планирование",
                       f"Расчёт-{i:02d}", f"Сценарий-{i:02d}",
                       date(2026, 7, (i % 27) + 1), "месячный",
                       f"PLS-{400 + i}", "Дежурный аналитик",
                       "Отдел планирования", ""])

    # Лист 1: правило-носители + дефекты + шум.
    rows1: list[tuple[str, str]] = []
    for rid, (name, text) in forms.XLSX.items():
        text = _apply_replace(text, "xlsx", rid)
        rows1.append((name, text))
        snippets[rid] = text[:60]
    for payload in _extras("xlsx"):
        rows1.append(("Дополнительно", payload))
    rows1 += [("Замечание", t) for t in XLSX_NOISE[:10]]

    fill(wb.active, "Журнал ОПЭ Полярис", rows1)
    ws2 = wb.create_sheet()
    fill(ws2, "Архив 2025",
         [("Замечание", t) for t in XLSX_NOISE[10:]])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return snippets


def build_jpg(path: Path) -> list[str]:
    """PR1-Pamyatka-Polyaris.jpg — изображение страницы памятки."""
    from PIL import Image, ImageDraw, ImageFont

    lines = ["ПАМЯТКА ПЛАНИРОВЩИКУ", "Авиакомпания «Полярис» (PLR)",
             ""]
    for rid, line in forms.PAMYATKA:
        lines.append("• " + _apply_replace(line, "jpg", rid))
    for payload in _extras("jpg"):
        lines.append("• " + payload)
    lines += ["", "Отдел планирования экипажей, 08.2026"]

    img = Image.new("RGB", (1240, 1754), "white")   # ~A4 150 dpi
    draw = ImageDraw.Draw(img)
    try:
        font_h = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf",
                                    44)
        font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 34)
    except OSError:
        font_h = font = ImageFont.load_default()
    y = 90
    for i, line in enumerate(lines):
        f = font_h if i == 0 else font
        draw.text((90, y), line, fill="black", font=f)
        y += 64 if i == 0 else 52
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(str(path), quality=92)
    return [ln[2:] for ln in lines if ln.startswith("• ")]


def freeze_and_trace(m0: dict, snippets: dict[str, dict[str, str]],
                     jpg_lines: list[str]) -> tuple[dict, dict]:
    """Штатная заморозка CDPM → сопоставление клауз → sources в M0."""
    import tempfile
    from cdpm import freeze as freeze_mod

    with tempfile.TemporaryDirectory() as td:
        kit = Path(td) / "kit"
        freeze_mod.freeze([DOCS / DOCX_NAME,
                           DOCS / XLSX_NAME], kit)
        entries = {}
        for sid in ("S1", "S2"):
            entries[sid] = json.loads(
                (kit / "clauses" / f"{sid}.json")
                .read_text(encoding="utf-8"))
        manifest = json.loads((kit / "manifest.json")
                              .read_text(encoding="utf-8"))

    for src in m0["sources"]:
        for f in manifest["files"]:
            if f["file"] == src["file"]:
                src["sha256"] = f["sha256"]

    anchors_map: dict[str, list[dict]] = {}
    for sid, doc_key in (("S1", "docx"), ("S2", "xlsx")):
        clauses = entries[sid]["clauses"]
        for rid, snippet in snippets[doc_key].items():
            hits = [c["anchor"] for c in clauses
                    if snippet in c["text"]]
            assert hits, f"{rid}: сниппет не найден в {sid}"
            anchors_map.setdefault(rid, []).append(
                {"sid": sid, "anchor": hits[0]})

    # Памятка (jpg): до ПИ-1 клауз нет — правила памятки уже
    # трассированы на docx (дубли по построению); фиксируем строки
    # для будущего сопоставления.
    anchors_map["_pamyatka_lines"] = [
        {"sid": "S3", "line": ln} for ln in jpg_lines]

    by_id = {r["id"]: r for r in m0["rules"]}
    for rid, refs in anchors_map.items():
        if rid.startswith("_"):
            continue
        by_id[rid]["sources"] = refs
    return m0, anchors_map


def _bind_soft(m0: dict) -> None:
    """SOFT-инвариант контракта: вес — только из ответа авиакомпании,
    поэтому каждое мягкое правило блокируется ЕДИНЫМ вопросом зоны о
    целевой и весах (Q-*-GOAL, склейка ЦФ-1 2026-08-21). Используется
    ШТАТНЫЙ строитель CDPM — эталон обязан выглядеть ровно как
    честный результат извлечения (он же создаёт вопрос с секцией
    таблицы целевых, когда целевых в документах нет)."""
    from cdpm import model as model_mod
    for part in ("3a", "3b"):
        model_mod.ensure_goal_question(m0, part)


def main() -> int:
    m0 = build_m0()
    snippets = {"docx": build_docx(DOCS / DOCX_NAME),
                "xlsx": build_xlsx(DOCS / XLSX_NAME)}
    jpg_lines = build_jpg(DOCS / JPG_NAME)
    m0, anchors_map = freeze_and_trace(m0, snippets, jpg_lines)
    _bind_soft(m0)

    ETALON.mkdir(parents=True, exist_ok=True)
    (ETALON / "M0-model.json").write_text(
        json.dumps(m0, ensure_ascii=False, indent=1),
        encoding="utf-8")
    (ETALON / "anchors-map.json").write_text(
        json.dumps(anchors_map, ensure_ascii=False, indent=1),
        encoding="utf-8")
    (ETALON / "DEFECTS-REGISTRY.json").write_text(
        json.dumps(defects_mod.registry(), ensure_ascii=False,
                   indent=1),
        encoding="utf-8")
    # Текст, фактически нанесённый на изображение памятки, — источник
    # истины для гейта (пиксели не грепаются) и эталон для будущего
    # распознавания ПИ-1.
    (ETALON / "jpg-lines.json").write_text(
        json.dumps(jpg_lines, ensure_ascii=False, indent=1),
        encoding="utf-8")

    from cdpm import model as model_mod
    problems = model_mod.validate_model(m0)
    traced = sum(1 for r in m0["rules"] if r["sources"])
    print(f"документы: {DOCS}")
    print(f"правил в M0: {len(m0['rules'])}, трассировано: {traced}; "
          f"дефектов в реестре: {len(defects_mod.DEFECTS)}")
    print(f"validate_model: нарушений {len(problems)}")
    for p in problems[:8]:
        print(" -", p)
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
