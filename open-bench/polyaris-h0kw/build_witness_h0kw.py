# -*- coding: utf-8 -*-
"""Расписание-свидетель набора «Полярис-H0-КИ-окна» — допустимый ростер по
книге комплекта, построенный конструктивно из замкнутых ротаций (без
позиционирования пассажирами), плюс отрицательные варианты с ровно одним
заведомым нарушением. Ожидание каждого свидетеля берётся из сценария и ВМП,
не из решателя. Свидетель — данные сценария для сертификатора COSMOS и
для старта ростеринга «от свидетеля»; выпускает CDPM.

Основное требование владельца к набору Н0 (05.09.2026): очень простой
демонстрационный набор маленькой авиакомпании, проход по обоим конвейерам
≈ 30 минут; всё, что мешает быстрому расчёту, меняется в наборе. Поэтому
программа рейсов книги сбалансирована (PL306 — возврат PL305), история по
умолчанию не задаётся (вариант книги nohist), а ночная пара PL401/PL402 —
замкнутая ротация PLY (PL401 суток d и PL402 суток d+1).

Варианты свидетеля (книга → ожидание сертификатора):
- nohist-positive — книга без «Истории» (умолчание комплекта; решение
  пользователя «продолжить с нулевой историей»): допустимо, нарушений нет,
  ночёвок вне базы нет;
- hist-positive — книга с «Историей» (сценарии A/B/C): допустимо;
- hist-neg-A — ровно одно нарушение VN-091: H03 (серия ночных 2 в истории,
  последняя смена ночная) на ночной ротации 04.09 → серия 3 > 2;
- hist-neg-B — ровно одно нарушение VN-023: H07 (конец последней смены
  03.09 22:00) на ротации PL301 04.09 06:00 → отдых дома 480 < 720.

Ротации суток (все замкнуты на базу экипажа; экипаж — КВС + Второй пилот
одной базы):
  PLY  P1 [PL101 08:00→SVG 09:30, PL102 SVG 10:20→PLY 11:50]        230 мин
       P3 [PL301 06:00→SVG 07:30] + дневной отдых в SVG 650 ≥ 600 +
          [PL306 SVG 18:20→PLY 19:50] — две смены одних суток
       P6 [PL401 d 23:30→SVG 01:10, PL402 d+1 02:10→PLY 03:50]        260 мин, ночная
  SVG  V1 [PL302 08:20→PLY 09:50, PL305 PLY 15:40→SVG 17:30]         550 мин
  GRT  G1 [PL304 13:10→PLY 14:50, PL201 PLY 19:00→GRT 20:40]         450 мин
       G2 [PL202 07:20→PLY 09:00, PL303 PLY 10:40→GRT 12:20]         300 мин
Ротации по дням: PLY — цикл 6 суток по 6 экипажам (экипаж X: P1 в день X,
P6 в день X+2, P3 в день X+4; каждые сутки заняты 4 экипажа — P1, P3, P6
этих суток и P6 предыдущих, свободны 2 = резерв VN-060 ровно по норме);
SVG — V1 по кругу из 3 экипажей; GRT — цикл 6 суток по 3 экипажам
[G2, G2, G1, G1, —, —] со сдвигом фаз.

Сироты ночной пары (PL402 первых суток без своего PL401 накануне, PL401
последних суток без PL402 в горизонте) свидетель НЕ покрывает: их не
должно быть в простой книге (ночная программа = замкнутые пары внутри
горизонта); при их наличии самопроверка называет их нарушением «книга».

    python open-bench/polyaris-h0kw/build_witness_h0kw.py            # проверка
    python open-bench/polyaris-h0kw/build_witness_h0kw.py --write    # witness/

Формат — вход сертификатора ростеринга (vmp/run/roster_certify): JSON
{assignments: [{employee, flight, date, dep_airport, arr_airport, dep_time,
arr_time, role, base, duty_id, deadhead}], duties: [{employee, id, start,
end, base, last_arr_airport, extension_consent}]} и та же книга xlsx с
листами «Назначения» и «Смены».
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "open-bench" / "polyaris-h0"))
sys.path.insert(0, str(HERE))
import build_data_h0 as h0  # noqa: E402
import build_data_h0kw as bw  # noqa: E402

OUT_DIR = HERE / "witness"
DAYS = bw.DAYS
N = len(DAYS)
# Ревизия 06.09.2026: предел смены комплекта — 600 минут без продления
# (build_kit_h0kw.DUTY_LIMIT_MIN); самая длинная ротация свидетеля — 550.
LIM = dict(h0.LIMITS, max_duty_duration=600)
#: Пороги правил окон истории (VN-091/092) — как в эталоне комплекта.
MAX_NIGHT_SERIES = 2
ROLLING_WINDOW_MIN, ROLLING_REST_MIN = 7 * 1440, 36 * 60
#: Норма выходных месяца (VN-093) — судится только при закрытом месяце.
MIN_DAYS_OFF_MONTH = getattr(bw, "MIN_DAYS_OFF_MONTH", 8)

#: Экипажи (КВС, Второй пилот) по базам; порядок фиксирует ротацию.
#: PLY: c0 несёт H03 (сценарий A: первая смена дневная P1 04.09), c5 — H07
#: (сценарий B: первая смена P1 09.09), c3 — H11 (сценарий C).
CREWS = {
    "PLY": [("H04", "H03"), ("H02", "H01"), ("H06", "H05"),
            ("H12", "H11"), ("H10", "H09"), ("H08", "H07")],
    "SVG": [("H14", "H13"), ("H16", "H15"), ("H18", "H17")],
    "GRT": [("H20", "H19"), ("H22", "H21"), ("H24", "H23")],
}
#: Шаблоны смен: (имя, база, [(рейс, смещение суток)]).
P1 = ("P1", "PLY", [("PL101", 0), ("PL102", 0)])
P3A = ("P3a", "PLY", [("PL301", 0)])
P3B = ("P3b", "PLY", [("PL306", 0)])
P6 = ("P6", "PLY", [("PL401", 0), ("PL402", 1)])
V1 = ("V1", "SVG", [("PL302", 0), ("PL305", 0)])
G1 = ("G1", "GRT", [("PL304", 0), ("PL201", 0)])
G2 = ("G2", "GRT", [("PL202", 0), ("PL303", 0)])

VARIANTS = ("nohist-positive", "hist-positive", "hist-neg-A", "hist-neg-B")
#: Книга каждого варианта (вариант генератора книги) и ожидания сценария.
BOOK_OF = {"nohist-positive": "nohist", "hist-positive": "hist",
           "hist-neg-A": "hist", "hist-neg-B": "hist"}
EXPECT_VIOLATIONS = {
    "nohist-positive": [],
    "hist-positive": [],
    "hist-neg-A": [("VN-091", "H03")],
    "hist-neg-B": [("VN-023", "H07")],
}


# ------------------------------------------------------------------ книга

def _mins(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def legs_by_key() -> dict:
    """(рейс, дата) → рейс книги с datetime вылета/прилёта (прилёт раньше
    вылета — следующие сутки, как у приёмника webapp/task_package)."""
    out = {}
    for r in bw.flights():
        d = dt.date.fromisoformat(r["Дата"])
        dep = dt.datetime.combine(d, dt.time(*divmod(_mins(r["Время вылета"]), 60)))
        arr = dt.datetime.combine(d, dt.time(*divmod(_mins(r["Время прилёта"]), 60)))
        if arr < dep:
            arr += dt.timedelta(days=1)
        out[(r["Номер рейса"], r["Дата"])] = dict(r, dep_dt=dep, arr_dt=arr)
    return out


def crew_table() -> dict:
    return {r["Табельный номер"]: r for r in h0.tables()["Экипажи"]["rows"]}


def history_records(book: str) -> dict:
    """Записи листа «История» книги варианта hist; у книги без истории —
    пусто (остатки по решению пользователя нулевые)."""
    if book != "hist":
        return {}
    out = {}
    crew = h0.tables()["Экипажи"]["rows"]
    for row in bw.history_rows(crew):
        end = dt.datetime.strptime(row["Последняя смена: конец"], "%d.%m.%Y %H:%M")
        out[row["Табельный номер"]] = {
            "last_end": end,
            "last_night": str(row["Последняя смена: ночная"]).strip().lower() == "да",
            "night_series": int(row["Серия ночных смен"]),
            "days_worked": int(row["Дней отработано"]),
            "station": row["Последняя смена: станция"],
        }
    return out


# --------------------------------------------------------------- ротации

def plan(variant: str) -> list[tuple[int, tuple, tuple]]:
    """Список (день k, шаблон смены, экипаж) по варианту."""
    ply, svg, grt = CREWS["PLY"], CREWS["SVG"], CREWS["GRT"]
    items: list[tuple[int, tuple, tuple]] = []
    for k in range(N):
        items.append((k, P1, ply[k % 6]))
        items.append((k, P3A, ply[(k + 2) % 6]))
        items.append((k, P3B, ply[(k + 2) % 6]))
        items.append((k, P6, ply[(k + 4) % 6]))
        items.append((k, V1, svg[k % 3]))
        m = k % 6
        g2 = grt[0] if m in (0, 1) else grt[2] if m in (2, 3) else grt[1]
        g1 = grt[1] if m in (0, 1) else grt[0] if m in (2, 3) else grt[2]
        items.append((k, G2, g2))
        items.append((k, G1, g1))
    if variant == "hist-neg-A":
        # H03 (c0) на ночную 04.09: обмен c0↔c4 в день 0 (P1↔P6) и в день 2
        # (P6↔P3) — серия сбрасывается дневной сменой, нарушение одно.
        c0, c4 = ply[0], ply[4]
        swap = {c0: c4, c4: c0}
        items = [(k, tpl, swap.get(crew, crew))
                 if (k in (0, 2) and crew in swap and tpl[1] == "PLY")
                 else (k, tpl, crew) for k, tpl, crew in items]
    elif variant == "hist-neg-B":
        # H07 (c5) на PL301 04.09 06:00 вместо c2 (обе смены P3 этих суток):
        # отдых от 03.09 22:00 = 480 < 720.
        items = [(k, tpl, ply[5]) if (k == 0 and tpl[0] in ("P3a", "P3b"))
                 else (k, tpl, crew) for k, tpl, crew in items]
    elif variant not in ("nohist-positive", "hist-positive"):
        raise ValueError(variant)
    return items


# ---------------------------------------------------------------- решение

def build_solution(variant: str) -> dict:
    legs = legs_by_key()
    crew = crew_table()
    assignments, duties = [], []
    for k, (name, base, tpl_legs), pilots in plan(variant):
        rows = []
        for flight, shift in tpl_legs:
            key = (flight, DAYS[k + shift]) if k + shift < N else None
            if key is None or key not in legs:
                continue  # рейса нет в книге (граница горизонта)
            rows.append(legs[key])
        if not rows:
            continue
        start, end = rows[0]["dep_dt"], rows[-1]["arr_dt"]
        for emp in pilots:
            duty_id = f"{emp}-{DAYS[k]}-{name}"
            for leg in rows:
                assignments.append({
                    "employee": emp, "flight": leg["Номер рейса"], "date": leg["Дата"],
                    "dep_airport": leg["Вылет"], "arr_airport": leg["Прилёт"],
                    "dep_time": leg["Время вылета"], "arr_time": leg["Время прилёта"],
                    "role": crew[emp]["Должность"], "base": crew[emp]["База"],
                    "duty_id": duty_id, "deadhead": "нет",
                })
            duties.append({
                "employee": emp, "id": duty_id,
                "start": start.strftime("%Y-%m-%dT%H:%M"),
                "end": end.strftime("%Y-%m-%dT%H:%M"),
                "base": crew[emp]["База"], "last_arr_airport": rows[-1]["Прилёт"],
                "extension_consent": "нет",
            })
    return {"variant": variant, "book": BOOK_OF[variant],
            "assignments": assignments, "duties": duties}


# -------------------------------------------------------------- самопроверка

def check(solution: dict) -> tuple[list[tuple[str, str, str]], dict]:
    """Жёсткие правила ВМП H0 с учётом истории (если книга её несёт).
    Возврат: (нарушения [(правило, табельный, слова)], числа)."""
    legs = legs_by_key()
    crew = crew_table()
    hist = history_records(solution["book"])
    horizon_start = dt.datetime.fromisoformat(DAYS[0] + "T00:00")
    horizon_end = dt.datetime.fromisoformat(DAYS[-1] + "T23:59") + dt.timedelta(minutes=1)
    viol: list[tuple[str, str, str]] = []
    facts: dict = {}
    # Покрытие и состав экипажа (VN-081): ровно 2 действующих пилота на рейс.
    operating: dict[tuple, list] = {key: [] for key in legs}
    for a in solution["assignments"]:
        key = (a["flight"], a["date"])
        if key not in legs:
            viol.append(("книга", a["employee"], f"рейса {key} в книге нет"))
            continue
        if a["deadhead"] == "нет":
            operating[key].append(a["employee"])
    for key, ppl in operating.items():
        if len(ppl) != LIM["crew_size"]:
            who = ",".join(ppl) or "—"
            note = ""
            if not ppl and key[0] in ("PL401", "PL402"):
                note = " (сирота ночной пары на границе горизонта — уберите из книги)"
            viol.append(("VN-081" if ppl else "книга", who,
                         f"{key}: пилотов {len(ppl)}{note}"))
    facts["flights"] = len(legs)
    facts["assignments"] = len(solution["assignments"])
    facts["deadheads"] = sum(1 for a in solution["assignments"] if a["deadhead"] == "да")
    # Допуски (VN-070): действуют на дату рейса.
    quals = {r["Табельный номер"]: dt.date.fromisoformat(str(r["Действует до"]))
             for r in h0.tables()["Допуски"]["rows"]}
    for a in solution["assignments"]:
        if a["deadhead"] == "нет" and quals[a["employee"]] < dt.date.fromisoformat(a["date"]):
            viol.append(("VN-070", a["employee"], f"допуск истёк до {a['date']}"))
    # По сотрудникам: смены в хронологии.
    by_emp: dict[str, list] = {}
    for d in solution["duties"]:
        by_emp.setdefault(d["employee"], []).append(d)
    legs_by_duty: dict[str, list] = {}
    for a in solution["assignments"]:
        legs_by_duty.setdefault(a["duty_id"], []).append(a)
    days_off_total: dict[str, int] = {}
    busy_days: dict[str, set] = {}
    layovers: dict[str, int] = {}
    open_ends: dict[str, int] = {}
    windows_unjudged = 0
    pre_days = dt.date.fromisoformat(DAYS[0]).day - 1
    for emp, ds in by_emp.items():
        ds.sort(key=lambda d: d["start"])
        base = crew[emp]["База"]
        h = hist.get(emp)
        location = h["station"] if h else base
        prev_end = h["last_end"] if h else None
        prev_home = (h["station"] == base) if h else True
        series = h["night_series"] if (h and h["last_night"]) else 0
        intervals = []
        for i, d in enumerate(ds):
            start = dt.datetime.fromisoformat(d["start"])
            end = dt.datetime.fromisoformat(d["end"])
            intervals.append((start, end))
            # Непрерывность маршрута и стыковки внутри смены (VN-020).
            alegs = sorted(legs_by_duty[d["id"]],
                           key=lambda a: legs[(a["flight"], a["date"])]["dep_dt"])
            if alegs[0]["dep_airport"] != location:
                viol.append(("маршрут", emp, f"смена {d['id']} стартует в "
                                             f"{alegs[0]['dep_airport']}, пилот в {location}"))
            for x, y in zip(alegs, alegs[1:]):
                lx, ly = legs[(x["flight"], x["date"])], legs[(y["flight"], y["date"])]
                if x["arr_airport"] != y["dep_airport"]:
                    viol.append(("маршрут", emp, f"{x['flight']}→{y['flight']}: разрыв аэропортов"))
                turn = (ly["dep_dt"] - lx["arr_dt"]).total_seconds() / 60
                if turn < LIM["min_turnaround"]:
                    viol.append(("VN-020", emp, f"{x['flight']}→{y['flight']}: стыковка {turn:.0f}"))
            location = alegs[-1]["arr_airport"]
            # Длительность смены (VN-010, без продления).
            dur = (end - start).total_seconds() / 60
            if dur > LIM["max_duty_duration"]:
                viol.append(("VN-010", emp, f"смена {d['id']}: {dur:.0f} мин"))
            # Отдых между сменами (VN-023 дома / VN-024 вне базы), включая
            # отдых от конца последней смены истории, если история есть.
            if prev_end is not None:
                rest = (start - prev_end).total_seconds() / 60
                floor = LIM["min_rest_home"] if prev_home else LIM["min_rest_outstation"]
                rule = "VN-023" if prev_home else "VN-024"
                if rest < floor:
                    viol.append((rule, emp, f"отдых перед {d['id']}: {rest:.0f} < {floor}"))
            # Серия ночных смен (VN-091) с учётом истории.
            night = bw.is_night_minute(start.hour * 60 + start.minute)
            series = series + 1 if night else 0
            if series > MAX_NIGHT_SERIES:
                viol.append(("VN-091", emp, f"смена {d['id']}: серия ночных {series}"))
            # Ночёвка вне базы доказана, когда смена кончается не на базе и
            # следующая стартует там же в ДРУГИЕ календарные сутки (отдых
            # внутри тех же суток — не ночёвка); конец вне базы без
            # продолжения в горизонте — открытый конец.
            if d["last_arr_airport"] != base:
                nxt = ds[i + 1] if i + 1 < len(ds) else None
                if nxt is None:
                    open_ends[emp] = open_ends.get(emp, 0) + 1
                elif dt.datetime.fromisoformat(nxt["start"]).date() > end.date():
                    layovers[emp] = layovers.get(emp, 0) + 1
            prev_end, prev_home = end, (d["last_arr_airport"] == base)
        # Окно 7 суток (VN-092): непрерывный отдых ≥ 36 ч в каждом окне
        # скользящего домена (любые 10080 минут подряд). Судимый домен — от
        # конца последней смены истории (без истории — от начала горизонта)
        # до границы горизонта; окна за границей — несудимые.
        w = dt.timedelta(minutes=ROLLING_WINDOW_MIN)
        domain_start = h["last_end"] if h else horizon_start
        points = {domain_start}
        for a, b in intervals:
            points.update({a, b, a - w, b - w})
        for s in sorted(points):
            if s < domain_start:
                continue
            w_end = s + w
            if w_end > horizon_end:
                windows_unjudged += 1
                continue
            cursor, best = s, 0.0
            for a, b in intervals:
                if b <= s or a >= w_end:
                    continue
                best = max(best, (max(a, s) - cursor).total_seconds() / 60)
                cursor = max(cursor, b)
            best = max(best, (w_end - cursor).total_seconds() / 60)
            if best < ROLLING_REST_MIN:
                viol.append(("VN-092", emp, f"окно от {s:%d.%m %H:%M}: макс. отдых "
                                            f"{best:.0f} < {ROLLING_REST_MIN}"))
                break
        # Выходные месяца (VN-093): сутки горизонта без единой минуты смены
        # + сутки месяца до среза без работы (по истории; без истории —
        # нулевые остатки: отработано 0).
        touched = set()
        for a, b in intervals:
            day = a.date()
            while day <= b.date():
                touched.add(day.isoformat())
                day += dt.timedelta(days=1)
        busy_days[emp] = touched
        off = sum(1 for d0 in DAYS if d0 not in touched)
        pre = pre_days - (h["days_worked"] if h else 0)
        days_off_total[emp] = off + pre
        if (dt.date.fromisoformat(DAYS[-1]) + dt.timedelta(days=1)).day == 1 \
                and off + pre < MIN_DAYS_OFF_MONTH:
            viol.append(("VN-093", emp, f"выходных в месяце {off + pre} < {MIN_DAYS_OFF_MONTH}"))
    for emp in crew:
        if emp not in by_emp:
            busy_days[emp] = set()
            days_off_total[emp] = N + pre_days - (hist[emp]["days_worked"] if emp in hist else 0)
    # Резерв по базам (VN-060): свободных за сутки полных экипажей ≥ нормы.
    need = LIM["reserve_crews"]
    reserve_min = {}
    for base, n_crews in need.items():
        pilots = [e for e, r in crew.items() if r["База"] == base]
        worst = None
        for d0 in DAYS:
            free = sum(1 for e in pilots if d0 not in busy_days[e])
            worst = free if worst is None else min(worst, free)
            if free < n_crews * LIM["crew_size"]:
                viol.append(("VN-060", base, f"{d0}: свободно {free} пилотов, "
                                             f"норма {n_crews * LIM['crew_size']}"))
        reserve_min[base] = worst
    facts["reserve_free_min"] = reserve_min
    facts["days_off_min"] = min(days_off_total.values())
    facts["layovers"] = dict(sorted(layovers.items()))
    facts["open_ends"] = dict(sorted(open_ends.items()))
    facts["windows_unjudged_beyond_horizon"] = windows_unjudged
    facts["duties"] = len(solution["duties"])
    facts["night_duties"] = sum(1 for d in solution["duties"]
                                if bw.is_night_minute(_mins(d["start"][11:16])))
    return viol, facts


# ------------------------------------------------------------------ вывод

def _rows_for_xlsx(solution: dict) -> dict:
    a_cols = [("Табельный номер", "employee"), ("Рейс", "flight"), ("Дата", "date"),
              ("Аэропорт вылета", "dep_airport"), ("Аэропорт прилёта", "arr_airport"),
              ("Время вылета", "dep_time"), ("Время прилёта", "arr_time"), ("Роль", "role"),
              ("База приписки", "base"), ("Смена", "duty_id"), ("Перегон", "deadhead")]
    d_cols = [("Табельный номер", "employee"), ("Идентификатор", "id"), ("Начало", "start"),
              ("Конец", "end"), ("База приписки", "base"), ("Место отдыха", "last_arr_airport"),
              ("Согласие на продление", "extension_consent")]
    return {
        "Назначения": ([c for c, _ in a_cols],
                       [[a[k] for _, k in a_cols] for a in solution["assignments"]]),
        "Смены": ([c for c, _ in d_cols],
                  [[d[k] for _, k in d_cols] for d in solution["duties"]]),
    }


def write(variant: str, solution: dict, viol: list, facts: dict) -> list[Path]:
    import openpyxl
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = {
        "form_version": 1,
        "kit": "Полярис-H0-КИ-окна",
        "variant": variant,
        "book": bw.VARIANTS[BOOK_OF[variant]],
        "expectation": ("допустимо: жёстких нарушений нет" if not EXPECT_VIOLATIONS[variant]
                        else "ровно одно нарушение: " + ", ".join(
                            f"{r} у {e}" for r, e in EXPECT_VIOLATIONS[variant])),
        "expected_violations": [{"rule": r, "employee": e} for r, e in EXPECT_VIOLATIONS[variant]],
        "self_check": {"violations": [{"rule": r, "employee": e, "text": t} for r, e, t in viol],
                       "facts": facts},
        "assignments": solution["assignments"],
        "duties": solution["duties"],
    }
    jpath = OUT_DIR / f"witness-{variant}.json"
    jpath.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    wb = openpyxl.Workbook()
    first = True
    for name, (cols, rows) in _rows_for_xlsx(solution).items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        first = False
        ws.append(cols)
        for row in rows:
            ws.append(row)
    xpath = OUT_DIR / f"witness-{variant}.xlsx"
    wb.save(str(xpath))
    return [jpath, xpath]


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--write", action="store_true", help="записать witness/*.json|xlsx")
    a = ap.parse_args(argv)
    red = False
    for variant in VARIANTS:
        sol = build_solution(variant)
        viol, facts = check(sol)
        got = sorted({(r, e) for r, e, _ in viol})
        want = sorted(EXPECT_VIOLATIONS[variant])
        ok = got == want and not facts["layovers"]
        red = red or not ok
        print(f"[{variant}] книга {bw.VARIANTS[BOOK_OF[variant]]}: рейсов {facts['flights']}, "
              f"смен {facts['duties']}, назначений {facts['assignments']}, ночных смен "
              f"{facts['night_duties']}; резерв min {facts['reserve_free_min']}; выходных min "
              f"{facts['days_off_min']}; ночёвок {facts['layovers'] or 0}, открытых концов "
              f"{sum(facts['open_ends'].values())}; нарушения {got} — "
              f"{'как в сценарии' if ok else 'НЕ КАК В СЦЕНАРИИ, ожидалось ' + str(want)}")
        for r, e, t in viol[:12]:
            print(f"    {r} {e}: {t}")
        if a.write:
            for p in write(variant, sol, viol, facts):
                print("   записано:", p)
    print("самопроверка чиста" if not red else "САМОПРОВЕРКА КРАСНАЯ")
    return 1 if red else 0


if __name__ == "__main__":
    raise SystemExit(main())
