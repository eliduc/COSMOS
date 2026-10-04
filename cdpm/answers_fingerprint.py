# -*- coding: utf-8 -*-
"""Отпечаток ревизии ответов ПО СОДЕРЖАНИЮ (пункт 11.7 проекта «Запись о
прогоне», раздел 12.6 шага 3).

Зачем не байты файла. Корень дерева прогонов — отпечатки ИДАК, данных и
ОТВЕТОВ на анкеты (решение владельца 10.09.2026). Пока отпечатком
ревизии был sha256 БАЙТОВ книги, две книги с одними и теми же ответами
(пересохранённая в Excel, собранная другой сборкой openpyxl, с иным
порядком листов) давали РАЗНЫЕ корни: дерево видело две ветви там, где
ответы одни. И наоборот: корень с ответами у набора нельзя было узнать
до старта, потому что байты книги набора и байты книги, положенной
панелью в ревизию, совпадали не всегда.

Отпечаток по содержанию снимает обе беды: судятся ПАРЫ «вопрос →
ответ», а не носитель. Обе системы (CDPM и COSMOS) считают его одним
каноном — иначе один и тот же набор ответов дал бы два разных корня и
дерево распалось бы надвое.

КАНОН (повторяется у COSMOS буква в букву):

  1. Читаются только листы ЧАСТЕЙ анкеты (`anketa_xlsx.SHEET_TITLES`).
     Служебные листы «Сводка» и «Принятые ответы» и любые чужие листы
     (пометки, расчёты) в отпечаток не входят — панель ответов с них
     тоже НЕ принимает (`answers.read_book`).
  2. Первая строка листа — шапка. Колонка ответа — по заголовку
     «ОТВЕТ АВИАКОМПАНИИ» (`anketa_xlsx.ANSWER_HEADER`), а если
     заголовка нет — шестая (`anketa_xlsx.ANSWER_COL`). Идентификатор
     вопроса — первая колонка строки.
  3. Ответ нормализуется: `str(значение)` → все подряд идущие
     пробельные знаки схлопываются в один пробел → обрезка по краям.
     Пустой ответ ОПУСКАЕТСЯ (пустая клетка — не ответ).
  4. Один вопрос с ДВУМЯ РАЗНЫМИ нормализованными ответами опускается
     целиком: правило чтения панели — «ни один не принят, оставьте
     один». Повтор ОДНОГО и того же ответа (две строки, два листа)
     схлопывается в одну пару.
  5. Пары сортируются по идентификатору вопроса и кладутся списком
     пар: `[[идентификатор, ответ], …]`.
  6. `content_sha256 = sha256(canonical(пары))`, где `canonical` —
     канон словаря записи (`vmp.run_record.canonical`):
     `json.dumps(obj, ensure_ascii=False, sort_keys=True,
     separators=(",", ":")).encode("utf-8")`.
  7. Книга не читается (не .xlsx, обрывок, ревизия без книги —
     «человек») либо ответов в ней нет — отпечатка НЕТ (`None`).
     Отпечаток пустого перечня не выдумывается: «ответов не прочитано»
     и «ответов нет» — разные вещи, и вторая в ревизию не попадает
     (панель отказывает книге без ответов).

Чего модуль НЕ делает: не судит, тому ли изданию вопроса дан ответ
(отпечаток строки, ревизия книги — дело `answers.read_book` на пути
применения) и не чинит книгу. Отпечаток обязан считаться ДО СТАРТА, по
одной лишь книге набора, когда никакой модели свода ещё нет."""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

from cdpm import anketa_xlsx
from vmp.run_record import canonical

#: Пробельные знаки схлопываются в один пробел: перенос строки внутри
#: клетки, двойной пробел и неразрывный пробел не меняют ОТВЕТА.
_SPACES = re.compile(r"\s+", re.UNICODE)


def normalized_answer(value) -> str:
    """Нормализованный ответ канона: пробельные знаки схлопнуты, края
    обрезаны. Пустая клетка (`None`) даёт пустую строку — такая пара в
    отпечаток не входит."""
    if value is None:
        return ""
    return _SPACES.sub(" ", str(value)).strip()


def read_pairs(path) -> dict[str, str]:
    """Пары «вопрос → нормализованный ответ» из заполненной книги
    анкеты. Правила — канон модуля (пункты 1–4).

    Книга не читается — исключение (`openpyxl` / OSError): решение,
    что делать с нечитаемой книгой, принимает вызывающая сторона.
    Отказ строки (два разных ответа) не роняет чтение: вопрос
    опускается, остальные пары целы."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    part_sheets = set(anketa_xlsx.SHEET_TITLES.values())
    raw: dict[str, list[str]] = {}
    try:
        for ws in wb.worksheets:
            if ws.title not in part_sheets:
                continue
            rows = ws.iter_rows(values_only=True)
            head = next(rows, None) or ()
            column = _answer_column(head)
            for row in rows:
                if not row or row[0] is None:
                    continue
                qid = str(row[0]).strip()
                if not qid:
                    continue
                answer = normalized_answer(
                    row[column] if len(row) > column else None)
                if not answer:
                    continue
                if answer not in raw.setdefault(qid, []):
                    raw[qid].append(answer)
    finally:
        wb.close()
    # Пункт 4: два РАЗНЫХ ответа на один вопрос — ни один не принят.
    return {qid: answers[0] for qid, answers in raw.items()
            if len(answers) == 1}


def _answer_column(head) -> int:
    """Номер (с нуля) колонки ответа: по заголовку, иначе шестая."""
    want = anketa_xlsx.ANSWER_HEADER.strip().casefold()
    for j, cell in enumerate(head):
        if str(cell or "").strip().casefold() == want:
            return j
    return anketa_xlsx.ANSWER_COL - 1


def content_sha256(pairs: dict) -> str | None:
    """Отпечаток содержания по парам «вопрос → ответ» (пункты 3, 5, 6).
    Пар нет — отпечатка нет."""
    # Формула — ОДИН носитель у обеих систем: `vmp.run_record.answers_fingerprint`
    # (словарь COSMOS, main 5f9e49bc): те же нормализация, отбрасывание
    # пустых, сортировка по номеру вопроса и sha256 канонического JSON.
    # Здесь — только чтение книги в пары; хеш не копируется.
    from vmp.run_record import answers_fingerprint
    rows = [[str(qid), normalized_answer(answer)]
            for qid, answer in (pairs or {}).items()
            if normalized_answer(answer)]
    return answers_fingerprint(rows) if rows else None


def book_content_sha256(path) -> str | None:
    """Отпечаток содержания книги ответов: пары из книги → канон.
    Книга не читается (не .xlsx, обрывок, файла нет) либо ответов в ней
    нет — `None` со словами в журнал не пишется: читатель (запись о
    прогоне) кладёт `null`, как у ревизии без книги."""
    p = Path(path)
    if not p.is_file():
        return None
    try:
        return content_sha256(read_pairs(p))
    except Exception:                                     # noqa: BLE001
        # Чужая книга, обрывок, книга другой сборки: отпечатка нет.
        # Ронять сборку записи о прогоне из-за одной книги нельзя.
        return None


def cached_book_content_sha256(path) -> str | None:
    """То же с памяткой в процессе: запись о прогоне пересобирается
    после каждого действия панели, а книга ревизии не меняется — по
    второму разу книга не открывается. Ключ памятки — путь, момент
    правки и размер: подменённый файл памяткой не прикрыт."""
    p = Path(path)
    try:
        st = p.stat()
        key = (str(p.resolve()), st.st_mtime_ns, st.st_size)
    except OSError:
        return None
    return _memo(key)


def _memo(key, _cache: dict = {}) -> str | None:      # noqa: B006
    if key not in _cache:
        if len(_cache) > 256:
            _cache.clear()
        _cache[key] = book_content_sha256(key[0])
    return _cache[key]
