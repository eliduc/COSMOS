# -*- coding: utf-8 -*-
"""Excel-издание анкеты CDPM (пакет ЭА-1).

Второе издание анкеты вопросов к авиакомпании — рабочая книга xlsx
для заполнения: лист «Сводка» плюс лист на каждую непустую часть
анкеты, строка на вопрос, пустая колонка «ОТВЕТ АВИАКОМПАНИИ».
Источник истины тот же, что у md-анкеты, — машинная модель
(`cdpm.model`): состав листов, порядок строк и содержимое каждой
ячейки — чистая функция модели. Побайтная стабильность файла xlsx
недостижима (zip-контейнер несёт метки времени), поэтому гейт
(`cdpm.gates.gate_anketa_xlsx`) сверяет содержимое ячеек, а не байты.

Запуск отдельно от приложения:
    python -m cdpm.anketa_xlsx MODEL.json --out КНИГА.xlsx
"""
from __future__ import annotations

import argparse
import hashlib
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cdpm import model as m

# Имя листа xlsx ограничено 31 знаком (Excel не открывает книгу с
# более длинным именем), поэтому тире в именах — без пробелов: полное
# «Часть 3б — постановка ростеринга» в лимит не помещается.
SHEET_TITLES = {
    "1": "Часть 1—противоречия",
    "2": "Часть 2—неоднозначности",
    "3a": "Часть 3а—постановка пейринга",
    "3b": "Часть 3б—постановка ростеринга",
    "4": "Часть 4—несоответствия",
}
SUMMARY_TITLE = "Сводка"
ANSWERED_TITLE = "Принятые ответы"
ANSWER_HEADER = "ОТВЕТ АВИАКОМПАНИИ"
FINGERPRINT_HEADER = "Отпечаток вопроса"
REVISION_LABEL = "Ревизия анкеты"
HEADERS = ("ID вопроса", "Вопрос", "Формат ответа", "Строки свода",
           "Источник", ANSWER_HEADER, FINGERPRINT_HEADER)
# Номера колонок (с единицы): ответ авиакомпании и отпечаток вопроса.
# Отпечаток стоит ПОСЛЕ колонки ответа: положение колонки ответа —
# внешний контракт (заполнители книг, машинная книга регрессии).
ANSWER_COL = 6
FINGERPRINT_COL = 7
ANSWERED_HEADERS = ("ID вопроса", "Вопрос", "Ответ", "Статус",
                    "Ревизия")
# Ширины колонок в знаках; колонка ответа — самая широкая.
COLUMN_WIDTHS = (14, 52, 44, 24, 34, 60, 14)
ANSWERED_WIDTHS = (14, 60, 44, 22, 10)

INSTRUCTION = (
    "Как заполнять: ответы вносите только в колонку «ОТВЕТ "
    "АВИАКОМПАНИИ» на листах частей; для вопроса с закрытым перечнем "
    "вариантов укажите букву выбранного варианта (а, б, …), при "
    "варианте «иное» — добавьте значение и основание; для остальных "
    "форматов пишите ответ в форме, названной в колонке «Формат "
    "ответа». Остальные колонки не редактируйте: их содержимое "
    "собрано из машинной модели и сверяется гейтом. Скрытая колонка "
    "«Отпечаток вопроса» и строка «Ревизия анкеты» сводки привязывают "
    "ответ к изданию вопроса: ответ на переизданный вопрос или из "
    "книги другой ревизии не принимается — скачайте книгу заново.")

# Фиксированные метки docProps: содержимое книги не должно зависеть
# от момента сборки (метки времени zip-архива вне нашей власти).
_FIXED_TS = datetime(2000, 1, 1)

_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _files_by_sid(doc: dict) -> dict[str, str]:
    return {s["sid"]: s["file"] for s in doc.get("sources", [])}


def _answer_format_cell(q: dict) -> str:
    """Колонка «Формат ответа»: формулировки повторяют md-издание
    (`cdpm.render`), чтобы оба издания говорили одними словами."""
    fmt = q["answer_format"]
    ftype = fmt["type"]
    if ftype == "choice":
        lines = ["закрытый перечень — выбрать один вариант:"]
        lines += [f"{o['key']}) {o['text']}"
                  for o in fmt.get("options", [])]
        if fmt.get("note"):
            lines.append(f"примечание: {fmt['note']}")
        return "\n".join(lines)
    label = {"value": "значение",
             "registry": "реестр (приложить перечень)",
             "text": "свободный текст"}[ftype]
    return f"{label} — {fmt['note']}" if fmt.get("note") else label


def question_fingerprint(q: dict) -> str:
    """Отпечаток издания вопроса (аудит 04.09, п. A3.1): идентификатор,
    нормализованные заголовок и текст, тип формата и перечень
    вариантов (ключ + нормализованный текст). Ответ в книге привязан
    к отпечатку: переформулированный вопрос или изменённый перечень
    дают другой отпечаток, и прежний ответ к нему не применяется."""
    if "edition" in q:
        from cdpm import question_editions
        return question_editions.fingerprint(q)
    fmt = q.get("answer_format") or {}
    parts = [str(q.get("id", "")),
             m.normalized_question_text(str(q.get("title", ""))),
             m.normalized_question_text(str(q.get("text", ""))),
             str(fmt.get("type", ""))]
    for o in fmt.get("options", []) or []:
        parts.append(f"{o.get('key', '')}="
                     f"{m.normalized_question_text(str(o.get('text', '')))}")
    digest = hashlib.sha256("\x1f".join(parts).encode("utf-8"))
    return digest.hexdigest()[:12]


def question_row(q: dict, files: dict[str, str], *, doc: dict | None = None) -> tuple:
    """Ячейки строки вопроса: пять заполняемых модулем, шестая («ОТВЕТ
    АВИАКОМПАНИИ») остаётся пустой — она принадлежит авиакомпании,
    седьмая — отпечаток издания вопроса (скрытая колонка)."""
    src = "; ".join(f"{files.get(r['sid'], r['sid'])}, {r['anchor']}"
                    for r in q["source_refs"])
    text = f"{q['title']}\n{q['text']}"
    answer = q.get("answer") or {}
    if answer.get("status") == m.REVISION_CONFLICT_STATUS:
        text += (f"\nПересмотр: конфликт. {answer.get('note', '')}"
                 f"\nНеприменённый ответ сохранён: «{answer.get('raw', '')}».")
    from cdpm import answer_feedback
    for feedback in answer_feedback.lines(q, doc=doc):
        text += "\n" + feedback
    return (q["id"],
            text,
            _answer_format_cell(q),
            ", ".join(q["rule_refs"]),
            src,
            None,
            question_fingerprint(q))


def part_questions(doc: dict, part: str) -> list[dict]:
    """Вопросы части для листа книги: ОТКРЫТЫЕ вопросы текущей ревизии
    (принятые ответы — на своём листе; ИР-1). Для моделей без ответов —
    все вопросы части, как прежде."""
    return m.open_questions(doc, part)


def published_question_row(doc: dict, q: dict, files: dict[str, str]) -> tuple:
    """The complete issued row, shared by the writer and the book gate."""
    from cdpm import question_editions, question_resolutions
    row = list(question_row(q, files, doc=doc))
    for note in question_editions.prior_answer_notes(doc, q):
        row[4] += "; " + note
    for note in question_editions.scope_change_notes(doc, q):
        row[4] += "; " + note
    for note in question_editions.source_change_notes(doc, q, files):
        row[4] += "; " + note
    for note in question_resolutions.notes(q, doc=doc):
        row[4] += "; " + note
    return tuple(row)


def sheet_plan(doc: dict) -> list[tuple[str, list[tuple]]]:
    """Листы частей в порядке `model.PARTS`: (имя листа, строки
    вопросов). Пустая часть листа не получает."""
    files = _files_by_sid(doc)
    plan = []
    for part in m.PARTS:
        rows = []
        for q in part_questions(doc, part):
            rows.append(published_question_row(doc, q, files))
        if rows:
            plan.append((SHEET_TITLES[part], rows))
    return plan


def answered_rows(doc: dict) -> list[tuple]:
    """Строки листа «Принятые ответы» (появляется только у моделей с
    ответами — прежние книги собираются без изменений)."""
    rows = []
    for q in doc.get("questions", []):
        if m.is_open_question(q, doc=doc):
            continue
        from cdpm import question_editions, question_resolutions
        prior = question_editions.prior_answer_notes(doc, q)
        question_text = f"{q['title']}\n{q['text']}"
        if prior:
            question_text += "\nИсточник прежнего ответа: " + "; ".join(prior)
        if question_resolutions.notes(q, doc=doc):
            question_text += "\n" + "; ".join(question_resolutions.notes(q, doc=doc))
        ans = question_resolutions.display_answer(q, doc=doc)
        shown = ans.get("key") or ans.get("raw", "")
        if ans.get("key") and ans.get("note"):
            shown = f"{ans['key']} ({ans['note']})"
        from cdpm import answer_feedback
        feedback = answer_feedback.lines({"answer": ans})
        if feedback:
            shown += "; " + "; ".join(feedback)
        rows.append((q["id"], question_text, shown,
                     ans["status"], ans["revision"]))
    return rows


def summary_rows(doc: dict) -> list[tuple]:
    """Строки листа «Сводка»; пустой кортеж — пустая строка. По этому
    же плану гейт сверяет сводку с моделью."""
    qs = doc.get("questions", [])
    rows: list[tuple] = [
        ("Анкета вопросов к авиакомпании",),
        ("Авиакомпания", doc["meta"]["customer"]),
        # Ревизия модели, для которой издана книга (A3.1): приёмник
        # ответов сверяет её с meta.revision свода.
        (REVISION_LABEL, int(doc.get("meta", {}).get("revision", 1))),
        (),
        ("Вопросов по частям",),
    ]
    for part in m.PARTS:
        n = len(part_questions(doc, part))
        if n:
            rows.append((SHEET_TITLES[part], n))
    n_open = sum(len(part_questions(doc, part)) for part in m.PARTS)
    n_answered = sum(1 for q in qs if not m.is_open_question(q, doc=doc))
    if n_answered or n_open != len(qs):
        from cdpm import question_editions, question_resolutions
        n_code = sum(question_resolutions.complete(q) and not m.is_open_question(q, doc=doc) for q in qs)
        n_retired = sum((q.get("answer") or {}).get("status") == "ожидает обработчика"
                        and not question_resolutions.complete(q)
                        and not m.is_open_question(q, doc=doc)
                        and question_editions.is_superseded(doc, q) for q in qs)
        rows.append(("Открытых вопросов", n_open))
        rows.append(("Принятых ответов", n_answered - n_code - n_retired))
        if n_retired:
            rows.append(("Прежних неприменённых ответов", n_retired))
        if n_code:
            rows.append(("Снято кодом (не ответы авиакомпании)", n_code))
        historical = len(qs) - n_open - n_answered
        if historical:
            rows.append(("Прежних изданий без ответа", historical))
    rows.append(("Всего вопросов", len(qs)))
    if any(q.get("edition_of") for q in qs):
        rows.append(("Смысл счётчиков", "Счётчики учитывают издания вопросов, а не число улаженных предметов; готовность определяется отдельными проверками."))
    rows.append(())
    rows.append((INSTRUCTION,))
    return rows


def _fill_part_sheet(ws, rows: list[tuple]) -> None:
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for row in rows:
        ws.append(row)
    for r in range(2, len(rows) + 2):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=c).alignment = _WRAP
    for i, width in enumerate(COLUMN_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    # Отпечаток — служебная колонка: скрыта от заполняющего, в фильтр
    # не входит; при пересохранении книги Excel её сохраняет.
    ws.column_dimensions[get_column_letter(FINGERPRINT_COL)].hidden = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"


def _fill_summary(ws, doc: dict) -> None:
    rows = summary_rows(doc)
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 44
    last = len(rows)
    # Инструкция — один абзац в объединённой ячейке с переносом;
    # высота задана явно: авто-высоты у объединённых ячеек в Excel
    # нет.
    ws.merge_cells(start_row=last, start_column=1,
                   end_row=last, end_column=2)
    ws.cell(row=last, column=1).alignment = _WRAP
    ws.row_dimensions[last].height = 90


def _fill_answered_sheet(ws, rows: list[tuple]) -> None:
    ws.append(ANSWERED_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for row in rows:
        ws.append(row)
    for r in range(2, len(rows) + 2):
        for c in range(1, len(ANSWERED_HEADERS) + 1):
            ws.cell(row=r, column=c).alignment = _WRAP
    for i, width in enumerate(ANSWERED_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def render_anketa_xlsx(doc: dict, path: str | Path) -> None:
    """Собирает книгу анкеты и пишет её в `path`. Порядок листов:
    «Сводка», непустые части в порядке `model.PARTS`, затем
    «Принятые ответы» (только у моделей с ответами)."""
    wb = Workbook()
    ws = wb.active
    ws.title = SUMMARY_TITLE
    _fill_summary(ws, doc)
    for title, rows in sheet_plan(doc):
        _fill_part_sheet(wb.create_sheet(title), rows)
    answered = answered_rows(doc)
    if answered:
        _fill_answered_sheet(wb.create_sheet(ANSWERED_TITLE), answered)
    props = wb.properties
    props.creator = "CDPM"
    props.created = _FIXED_TS
    props.modified = _FIXED_TS
    wb.save(str(path))


def main(argv: list[str] | None = None) -> int:
    # Печать кириллицы не должна ронять процесс на консоли cp1251
    # (У-14).
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(
        description="Excel-издание анкеты CDPM из машинной модели")
    ap.add_argument("model", help="путь к model.json")
    ap.add_argument("--out", required=True, help="путь к книге .xlsx")
    args = ap.parse_args(argv)
    doc = m.load_model(args.model)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    m.publish_model(doc, args.model)
    render_anketa_xlsx(doc, out)
    plan = sheet_plan(doc)
    n_questions = sum(len(rows) for _, rows in plan)
    print(f"записано: {out.resolve()} (листов {len(plan) + 1}, "
          f"вопросов {n_questions})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
