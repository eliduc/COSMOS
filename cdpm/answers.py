# -*- coding: utf-8 -*-
"""Приёмник ответов анкеты (интерактивный цикл, пакет ИР-1).

Директива владельца 2026-08-19: после получения ответов авиакомпании
должна быть возможность загрузить их и снова проанализировать —
скорректировать формализованный свод и сформировать анкету следующей
ревизии. Этот модуль — первое звено цикла: чтение заполненной
Excel-книги анкеты, разбор ответов и МАШИННО-ОДНОЗНАЧНОЕ применение.

Разделение труда прежнее (fail-closed):

- КОД читает книгу (колонка «ОТВЕТ АВИАКОМПАНИИ») и принимает ответ
  только на то издание вопроса, для которого книга собрана: ревизия
  книги («Сводка») равна ревизии свода, отпечаток строки равен
  отпечатку вопроса модели; иначе, как и при двух разных ответах на
  один вопрос, ответ ложится «не разобран» с причиной и вопрос
  задаётся повторно (аудит 04.09, A3.1/A3.2);
- КОД сопоставляет ответы развилок с вариантами детерминированно
  (буква, «вариант а», текст варианта) и применяет только то, что
  однозначно: разрешение УСЛОВНОГО типа строки (ветвь по выбранному
  ключу, снятие блокировки); ключ с оговоркой («а, но только для…»)
  машинно не применяется — уходит обработчику с сохранением
  блокировок (A3.3);
- МОДЕЛЬ зовётся только для нечётких ответов развилок — предлагает
  ключ варианта, КОД валидирует по перечню; неразборчивый ответ —
  честный статус «не разобран», вопрос в новой анкете задаётся
  повторно;
- ответы, требующие правки текста строк (выбор чтения, значения
  частей 1 и 3а/3б), получают статус «ожидает обработчика» —
  их вносит модельный патч под контрактом (пакет ИР-2); блокировки
  до этого НЕ снимаются;
- каждый принятый ответ записывается в машинную модель
  (question.answer) с номером ревизии; meta.revision растёт за проход.
  Номер может пропускать сохранённые ревизии прежнего извлечения.

Запуск отдельно от приложения:
    python -m cdpm.answers MODEL.json КНИГА.xlsx --out-dir КАТАЛОГ \
        [--llm cloud|local|auto]
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from cdpm import anketa_xlsx
from cdpm import answer_versions, answer_parsing
from cdpm import family_params as family_params_mod
from cdpm import history_section as _hist
from cdpm import llm as llm_mod
from cdpm import model as m
from cdpm.diagnostics import summarize_items

MATCH_CHUNK = 12         # нечётких ответов в одном обращении к модели

class AnswersError(ValueError):
    """Честный отказ приёмника: причина всегда названа."""


def reserve_revision_directory(root: str | Path, current_revision: int
                               ) -> tuple[int, Path]:
    """Reserve a fresh answer-history slot, never reuse an occupied number.

    Re-extraction can reset model revision while the old answer history stays.
    Include every numeric rev-* entry (even empty directories/failed uploads
    or files); exclusive mkdir arbitrates a concurrent claim of the same slot.
    Reserving does not change the model: input-book validation still uses its
    current revision. A failed upload leaves its reserved evidence in place.
    """
    if type(current_revision) is not int or current_revision < 1:
        raise AnswersError("неверная текущая ревизия свода")
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    for _ in range(100):
        latest = current_revision
        for entry in root.iterdir():
            match = re.fullmatch(r"rev-([0-9]+)", entry.name, re.IGNORECASE)
            if match:
                latest = max(latest, int(match.group(1)))
        revision = latest + 1
        directory = root / f"rev-{revision:03d}"
        try:
            directory.mkdir(exist_ok=False)
        except FileExistsError:
            continue
        return revision, directory
    raise AnswersError("не удалось зарезервировать новую ревизию ответов")


_PROMPT_MATCH = """Ты — разборщик ответов авиакомпании на вопросы
анкеты. У каждого вопроса закрытый перечень вариантов; авиакомпания
ответила свободным текстом. По КАЖДОМУ ответу определи выбранный
вариант. Верни СТРОГО один JSON-объект без пояснений:

{{"matches": [{{
  "question": "<id вопроса, ровно как в списке>",
  "key": "<ТОЛЬКО ключ варианта — то, что стоит ПЕРЕД скобкой в
          перечне (например «а»), БЕЗ текста варианта>",
  "note": "<оговорки ответа помимо выбора, если есть>"
}} | {{
  "question": "<id>",
  "unparsed": "<почему ответ нельзя свести к одному варианту>"
}}, …]}}

Правила: ровно один объект на каждый вопрос списка; в key — ТОЛЬКО
сам ключ из перечня (буква или слово перед скобкой), НЕ текст
варианта и НЕ форма «ключ) текст»: живой прогон run-0022 — ключ
«а) Полные экипажи…» с текстом отбраковывается, ключ «а» принимается;
выбор за авиакомпанию не делай — если ответ
не называет вариант однозначно, честно верни unparsed. Исключение:
если ответ прямо говорит, что НИ ОДИН вариант перечня не действует
(строка неприменима, перечень ложный), верни ключ "неприменимо" —
это законный служебный ключ.

Ответы ({n}):
{answers}""" + m.ID_IMMUTABILITY_RULE


# ---------------------------------------------------------------------------
# Чтение книги
# ---------------------------------------------------------------------------

@dataclass
class BookReading:
    """Итог чтения книги ответов. `answers` — принятые к разбору
    ответы {id вопроса: дословный ответ}; `refused` — ответы, которые
    приёмник ОТКАЗАЛСЯ разбирать, с причиной {id: {raw, reason}}:
    книга другой ревизии, переизданный вопрос (отпечаток не сошёлся),
    противоречащие дубли; они попадают в модель статусом «не разобран»
    с причиной — вопрос задаётся повторно. `problems` — все замечания
    чтения словами (для сводки прохода и гейта ответов)."""
    answers: dict[str, str] = field(default_factory=dict)
    problems: list[str] = field(default_factory=list)
    refused: dict[str, dict] = field(default_factory=dict)
    received_at: str | None = None
    book_sha256: str | None = None


def _header_index(head: tuple, title: str) -> int | None:
    """Номер (с нуля) колонки шапки по заголовку; None, если её нет."""
    want = title.strip().casefold()
    for j, h in enumerate(head):
        if str(h or "").strip().casefold() == want:
            return j
    return None


def _book_revision(wb) -> int | None:
    """Ревизия анкеты из листа «Сводка» (строка «Ревизия анкеты»);
    None у книг старого издания, где строки нет."""
    if anketa_xlsx.SUMMARY_TITLE not in wb.sheetnames:
        return None
    ws = wb[anketa_xlsx.SUMMARY_TITLE]
    for row in ws.iter_rows(values_only=True):
        if row and str(row[0] or "").strip() == anketa_xlsx.REVISION_LABEL:
            try:
                return int(str(row[1]).strip()) if len(row) > 1 \
                    and row[1] is not None else None
            except ValueError:
                return None
    return None


def read_book(path: str | Path | bytes, doc: dict) -> BookReading:
    """Чтение заполненной книги анкеты (аудит 04.09, п. A3.1–A3.2).

    Ответ принимается к разбору только если он дан на ТО САМОЕ
    издание вопроса: ревизия книги («Сводка», строка «Ревизия анкеты»)
    равна meta.revision свода и отпечаток строки совпадает с
    отпечатком вопроса модели (`anketa_xlsx.question_fingerprint`).
    Книга старого издания — без строки ревизии или без колонки
    отпечатков — отклоняется целиком: без отпечатка нет
    доказательства, что ответ дан на этот вопрос (fail-closed).
    Два РАЗНЫХ ответа на один вопрос — ни один не принимается.
    Отказы возвращаются с причиной и словами повторяются в
    `problems`; молчаливого применения нет."""
    from openpyxl import load_workbook
    # Bytes input lets the package reader apply these identical guards without
    # writing an imported source book to a temporary file.
    book_bytes = path if isinstance(path, bytes) else Path(path).read_bytes()
    receipt = answer_versions.context(book_sha256=hashlib.sha256(book_bytes).hexdigest())
    wb = load_workbook(BytesIO(book_bytes), read_only=True, data_only=True)
    by_id = {q.get("id"): q for q in doc.get("questions", [])}
    part_sheets = set(anketa_xlsx.SHEET_TITLES.values())
    out = BookReading(**receipt)
    model_rev = int(doc.get("meta", {}).get("revision", 1))
    book_rev = _book_revision(wb)
    if book_rev is None:
        stale = ("книга старого издания: в «Сводке» нет строки "
                 f"«{anketa_xlsx.REVISION_LABEL}» — скачайте книгу "
                 "анкеты заново")
    elif book_rev != model_rev:
        stale = (f"книга собрана для ревизии {book_rev}, а свод — "
                 f"ревизии {model_rev}: вопросы переизданы — скачайте "
                 "книгу анкеты заново")
    else:
        stale = None
    if stale:
        out.problems.append(stale)

    # Проход 1 — сбор строк с ответами по вопросам: решение по
    # вопросу принимается по ВСЕМ его строкам сразу, а не по порядку
    # листов (иначе итог зависел бы от того, какой лист в книге
    # раньше). Сохраняем также напечатанный контекст: v1-отпечаток
    # не охватывал область, источники и примечание формата.
    rows_by_q: dict[str, list[tuple[str, str, str, tuple]]] = {}
    order: list[str] = []
    for ws in wb.worksheets:
        if ws.title in (anketa_xlsx.SUMMARY_TITLE,
                        anketa_xlsx.ANSWERED_TITLE):
            # Штатные листы книги: «Сводка» и «Принятые ответы»
            # (последний панель отдаёт со второго круга — он несёт
            # «ID вопроса», но не колонку ответа и не ответный).
            continue
        rows_iter = ws.iter_rows(values_only=True)
        head = next(rows_iter, None) or ()
        ans_col = _header_index(head, anketa_xlsx.ANSWER_HEADER)
        if ws.title not in part_sheets:
            # Молчаливый пропуск листа — сам по себе находка (живой
            # случай run-0002 ночи 30→31.08: лист «Переподача» с
            # ответами исчез без слов, received молча просел).
            # Ответный лист — тот, что несёт колонку «ОТВЕТ
            # АВИАКОМПАНИИ»; прочие чужие листы (пометки, расчёты)
            # не ответные.
            if ans_col is not None:
                extra = sum(1 for row in rows_iter
                            if row and row[0])
                if extra:
                    out.problems.append(
                        f"лист {ws.title!r} вне частей анкеты — "
                        f"{extra} строк(и) пропущено: ответы "
                        "подаются только на листах частей")
            continue
        if ans_col is None:
            ans_col = anketa_xlsx.ANSWER_COL - 1
        fp_col = _header_index(head, anketa_xlsx.FINGERPRINT_HEADER)
        context_cols = [_header_index(head, name) for name in anketa_xlsx.HEADERS[1:5]]
        for row in rows_iter:
            if not row or row[0] is None:
                continue
            qid = str(row[0]).strip()
            raw = str(row[ans_col]).strip() \
                if len(row) > ans_col and row[ans_col] else ""
            if not raw:
                continue
            if qid not in by_id:
                out.problems.append(f"{ws.title}: ответ на неизвестный "
                                    f"вопрос {qid!r} — пропущен")
                continue
            got_fp = str(row[fp_col]).strip() \
                if fp_col is not None and len(row) > fp_col \
                and row[fp_col] else ""
            if qid not in rows_by_q:
                order.append(qid)
            context = tuple(str(row[c] or "") if c is not None and c < len(row)
                            else None for c in context_cols)
            rows_by_q.setdefault(qid, []).append((ws.title, raw, got_fp, context))
    wb.close()

    def refuse(qid: str, raw: str, reason: str, sheet: str) -> None:
        out.refused[qid] = {"raw": raw, "reason": reason}
        if reason != stale:
            suffix = "" if "не принят" in reason else " — ответ не принят"
            out.problems.append(f"{sheet}: {qid}: {reason}{suffix}")

    # Проход 2 — решение по каждому вопросу.
    for qid in order:
        rows = rows_by_q[qid]
        if stale:
            refuse(qid, rows[0][1], stale, rows[0][0])
            continue
        from cdpm import question_editions
        try:
            want_fp = anketa_xlsx.question_fingerprint(by_id[qid])
        except question_editions.EditionRefusal as exc:
            refuse(qid, rows[0][1], str(exc), rows[0][0])
            continue
        current = [r for r in rows if r[2] == want_fp]
        foreign = [r for r in rows if r[2] != want_fp]
        if not current:
            sheet, raw, got_fp, _context = foreign[0]
            if not got_fp:
                refuse(qid, raw, "в строке нет отпечатка вопроса — "
                                 "книга старого издания, скачайте "
                                 "книгу анкеты заново", sheet)
            else:
                refuse(qid, raw, f"вопрос {qid} переиздан (отпечаток "
                                 f"{got_fp} вместо {want_fp}) — "
                                 f"скачайте книгу анкеты заново", sheet)
            continue
        q = by_id[qid]
        if "edition" not in q or "fingerprint_v1" in q.get("edition", {}):
            # Compatibility with v1 cannot authorize a different printed
            # scope just because its old, narrower fingerprint still matches.
            files = {s["sid"]: s["file"] for s in doc.get("sources", [])}
            expected = anketa_xlsx.published_question_row(doc, q, files)[1:5]
            # Edge whitespace introduced by Excel is not another context.
            # Internal wording, scope and source differences still refuse.
            expected = [str(cell or "").strip() for cell in expected]
            mismatched = [r for r in current
                          if [str(cell or "").strip() for cell in r[3]] != expected]
            if mismatched:
                row = mismatched[0]
                refuse(qid, row[1], "контекст старого издания в книге не совпадает "
                       "с моделью — скачайте совместимую книгу заново", row[0])
                continue
        distinct = []
        for _sheet, raw, _fp, _context in current:
            if raw not in distinct:
                distinct.append(raw)
        if len(distinct) > 1:
            listed = " и ".join(f"«{r}»" for r in distinct)
            refuse(qid, distinct[0], f"в книге два разных ответа "
                                     f"({listed}) — ни один не принят, "
                                     f"оставьте один", current[0][0])
            continue
        out.answers[qid] = distinct[0]
        for sheet, _raw, got_fp, _context in foreign:
            # Лишняя строка того же вопроса из другого издания
            # (скопирована из старой книги): принятый ответ строки
            # текущего издания цел, чужая строка не принята — с
            # причиной, а не «вопрос переиздан» (книга-то свежая).
            out.problems.append(
                f"{sheet}: {qid}: в книге две строки вопроса, одна из "
                f"них — из другого издания (отпечаток "
                f"{got_fp or 'пуст'} вместо {want_fp}); принят ответ "
                f"строки текущего издания, лишнюю строку уберите")
    return out


def read_answer_book(path: str | Path,
                     doc: dict) -> tuple[dict[str, str], list[str]]:
    """Прежняя форма чтения: (принятые ответы, замечания чтения).
    Отказанные ответы (см. `read_book`) в первый элемент не входят,
    их причины — во втором; путь панели пользуется `read_book`, чтобы
    отказы легли в модель статусом «не разобран»."""
    got = read_book(path, doc)
    return got.answers, got.problems


# Вопросы, рождаемые ЛИНИЕЙ ДАННЫХ (после книги данных и сверки): метка
# `-дл-` — промоция проверки в вопрос (`stable_question_id("2",
# "дл-<проверки>")` → `Q-2-дл-…`, data_to_anketa.promote_to_question);
# префиксы — вопросы кода по книге данных (нет данных / значения вне
# перечня / история: history_questions.ensure_history_questions на
# сверке данных). Вопросы о целевой (`Q-3a/3b-GOAL`) тоже строит код, но
# на извлечении — линии данных они не принадлежат (см.
# exchange.CODE_BUILT_QUESTION_PREFIXES).
DATA_LINE_QUESTION_MARK = "-дл-"
DATA_LINE_QUESTION_PREFIXES = ("Q-3a-нет-данных-", "Q-3b-нет-данных-",
                               "Q-3a-данные-", "Q-3b-данные-",
                               "Q-3b-история-")


def is_data_line_question(qid) -> bool:
    """Вопрос рождён линией данных — книга с ответом на него заведомо
    подана после сверки данных (живой случай 08.09 на Test: повтор
    «заново» run-0012 пропускал ответ на `Q-3b-история-нет` как «ответ на
    неизвестный вопрос», и готовность потомка была «б» при «в» у
    источника)."""
    s = str(qid or "")
    return DATA_LINE_QUESTION_MARK in s or s.startswith(DATA_LINE_QUESTION_PREFIXES)


def answered_question_ids(path: str | Path) -> set[str]:
    """Идентификаторы вопросов, на которые книга даёт ответ (непустая
    колонка ответа на листах частей), БЕЗ сверки со сводом и ревизией.
    Нужно восстановлению цепочки повтора (`replay.reconstruct_actions`):
    книга с ответом на вопрос линии данных (`Q-2-дл-…`) заведомо подана
    ПОСЛЕ сверки данных — иначе такого вопроса в своде ещё не было.
    Книга, которую openpyxl не читает, даёт пустое множество: порядок
    остаётся порядком конвейера, а сама книга откажется словами при
    применении."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — любая нечитаемая книга: см. док-строку
        return set()
    part_sheets = set(anketa_xlsx.SHEET_TITLES.values())
    out: set[str] = set()
    for ws in wb.worksheets:
        if ws.title not in part_sheets:
            continue
        rows_iter = ws.iter_rows(values_only=True)
        head = next(rows_iter, None) or ()
        ans_col = _header_index(head, anketa_xlsx.ANSWER_HEADER)
        if ans_col is None:
            ans_col = anketa_xlsx.ANSWER_COL - 1
        for row in rows_iter:
            if not row or row[0] is None:
                continue
            raw = str(row[ans_col]).strip() \
                if len(row) > ans_col and row[ans_col] else ""
            if raw:
                out.add(str(row[0]).strip())
    return out


# ---------------------------------------------------------------------------
# Сопоставление ответов развилок
# ---------------------------------------------------------------------------

def _norm(text: str) -> str:
    return unicodedata.normalize("NFC", (text or "")).strip().casefold()


# Штатный ответ «перечень ложный» (запрос одной из линий
# 2026-08-27):
# закрытый перечень может ЦЕЛИКОМ не действовать (живой случай —
# FRMS-зависимые значения при невнедрённой FRMS); канонические формы
# распознаются кодом, строка-носитель уходит «вне контура» (OUT).
INAPPLICABLE_KEY = "неприменимо"
_INAPPLICABLE_PREFIXES = (
    "строка целиком неприменима", "строка неприменима",
    "ни одно из значений не действует",
    "ни один вариант не действует",
    "ни один из вариантов не действует",
    "перечень ложный", "строка снимается",
    "неприменимо", "неприменима")


def match_choice(raw: str, options: list[dict]) -> tuple[str, str] | None:
    """Детерминированное сопоставление ответа развилки: (ключ,
    оговорка) либо None, если однозначно не выходит. Канонический
    ответ «строка неприменима» (см. _INAPPLICABLE_PREFIXES) даёт
    служебный ключ INAPPLICABLE_KEY."""
    text = _norm(raw)
    for ph in _INAPPLICABLE_PREFIXES:
        if text.startswith(ph):
            note = raw.strip()[len(ph):].strip(" .,:;—-")
            return INAPPLICABLE_KEY, note
    keys = {_norm(o.get("key", "")): o.get("key") for o in options}
    if text in keys:
        return keys[text], ""
    for o in options:
        if text == _norm(o.get("text", "")):
            return o.get("key"), ""
    for o in options:
        k = _norm(o.get("key", ""))
        for prefix in (f"{k})", f"{k}.", f"{k}:", f"вариант {k}",
                       f"({k})"):
            if text.startswith(prefix):
                note = raw.strip()[len(prefix):].strip(" .,;—-")
                return o.get("key"), _tail(note, o)
    # «а — потому что…» / «а, но…»: ключ + разделитель + хвост.
    for o in options:
        k = _norm(o.get("key", ""))
        if text.startswith(k + " ") or text.startswith(k + ",") \
                or text.startswith(k + " —"):
            note = raw.strip()[len(o.get("key", "")):].strip(" .,;—-")
            return o.get("key"), _tail(note, o)
    return None


def _tail(note: str, option: dict) -> str:
    """Хвост ответа после ключа. Хвост, повторяющий текст самого
    варианта («а) допущение верно» — пересказ перечня), — пустой;
    остальное возвращается дословно: пояснение или оговорка, разницу
    между ними ставит `_caveat`. Сравнение детерминированное:
    нормализованное равенство хвоста (без обрамляющих скобок) тексту
    варианта."""
    tail = _norm(note).strip("()[] ")
    if not tail or tail == _norm(option.get("text", "")).strip("()[] "):
        return ""
    return note


# Оговорка — ОГРАНИЧЕНИЕ применимости выбранного варианта (аудит
# 04.09, A3.3): её узнают маркеры ограничения в хвосте ответа. Хвост
# без маркеров — пояснение (обоснование выбора: «а (чтение
# аудитора)», «г (баланс величин: код — справочная величина)»): ключ
# применяется машинно, хвост ложится заметкой ответа. Принцип, не
# заплатка под набор: сформулирован без упоминания вопросов и книг.
_CAVEAT_MARKERS = re.compile(
    r"(?<![\w-])(?:только|кроме|за исключением|но|если|при условии|"
    r"начиная с|в случае|лишь)(?![\w-])")
# Предлоги времени «до»/«после» ограничивают ключ, когда возглавляют
# хвост («а, до 31 декабря», «б (после ввода FRMS)»); внутри
# пояснения они — его собственная грамматика («вплоть до 0», «считать
# до получения данных») и ограничением ответа не являются.
_CAVEAT_HEAD = re.compile(r"^(?:до|после)(?![\w-])")


def _caveat(note: str) -> str:
    """Оговорка из хвоста ответа: сам хвост, если он несёт маркер
    ограничения применимости (см. `_CAVEAT_MARKERS`, `_CAVEAT_HEAD`),
    иначе пустая строка (хвост — пояснение)."""
    tail = _norm(note).strip("()[] ")
    if not tail:
        return ""
    if _CAVEAT_MARKERS.search(tail) or _CAVEAT_HEAD.match(tail):
        return note
    return ""


def _canon_option_key(raw_key, options) -> str | None:
    """Ключ разборщика → канонический ключ перечня БЕЗ гадания
    (живой класс run-0019/0022, тот же, что фантомные номера
    замечаний: модель приписывает к идентификатору лишнее). Формы:
    «а) полный текст варианта» — буква известна и текст сверяется с
    вариантом этой буквы; «полный текст варианта» — однозначное
    совпадение ровно с одним вариантом. Всё прочее — None (честный
    unparsed, fail-closed)."""
    if not isinstance(raw_key, str):
        return None
    k = raw_key.strip()
    keys = [str(o.get("key")) for o in options]

    def norm(s) -> str:
        return " ".join(str(s or "").split()).casefold() \
            .strip(" .,;:—–-")

    mk = re.match(r"^\s*(\S{1,8}?)\s*\)\s*(.*)$", k, re.S)
    if mk and mk.group(1).strip() in keys:
        key = mk.group(1).strip()
        o = next(o for o in options if str(o.get("key")) == key)
        tail = mk.group(2).strip()
        if not tail or norm(tail) == norm(o.get("text")):
            return key
        return None
    hits = [str(o.get("key")) for o in options
            if norm(o.get("text")) == norm(k)]
    return hits[0] if len(hits) == 1 else None


def _match_with_llm(pending: list[tuple[dict, str]], llm,
                    control=None) -> dict[str, dict]:
    """Нечёткие ответы развилок — модели чанками; КОД валидирует ключи
    по перечню вопроса. Возвращает {id вопроса: {key, note} |
    {unparsed}}. Молчание канала — честный unparsed (fail-closed)."""
    out: dict[str, dict] = {}
    for start in range(0, len(pending), MATCH_CHUNK):
        if control:
            control()          # пауза/остановка между чанками (ПС-1)
        batch = pending[start:start + MATCH_CHUNK]
        listing = "\n".join(
            f"{q['id']} | варианты: "
            + "; ".join(f"{o['key']}) {o['text']}"
                        for o in q['answer_format'].get('options', []))
            + f" | ответ: {raw}"
            for q, raw in batch)
        prompt = _PROMPT_MATCH.format(n=len(batch), answers=listing)
        reply = None
        for attempt in (1, 2):
            ask = prompt if attempt == 1 else (
                prompt + "\n\nПредыдущий ответ не разобрался. Верни "
                         "СТРОГО один JSON-объект и ничего больше.")
            try:
                reply = llm_mod.parse_json_reply(llm(ask))
                break
            except llm_mod.LLMReplyError:
                continue
        got = {r.get("question"): r
               for r in (reply or {}).get("matches", [])}
        for q, raw in batch:
            r = got.get(q["id"])
            keys = [o.get("key")
                    for o in q["answer_format"].get("options", [])]
            if r is None:
                out[q["id"]] = {"unparsed": "разборщик не дал ответа "
                                            "по вопросу"}
            elif r.get("unparsed"):
                out[q["id"]] = {"unparsed": str(r["unparsed"])}
            elif r.get("key") in keys \
                    or r.get("key") == INAPPLICABLE_KEY:
                out[q["id"]] = {"key": r["key"],
                                "note": (r.get("note") or "").strip()}
            else:
                fixed = _canon_option_key(
                    r.get("key"),
                    q["answer_format"].get("options", []))
                if fixed is not None:
                    out[q["id"]] = {"key": fixed,
                                    "note": (r.get("note")
                                             or "").strip()}
                else:
                    out[q["id"]] = {"unparsed": f"разборщик назвал "
                                                f"ключ {r.get('key')!r}"
                                                f" вне вариантов "
                                                f"{keys}"}
    return out


# ---------------------------------------------------------------------------
# Применение
# ---------------------------------------------------------------------------

def _bind_soft_to_objective(doc: dict, rule: dict) -> None:
    """Строка стала SOFT и осталась без блокировок: вес — только из
    ответа, поэтому привязка к ЕДИНОМУ вопросу зоны о целевой и весах
    (Q-*-GOAL, склейка ЦФ-1 2026-08-21; создаётся общим строителем,
    если его ещё нет) — как в сборке извлечения."""
    parts = []
    if rule["zone"] in ("PAIRING", "BOTH"):
        parts.append("3a")
    if rule["zone"] in ("ROSTERING", "BOTH"):
        parts.append("3b")
    for part in parts:
        m.ensure_goal_question(doc, part)


# ── Журнал следствий ответа (аудит 04.09, п. A3.4) ──
#
# Следствие ответа — разница состояния свода до и после его применения,
# поле за полем по строкам и целевым: журнал не зависит от ветви
# применения (развилка, «неприменимо», «нет данных», оговорка, патч
# обработчика) и потому полон по построению. По журналу пересмотр
# возвращает строки к виду «до», а не пере-применяет ответ поверх.
_ABSENT = object()
_OBJECTIVES_KEY = "__objectives__"


def _snapshot(doc: dict) -> dict:
    rules = {r.get("id"): copy.deepcopy(r)
             for r in doc.get("rules", []) if isinstance(r, dict)}
    objs = {(o.get("id") or f"#{i}"): copy.deepcopy(o)
            for i, o in enumerate(doc.get("objectives") or [])
            if isinstance(o, dict)}
    return {"rules": rules, "objectives": objs}


def _diff_entries(before: dict, after: dict, rid: str) -> list[dict]:
    out = []
    for field in sorted(set(before) | set(after), key=str):
        b, a = before.get(field, _ABSENT), after.get(field, _ABSENT)
        if b is not _ABSENT and a is not _ABSENT and b == a:
            continue
        e = {"rule": rid, "field": str(field),
             "before": None if b is _ABSENT else copy.deepcopy(b),
             "after": None if a is _ABSENT else copy.deepcopy(a)}
        if b is _ABSENT:
            e["before_absent"] = True
        if a is _ABSENT:
            e["after_absent"] = True
        out.append(e)
    return out


def _diff_effects(snap: dict, doc: dict) -> list[dict]:
    """effects = [{rule, field, before, after}] — что ответ изменил в
    строках свода и целевых. Поле «*» — строка появилась или исчезла
    целиком; строка «__objectives__» — целевые (поле = id целевой)."""
    effects: list[dict] = []
    now = {r.get("id"): r for r in doc.get("rules", []) if isinstance(r, dict)}
    for rid in sorted(set(snap["rules"]) | set(now), key=str):
        before, after = snap["rules"].get(rid), now.get(rid)
        if before is None or after is None:
            e = {"rule": str(rid), "field": "*",
                 "before": copy.deepcopy(before), "after": copy.deepcopy(after)}
            if before is None:
                e["before_absent"] = True
            if after is None:
                e["after_absent"] = True
            effects.append(e)
            continue
        effects += _diff_entries(before, after, str(rid))
    now_o = {(o.get("id") or f"#{i}"): o
             for i, o in enumerate(doc.get("objectives") or []) if isinstance(o, dict)}
    effects += _diff_entries(snap["objectives"], now_o, _OBJECTIVES_KEY)
    return effects


def _journal(q: dict, snap: dict, doc: dict, revision: int) -> None:
    """Записать следствия в ответ вопроса, если ответ поставлен этой
    ревизией и что-то применял (статусы «применён» / «ожидает
    обработчика» — блокировки тоже следствие)."""
    ans = q.get("answer")
    if not isinstance(ans, dict) or ans.get("revision") != revision:
        return
    if ans.get("status") not in ("применён", "ожидает обработчика", m.PARTIAL_ANSWER_STATUS):
        return
    ans["effects"] = _diff_effects(snap, doc)


def _restore_effects(doc: dict, effects: list[dict], qid: str) -> None:
    """Вернуть свод к виду «до» по журналу: записи в обратном порядке;
    строка из журнала, которой нет в своде, — честный отказ."""
    for e in reversed(effects):
        rid, field = e.get("rule"), e.get("field")
        before_absent = bool(e.get("before_absent"))
        before = copy.deepcopy(e.get("before"))
        if rid == _OBJECTIVES_KEY:
            objs = doc.setdefault("objectives", [])
            idx = next((i for i, o in enumerate(objs)
                        if isinstance(o, dict) and (o.get("id") or f"#{i}") == field), None)
            if before_absent:
                if idx is not None:
                    objs.pop(idx)
            elif idx is None:
                objs.append(before)
            else:
                objs[idx] = before
            continue
        by_id = {r.get("id"): r for r in doc.get("rules", []) if isinstance(r, dict)}
        if field == "*":
            if before_absent:
                doc["rules"] = [r for r in doc.get("rules", []) if r.get("id") != rid]
            elif rid in by_id:
                by_id[rid].clear()
                by_id[rid].update(before)
            else:
                doc.setdefault("rules", []).append(before)
            continue
        rule = by_id.get(rid)
        if rule is None:
            raise AnswersError(f"пересмотр {qid}: строки {rid!r} из журнала "
                               f"следствий в своде нет — восстановление «до» "
                               f"невозможно")
        if before_absent:
            rule.pop(field, None)
        else:
            rule[field] = before


def _answer_targets(q: dict, doc: dict | None = None) -> set:
    """Идентификаторы, за которые действует ответ на вопрос: сам вопрос,
    у вопроса-пересмотра — корневой вопрос и вся цепочка его пересмотров
    (после отката в blockers строки стоят и корень, и прежние -R):
    развилка корня и все эти блокировки — цели ответа."""
    out = {q.get("id"), q.get("revision_of")} - {None}
    if doc is not None:
        root = q.get("revision_of") or q.get("id")
        out |= {x.get("id") for x in doc.get("questions", [])
                if isinstance(x, dict) and x.get("revision_of") == root}
    return out


def _root_question_id(q: dict) -> str:
    return str(q.get("revision_of") or q.get("id"))


def _revision_unsupported(root: str) -> str | None:
    """Семьи, чей ответ пересматривается не машинным откатом, а добором
    своей семьи (вопрос следующей ревизии): целевая и веса (GOAL, GOAL-T),
    машинный критерий пейринга, носители текста условий расчёта. Их
    следствия живут вне строк свода (целевые, экспорт условий), и откат по
    журналу строк был бы ложным."""
    if m.is_goal_family(root):
        return ("ответ о целевой и весах пересматривается добором семьи "
                "GOAL (вопрос следующей ревизии), а не машинным откатом")
    if root == m.CRITERION_QUESTION_ID or root.startswith("Q-3a-критерий-пейринга"):
        return ("ответ о машинном критерии пересматривается добором "
                "критерия, а не машинным откатом")
    if _hist.is_history_question(root):
        return ("решение по истории сотрудников пересматривается повторной "
                "загрузкой данных и новым вопросом линии данных, а не машинным "
                "откатом (следствия живут в пакете, не в строках)")
    if root.startswith(("Q-3a-условия-", "Q-3b-условия-")):
        return ("ответ-носитель текста условий расчёта читает экспорт "
                "условий напрямую — пересмотр таким же вопросом следующей "
                "ревизии через обработчика, машинного отката нет")
    return None


def _revert_conflicts(doc: dict, effects: list[dict]) -> list[str]:
    """Сухой прогон отката: поля журнала, которые с момента ответа
    изменил кто-то ещё (текущее значение ≠ «после» журнала), и строки,
    которых в своде нет. Непусто — откатывать нельзя молча."""
    trial = copy.deepcopy(doc)
    conflicts: list[str] = []
    for e in reversed(effects):
        rid, field = e.get("rule"), e.get("field")
        if rid == _OBJECTIVES_KEY:
            objs = trial.get("objectives") or []
            cur = next((o for i, o in enumerate(objs) if isinstance(o, dict)
                        and (o.get("id") or f"#{i}") == field), _ABSENT)
        else:
            by_id = {r.get("id"): r for r in trial.get("rules", []) if isinstance(r, dict)}
            if field == "*":
                cur = by_id.get(rid, _ABSENT)
            elif rid not in by_id:
                conflicts.append(f"строки {rid} нет в своде")
                continue
            else:
                cur = by_id[rid].get(field, _ABSENT)
        expect_absent = bool(e.get("after_absent"))
        if expect_absent != (cur is _ABSENT) \
                or (cur is not _ABSENT and cur != e.get("after")):
            conflicts.append(f"{rid}.{field}: значение изменено позже другим "
                             f"ответом или обработчиком")
            continue
        try:
            _restore_effects(trial, [e], "сухой прогон")
        except AnswersError as exc:
            conflicts.append(str(exc))
    return conflicts


def open_revision(doc: dict, qid: str) -> dict:
    """Пересмотр применённого ответа (право авиакомпании передумать):
    следствия прежнего ответа снимаются по журналу (строки — к виду
    «до», развилка возвращается), строки блокируются вопросом-пересмотром
    `<корень>-R`, прежний ответ получает статус «пересмотрен» со ссылкой на
    него; ответ на -R применяется штатными ветвями с реальным эффектом
    (`revision_of` указывает КОРНЕВОЙ вопрос — и у пересмотра пересмотра).
    Откат не делается, а -R создаётся с честной пометкой, если журнала
    нет (ответ принят до его появления) или поля журнала с тех пор менял
    кто-то ещё (следствия перекрыты) — тогда следствия заменяет
    обработчик под контрактом. Семьи целевой/критерия/условий машинно не
    пересматриваются — отказ словами."""
    q = next((x for x in doc.get("questions", []) if x.get("id") == qid), None)
    if q is None:
        raise AnswersError(f"вопроса {qid!r} нет")
    if "parameter_contract" in q:
        from cdpm import question_editions
        try:
            question_editions.require_valid(q)
        except question_editions.EditionRefusal as exc:
            raise AnswersError(str(exc)) from exc
    ans = q.get("answer") or {}
    retry_conflict = (m.is_revision_question(q)
                      and ans.get("status") == m.REVISION_CONFLICT_STATUS)
    if ans.get("status") not in ("применён", m.PARTIAL_ANSWER_STATUS) and not retry_conflict:
        raise AnswersError(f"{qid}: пересматривать можно только ПРИМЕНЁННЫЙ "
                           f"ответ — открытый вопрос отвечается прямо")
    root = _root_question_id(q)
    why_not = _revision_unsupported(root)
    if why_not:
        raise AnswersError(f"{qid}: {why_not}")
    effects = None if retry_conflict else ans.get("effects")
    conflicts: list[str] = []
    restored = False
    if isinstance(effects, list):
        conflicts = _revert_conflicts(doc, effects)
        if not conflicts:
            _restore_effects(doc, effects, qid)
            restored = True
    existing = {x.get("id") for x in doc.get("questions", [])}
    base = f"{root}-R"
    rid_new, n = base, 1
    while rid_new in existing:
        n += 1
        rid_new = f"{base}-{n}"
    touched = sorted({str(e["rule"]) for e in (effects or [])
                      if e.get("rule") != _OBJECTIVES_KEY
                      and (e.get("field") != "*" or not e.get("before_absent"))}
                     | {str(r) for r in (q.get("rule_refs") or [])})
    blocked = []
    for rule in doc.get("rules", []):
        if rule.get("id") in touched:
            bl = rule.setdefault("blockers", [])
            if rid_new not in bl:
                bl.append(rid_new)
            blocked.append(rule["id"])
    prev_raw = str(ans.get("raw") or "")
    if retry_conflict:
        tail = (" Предыдущий пересмотр встретил конфликт. Текущие значения "
                "сохранены без отката; новый вопрос фиксирует их нынешнее "
                "состояние. Перед ответом сверьте поздние изменения.")
    elif restored:
        tail = (f" Следствия прежнего ответа сняты по журналу ({len(effects)} "
                f"изменений): строки возвращены к виду до него и заблокированы "
                f"этим вопросом; новый ответ применится с реальным эффектом.")
    elif conflicts:
        tail = (" Следствия прежнего ответа перекрыты позднейшими изменениями ("
                + summarize_items(conflicts, limit=3, separator="; ") + "): вернуть строки к виду «до» "
                "машинно нельзя — строки заблокированы этим вопросом, новый "
                "ответ применит обработчик, заменив следствия под контрактом.")
    else:
        tail = (" Прежний ответ применён без журнала следствий (принят до его "
                "появления): вернуть строки к виду «до» машинно нельзя — новый "
                "ответ применит обработчик, заменив следствия под контрактом.")
    previous = (f"Прежний ответ на {qid} НЕ применён из-за конфликта («{prev_raw}»). "
                if retry_conflict else
                f"Прежний ответ на {qid} принят и применён («{prev_raw}»). ")
    doc.setdefault("questions", []).append({
        "id": rid_new, "part": q.get("part"),
        "title": f"Пересмотр: {q.get('title')}",
        "text": (previous +
                 f"Авиакомпания пересматривает решение — дайте новый ответ; "
                 f"его применение заменит следствия прежнего установленным "
                 f"порядком." + tail),
        "answer_format": copy.deepcopy(q.get("answer_format") or {"type": "value"}),
        "rule_refs": list(q.get("rule_refs") or []),
        "source_refs": [dict(s) for s in (q.get("source_refs") or [])],
        "revision_of": root,
    })
    ans["status"] = "пересмотрен"
    if "parameter_contract" in q:
        # A reconsideration must not silently reinterpret the old question
        # under today's catalogue. Reissuing meaning is a separate operation.
        doc["questions"][-1]["parameter_contract"] = copy.deepcopy(q["parameter_contract"])
        question_editions.freeze(doc["questions"][-1])
    ans["revised_by"] = rid_new
    # K04 (§2.4 (а)): ячейки посадок, внесённые прежним ответом и
    # ПЕРЕЖИВШИЕ откат (журнала нет либо он перекрыт), помечаются
    # открытым пересмотром. Пока признак стоит, прежнее значение
    # решением авиакомпании не является: готовность затронутого режима
    # ниже «в», выпуск отказывает, гейт производителя красный — и
    # синхронизация не снимает пересмотр как лишний.
    pending = m.mark_pending_revision(doc, root, rid_new)
    from cdpm import revision_cells
    revision_cells.capture(doc, doc["questions"][-1])
    return {"created": rid_new, "restored": restored, "conflicts": conflicts,
            "effects_reverted": len(effects or []) if restored else 0,
            "blocked": blocked, "pending_cells": pending}


def resolve_source(doc: dict, cell_source: dict, applying=()) -> tuple:
    """Обратный ход происхождения значения ячейки посадки (контракт П0
    §2.3 п. 1): по объекту `source` возвращает (вопрос, дословный `raw`
    ТОЙ САМОЙ версии ответа), из которой значение прочитано.

    Поиск — по ТОЧНОЙ идентичности (идентификатор издания вопроса,
    отпечаток издания, отпечаток версии ответа), а не «первый
    применённый по порядку списка»; любое несовпадение — отказ
    `provenance.ProvenanceRefusal` с кодом `answer-provenance` и
    словами, называющими вопрос и оба отпечатка. Закон живёт в
    `cdpm.provenance`; здесь — имя, которым его зовёт контракт."""
    from cdpm import provenance as _prov
    return _prov.resolve_source(doc, cell_source, applying=applying)


def apply_answers(doc: dict, raw_map: dict[str, str], llm=None,
                  control=None, refused: dict[str, dict] | None = None,
                  *, received_at: str | None = None, book_sha256: str | None = None,
                  revision: int | None = None, restore: bool = False
                  ) -> dict:
    """Разбор и применение ответов к модели (на месте). Возвращает
    сводку прохода. Машинно применяется только однозначное: разрешение
    условного типа по выбранной ветви БЕЗ оговорки; ответ с оговоркой
    («а, но только для…») и остальное — «ожидает обработчика» с
    сохранением блокировок. `refused` — ответы, отклонённые чтением
    книги (`read_book`: чужая ревизия, переизданный вопрос, дубли):
    ложатся в модель статусом «не разобран» с причиной, вопрос
    задаётся повторно.

    `restore` — восстановление ответа ПРЕЖНЕЙ ревизии сквозь
    переизвлечение (`answer_carry`, РВ-140 п. 1): ревизия и квитанция
    книги — исходные (их сверяет гейт выпуска `source-answer-evidence`),
    поэтому ревизия может не превышать текущую — ответ «Истории» той же
    ревизии ложится позже, на линии данных."""
    invalid_versions = (answer_versions.document_problems(doc)
                        + answer_parsing.document_problems(doc))
    if invalid_versions:
        raise AnswersError("; ".join(invalid_versions))
    current_revision = int(doc.get("meta", {}).get("revision", 1))
    if revision is None:
        revision = current_revision + 1
    if type(revision) is not int or (
            revision < 2 if restore else revision <= current_revision):
        raise AnswersError("новая ревизия ответов должна быть больше текущей")
    receipt = answer_versions.context(received_at, book_sha256)
    questions = {q.get("id"): q for q in doc.get("questions", [])}
    previous_answers = {qid: q.get("answer") for qid, q in questions.items()}
    refused = dict(refused or {})

    parsed: dict[str, dict] = {}
    pending_llm: list[tuple[dict, str]] = []
    for qid, raw in raw_map.items():
        q = questions.get(qid)
        if q is None:
            continue
        from cdpm import question_editions
        edition_problems = question_editions.application_problems(doc, q)
        if edition_problems:
            refused[qid] = {"raw": raw, "reason": "; ".join(edition_problems)}
            continue
        if (q.get("answer") or {}).get("status") == m.PARTIAL_ANSWER_STATUS:
            parsed[qid] = {"partial_reanswer": (
                "ответ применён частично; прежний текст и его следствия сохранены. "
                "Продолжите обработку оставшихся строк; для изменения ответа "
                "сначала откройте пересмотр вопроса")}
            continue
        if q.get("answer") and not m.is_open_question(q):
            parsed[qid] = {"unparsed": "вопрос уже отвечен в ревизии "
                                       f"{q['answer'].get('revision')} "
                                       f"— повторный ответ не применён"}
            continue
        fmt = q.get("answer_format", {})
        if fmt.get("type") == "choice":
            got = match_choice(raw, fmt.get("options", []))
            if got is not None:
                parsed[qid] = {"key": got[0], "note": got[1]}
            else:
                pending_llm.append((q, raw))
        else:
            parsed[qid] = {"freeform": True}
    if pending_llm:
        if llm is None:
            for q, _raw in pending_llm:
                parsed[q["id"]] = {"unparsed": "ответ не сопоставился "
                                               "с вариантами, а канал "
                                               "разборщика не задан"}
        else:
            parsed.update(_match_with_llm(pending_llm, llm, control))

    summary = {"revision": revision, "source_revision": current_revision,
               "received": len(set(raw_map) | set(refused)),
               "applied": 0, "awaiting": 0, "unparsed": 0,
               "resolved_rules": [], "already_answered": 0,
               "refused": []}
    # Отказы чтения книги (A3.1/A3.2): в модель — «не разобран» с
    # причиной; уже принятый ответ прежней ревизии не затирается.
    for qid, rej in sorted(refused.items()):
        q = questions.get(qid)
        if q is None or qid in parsed:
            continue
        from cdpm import question_resolutions
        if question_resolutions.complete(q) or any(x.get("edition_of") == qid for x in questions.values()):
            # A stale book must not add/replace an answer in historical scope.
            summary["unparsed"] += 1
            summary["refused"].append({"question": qid, "reason": str(
                rej.get("reason") or "вопрос уже переиздан")})
            continue
        if q.get("answer") and (not m.is_open_question(q)
                or q["answer"].get("status") == m.PARTIAL_ANSWER_STATUS):
            summary["already_answered"] += 1
            if q["answer"].get("status") == m.PARTIAL_ANSWER_STATUS:
                summary["refused"].append({"question": qid, "reason": (
                    "ответ применён частично; прежний текст и следствия сохранены. "
                    "Для изменения ответа сначала откройте пересмотр вопроса; "
                    + str(rej.get("reason") or "ответ отклонён чтением книги"))})
            continue
        reason = str(rej.get("reason") or "ответ отклонён чтением книги")
        q["answer"] = {"raw": str(rej.get("raw") or "") or "—",
                       "status": "не разобран", "revision": revision,
                       "note": reason}
        summary["unparsed"] += 1
        summary["refused"].append({"question": qid, "reason": reason})
    # Журнал следствий (A3.4): снимок свода перед каждым вопросом, разница
    # после — в его ответ; закрывается перед следующим вопросом и в конце.
    pending_journal: list = []

    def _close_journal() -> None:
        if pending_journal:
            pq, snap = pending_journal.pop()
            _journal(pq, snap, doc, revision)

    for qid, verdict in sorted(parsed.items()):
        q = questions[qid]
        raw = raw_map[qid]
        _close_journal()
        if "partial_reanswer" in verdict:
            summary["already_answered"] += 1
            summary["refused"].append({"question": qid,
                                       "reason": verdict["partial_reanswer"]})
            continue
        pending_journal.append((q, _snapshot(doc)))
        # Вопрос-пересмотр (-R) действует за корневой: развилки и
        # блокировки корня и прежних пересмотров — его цели.
        targets = _answer_targets(q, doc)
        if "unparsed" in verdict:
            if "уже отвечен" in verdict["unparsed"]:
                summary["already_answered"] += 1
                continue
            q["answer"] = {"raw": raw, "status": "не разобран",
                           "revision": revision}
            if verdict["unparsed"]:
                q["answer"]["note"] = verdict["unparsed"]
            summary["unparsed"] += 1
            continue
        if family_params_mod.is_param_question(
                family_params_mod._root_id(q)):
            # Полнота параметров семьи (09.09.2026): вопрос строит код,
            # ответ применяет код — значение в единице параметра садится
            # в посадку первого носителя семьи, «не применяется» —
            # решение ответом; неразборчивое — «не разобран» (спросить
            # снова), догадок нет. Внутри журнала следствий: изменение
            # посадки — следствие ответа, пересмотр откатывает его.
            res = family_params_mod.apply_param_answer(
                doc, q, raw, revision, targets)
            if res.get("applied"):
                summary["applied"] += 1
                summary["resolved_rules"] += list(res.get("touched") or [])
            elif res.get("partial"):
                summary["partial"] = summary.get("partial", 0) + 1
                summary["resolved_rules"] += [
                    r["id"] for r in _affected_rules(doc, q, include_processed=True)
                    if not targets.intersection(r.get("blockers", []))]
            elif (q.get("answer") or {}).get("status") == m.REVISION_CONFLICT_STATUS:
                summary["conflicted"] = summary.get("conflicted", 0) + 1
            else:
                summary["unparsed"] += 1
            continue
        if verdict.get("freeform"):
            # Вопросы условий расчёта — НОСИТЕЛИ ТЕКСТА: их ответы
            # читает экспорт условий напрямую из raw строгим
            # разбором, модельный обработчик им не нужен вовсе —
            # применяются машинно (живой урок 27-28.08: применение
            # ответа о размахе трижды падало на мёртвом канале
            # локальной модели, хотя канал ему не был нужен).
            if str(qid).startswith("Q-3a-условия-") \
                    or str(qid).startswith("Q-3a-критерий-пейринга"):
                q["answer"] = {"raw": raw, "status": "применён",
                               "revision": revision,
                               "note": "носитель текста: значение "
                                       "читает экспорт условий "
                                       "напрямую, обработчик не "
                                       "требуется"}
                summary["applied"] += 1
                continue
            # Значения и перечни (части 1/3а/3б и не-развилки): вносит
            # обработчик под контрактом (ИР-2); блокировки на месте.
            q["answer"] = {"raw": raw, "status": "ожидает обработчика",
                           "revision": revision}
            summary["awaiting"] += 1
            continue
        key, note = verdict["key"], verdict.get("note", "")
        if _hist.is_history_question(qid):
            # Три исхода по истории (1.17.2, решение владельца 05.09): ответ —
            # решение авиакомпании, строки не меняет; его читают экспорт и гейт
            # пакета (manifest.history.user_decision), потребитель называет
            # допущение словами. Ветка стоит ДО «перечень неприменим»: ответ
            # «неприменимо» на вопрос о данных строки из контура не выводит и
            # решением не является — «не разобран», вопрос задаётся снова.
            if key not in ("а", "б"):
                q["answer"] = {"raw": raw, "status": "не разобран",
                               "revision": revision,
                               "note": "ответ на вопрос об истории — только «а» "
                                       "или «б»; «неприменимо» решением не является"}
                summary["unparsed"] += 1
                continue
            q["answer"] = {"raw": raw, "status": "применён", "key": key,
                           "revision": revision}
            if note:
                q["answer"]["note"] = note
            summary["applied"] += 1
            continue
        data_scope_choice = str(qid).startswith(("Q-3a-данные-", "Q-3b-данные-"))
        missing_data_choice = str(qid).startswith(("Q-3a-нет-данных-", "Q-3b-нет-данных-"))
        if (data_scope_choice or missing_data_choice) and key not in ("а", "б"):
            q["answer"] = {"raw": raw, "status": "не разобран",
                           "revision": revision,
                           "note": ("для решения о недостающих данных выберите «а» или «б»"
                                    if missing_data_choice else "для значений вне перечня выберите «а» или «б»")}
            summary["unparsed"] += 1
            continue
        if key == INAPPLICABLE_KEY:
            # «Перечень ложный»: строка-носитель уходит вне контура
            # (OUT — экспорт её не шлёт), блокировки этого вопроса
            # снимаются; всё со следом (запрос одной из линий
            # 2026-08-27).
            removed = []
            for rule in doc.get("rules", []):
                cond = rule.get("conditional")
                if cond and cond.get("question") in targets:
                    rule["kind"] = "OUT"
                    del rule["conditional"]
                    m.align_fit_with_kind(doc, rule,
                                          "перечень неприменим")
                    rule.pop("weight", None)
                    extra = (f"строка снята ответом авиакомпании на "
                             f"{qid}: перечень признан неприменимым")
                    rule["note"] = ((rule.get("note", "") + "; "
                                     + extra).strip("; "))
                    removed.append(rule["id"])
                if targets & set(rule.get("blockers", [])):
                    rule["blockers"] = [b for b in rule["blockers"]
                                        if b not in targets]
                    if not rule["blockers"]:
                        rule.pop("blockers", None)
            q["answer"] = {"raw": raw, "status": "применён",
                           "key": INAPPLICABLE_KEY,
                           "revision": revision}
            if note:
                q["answer"]["note"] = note
            summary["applied"] += 1
            summary["resolved_rules"] += removed
            continue
        if missing_data_choice and not _caveat(note):
            # Правило без данных (решение владельца 2026-09-02): «а» —
            # считать до получения данных: строки-носители получают
            # пометку excluded_until_data и разблокируются (сертификат
            # помечает их непроверенными); «б» — ждать данных: блокировка
            # остаётся (fail-closed), расчёт не начинается.
            touched = []
            from cdpm import question_resolutions
            active_refs = set(q.get("rule_refs") or []) - question_resolutions.resolved_targets(q)
            for rule in doc.get("rules", []):
                if rule.get("id") not in active_refs:
                    continue
                if key == "а":
                    rule["excluded_until_data"] = {"question": qid,
                                                   "revision": revision}
                    rule["blockers"] = [b for b in rule.get("blockers", [])
                                        if b not in targets]
                    if not rule["blockers"]:
                        rule.pop("blockers", None)
                    extra = (f"исключено из расчёта до получения данных "
                             f"по решению авиакомпании ({qid})")
                    rule["note"] = ((rule.get("note", "") + "; " + extra)
                                    .strip("; "))
                    touched.append(rule["id"])
            q["answer"] = {"raw": raw, "status": "применён", "key": key,
                           "revision": revision}
            if key == "а":
                q["answer"]["note"] = ("исключены из расчёта до получения данных: "
                                       + ", ".join(touched)
                                       + "; сертификат помечает эти правила непроверенными")
            else:
                q["answer"]["note"] = ("ждём данных: строки-носители "
                                       "остаются заблокированными до "
                                       "поступления данных")
            if note:
                q["answer"]["note"] = (q["answer"].get("note", "") + "; "
                                       + note).strip("; ")
            summary["applied"] += 1
            summary["resolved_rules"] += touched
            continue
        caveat = _caveat(note)
        if caveat or data_scope_choice:
            # Choosing to extend a closed list or exclude data has an effect;
            # a bare choice key is not proof that this effect was implemented.
            # Аудит 04.09, п. A3.3: ключ с оговоркой («а, но только
            # для дальних рейсов») машинно НЕ применяется — оговорка
            # есть ограничение применимости, которое код истолковать
            # не может, а заметкой в листе «Принятые ответы» оно
            # терялось для ВМП. Развилка и блокировки остаются, строки
            # вопроса блокируются им (для вопросов, не блокировавших
            # строки по построению, — вопросов посадки); обработчик
            # получает «ключ (оговорка)» и вносит её в строки под
            # контрактом. Хвост-пояснение (без маркеров ограничения)
            # оговоркой не считается — ключ применяется ниже, хвост
            # ложится заметкой ответа.
            from cdpm import question_resolutions
            active_refs = set(q.get("rule_refs") or []) - question_resolutions.resolved_targets(q)
            for rule in doc.get("rules", []):
                if rule.get("id") in active_refs:
                    bl = rule.setdefault("blockers", [])
                    if qid not in bl:
                        bl.append(qid)
            q["answer"] = {"raw": raw, "status": "ожидает обработчика",
                           "key": key, "revision": revision,
                           "note": caveat or note}
            if data_scope_choice and key == "б":
                from cdpm.answer_feedback import DATA_EXCLUSION_PENDING
                q["answer"]["processing_note"] = DATA_EXCLUSION_PENDING
            summary["awaiting"] += 1
            continue
        # Разрешение условного типа — единственное машинно-однозначное
        # следствие: ветвь названа контрактом для каждого ключа.
        resolved_here = []
        still_blocking = False
        overruled = []
        for rule in doc.get("rules", []):
            cond = rule.get("conditional")
            if cond and cond.get("question") in targets:
                chosen = cond["branches"][key]
                rule["kind"] = chosen
                # П7: вес, названный клаузой у условного правила, ждал
                # в развилке — вступает при ветви SOFT.
                if chosen == "SOFT" and cond.get("weight") is not None \
                        and rule.get("weight") is None:
                    rule["weight"] = cond["weight"]
                    if cond.get("weight_quote"):
                        rule["weight_quote"] = cond["weight_quote"]
                del rule["conditional"]
                m.align_fit_with_kind(doc, rule,
                                      "вид разрешён ответом развилки")
                if rule.get("kind") != chosen:
                    # Ярлыки кодом (ревью r54): семья посадки задаёт
                    # вид; ответ развилки не подменяется молча — след
                    # в ответе анкеты (спорные развилки на семьях с
                    # фиксированным видом код снимает ещё до анкеты).
                    overruled.append(f"{rule['id']}: вид {rule['kind']} "
                                     f"задаёт семья посадки, ответ "
                                     f"«{chosen}» учтён как справка")
                rule["blockers"] = [b for b in rule.get("blockers", [])
                                    if b not in targets]
                if rule["kind"] == "SOFT" and not rule["blockers"]:
                    _bind_soft_to_objective(doc, rule)
                if not rule["blockers"]:
                    rule.pop("blockers", None)
                resolved_here.append(rule["id"])
            elif targets & set(rule.get("blockers", [])):
                still_blocking = True
        # B-вопрос посадки (пакет посадки, живой прогон H0 run-0005
        # 2.09): ответ «неверно» закрывался статусом без действия —
        # ложное покрытие гейта. Отрицание допущения отдаёт строки
        # обработчику с уточнением авиакомпании (пересадка патчем).
        if str(qid).startswith("Q-2-fit-") and key != "а":
            for rule in doc.get("rules", []):
                if rule.get("id") in (q.get("rule_refs") or []):
                    bl = rule.setdefault("blockers", [])
                    if qid not in bl:
                        bl.append(qid)
                    still_blocking = True
        status = "применён" if resolved_here and not still_blocking \
            else ("ожидает обработчика" if still_blocking
                  else "применён")
        q["answer"] = {"raw": raw, "status": status, "key": key,
                       "revision": revision}
        # Пояснение авиакомпании (хвост без ограничения) — след в
        # ответе; ярлык-справка «вид задаёт семья посадки» — рядом.
        notes = ([note] if note else []) + overruled
        if notes:
            q["answer"]["note"] = "; ".join(notes)
        if status == "применён":
            summary["applied"] += 1
            summary["resolved_rules"] += resolved_here
        else:
            summary["awaiting"] += 1
    _close_journal()
    # Every actual new receipt (including an unparsed/refused answer) gets
    # one identity. A retry/worker/status transition must not re-date it.
    for qid in set(raw_map) | set(refused):
        q = questions.get(qid)
        answer = q.get("answer") if q is not None else None
        previous = previous_answers.get(qid)
        if isinstance(answer, dict) and answer is not previous:
            answer_versions.stamp(answer, previous, receipt)
    # K04 (§2.4 (б)): человек применил ответ на пересмотр — признак
    # `pending_revision` этого вопроса снимается. Второй законный путь
    # снятия (кодовый исход «вопрос больше не требуется») живёт в
    # `model.close_question_by_code`; третьего нет.
    summary["settled_revisions"] = m.settle_applied_revisions(doc)
    doc.setdefault("meta", {})["revision"] = max(revision, current_revision)
    return summary


def open_questions(doc: dict) -> list[dict]:
    """Вопросы, которые войдут в анкету следующей ревизии."""
    return m.open_questions(doc)


def remaining_awaiting(doc: dict) -> int:
    """Ответы, всё ещё ожидающие обработчика, — пересчёт по модели.
    Сводка прохода обязана сходиться с ним ПОСЛЕ обработчика (живой
    прогон 2026-08-19: awaiting прохода лгал после применения патча,
    гейт поймал)."""
    from cdpm import question_editions, question_resolutions
    count = 0
    for q in doc.get("questions", []):
        if (q.get("answer") or {}).get("status") != "ожидает обработчика":
            continue
        if question_editions.is_superseded(doc, q):
            continue
        if (not question_editions.problems(q)
                and question_resolutions.complete(q)
                and not question_resolutions.problems(doc, q)):
            continue
        count += 1
    return count


# ---------------------------------------------------------------------------
# Модельные патчи обработчика (ИР-2): части 1 и 3а/3б, чтения части 2
# ---------------------------------------------------------------------------

WORKER_CHUNK = 5         # вопросов в одном обращении к обработчику

_PROMPT_WORKER = """Ты — формализатор свода правил авиакомпании
(обработчик). Авиакомпания ОТВЕТИЛА на вопросы анкеты — внеси каждый
ответ в затронутые строки свода. Верни СТРОГО один JSON-объект:

{{"patches": [{{
  "question": "<id вопроса, ровно как в списке>",
  "change": "<одно предложение: что изменено>",
  "rules": [{{<полная строка свода в машинной форме>}}, …]
}}, …]}}

Строка свода в машинной форме — объект с полями РОВНО такого состава:
id, title, statement, zone ("PAIRING" | "ROSTERING" | "BOTH"), kind
("HARD" | "SOFT" | "DEF" | "DATA" | "CASE" | "OUT"), sources (как в
исходной строке, НЕ менять — это трассировка), note (необязательно),
blockers (необязательно, список id вопросов), weight (ТОЛЬКО у kind
"SOFT" и ТОЛЬКО когда ответ называет вес/штраф: целое неотрицательное
ЧИСЛО — вес обязан быть машинным полем, не только словами в
statement), params (необязательно: машинные параметры правила ровно
как в исходной строке; если ответ МЕНЯЕТ или ЗАДАЁТ порог — обнови
value/as_written записи: value длительностей — В МИНУТАХ, имя ТОЛЬКО
из словаря ниже, чужие имена и единицы будут отброшены), template_fit
(посадка строки на шаблон базы; ОБЯЗАТЕЛЬНО переложи её, если ответ
авиакомпании отвергает допущение посадки (вопрос Q-2-fit-…):
{{"category": "A"|"B"|"C"|"D", "family", "params": {{имя: {{value,
quote — дословная цитата клаузы}}}}, "assumption" при B, "gap" при C,
"candidates" при D}}; registry_rev не пиши — поставит код; иначе
посадку не трогай — она наследуется). Других полей НЕТ.

{PARAM_DICT}

Правила работы:
- ровно один объект patches на каждый вопрос списка; в rules — только
  строки из перечня затронутых этим вопросом, каждая целиком;
- по КАЖДОЙ затронутой строке нужен результат: изменённая полная строка
  либо {{"id": "<id строки>", "unchanged": {{"basis": "<дословная цитата
  из ответа, обосновывающая отсутствие изменений>"}}}}. В объект unchanged
  не добавляй другие поля строки. Если результата нет, строка остаётся
  заблокированной, а ответ применяется только частично;
- unchanged.basis — непустая дословная подстрока поля «ДОСЛОВНЫЙ ОТВЕТ
  (raw)» ниже. Это поле записано JSON-строкой: цитируй её декодированное
  содержимое, без служебных кавычек и экранирования. Если raw равен "б",
  цитата basis может быть только "б". Пояснение «выбранный вариант» —
  контекст вопроса, не цитата ответа. Не подставляй его вместо raw;
- ответ вносится содержательно: выбранное чтение или значение — в
  statement; противоречие (часть 1): действующее требование остаётся,
  строка отменённого требования получает kind "OUT" и note с причиной
  (удалять строки нельзя);
- id и sources строки не меняются; новые строки не создаются;
- в blockers НЕ указывай отвеченный вопрос — его блокировка снята
  самим применением;
- посадку строки с ячейками из ответов не повторяй — такие ячейки
  хранит и судит код;
- дословный ответ и оговорки — ниже; выдумывать сверх ответа нельзя.

Вопросы с ответами ({n}):
{items}""" + m.ID_IMMUTABILITY_RULE

_PATCH_FIELDS = {"id", "title", "statement", "zone", "kind",
                 "sources", "note", "blockers", "weight", "weight_quote",
                 "params", "logic", "template_fit"}


def _affected_rules(doc: dict, q: dict, *, include_processed: bool = False) -> list[dict]:
    """Строки, затронутые вопросом: названные в rule_refs и
    заблокированные им. include_processed нужен только для показа прогресса;
    обработчик по умолчанию получает только ещё не обработанные строки."""
    if "edition" in q:
        from cdpm import question_editions, question_resolutions
        scope = question_editions.targets(q) - question_resolutions.resolved_targets(q)
        out = [r for r in doc.get("rules", []) if r.get("id") in scope]
        if not include_processed and (q.get("answer") or {}).get("status") == m.PARTIAL_ANSWER_STATUS:
            ids = _answer_targets(q, doc)
            out = [r for r in out if ids.intersection(r.get("blockers", []))]
        return out
    targets = _answer_targets(q, doc)
    out = []
    for r in doc.get("rules", []):
        if r.get("id") in q.get("rule_refs", []) \
                or targets & set(r.get("blockers", [])):
            if r not in out:
                out.append(r)
    if not include_processed and (q.get("answer") or {}).get("status") == m.PARTIAL_ANSWER_STATUS:
        out = [r for r in out if targets.intersection(r.get("blockers", []))]
    return out


def _check_patch(q: dict, patch: dict,
                 affected: dict[str, dict]) -> str | None:
    """Причина отбраковки патча вопроса либо None (fail-closed)."""
    rules = patch.get("rules")
    if not isinstance(rules, list) or not rules:
        return "патч без строк"
    seen = set()
    for r in rules:
        if not isinstance(r, dict):
            return "результат по строке не объект"
        rid = r.get("id")
        if not isinstance(rid, str):
            return "результат по строке без строкового id"
        if rid not in affected:
            return f"строка {rid!r} не затронута вопросом — патчить " \
                   f"можно только перечисленные строки"
        if rid in seen:
            return f"строка {rid!r} в патче дважды"
        seen.add(rid)
        if "unchanged" in r:
            unchanged = r["unchanged"]
            if (set(r) != {"id", "unchanged"} or not isinstance(unchanged, dict)
                    or set(unchanged) != {"basis"}):
                return f"строка {rid}: unchanged содержит только basis, без правок строки"
            from cdpm.provenance import norm_raw
            basis = unchanged.get("basis")
            if (not isinstance(basis, str) or not norm_raw(basis)
                    or norm_raw(basis) not in norm_raw((q.get("answer") or {}).get("raw"))):
                return f"строка {rid}: unchanged требует дословную непустую цитату ответа"
            if affected[rid].get("conditional"):
                return f"строка {rid}: unchanged не разрешает условный тип строки"
            continue
        extra = set(r) - _PATCH_FIELDS
        if extra:
            return f"строка {rid}: поля {sorted(extra)} вне схемы"
        if r.get("sources") != affected[rid].get("sources"):
            return f"строка {rid}: изменена трассировка (sources)"
        if affected[rid].get("conditional") \
                and (affected[rid]["conditional"].get("question")
                     != _root_question_id(q)):
            return f"строка {rid}: условный тип другого вопроса — " \
                   f"обработчику её патчить нельзя"
        if any(not isinstance(r.get(field), str) or not r[field].strip()
               for field in ("statement", "title")):
            return f"строка {rid}: пустое имя или формальная запись"
    return None


_PROMPT_GOAL = """Ты — формализатор свода правил авиакомпании
(обработчик). Авиакомпания ответила на вопрос о целевой функции и
весах мягких правил (ЦФ-1). Разбери ответ в машинную форму. Верни
СТРОГО один JSON-объект:

{{"decline": true | false,
  "objectives": [{{"title": "<имя целевой>",
    "statement": "<формальная запись метрики>",
    "zone": {zones_json},
    "direction": "minimize" | "maximize",
    "metric": "<ОБЯЗАТЕЛЬНОЕ имя из закрытого словаря:
total_pairing_cost, layover_nights_total, credit_hours_spread,
positioning_cost, uncovered_seats,
preferred_days_off_violations>"}}, …],
  "weights": [{{"rule": "<id строки>",
    "weight": <целое неотрицательное ЧИСЛО веса/штрафа>,
    "as_written": "<ДОСЛОВНЫЙ фрагмент ответа авиакомпании с этим числом>",
    "statement": "<ПОЛНАЯ формальная запись строки с внесённым "
                 "весом/штрафом>"}}, …]}}

Правила:
- "decline": true ТОЛЬКО если авиакомпания явно подтвердила отказ от
  оптимизации: режим «любое допустимое решение» ЛИБО прямо сказала,
  что критерия оптимизации нет / оптимизация не требуется (тогда
  решение ищется в режиме взвешенных штрафов); одновременно
  decline=true и непустой objectives — противоречие, так нельзя;
- каждая целевая ОБЯЗАНА нести machine-метрику из словаря выше:
  целевая без метрики неисполнима и будет отклонена кодом;
  ПОЯСНЕНИЯ об уже взвешенных строках («… уже взвешено ранее»,
  «вес такой-то назначен») целевыми НЕ являются — не включай их
  в objectives;
- каждая строка таблицы ответа с компонентом/метрикой/направлением —
  ОТДЕЛЬНАЯ целевая в objectives (каждая = отдельная задача
  оптимизации); допустимые зоны: {zones_list};
- "weights" — только для перечисленных затронутых мягких строк и
  только если ответ называет вес/штраф/разряд этой строки; statement
  переписывается целиком с весом; выдумывать веса нельзя;
- ничего сверх дословного ответа не вносить.

Вопрос {qid}: {text}

ОТВЕТ АВИАКОМПАНИИ: {raw}

Затронутые мягкие строки (кандидаты weights):
{rules}""" + m.ID_IMMUTABILITY_RULE


def _quote_around_number(raw: str, number) -> str:
    """Дословный фрагмент текста вокруг числа как целого токена: до
    двух слов слева и до трёх справа; пусто — числа в тексте нет."""
    import re as _re
    toks = str(raw or "").split()
    sval = str(number)
    for i, tk in enumerate(toks):
        if _re.fullmatch(r"\W*" + _re.escape(sval) + r"\W*", tk):
            return " ".join(toks[max(0, i - 2): i + 4])
    return ""


def _apply_goal_answer(doc: dict, q: dict, llm) -> list[str]:
    """ЦФ-1 фаза Б: применение ответа на Q-*-GOAL — обработчик
    разбирает таблицу ответа, КОД проверяет и вносит: целевые — в
    коллекцию objectives (источник — трассировка вопроса, происхождение
    отмечено origin), веса — в statement заблокированных мягких строк
    (снятие блокировки — решение кода и только по строкам с весом).
    Контракт П-1 гоняется на копии целиком; красный — откат вопроса.
    Возвращает [] при успехе либо причины отказа (fail-closed)."""
    from cdpm import question_editions
    edition_problems = question_editions.application_problems(doc, q)
    if edition_problems:
        return edition_problems
    qid = q["id"]
    part = q.get("part")
    zones = ("PAIRING", "BOTH") if part == "3a" \
        else ("ROSTERING", "BOTH")
    ans = q["answer"]
    # Кандидаты весов — ТОЛЬКО действующие мягкие строки: строка,
    # ставшая OUT/HARD после блокировки (конфликт «5 подряд против
    # 10 между парами», run-0046), вес не носит — её блокировка
    # снимается без веса ниже, а не роняет весь разбор целевой.
    blocked_all = [r for r in doc.get("rules", [])
                   if qid in (r.get("blockers") or [])]
    blocked = [r for r in blocked_all if r.get("kind") == "SOFT"]
    rules_listing = "\n".join(
        f"{r['id']} | {r['zone']} | {r['statement']}"
        for r in blocked) or "(заблокированных строк нет)"
    prompt = _PROMPT_GOAL.format(
        zones_json=" | ".join(f'"{z}"' for z in zones),
        zones_list=", ".join(zones), qid=qid,
        text=q.get("text", ""), raw=ans.get("raw", ""),
        rules=rules_listing)
    reply = None
    for attempt in (1, 2):
        ask = prompt if attempt == 1 else (
            prompt + "\n\nПредыдущий ответ не разобрался. Верни "
                     "СТРОГО один JSON-объект и ничего больше.")
        try:
            reply = llm_mod.parse_json_reply(llm(ask))
            break
        except llm_mod.LLMReplyError:
            continue
    if reply is None:
        return ["обработчик не дал разбираемого ответа по вопросу "
                "целевой"]
    # The question may have been reissued while the worker was running.
    # Recheck before constructing or committing any answer effects.
    edition_problems = question_editions.application_problems(doc, q)
    if edition_problems:
        return edition_problems
    decline = reply.get("decline")
    objectives = reply.get("objectives") or []
    weights = reply.get("weights") or []
    if not isinstance(decline, bool) \
            or not isinstance(objectives, list) \
            or not isinstance(weights, list):
        return ["разбор целевой: поля decline/objectives/weights "
                "неверных типов"]
    if decline and objectives:
        return ["разбор целевой: одновременно отказ от оптимизации и "
                "целевые — противоречие"]
    if not decline and not objectives and not weights:
        return ["разбор целевой: ни отказа, ни целевых, ни весов — "
                "ответ не применим машинно"]
    for i, o in enumerate(objectives):
        if not isinstance(o, dict):
            return [f"целевая [{i}] не объект"]
        if not str(o.get("title") or "").strip() \
                or not str(o.get("statement") or "").strip():
            return [f"целевая [{i}] без имени или формальной записи"]
        if o.get("zone") not in zones:
            return [f"целевая [{i}]: зона {o.get('zone')!r} вне "
                    f"допустимых {zones}"]
        if o.get("direction") not in m.DIRECTIONS:
            return [f"целевая [{i}]: направление "
                    f"{o.get('direction')!r} вне перечня"]
    blocked_ids = {r["id"] for r in blocked}
    seen_w = set()
    for i, w in enumerate(weights):
        if not isinstance(w, dict):
            return [f"вес [{i}]: элемент не объект: {w!r}"]
        rid = w.get("rule")
        if not isinstance(rid, str) or rid not in blocked_ids:
            return [f"вес [{i}]: строка {rid!r} не заблокирована "
                    f"этим вопросом"]
        if rid in seen_w:
            return [f"вес [{i}]: строка {rid!r} дважды"]
        seen_w.add(rid)
        if not str(w.get("statement") or "").strip():
            return [f"вес [{i}]: пустая формальная запись"]
        wv = w.get("weight")
        if not isinstance(wv, int) or isinstance(wv, bool) or wv < 0:
            # Третий живой урок класса «число в свободном тексте»:
            # вес обязан прийти МАШИННЫМ числом, не только в записи.
            return [f"вес [{i}]: поле weight не целое "
                    f"неотрицательное: {wv!r}"]
        # Цитата веса — из дословного ответа авиакомпании (доказательный
        # слой веса на пути ответов): число обязано стоять в ответе.
        aw = " ".join(str(w.get("as_written") or "").split())
        raw_answer = str(ans.get("raw") or "")
        if not aw:
            # Модель цитату не дала — её строит код из дословного
            # ответа: окрестность числа (два слова до, три после).
            aw = _quote_around_number(raw_answer, wv)
        if not aw or not m.value_quote_in_text(aw, wv, raw_answer):
            return [f"вес [{i}]: цитата {aw[:40]!r} с числом {wv} не из "
                    f"дословного ответа авиакомпании"]
        w["as_written"] = aw

    trial = copy.deepcopy(doc)
    trial.setdefault("objectives", [])
    src_refs = [dict(s) for s in q.get("source_refs", [])]
    for o in objectives:
        first = src_refs[0] if src_refs else {"sid": "?", "anchor": "?"}
        oid = "O-" + m.stable_rule_id(first["sid"],
                                      first["anchor"])[2:]
        n, base = 1, oid
        while any(x.get("id") == oid for x in trial["objectives"]):
            n += 1
            oid = f"{base}-{n}"
        entry = {
            "id": oid, "title": str(o["title"]).strip(),
            "statement": str(o["statement"]).strip(),
            "zone": o["zone"], "direction": o["direction"],
            "sources": [dict(s) for s in src_refs],
            "origin": {"question": qid,
                       "revision": int(ans.get("revision", 1))},
        }
        # Живой затык run-0002 (ночь 30→31.08): целевая БЕЗ машинной
        # метрики рождалась «голой», чистка осколков её снимала, а
        # сквозная проверка того же прохода ставила второй добор —
        # контракт дублей валил применение целиком. Метрика теперь
        # ОБЯЗАТЕЛЬНА (fail-closed с перечнем): целевая без метрики
        # неисполнима по построению — пустышка это доказала.
        metric = o.get("metric")
        if metric in m.METRICS and (
                set(m._ZONE_TASKS.get(o["zone"], ()))
                & set(m.METRICS[metric]["tasks"])):
            entry["metric"] = metric
        else:
            zone_metrics = sorted(
                name for name, spec in m.METRICS.items()
                if set(m._ZONE_TASKS.get(o["zone"], ()))
                & set(spec["tasks"]))
            return [
                f"целевая {str(o.get('title'))[:60]!r} без машинной "
                f"метрики из закрытого словаря (дано: "
                f"{str(metric)[:40]!r}); словарь зоны {o['zone']}: "
                + ", ".join(zone_metrics)
                + "; пояснения об уже взвешенных строках целевыми "
                  "не являются"]
        trial["objectives"].append(entry)
    trial_rules = {r["id"]: r for r in trial["rules"]}
    for w in weights:
        tr = trial_rules[w["rule"]]
        tr["statement"] = str(w["statement"]).strip()
        tr["weight"] = int(w["weight"])
        tr["weight_quote"] = " ".join(str(w.get("as_written") or "").split())
        # Снятие блокировки — решение кода и ТОЛЬКО по строке с
        # внесённым весом.
        tr["blockers"] = [b for b in tr.get("blockers", [])
                          if b != qid]
        if not tr["blockers"]:
            tr.pop("blockers", None)
    # Строки, чей вес ответ НЕ назвал (в т.ч. при отказе от
    # оптимизации), не оставляются на применённом вопросе (дедлок из
    # ревью фазы Б: применённый вопрос в анкету не попадает, а
    # контракт не даёт его применить) — они перевешиваются на НОВЫЙ
    # вопрос весов, который войдёт в анкету следующей ревизии.
    leftover = [trial_rules[rid] for rid in sorted(blocked_ids)
                if rid not in seen_w]
    # Немягкие строки под этим вопросом: блокировка снимается без
    # веса (отменённой/жёсткой строке вес не положен), со следом.
    for r0 in blocked_all:
        if r0.get("kind") == "SOFT":
            continue
        tr = trial_rules.get(r0["id"])
        if tr is None:
            continue
        tr["blockers"] = [b for b in tr.get("blockers", [])
                          if b != qid]
        if not tr["blockers"]:
            tr.pop("blockers", None)
        extra = (f"вопрос {qid} снят без веса: вид строки "
                 f"{tr.get('kind')} веса не носит")
        tr["note"] = ((tr.get("note", "") + "; " + extra)
                      .strip("; "))
    if leftover:
        for tr in leftover:
            tr["blockers"] = [b for b in tr.get("blockers", [])
                              if b != qid]
        m.ensure_weights_question(trial, part, leftover)
    trial_q = next(x for x in trial["questions"] if x["id"] == qid)
    trial_q["answer"]["status"] = "применён"
    trial_q["answer"]["goal_decline"] = decline
    # Ярлыки кодом: своды до детерминизации самолечатся до контракта.
    m.derive_all_rule_traits(trial, "разбор целевой")
    problems = m.validate_schema(trial) + m.validate_model(trial)
    if problems:
        return [f"разбор целевой не проходит контракт: {problem}" for problem in problems]
    # Слияние на месте: строки и вопросы сохраняют идентичность;
    # новые вопросы (повторный запрос весов) переносятся целиком.
    by_id = {r["id"]: r for r in doc["rules"]}
    for tr in trial["rules"]:
        orig = by_id[tr["id"]]
        orig.clear()
        orig.update(tr)
    have_q = {x.get("id") for x in doc["questions"]}
    for tq in trial["questions"]:
        if tq.get("id") not in have_q:
            doc["questions"].append(tq)
    doc.setdefault("objectives", [])
    have = {x.get("id") for x in doc["objectives"]}
    for x in trial["objectives"]:
        if x.get("id") not in have:
            doc["objectives"].append(x)
    q["answer"]["status"] = "применён"
    q["answer"]["goal_decline"] = decline
    return []


def apply_worker_patches(doc: dict, llm, control=None) -> dict:
    """Внесение принятых ответов со статусом «ожидает обработчика» в
    строки свода модельными патчами: модель предлагает полные строки,
    код валидирует схему, принадлежность, трассировку и контракт П-1
    целиком; невалидный патч — откат вопроса с причиной (fail-closed
    повопросно). Вопросы о целевой (Q-*-GOAL) идут отдельной веткой
    _apply_goal_answer: их ответ рождает целевые и веса, а не патчи
    строк. Частичный результат оставляет незавершённые строки заблокированными.
    Возвращает {worker_applied, worker_rejected}, а при частичном результате
    дополнительно worker_partial — число частично обработанных вопросов."""
    from cdpm import question_editions
    rejected: list[dict] = []
    pending_all = []
    for q in doc.get("questions", []):
        if (q.get("answer") or {}).get("status") not in (
                "ожидает обработчика", m.PARTIAL_ANSWER_STATUS):
            continue
        if question_editions.is_superseded(doc, q):
            # Keep the original answer unchanged as history, not a new
            # failure on every worker pass. Mid-call races still refuse below.
            continue
        issues = question_editions.application_problems(doc, q)
        if issues:
            rejected.append({"question": q["id"], "reasons": issues})
        else:
            pending_all.append(q)
    pending = [q for q in pending_all
               if not m.is_goal_family(q.get("id"))]
    goal_pending = [q for q in pending_all
                    if m.is_goal_family(q.get("id"))]
    applied = partial = 0
    for start in range(0, len(pending), WORKER_CHUNK):
        if control:
            control()          # пауза/остановка между чанками (ПС-1)
        batch = pending[start:start + WORKER_CHUNK]
        items = []
        sent_rows = {}
        for q in batch:
            ans = q["answer"]
            shown = ans.get("key") or ans.get("raw", "")
            if ans.get("key") and ans.get("note"):
                shown = f"{ans['key']} ({ans['note']})"
            elif not ans.get("key"):
                shown = ans.get("raw", "")
            # A choice key alone carries no meaning outside its issued options.
            # Supply only the unambiguous selected text as worker context; never
            # rewrite raw/note or use this context as answer-quote provenance.
            fmt = q.get("answer_format") or {}
            options = fmt.get("options") if isinstance(fmt, dict) else None
            if (isinstance(fmt, dict) and fmt.get("type") == "choice"
                    and ans.get("key") and isinstance(options, list)):
                matches = [o for o in options if isinstance(o, dict)
                           and o.get("key") == ans["key"]]
                if (len(matches) == 1 and isinstance(matches[0].get("text"), str)
                        and matches[0]["text"]):
                    shown += f" — выбранный вариант: {matches[0]['text']}"
            if str(q["id"]).startswith(("Q-3a-данные-", "Q-3b-данные-")):
                if ans.get("key") == "а" and q.get("data_scope"):
                    shown += ("\nТочный контракт перечня: "
                              + json.dumps(q["data_scope"], ensure_ascii=False)
                              + "\nДобавьте только перечисленные значения, сохраните прежние "
                              "и текст условия вне перечня. Не меняйте регистр и символы "
                              "обозначений. Согласованно обновите существующие машинные "
                              "перечни. В явном перечне разделяйте значения запятыми; "
                              "значение с собственной запятой заключите в кавычки.")
            affected = _affected_rules(doc, q)
            sent_rows[q["id"]] = {r["id"]: copy.deepcopy(r) for r in affected}
            items.append(
                f"{q['id']} | {q['title']}\n"
                f"Вопрос: {q['text']}\n"
                f"ДОСЛОВНЫЙ ОТВЕТ (raw): {json.dumps(ans.get('raw', ''), ensure_ascii=False)}\n"
                f"ОТВЕТ АВИАКОМПАНИИ: {shown}\n"
                f"Затронутые строки: "
                + json.dumps(affected, ensure_ascii=False))
        prompt = _PROMPT_WORKER.format(
            PARAM_DICT=m.machine_dictionary_listing(),
            n=len(batch),
            items="\n\n".join(items))
        reply = None
        for attempt in (1, 2):
            ask = prompt if attempt == 1 else (
                prompt + "\n\nПредыдущий ответ не разобрался. Верни "
                         "СТРОГО один JSON-объект и ничего больше.")
            try:
                reply = llm_mod.parse_json_reply(llm(ask))
                break
            except llm_mod.LLMReplyError:
                continue
        rows = reply.get("patches") if isinstance(reply, dict) else None
        got, repeated = {}, set()
        for patch in rows if isinstance(rows, list) else []:
            if not isinstance(patch, dict) or not isinstance(patch.get("question"), str):
                continue
            key = patch["question"]
            if key in got:
                repeated.add(key)
            else:
                got[key] = patch
        for q in batch:
            qid = q["id"]
            issues = question_editions.application_problems(doc, q)
            if issues:
                rejected.append({"question": qid, "reasons": issues})
                continue
            if qid in repeated:
                rejected.append({"question": qid, "reasons": [
                    "обработчик дал несколько патчей одного вопроса — выбор одного запрещён"]})
                continue
            if str(qid).startswith(("Q-3a-данные-", "Q-3b-данные-")) \
                    and q["answer"].get("key") == "б":
                # This worker can patch rules, not operational records. No
                # rule edit can prove the selected data-exclusion effect.
                from cdpm.answer_feedback import DATA_EXCLUSION_PENDING
                rejected.append({"question": qid, "reasons": [DATA_EXCLUSION_PENDING]})
                continue
            patch = got.get(qid)
            if patch is None:
                rejected.append({"question": qid,
                                 "reasons": ["обработчик не дал патча "
                                             "по вопросу"]})
                continue
            affected = {r["id"]: r for r in _affected_rules(doc, q)}
            reason = _check_patch(q, patch, affected)
            if reason:
                rejected.append({"question": qid, "reasons": [reason]})
                continue
            # K04 applies to the second (worker) stage too. A partial
            # deterministic answer is not permission to overwrite a later
            # change in a remaining row. Keep the original receipt/journal.
            from cdpm import revision_cells
            conflicts = revision_cells.worker_conflicts(
                doc, q, [affected[r["id"]] for r in patch["rules"]])
            conflicts += revision_cells.preserve_worker_neighbors(
                q, patch, sent_rows[qid], affected)
            if conflicts:
                q["answer"]["status"] = m.REVISION_CONFLICT_STATUS
                q["answer"]["conflicts"] = conflicts
                q["answer"]["processing_note"] = (
                    "Пересмотр: конфликт — поздние изменения сохранены; "
                    "сверьте значения и явно откройте новый пересмотр через API reconsider.")
                rejected.append({"question": qid, "reasons": conflicts})
                continue
            # Применение к копии + полный контракт; красный — откат.
            trial = copy.deepcopy(doc)
            trial_rules = {r["id"]: r for r in trial["rules"]}
            patch_problem = None
            for r in patch["rules"]:
                if "unchanged" in r:
                    # Keep the rule byte-for-byte; only this question's lock
                    # is removed below after the quote has been checked.
                    continue
                new_rule = {k: v for k, v in r.items()
                            if v is not None}
                new_rule.pop("conditional", None)
                # Машинные поля не исчезают МОЛЧА (ревью 1.3): патч
                # без непустого params/weight наследует значения
                # строки — обработчик, не повторивший поле, стирал
                # только что ставшие машинными пороги и веса.
                old_rule = trial_rules[r["id"]]
                prev_kind = old_rule.get("kind")
                if not new_rule.get("params"):
                    new_rule.pop("params", None)
                    if old_rule.get("params"):
                        new_rule["params"] = copy.deepcopy(
                            old_rule["params"])
                else:
                    # Params из патча — той же очисткой, что при
                    # извлечении: выдуманное имя/единица отбрасывается
                    # со следом, а не топит весь патч (живой Лайт
                    # run-0040: 'min_rest_away'/'minutes' валили
                    # ответ, строка оставалась заблокированной).
                    kept_p, notes_p = m.sanitize_params(
                        new_rule.get("params"), new_rule.get("kind"),
                        rule_text=(str(new_rule.get("formal") or "")
                                   + " "
                                   + str(new_rule.get("statement")
                                         or "")))
                    if kept_p:
                        new_rule["params"] = kept_p
                    else:
                        new_rule.pop("params", None)
                        if old_rule.get("params"):
                            new_rule["params"] = copy.deepcopy(
                                old_rule["params"])
                    for extra in notes_p:
                        new_rule["note"] = ((new_rule.get("note", "")
                                             + "; " + extra)
                                            .strip("; "))
                if "weight" not in new_rule \
                        and old_rule.get("weight") is not None \
                        and new_rule.get("kind") == "SOFT":
                    new_rule["weight"] = old_rule["weight"]
                if not new_rule.get("logic"):
                    new_rule.pop("logic", None)
                    if old_rule.get("logic"):
                        # 1.5: logic наследуется как params и
                        # weight — молчаливое стирание закрыто
                        # классом (гл. 3 дизайна ступени 3).
                        new_rule["logic"] = copy.deepcopy(
                            old_rule["logic"])
                # Пакет посадки (ревью, блокер): патч обработчика
                # посадку не пишет (_check_patch бракует чужое поле)
                # — наследование, как params/weight/logic, иначе при
                # контракте посадки гиб весь путь внесения ответов.
                if not isinstance(new_rule.get("template_fit"),
                                  dict) \
                        and isinstance(old_rule.get("template_fit"),
                                       dict):
                    new_rule["template_fit"] = copy.deepcopy(
                        old_rule["template_fit"])
                elif isinstance(new_rule.get("template_fit"), dict) \
                        and not new_rule["template_fit"].get(
                            "registry_rev"):
                    # Посадку патча обработчик пересадил — registry_rev
                    # ставит код (прогон №2 run-0007: патч отбракован
                    # схемой «registry_rev is required»).
                    new_rule["template_fit"]["registry_rev"] = \
                        m.registry_fingerprint()
                # Замки 4П на пути ответов — те же, что в цикле (план
                # «прямой ход» 2.8; проверка К1): смена вида только по
                # перечню миграций либо по семье посадки A/B; заплата не
                # уменьшает информацию посаженной строки (понижение
                # посадки с цитатами, потеря параметров, снятие веса).
                if not m.kind_change_allowed(prev_kind, new_rule.get("kind"),
                                             new_rule.get("template_fit")):
                    patch_problem = (
                        f"смена вида {prev_kind!r}→{new_rule.get('kind')!r} "
                        f"патчем ответа вне перечня миграций и не "
                        f"подтверждена посадкой A/B на семью этого вида "
                        f"(конвенция ВМП)")
                    break
                from cdpm import exchange as _ex
                integrity = _ex.fit_integrity_problems(old_rule, new_rule, None)
                if integrity:
                    patch_problem = "; ".join(integrity)
                    break
                # Точное происхождение (A03-мин/K05-мин, контракт П0
                # §2.3 п. 1 и 3, закон 5, §12.3). Сперва закон заплаты:
                # ячейка, помеченная ответом, заводится и меняется
                # только вместе с точным происхождением — иначе форма
                # ячейки (одна строка `answer`) была бы обходом закона
                # точной версии; и строка с такими ячейками не меняет
                # вид (OUT/DATA/DEF/CASE) — иначе посадку у неё снимет
                # ниже `align_fit_with_kind`, и значение авиакомпании
                # исчезнет руками кода. Тот же расчёт судит заплату
                # обмена, включая заплату вида (общий закон П0 № 4):
                # смена вида — разрешение `kind_change_allowed` выше
                # говорит о ФОРМЕ перехода, а не о праве снять ответ.
                from cdpm import provenance as _prov
                prov_patch = _prov.patch_problems(old_rule, new_rule)
                if prov_patch:
                    patch_problem = "; ".join(prov_patch)
                    break
                # Затем сама ячейка: кандидат заплаты и эффект до
                # фиксации — статус «применён» ещё не проставлен РОВНО у
                # отвечаемого вопроса, только он и назван в `applying`
                # (корень цепочки и соседний пересмотр применёнными не
                # считаются: иначе заплата на `X-R` оправдывала бы
                # значение пересмотренным ответом на `X`). Цитата —
                # дословная подстрока raw точной версии, значение —
                # читается из этой цитаты. ОГОВОРКА: минимальная часть
                # судит подлинность подстроки и читаемость числа из неё,
                # а не смысл основания — «без изменений» не обосновывает
                # никакого числа (числа в цитате нет), но цитата «660
                # минут», вырезанная из ответа «…600 минут, не 660
                # минут», подстрокой проходит: смысл и область основания
                # судит полный A02/A03 (волна 17б, контракт §12.3).
                prov_problems = _prov.rule_problems(
                    trial, new_rule, applying=(qid,))
                if prov_problems:
                    patch_problem = "; ".join(prov_problems)
                    break
                trial_rules[r["id"]].clear()
                trial_rules[r["id"]].update(new_rule)
                requested_kind = new_rule.get("kind")
                changes_k = m.align_fit_with_kind(
                    trial, trial_rules[r["id"]],
                    "строка обновлена ответом анкеты")
                if requested_kind in ("HARD", "SOFT") \
                        and requested_kind != prev_kind \
                        and trial_rules[r["id"]].get("kind") \
                        != requested_kind:
                    # П4 (ревью r54/r55): патч, МЕНЯЮЩИЙ вид против
                    # семьи посадки, бракуется словами, как в обмене, —
                    # не проходит «применённым» без эффекта; патч, вид
                    # не трогавший (устаревший ярлык свода выровнял
                    # код), не виноват.
                    patch_problem = (
                        f"вид {requested_kind!r} отменён ярлыками кода "
                        f"(итог {trial_rules[r['id']].get('kind')!r}: "
                        f"{'; '.join(changes_k)}) — смените посадку, "
                        f"не вид")
                    break
            if patch_problem:
                rejected.append({"question": qid,
                                 "reasons": [patch_problem]})
                continue
            # Снятие блокировки отвеченного вопроса — решение кода, не
            # модели: qid снимается только со строк, по которым дан результат.
            # Вопрос-пересмотр действует за корневой: снимается вся
            # цепочка (после отката корневой и прежние -R снова в blockers).
            targets = _answer_targets(q, doc)
            handled = {r["id"] for r in patch["rules"]}
            for rid in affected:
                tr = trial_rules[rid]
                if rid not in handled:
                    if not targets.intersection(tr.get("blockers", [])):
                        tr.setdefault("blockers", []).append(qid)
                    continue
                # A patch cannot erase blockers belonging to other questions.
                tr["blockers"] = list(dict.fromkeys([
                    *affected[rid].get("blockers", []), *tr.get("blockers", [])]))
                tr["blockers"] = [b for b in tr.get("blockers", [])
                                  if b not in targets]
                if not tr["blockers"]:
                    tr.pop("blockers", None)
            trial_q = next(x for x in trial["questions"]
                           if x["id"] == qid)
            if "edition" not in trial_q:
                # A legacy draft retains unfinished scope before publication.
                # A completed blocker-only effect does NOT make that row an
                # addressee of the answer (A03). A substantive result, including
                # explicitly quoted unchanged, does (A02). Never alter an edition.
                content_rows = [r["id"] for r in patch["rules"]
                                if "unchanged" in r or any(
                                    key not in ("id", "blockers")
                                    and r[key] != affected[r["id"]].get(key)
                                    for key in r)]
                unfinished = [rid for rid in affected if rid not in handled]
                trial_q["rule_refs"] = list(dict.fromkeys([
                    *(trial_q.get("rule_refs") or []), *content_rows, *unfinished]))
            status = "применён" if handled == set(affected) else m.PARTIAL_ANSWER_STATUS
            trial_q["answer"]["status"] = status
            unchanged = {r["id"]: r["unchanged"]["basis"]
                         for r in patch["rules"] if "unchanged" in r}
            if unchanged:
                trial_q["answer"].setdefault("worker_unchanged", {}).update(unchanged)
            # Путь к весу — тем же сквозным строителем, что в цикле
            # обмена (живой прогон H0 run-0005, 2.09): патч снимал
            # блокировку вопроса единицы счёта у SOFT-строки без веса,
            # и контракт «SOFT без пути к весу» ронял патч — вопрос
            # висел «ожидает обработчика» вечно.
            # Ярлыки кодом (2026-09-02): своды до детерминизации
            # самолечатся на пути патча, а не роняют его контрактом.
            m.derive_all_rule_traits(trial, "патч обработчика")
            m.ensure_goal_questions_sweep(trial)
            problems = m.validate_schema(trial) + m.validate_model(trial)
            if problems:
                rejected.append(
                    {"question": qid,
                     "reasons": [f"патч не проходит контракт: {problem}"
                                 for problem in problems]})
                continue
            if str(qid).startswith(("Q-3a-данные-", "Q-3b-данные-")):
                if any(trial_rules[rid].get("kind") != old.get("kind")
                       for rid, old in affected.items()):
                    rejected.append({"question": qid, "reasons": [
                        "выбранное расширение перечня не меняет вид правила; "
                        "сохраните вид и внесите выбранные значения"]})
                    continue
                from cdpm.data_scope import extension_problems
                issues = extension_problems(q, affected, trial_rules)
                if issues:
                    rejected.append({"question": qid, "reasons": issues})
                    continue
            # Слияние на месте: объекты строк и вопросов сохраняют
            # идентичность — внешние ссылки не протухают.
            snap = _snapshot(doc)
            by_id = {r["id"]: r for r in doc["rules"]}
            for tr in trial["rules"]:
                orig = by_id[tr["id"]]
                orig.clear()
                orig.update(tr)
            q["answer"]["status"] = status
            if "edition" not in q:
                q["rule_refs"] = copy.deepcopy(trial_q["rule_refs"])
            if "worker_unchanged" in trial_q["answer"]:
                q["answer"]["worker_unchanged"] = copy.deepcopy(
                    trial_q["answer"]["worker_unchanged"])
            q["answer"].pop("processing_note", None)
            # Журнал следствий патча (A3.4): та же разница «до/после»,
            # ДОБАВЛЯЕТСЯ к следствиям первой стадии (блокировки, ярлыки
            # ответа с оговоркой) — откат идёт по всей цепочке до «до».
            q["answer"]["effects"] = ((q["answer"].get("effects") or [])
                                      + _diff_effects(snap, doc))
            if status == m.PARTIAL_ANSWER_STATUS:
                partial += 1
            else:
                applied += 1
    for q in goal_pending:
        if control:
            control()
        snap = _snapshot(doc)
        reasons = _apply_goal_answer(doc, q, llm)
        if reasons:
            rejected.append({"question": q["id"], "reasons": reasons})
        else:
            if isinstance(q.get("answer"), dict):
                q["answer"]["effects"] = ((q["answer"].get("effects") or [])
                                          + _diff_effects(snap, doc))
            applied += 1
    # Keep machine feedback separate from raw/note (the airline's answer).
    # Never rewrite an invalid or superseded edition while reporting refusal.
    by_question = {q["id"]: q for q in pending_all}
    for failure in rejected:
        q = by_question.get(failure["question"])
        if q is not None and str(q["id"]).startswith(("Q-3a-данные-", "Q-3b-данные-")) \
                and q["answer"].get("status") == "ожидает обработчика" \
                and not question_editions.application_problems(doc, q):
            q["answer"]["processing_note"] = "; ".join(failure["reasons"])
    # K04 (§2.4 (б)): применение ответа на пересмотр патчем обработчика —
    # тот же законный путь снятия признака `pending_revision`.
    m.settle_applied_revisions(doc)
    result = {"worker_applied": applied, "worker_rejected": rejected}
    if partial:
        result["worker_partial"] = partial
    return result


def sha256_of(path: str | Path) -> str:
    h = hashlib.sha256()
    h.update(Path(path).read_bytes())
    return h.hexdigest()


def main(argv: list[str] | None = None) -> int:
    import argparse
    import sys
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")   # У-14
    from cdpm import render
    ap = argparse.ArgumentParser(
        description="Приёмник ответов анкеты CDPM (ИР-1)")
    ap.add_argument("model", help="model.json свода")
    ap.add_argument("book", help="заполненная книга анкеты (.xlsx)")
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--llm", default="cloud",
                    choices=("cloud", "local", "auto"),
                    help="канал разборщика нечётких ответов")
    args = ap.parse_args(argv)
    doc = m.load_model(args.model)
    reading = read_book(args.book, doc)
    raw_map, problems = reading.answers, reading.problems
    for p in problems:
        print(f"ЗАМЕЧАНИЕ: {p}")
    if not raw_map and not reading.refused:
        print("в книге нет ни одного ответа — применять нечего")
        return 1
    llm = llm_mod.cosmos_llm(requested=args.llm)
    summary = apply_answers(doc, raw_map, llm, refused=reading.refused,
                            received_at=reading.received_at, book_sha256=reading.book_sha256)
    summary["read_problems"] = problems
    # Ответы, требующие правки текста строк, вносит обработчик (ИР-2).
    summary.update(apply_worker_patches(doc, llm))
    summary["awaiting"] = remaining_awaiting(doc)
    bad = m.validate_schema(doc) + m.validate_model(doc)
    if bad:
        # Do not silently hide remaining errors (one record can yield several).
        for p in bad:
            print(f"НАРУШЕНИЕ: {p}")
        return 1
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    m.publish_model(doc, out / "model.json")
    (out / "svod.md").write_text(render.render_svod(doc),
                                 encoding="utf-8")
    (out / "anketa.md").write_text(render.render_anketa(doc),
                                   encoding="utf-8")
    anketa_xlsx.render_anketa_xlsx(doc, out / "anketa.xlsx")
    summary["book_sha256"] = reading.book_sha256
    # Отпечаток ревизии по СОДЕРЖАНИЮ (11.7 проекта «Запись о
    # прогоне»): корень дерева прогонов судится парами «вопрос →
    # ответ», а не байтами книги.
    from cdpm import answers_fingerprint
    summary["content_sha256"] = \
        answers_fingerprint.book_content_sha256(args.book)
    (out / "answers-summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8")
    print(f"ревизия {summary['revision']}: получено "
          f"{summary['received']}, применено {summary['applied']}, "
          f"обработчиком {summary.get('worker_applied', 0)}, "
          f"отброшено патчей {len(summary.get('worker_rejected', []))},"
          f" не разобрано {summary['unparsed']}"
          + (f", конфликтов пересмотра {summary['conflicted']}"
             if summary.get("conflicted") else "")
          + (f", частично обработано {summary['worker_partial']}"
             if summary.get("worker_partial") else ""))
    print(f"модель: {out / 'model.json'}")
    print(f"сводка: {out / 'answers-summary.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
