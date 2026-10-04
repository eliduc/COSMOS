"""Authentic applied questionnaire sources for the blind source checker.

The model supplies identity and application status, NEVER the expected value.
The reader sees the issued workbook row and the airline's literal answer.
No source clause is rewritten. Missing/corrupt evidence is not a legacy fallback.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import re
from pathlib import Path

SCHEMA = "source-answer-evidence/1"
EVIDENCE = "source-check/answer-evidence.json"
PACKAGE_MANIFEST = "source-check/package-manifest.json"
PACKAGE_MODEL = "source-check/package-model.json"
RECOVERY_SCHEMA = "source-revision-recovery/1"


class EvidenceError(ValueError):
    pass


class _BookRevisionMismatch(EvidenceError):
    pass


def _fail(words):
    raise EvidenceError("источники ответов анкеты: " + words)


def _sha(body):
    return hashlib.sha256(body).hexdigest()


def _json(body):
    def unique(pairs):
        out = {}
        for key, value in pairs:
            if key in out:
                _fail(f"повтор ключа JSON {key}")
            out[key] = value
        return out
    try:
        return json.loads(body, object_pairs_hook=unique)
    except (TypeError, ValueError, UnicodeError) as exc:
        raise EvidenceError(f"источники ответов анкеты: нечитаемый JSON: {exc}") from exc


def _bytes(doc):
    return (json.dumps(doc, ensure_ascii=False, sort_keys=True,
                       separators=(",", ":")) + "\n").encode("utf-8")


def _ref(name, body):
    return {"name": name, "sha256": _sha(body), "bytes": len(body)}


def _applied(doc):
    from cdpm import answer_versions, provenance, question_editions
    if not isinstance(doc, dict) or not isinstance(doc.get("questions", []), list):
        _fail("модель / вопросы не того вида")
    # A retained answer-derived cell cannot hide by dropping/rejecting its question.
    for _rid, _family, _name, cell in provenance.answer_cells(doc):
        src = provenance.cell_source(cell)
        try:
            if src and src.get("kind") == "answer":
                q, reason = provenance.resolve_source(doc, src)
            else:
                q, reason = provenance._legacy_raw(doc, cell)
        except ValueError as exc:
            raise EvidenceError(f"источники ответов анкеты: {exc}") from exc
        if q is None:
            _fail(str(reason))
    result, seen = [], set()
    for q in doc.get("questions", []):
        if not isinstance(q, dict) or not isinstance(q.get("id"), str) or q["id"] in seen:
            _fail("вопрос отсутствует или его ID неоднозначен")
        seen.add(q["id"])
        a = q.get("answer")
        if not isinstance(a, dict) or a.get("closed_by") == "code":
            continue
        if a.get("status") not in provenance.APPLIED_STATUSES:
            continue  # rejected/pending/historical answers are not normative evidence
        if question_editions.is_superseded(doc, q):
            continue
        if a.get("status") != "применён":
            _fail(f"{q['id']}: частичный ответ не удостоверяет весь исходный текст")
        if (type(a.get("revision")) is not int or a["revision"] < 2
                or type((doc.get("meta") or {}).get("revision")) is not int
                or a["revision"] > doc["meta"]["revision"]):
            _fail(f"{q['id']}: неверная ревизия применения")
        if not isinstance(a.get("raw"), str) or not a["raw"].strip():
            _fail(f"{q['id']}: нет дословного ответа")
        problems = answer_versions.problems(a)
        version = a.get("answer_version")
        if problems or not isinstance(version, dict) or not re.fullmatch(
                r"[0-9a-f]{64}", str(version.get("book_sha256") or "")):
            _fail(f"{q['id']}: нет достоверного отпечатка исходной книги / версии")
        provenance.question_edition(q)  # validates the immutable question scope
        result.append(q)
    return sorted(result, key=lambda q: q["id"])


def _book_row(body, doc, q, source_revision=None):
    from openpyxl import load_workbook
    from cdpm import anketa_xlsx, provenance
    a = q["answer"]
    try:
        wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        try:
            revision = None
            if anketa_xlsx.SUMMARY_TITLE in wb.sheetnames:
                for row in wb[anketa_xlsx.SUMMARY_TITLE].iter_rows(values_only=True):
                    if row and row[0] == anketa_xlsx.REVISION_LABEL:
                        if revision is not None:
                            _fail("повтор ревизии книги")
                        revision = row[1] if len(row) > 1 else None
            expected_revision = a["revision"] - 1 if source_revision is None else source_revision
            if str(revision) != str(expected_revision):
                raise _BookRevisionMismatch(
                    f"источники ответов анкеты: {q['id']}: книга не того издания / ревизии")
            # Issued context, not parsed answer/effect or a current rule value.
            published_doc = copy.deepcopy(doc)
            pq = next(x for x in published_doc["questions"] if x["id"] == q["id"])
            pq.pop("answer", None)
            expected = anketa_xlsx.published_question_row(
                published_doc, pq, {s["sid"]: s["file"] for s in doc.get("sources", [])})
            found = []
            for ws in wb.worksheets:
                if ws.title not in anketa_xlsx.SHEET_TITLES.values():
                    continue
                rows = ws.iter_rows(values_only=True)
                head = next(rows, ())
                names = [str(x or "").strip().casefold() for x in head]
                required = [*anketa_xlsx.HEADERS[:5], anketa_xlsx.ANSWER_HEADER,
                            anketa_xlsx.FINGERPRINT_HEADER]
                if not all(names.count(x.strip().casefold()) == 1 for x in required):
                    _fail(f"{q['id']}: неполная / неоднозначная шапка книги")
                columns = [names.index(x.strip().casefold()) for x in required]
                for n, row in enumerate(rows, 2):
                    cells = [str(row[c] or "").strip() if c < len(row) else ""
                             for c in columns]
                    if cells[0] != q["id"] or not cells[5]:
                        continue
                    wanted = [str(x or "").strip() for x in expected[:5]]
                    tail = cells[1][len(wanted[1]):] if cells[1].startswith(wanted[1]) else None
                    # answer_feedback.lines appends the PREVIOUS rejected
                    # answer outside the immutable question. Preserve the full
                    # printed row as evidence, but do not turn that rejection
                    # into a current airline requirement for the blind reader.
                    feedback = (tail == "" or (tail is not None and re.fullmatch(
                        r"\nНеприменённый ответ сохранён: «[^\n]*»\."
                        r"(?:\nОтвет не разобран: [^\n]+)?", tail) is not None))
                    if (not feedback or cells[0] != wanted[0] or cells[2:5] != wanted[2:5]
                            or cells[6] != provenance.question_edition(q)["fingerprint"]
                            or cells[5] != a["raw"]):
                        _fail(f"{q['id']}: исходный вопрос / ответ / отпечаток книги изменён")
                    found.append({"sheet": ws.title, "row": n,
                                  "question": wanted[1], "printed_question": cells[1],
                                  "historical_feedback": tail, "format": cells[2],
                                  "targets": cells[3], "sources": cells[4],
                                  "raw": cells[5]})
            if not found:
                _fail(f"{q['id']}: применённого ответа нет в исходной книге")
            return found
        finally:
            wb.close()
    except EvidenceError:
        raise
    except Exception as exc:
        raise EvidenceError(f"источники ответов анкеты: книга не читается: {exc}") from exc


def _source_revision(value, applied_revision):
    if type(value) is not int or not 1 <= value < applied_revision:
        _fail("неверная исходная ревизия книги")
    return value


def _recovery_revision(read, source_run, q, book, summary_bytes):
    """Validate an explicitly posthoc witness, never infer missing receipt facts."""
    from cdpm import answers, model, provenance
    revision = q["answer"]["revision"]
    prefix = f"answers/rev-{revision:03d}/"
    witness_name = prefix + "source-revision-recovery.json"
    model_name = prefix + "source-model-recovery.json"
    try:
        witness_bytes, source_bytes = read(witness_name), read(model_name)
    except (OSError, KeyError, EvidenceError) as exc:
        raise EvidenceError("источники ответов анкеты: книга не того издания / ревизии; "
                            "нет полного свидетельства source-revision-recovery") from exc
    witness, source_doc = _json(witness_bytes), _json(source_bytes)
    if not isinstance(witness, dict) or not isinstance(source_doc, dict):
        _fail("свидетельство исходной ревизии / модель не объект")
    if (witness.get("schema") != RECOVERY_SCHEMA
            or witness.get("kind") != "posthoc-pre-apply-model"
            or witness.get("source_run") != source_run
            or type(witness.get("revision")) is not int or witness["revision"] != revision):
        _fail("свидетельство другой ревизии / прогона или неизвестного вида")
    for key, name, body in (("book", prefix + "anketa-otvety.xlsx", book),
                            ("summary", prefix + "answers-summary.json", summary_bytes),
                            ("source_model", model_name, source_bytes)):
        if witness.get(key) != _ref(name, body):
            _fail(f"свидетельство: {key} не совпадает с исходными байтами")
    origin = witness.get("origin")
    if (not isinstance(origin, dict) or not re.fullmatch(r"[0-9a-f]{64}", str(
            origin.get("backup_manifest_sha256") or ""))
            or origin.get("backup_member") != f"data/cdpm-runs/{source_run}/extract/model.json"):
        _fail("свидетельство: нет адреса проверенного резервного источника")
    source_revision = _source_revision(witness.get("source_revision"), revision)
    meta = source_doc.get("meta")
    if (not isinstance(meta, dict) or type(meta.get("revision")) is not int
            or meta["revision"] != source_revision
            or _json(summary_bytes).get("read_problems") != []):
        _fail("исходная модель / чтение книги не подтверждают ревизию")
    source_questions, editions = source_doc.get("questions"), witness.get("questions")
    if (not isinstance(source_questions, list) or not isinstance(editions, list)
            or not editions or any(not isinstance(x, dict) or not isinstance(x.get("id"), str)
                                   for x in source_questions + editions)):
        _fail("свидетельство: неверные исходные вопросы")
    by_id = {x["id"]: x for x in source_questions}
    if len(by_id) != len(source_questions) or len({x["id"] for x in editions}) != len(editions):
        _fail("свидетельство: неоднозначный исходный вопрос")
    for edition in editions:
        before = by_id.get(edition["id"])
        if before is None or not model.is_open_question(before):
            _fail("свидетельство: чужой или уже отвеченный исходный вопрос")
        try:
            actual = provenance.question_edition(before)
        except ValueError as exc:
            raise EvidenceError(f"свидетельство: неверное исходное издание: {exc}") from exc
        if actual != edition:
            _fail("свидетельство: исходное издание вопроса изменено")
    if provenance.question_edition(q) not in editions:
        _fail("свидетельство не относится к применённому изданию вопроса")
    try:
        reading = answers.read_book(book, source_doc)
    except Exception as exc:
        raise EvidenceError(f"свидетельство: книга не читается по исходной модели: {exc}") from exc
    if reading.problems or reading.refused or reading.answers.get(q["id"]) != q["answer"]["raw"]:
        _fail("свидетельство: исходный приёмник отказал книге / ответу")
    files = {witness_name: witness_bytes, model_name: source_bytes}
    return source_revision, files, _ref(witness_name, witness_bytes)


def build_recovery_witness(run_dir: Path, model_bytes: bytes, source_model_bytes: bytes, *,
                           applied_revision: int, expected_source_model_sha256: str,
                           backup_manifest_sha256: str, backup_member: str) -> dict[str, bytes]:
    """Prepare two additive payloads; NEVER write or repair a historical receipt.

    The caller must independently verify the named immutable backup and its
    expected model hash. Origin is an explicit posthoc attestation/reference,
    not a signature or a claim that this witness existed at answer application.
    Export/import revalidate included source bytes and the ordinary book reader.
    """
    from cdpm import provenance
    if (_sha(source_model_bytes) != expected_source_model_sha256
            or type(applied_revision) is not int or applied_revision < 2):
        _fail("свидетельство: неверный pinned источник / ревизия применения")
    doc, source_doc = _json(model_bytes), _json(source_model_bytes)
    if not isinstance(source_doc, dict) or not isinstance(source_doc.get("meta"), dict):
        _fail("свидетельство: неверный вид исходной модели")
    questions = [q for q in _applied(doc) if q["answer"]["revision"] == applied_revision]
    if not questions:
        _fail("свидетельство: нет применённых вопросов этой ревизии")
    prefix = f"answers/rev-{applied_revision:03d}/"
    book_name, summary_name = prefix + "anketa-otvety.xlsx", prefix + "answers-summary.json"
    book, summary_bytes = (run_dir / book_name).read_bytes(), (run_dir / summary_name).read_bytes()
    if "source_revision" in _json(summary_bytes):
        _fail("в квитанции уже есть исходная ревизия; posthoc замена запрещена")
    source_revision = _source_revision(source_doc["meta"].get("revision"), applied_revision)
    if source_revision == applied_revision - 1:
        _fail("последовательная legacy ревизия не требует posthoc свидетельства")
    model_name, witness_name = prefix + "source-model-recovery.json", prefix + "source-revision-recovery.json"
    witness = {"schema": RECOVERY_SCHEMA, "kind": "posthoc-pre-apply-model",
               "source_run": run_dir.name, "revision": applied_revision,
               "source_revision": source_revision,
               "book": _ref(book_name, book), "summary": _ref(summary_name, summary_bytes),
               "source_model": _ref(model_name, source_model_bytes),
               "questions": [provenance.question_edition(q) for q in questions],
               "origin": {"backup_manifest_sha256": backup_manifest_sha256,
                          "backup_member": backup_member}}
    files = {model_name: source_model_bytes, witness_name: _bytes(witness)}
    overlay = {**files, book_name: book, summary_name: summary_bytes}
    _assemble(model_bytes, lambda name: overlay[name] if name in overlay else
              (run_dir / name).read_bytes(), run_dir.name)
    return files


def _assemble(model_bytes, read, source_run):
    from cdpm import provenance
    doc = _json(model_bytes)
    questions = _applied(doc)
    if not questions:
        return {}, []
    if not isinstance(source_run, str) or not re.fullmatch(r"[\w-]+", source_run):
        _fail("неверный исходный прогон")
    files, rows, clauses = {}, [], []
    for q in questions:
        a, qid = q["answer"], q["id"]
        prefix = f"answers/rev-{a['revision']:03d}"
        book_name, summary_name = prefix + "/anketa-otvety.xlsx", prefix + "/answers-summary.json"
        try:
            book, summary_bytes = read(book_name), read(summary_name)
        except (OSError, KeyError) as exc:
            raise EvidenceError(f"источники ответов анкеты: {qid}: исходная книга / квитанция отсутствует") from exc
        summary = _json(summary_bytes)
        if not isinstance(summary, dict):
            _fail(f"{qid}: квитанция не объект")
        if (_sha(book) != a["answer_version"]["book_sha256"]
                or summary.get("book_sha256") != _sha(book)
                or summary.get("revision") != a["revision"]):
            _fail(f"{qid}: книга / квитанция не соответствуют применённой ревизии")
        applied = [summary.get("applied", 0), summary.get("worker_applied", 0)]
        if any(type(x) is not int or x < 0 for x in applied) or sum(applied) == 0:
            _fail(f"{qid}: квитанция не подтверждает применения")
        for key in ("refused", "worker_rejected"):
            if not isinstance(summary.get(key, []), list):
                _fail(f"{qid}: неверный журнал отказов {key}")
            for item in summary.get(key) or []:
                if ((isinstance(item, dict) and item.get("question") == qid)
                        or (isinstance(item, str) and qid in item)):
                    _fail(f"{qid}: ответ / его обработка были отклонены")
        source_revision = (_source_revision(summary["source_revision"], a["revision"])
                           if "source_revision" in summary else None)
        recovery = None
        try:
            printed = _book_row(book, doc, q, source_revision)
        except _BookRevisionMismatch:
            if "source_revision" in summary:
                raise  # a conflicting explicit receipt cannot borrow a witness
            source_revision, recovered, recovery = _recovery_revision(
                read, source_run, q, book, summary_bytes)
            printed = _book_row(book, doc, q, source_revision)
            files.update(recovered)
        files.update({book_name: book, summary_name: summary_bytes})
        entry = {"question": provenance.question_edition(q),
                 "answer": {"revision": a["revision"], **a["answer_version"]},
                 "book": _ref(book_name, book), "summary": _ref(summary_name, summary_bytes),
                 "rows": printed}
        if source_revision is not None:
            entry["answer"]["source_revision"] = source_revision
        if recovery is not None:
            entry["source_revision_recovery"] = recovery
        rows.append(entry)
        first = printed[0]
        anchor = f"{qid}; ответ rev-{a['revision']:03d}; {first['sheet']}:{first['row']}"
        clauses.append({"sid": "ANSWER-" + _sha(qid.encode())[:12], "anchor": anchor,
                        "text": (f"Исходный вопрос анкеты:\n{first['question']}\n"
                                 f"Формат и варианты ответа:\n{first['format']}\n"
                                 f"Область вопроса: {first['targets']}\n"
                                 f"Исходные документы: {first['sources']}\n"
                                 f"Дословный применённый ответ авиакомпании:\n{first['raw']}"),
                        "answer_source": entry})
    files[EVIDENCE] = _bytes({"schema": SCHEMA, "source_run": source_run,
                              "revision": doc["meta"]["revision"],
                              "model_sha256": _sha(model_bytes), "answers": rows})
    return files, clauses


def export_files(run_dir: Path, model_bytes: bytes, source_run: str) -> dict[str, bytes]:
    """Additional manifest.files payloads; only immutable source books/receipts."""
    files, _ = _assemble(model_bytes, lambda name: (run_dir / name).read_bytes(), source_run)
    return files


def package_sources(archive, source_run):
    """Validate all answer evidence BEFORE extraction; returns needed payloads.

    Legacy clause-only packages without any answers stay readable. A package
    declaring answers may not borrow local files, even if its entire evidence
    index and manifest entries were removed together.
    """
    names = archive.namelist()
    if len(names) != len(set(names)):
        _fail("повтор имени файла ZIP")
    manifest_bytes = archive.read("manifest.json")
    manifest = _json(manifest_bytes)
    if not isinstance(manifest, dict):
        _fail("манифест не объект")
    if manifest.get("source_run", source_run) != source_run:
        _fail("манифест другого прогона")
    listed = {}
    for item in manifest.get("files", []):
        if not isinstance(item, dict) or not isinstance(item.get("name"), str) or item["name"] in listed:
            _fail("неверный / повторный файл манифеста")
        listed[item["name"]] = item

    def read(name, *, required=True):
        try:
            body = archive.read(name)
        except KeyError as exc:
            raise EvidenceError(f"источники ответов анкеты: нет файла {name}") from exc
        rec = listed.get(name)
        if required and rec is None:
            _fail(f"{name}: нет привязки в manifest.files")
        if rec is not None and (rec.get("sha256") != _sha(body)
                                or rec.get("bytes", len(body)) != len(body)):
            _fail(f"{name}: SHA / размер не совпадает")
        return body

    model = read("model.json", required=False) if "model.json" in names or "model.json" in listed else b"{}"
    needed, _ = _assemble(model, read, source_run)
    if needed or EVIDENCE in names or EVIDENCE in listed:
        if not needed:
            _fail("есть evidence, но нет применённых ответов этой модели")
        read("model.json")
        if manifest.get("revision") != (_json(model).get("meta") or {}).get("revision"):
            _fail("ревизия манифеста не соответствует ответам / модели")
        actual = read(EVIDENCE)
        if _json(actual) != _json(needed[EVIDENCE]):
            _fail("evidence не соответствует текущему пакету / ревизии / ответам")
        needed[EVIDENCE] = actual
    clause_names = [n for n in names if n.startswith("freeze/clauses/") and n.endswith(".json")]
    if any(not re.fullmatch(r"freeze/clauses/[^/\\.][^/\\]*\.json", n)
           or re.search(r'[:*?"<>|\x00-\x1f]', n) for n in clause_names):
        _fail("небезопасное имя клаузы")
    if needed and not clause_names:
        _fail("применённые ответы без исходных клауз пакета")
    out = {n: read(n, required=bool(needed)) for n in clause_names}
    if clause_names:
        # Persist the package binding even for answer-free packages. Deleting
        # evidence after import cannot turn an answered package into legacy.
        out.update(needed)
        out[PACKAGE_MANIFEST] = manifest_bytes
        out[PACKAGE_MODEL] = model
    return out


def clauses_from_run(run_dir: Path) -> list[dict]:
    """Load only validated original question/answer sources, never parsed cells."""
    marker = run_dir / PACKAGE_MANIFEST
    if marker.exists():
        manifest = _json(marker.read_bytes())
        model = (run_dir / PACKAGE_MODEL).read_bytes()
        class LocalArchive:
            def namelist(self):
                return ["manifest.json", *[r["name"] for r in manifest.get("files", [])]]

            def read(self, name):
                if name == "manifest.json":
                    return marker.read_bytes()
                if name == "model.json":
                    return model
                return (run_dir / name).read_bytes()
        # Revalidate required sources on each call, including post-import loss.
        package_sources(LocalArchive(), run_dir.name)
        _, clauses = _assemble(model, lambda name: (run_dir / name).read_bytes(), run_dir.name)
        return clauses
    if (run_dir / EVIDENCE).exists() or (run_dir / PACKAGE_MODEL).exists():
        _fail("извлечённое evidence потеряло привязку к пакету")
    model = run_dir / "extract/model.json"
    if not model.is_file():
        return []  # historical answer-free direct clause tree
    _, clauses = _assemble(model.read_bytes(), lambda name: (run_dir / name).read_bytes(), run_dir.name)
    return clauses
