# -*- coding: utf-8 -*-
"""Заморозка источников CDPM (пакет работ П-2).

Превращает исходные документы правил (doc/docx/md/xlsx — объём первой
очереди, решение владельца 2026-08-17; jpg/jpeg/png — изображения
документов, пакет ПИ-1) в замороженный комплект:

- **машинная форма клауз** (`clauses/<sid>.json`) — источник истины для
  извлечения правил (пакет П-3): каждая клауза несёт якорь (адрес внутри
  документа), вид и текст с сохранённой смысловой разметкой;
- **каноническая md-проекция** (`projections/<sid>-<имя>.md`) —
  человекочитаемое отображение клауз; рендерится из машинной формы
  детерминированно, как свод и анкета из модели;
- **манифест** (`manifest.json`) — SHA-256 каждого оригинала и правило
  арбитража: проекция и конвертированный файл — проекции, при
  расхождении первичен оригинал.

Уроки, зашитые в контракт:

- нарезка на уровне КЛАУЗ, не абзацев целиком слитого текста: шесть
  пропущенных клауз исходного документа стоили седьмого обмена аудита;
- конвертация ТЕРЯЕТ смысловую разметку: жирный шрифт пункта П.1.9
  нёс норму — жирные фрагменты сохраняются маркерами `**…**` и в
  клаузах, и в проекции;
- вложенные таблицы docx извлекаются (дословное извлечение вложенных
  таблиц исходного документа появилось только ревизией v6 — здесь с
  первого дня);
- `.doc` конвертируется в `.docx` силами Microsoft Word (скрипт
  `cdpm/scripts/convert_doc.ps1`); Word отсутствует — честный отказ с
  названной причиной, а не тихий пропуск файла;
- изображение (`.jpg`/`.jpeg`/`.png`) распознаётся зрительной моделью
  (канал передаётся параметром `vision_llm`): текст переписывается
  ДОСЛОВНО абзацами, запись источника получает признак
  `"recognized": "image"`, манифест — SHA-256 самого изображения;
  канала нет — честный отказ, не тихий пропуск;
- `.pdf` (отмашка владельца 2026-08-26) читается ПОСТРАНИЧНО:
  текстовый слой — абзацы с якорями «стр.P, абз.N» (нумерованный
  пункт — якорь «п.X») и жирностью; страница-скан (без текстового
  слоя либо крохотный слой при картинке) рендерится в изображение и
  идёт зрительной моделью — запись получает `"recognized": "image"`
  и перечень распознанных страниц `recognized_pages`; зашифрованный
  PDF и скан без канала зрения — честные отказы;
- `.json` (отмашка владельца 2026-08-26): каждое скалярное значение —
  клауза с якорем-путём по ключам («rules[0].text», индексы с нуля);
  повторяющийся ключ объекта — честный отказ (штатный разбор молча
  теряет первое значение);
- замки чтения по аудиту 2026-09-04 (п. 2): в docx абзацы и таблицы
  идут в порядке тела документа (таблица стоит между своими абзацами,
  якоря прежние); в xlsx формула без сохранённого значения — честный
  отказ «книга сохранена без пересчёта формул» с адресом ячейки
  (формула с сохранённым результатом, пусть и пустой строкой,
  прочитана как значение); в pdf страница с текстовым слоем и
  рисунком, который слой не покрывает, помечается `mixed_pages`,
  рисунки при канале зрения распознаются с якорями «стр.P, рис.K,
  абз.N», рисунок без текста по ответу модели — след
  `mixed_empty_figures`; без канала зрения пометка идёт в статус
  прогона предупреждением (`mixed_warnings`).

Запуск отдельно от приложения:
    python -m cdpm.freeze ФАЙЛ... --out-dir КАТАЛОГ [--cloud-images]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

ARBITRATION = ("конвертированный файл и md-проекция — проекции; при "
               "расхождении первичен оригинал")

# Изображения документов (ПИ-1): распознаются зрительной моделью.
IMAGE_SUPPORTED = (".jpg", ".jpeg", ".png")

SUPPORTED = (".doc", ".docx", ".md", ".xlsx", ".pdf",
             ".json") + IMAGE_SUPPORTED

# PDF (отмашка владельца 2026-08-26): страница с текстовым слоем короче
# этого порога ПРИ НАЛИЧИИ картинки считается сканом с подписью —
# содержимое в картинке, текстовый слой его не покрывает; страница
# уходит зрительной модели целиком. Тот же порог решает, покрыт ли
# текстовым слоем отдельный рисунок текстовой страницы (смешанная
# страница, аудит 2026-09-04, п. 2).
PDF_SCAN_TEXT_MIN = 25

# Рисунок текстовой страницы PDF: изображение или связная область
# векторной графики не меньше этого размера (пунктов) в обе стороны;
# линейки, подчёркивания и рамки-полоски рисунками не считаются.
PDF_FIGURE_MIN_PT = 20

# Связность векторной графики: соседние контуры сливаются в одну
# область с этим допуском (пунктов) — таблица из сотен ячеек
# схлопывается в одну область, покрытую своим текстом, а не даёт
# сотни «рисунков» (замечание ревизии 2026-09-04).
PDF_FIGURE_JOIN_PT = 2

# Одиночная заливка или рамка (контур из одного элемента — фон ячейки
# таблицы, линейка, рамка) рисунком не считается; область векторной
# графики признаётся рисунком (схема, диаграмма) при таком суммарном
# числе элементов контуров, не меньше.
PDF_FIGURE_MIN_ITEMS = 8

_CONVERT_SCRIPT = Path(__file__).parent / "scripts" / "convert_doc.ps1"

# Номер пункта в начале клаузы: «1.9.», «2.3)», «П.1.4», «Р.1.14.» —
# становится якорем «п.1.9»; без номера якорь позиционный («абз.N»).
_NUM_RE = re.compile(r"^\s*(?:[ПпРр]\.\s*)?(\d+(?:\.\d+)+|\d+)[.)]\s+")


class FreezeError(ValueError):
    """Честный отказ заморозки: причина всегда названа."""


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------------------
# Извлечение клауз по форматам
# ---------------------------------------------------------------------------

def _wrap_bold(segments: list[tuple[str, bool]]) -> str:
    """Склейка прогонов текста с маркировкой жирных фрагментов `**…**`.
    Пробелы по краям жирного фрагмента выносятся наружу маркеров, иначе
    Markdown-разметка недействительна."""
    out = []
    for text, bold in segments:
        if not text:
            continue
        if bold and text.strip():
            lead = text[:len(text) - len(text.lstrip())]
            trail = text[len(text.rstrip()):]
            out.append(f"{lead}**{text.strip()}**{trail}")
        else:
            out.append(text)
    return "".join(out)


def _para_text(para) -> str:
    """Текст абзаца docx с сохранением жирности прогонов; соседние
    прогоны одной жирности сливаются, чтобы `**а****б**` не рвало
    разметку."""
    segments: list[tuple[str, bool]] = []
    for run in para.runs:
        bold = bool(run.bold)
        if segments and segments[-1][1] == bold:
            segments[-1] = (segments[-1][0] + run.text, bold)
        else:
            segments.append((run.text, bold))
    return _wrap_bold(segments)


class _Anchors:
    """Раздатчик якорей: повторный якорь получает суффикс « (2)», « (3)»
    — двусмысленный адрес хуже длинного (урок якорей Е1: 14 номеров на
    двух листах требуют суффикса листа)."""

    def __init__(self):
        self._seen: dict[str, int] = {}

    def take(self, anchor: str) -> str:
        n = self._seen.get(anchor, 0) + 1
        self._seen[anchor] = n
        return anchor if n == 1 else f"{anchor} ({n})"


def _numbered_anchor(text: str) -> str | None:
    m = _NUM_RE.match(text.replace("**", "", 2))
    return f"п.{m.group(1)}" if m else None


def _docx_body_items(doc):
    """Абзацы и таблицы тела документа В ПОРЯДКЕ ИСХОДНИКА (аудит
    2026-09-04, п. 2: раздельные обходы `doc.paragraphs` и `doc.tables`
    уносили таблицу в конец, и «исключение к таблице выше» теряло
    адресата по позиции). Обход детей `w:body` — штатный приём
    python-docx; состав элементов тот же, что у `doc.paragraphs` /
    `doc.tables` (прямые дети тела), поэтому якоря не меняются —
    меняется только взаимный порядок абзацев и таблиц в списке клауз."""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", Paragraph(child, doc._body)
        elif child.tag == qn("w:tbl"):
            yield "tbl", Table(child, doc._body)


def _extract_docx(path: Path) -> list[dict]:
    import docx
    doc = docx.Document(str(path))
    anchors = _Anchors()
    clauses: list[dict] = []
    n_para = 0
    n_table = 0

    def walk_table(table, base: str):
        for r, row in enumerate(table.rows, start=1):
            for c, cell in enumerate(row.cells, start=1):
                text = " ".join(
                    t for t in (_para_text(p).strip()
                                for p in cell.paragraphs) if t)
                if text:
                    clauses.append({
                        "anchor": anchors.take(
                            f"{base}, строка {r}, кол {c}"),
                        "kind": "cell", "text": text})
                for k, nested in enumerate(cell.tables, start=1):
                    # Вложенные таблицы исходного документа несли
                    # дословные нормы — извлекаются рекурсивно, не
                    # теряются.
                    walk_table(nested,
                               f"{base}, строка {r}, кол {c}, вл.табл.{k}")

    # Порядок клауз = порядок тела документа: таблица стоит между теми
    # абзацами, между которыми она стоит в оригинале. Нумерация якорей
    # («абз.N» — по непустым абзацам, «табл.T» — по таблицам) прежняя.
    for kind, item in _docx_body_items(doc):
        if kind == "p":
            text = _para_text(item).strip()
            if not text:
                continue
            n_para += 1
            anchor = _numbered_anchor(text) or f"абз.{n_para}"
            clauses.append({"anchor": anchors.take(anchor), "kind": "p",
                            "text": text})
        else:
            n_table += 1
            walk_table(item, f"табл.{n_table}")
    return clauses


def _xlsx_uncalculated_cells(path: Path, wb_values) -> list[str]:
    """Адреса ячеек с формулой, у которых в книге НЕТ сохранённого
    значения (аудит 2026-09-04, п. 2: книга, сохранённая без пересчёта
    формул, теряла значение ячейки молча). Формулы видны только при
    загрузке без `data_only`; сохранённое значение — при загрузке с
    `data_only`. Непересчитанной считается только ячейка БЕЗ
    типизированного кэша (тип `n` и значение None — так пишет книгу
    без пересчёта openpyxl: `<c><f>…</f><v></v></c>`); формула с
    сохранённым результатом любого типа — строкой, в том числе
    пустой (`t="str"`, `<v></v>`: IF(A1="","",…) в журналах и
    таблицах норм), числом, логическим, ошибкой — пересчитана и
    читается как значение (ревизия 2026-09-04: пустая строка —
    честный результат, а не потеря)."""
    import openpyxl
    wb_formulas = openpyxl.load_workbook(str(path), data_only=False)
    missing: list[str] = []
    for ws_f in wb_formulas.worksheets:
        ws_v = wb_values[ws_f.title]
        for row in ws_f.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                cached = ws_v[cell.coordinate]
                if cached.value is None and cached.data_type == "n":
                    missing.append(f"лист «{ws_f.title}», ячейка "
                                   f"{cell.coordinate}")
    return missing


def _extract_xlsx(path: Path) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    missing = _xlsx_uncalculated_cells(path, wb)
    if missing:
        shown = "; ".join(missing[:5])
        more = (f" и ещё {len(missing) - 5}" if len(missing) > 5
                else "")
        raise FreezeError(
            f"{path.name}: {shown}{more} — формула без сохранённого "
            f"значения: книга сохранена без пересчёта формул, значение "
            f"ячейки потерялось бы молча (тихая потеря нормы "
            f"запрещена); откройте книгу в Excel, пересчитайте, "
            f"сохраните и загрузите заново")
    anchors = _Anchors()
    clauses: list[dict] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            cells = []
            for cell in row:
                v = cell.value
                if v is None or str(v).strip() == "":
                    continue
                text = str(v).strip()
                if cell.font is not None and cell.font.b:
                    text = f"**{text}**"
                cells.append(text)
            if not cells:
                continue
            clauses.append({
                "anchor": anchors.take(
                    f"лист «{ws.title}», строка {row[0].row}"),
                "kind": "row", "text": " | ".join(cells)})
    return clauses


def _extract_json(path: Path) -> list[dict]:
    """Клаузы JSON-документа (отмашка владельца 2026-08-26): каждое
    скалярное значение — клауза с якорем-путём по ключам (индексы
    массивов с нуля, как в инструментах JSON: «rules[0].text»).
    Строки идут дословно; числа, истина/ложь и null — литералами
    JSON. Повторяющийся ключ объекта — честный отказ: штатный разбор
    молча теряет первое значение, а тихая потеря нормы запрещена.
    Пустые контейнеры пропускаются (нормы не несут)."""
    def no_dup(pairs):
        seen = {}
        for k, v in pairs:
            if k in seen:
                raise FreezeError(
                    f"{path.name}: повторяющийся ключ {k!r} в объекте "
                    f"— штатный разбор JSON молча потерял бы первое "
                    f"значение (тихая потеря нормы запрещена)")
            seen[k] = v
        return seen

    try:
        data = json.loads(path.read_text(encoding="utf-8"),
                          object_pairs_hook=no_dup)
    except FreezeError:
        raise
    except UnicodeDecodeError as e:
        raise FreezeError(f"{path.name}: файл не в UTF-8: {e}")
    except json.JSONDecodeError as e:
        raise FreezeError(
            f"{path.name}: JSON не разбирается — строка {e.lineno}, "
            f"позиция {e.colno}: {e.msg}")
    anchors = _Anchors()
    clauses: list[dict] = []

    def walk(node, trail: str):
        if isinstance(node, dict):
            for k, v in node.items():
                walk(v, f"{trail}.{k}" if trail else str(k))
        elif isinstance(node, list):
            for i, v in enumerate(node):
                walk(v, f"{trail}[{i}]" if trail else f"[{i}]")
        else:
            if isinstance(node, str):
                text = node.strip()
                if not text:
                    return
            elif isinstance(node, bool):
                text = "true" if node else "false"
            elif node is None:
                text = "null"
            else:
                text = json.dumps(node)
            clauses.append({"anchor": anchors.take(trail or "корень"),
                            "kind": "field", "text": text})

    walk(data, "")
    if not clauses:
        raise FreezeError(
            f"{path.name}: из JSON не извлеклось ни одной клаузы "
            f"(документ пуст?)")
    return clauses


def _extract_md(path: Path) -> list[dict]:
    anchors = _Anchors()
    clauses = []
    for i, line in enumerate(
            path.read_text(encoding="utf-8").splitlines(), start=1):
        text = line.strip()
        if not text:
            continue
        anchor = _numbered_anchor(text) or f"строка {i}"
        clauses.append({"anchor": anchors.take(anchor), "kind": "line",
                        "text": text})
    return clauses


# Промпт зрительной модели (ПИ-1): дословность — контракт, улучшение
# текста моделью = подмена оригинала.
_PROMPT_IMAGE = """Ты — распознаватель отсканированных документов
авиакомпании. Прочитай изображение и перепиши ВЕСЬ текст с него
ДОСЛОВНО: с опечатками, номерами пунктов и регистром оригинала — без
исправлений, улучшений и пересказа. Разбей текст на абзацы так, как
они расположены на изображении. Верни СТРОГО один JSON-объект без
пояснений:

{"paragraphs": ["<текст абзаца 1>", "<текст абзаца 2>", ...]}

Текста на изображении нет — верни {"paragraphs": []}."""


def _check_image_reply(got: dict,
                       allow_empty: bool = False) -> list[str] | str:
    """Список абзацев из ответа модели либо причина невалидности
    (строка). Контракт: список непустых строк; пустой список —
    невалиден для документа целиком (текст там есть по определению),
    но допустим для рисунка на текстовой странице (`allow_empty`:
    схема, эмблема — рисунок без текста нормы не несёт)."""
    paras = got.get("paragraphs")
    if not isinstance(paras, list):
        return "в ответе нет списка paragraphs"
    if not paras and not allow_empty:
        return "модель не нашла текста на изображении (пустой список)"
    if not all(isinstance(p, str) and p.strip() for p in paras):
        return "среди абзацев есть пустой или нестроковый"
    return [p.strip() for p in paras]


def _recognize_paragraphs(path: Path, llm, what: str,
                          allow_empty: bool = False) -> list[str]:
    """Дословные абзацы с изображения: зрительная модель переписывает
    текст, код валидирует («модель предлагает — код проверяет»).
    Невалидный ответ — одна повторная попытка, затем честный провал
    с именем источника `what`."""
    from cdpm import llm as llm_mod
    retry = _PROMPT_IMAGE + ("\n\nПредыдущий ответ не разобрался. "
                             "Верни СТРОГО один JSON-объект по схеме "
                             "выше, без какого-либо текста вокруг.")
    reason = "ответа не было"
    for ask in (_PROMPT_IMAGE, retry):
        try:
            got = llm_mod.parse_json_reply(llm(ask, path))
        except llm_mod.LLMReplyError as e:
            reason = str(e)
            continue
        paras = _check_image_reply(got, allow_empty=allow_empty)
        if isinstance(paras, str):
            reason = paras
            continue
        return paras
    raise FreezeError(
        f"{what}: зрительная модель не дала валидного ответа "
        f"после 2 попыток: {reason}")


def _extract_image(path: Path, llm) -> list[dict]:
    """Клаузы изображения: последовательные якоря «стр-1-абз-N»."""
    return [{"anchor": f"стр-1-абз-{n}", "kind": "p", "text": t}
            for n, t in enumerate(
                _recognize_paragraphs(path, llm, path.name), start=1)]


def _pdf_text_paragraphs(page) -> list[str]:
    """Абзацы текстового слоя страницы PDF: блок разметки = абзац
    (переносы строк внутри блока — вёрстка, склеиваются пробелом);
    жирные прогоны сохраняются маркерами `**…**` (урок П.1.9 — жирный
    шрифт нёс норму)."""
    paras: list[str] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        segments: list[tuple[str, bool]] = []
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                t = span.get("text", "")
                if not t:
                    continue
                bold = bool(span.get("flags", 0) & 16) \
                    or "bold" in (span.get("font") or "").lower()
                if segments and segments[-1][1] == bold:
                    segments[-1] = (segments[-1][0] + t, bold)
                else:
                    segments.append((t, bold))
            if segments and not segments[-1][0].endswith(" "):
                segments[-1] = (segments[-1][0] + " ",
                                segments[-1][1])
        text = _wrap_bold(segments).strip()
        if text:
            paras.append(text)
    return paras


def _pdf_figures(page) -> list:
    """Рисунки текстовой страницы, которые текстовый слой НЕ покрывает
    (аудит 2026-09-04, п. 2: картинка с дополнительной таблицей на
    странице с текстом молча проходила текстовым путём).

    Рисунок — встроенное изображение либо связная область векторной
    графики размером не меньше PDF_FIGURE_MIN_PT в обе стороны (линейки
    и подчёркивания — не рисунки); рисунок «не покрыт», если внутри
    его области текстового слоя меньше PDF_SCAN_TEXT_MIN знаков (тот же
    порог, что для скана с подписью). Изображение с распознанным
    поверх него слоем (скан после OCR) и таблица, начерченная
    рамками и заливками вокруг текста, рисунками не считаются — их
    содержимое уже в клаузах текстового слоя: одиночные заливки и
    рамки (контур из одного элемента) пропускаются, соседние контуры
    сливаются с допуском PDF_FIGURE_JOIN_PT, область признаётся
    рисунком при PDF_FIGURE_MIN_ITEMS элементах и больше (ревизия
    2026-09-04: экспорт Word/LibreOffice с таблицами давал сотни
    ложных рисунков). Возвращает прямоугольники в порядке чтения
    (сверху вниз, слева направо)."""
    import fitz
    page_rect = page.rect
    rects: list = []
    for info in page.get_images(full=True):
        for r in page.get_image_rects(info[0]):
            rects.append(fitz.Rect(r) & page_rect)
    # Векторная графика: связные области — слияние соседних (с
    # допуском) прямоугольников с подсчётом элементов, иначе
    # диаграмма из сотен штрихов даст сотни «рисунков», а таблица —
    # по «рисунку» на ячейку.
    join = PDF_FIGURE_JOIN_PT
    clusters: list = []                       # [(rect, items)]
    for d in page.get_drawings():
        n_items = len(d.get("items") or [])
        if n_items < 2:
            continue                          # фон ячейки, линейка, рамка
        r = fitz.Rect(d.get("rect") or fitz.Rect())
        if r.is_empty:
            continue
        r = r & page_rect
        merged = True
        while merged:
            merged = False
            grown = r + (-join, -join, join, join)
            for i, (c, c_items) in enumerate(clusters):
                if grown.intersects(c):
                    r = r | c
                    n_items += c_items
                    del clusters[i]
                    merged = True
                    break
        clusters.append((r, n_items))
    rects.extend(r for r, n in clusters if n >= PDF_FIGURE_MIN_ITEMS)
    figures = []
    for r in rects:
        if r.is_empty or r.width < PDF_FIGURE_MIN_PT \
                or r.height < PDF_FIGURE_MIN_PT:
            continue
        covered = len(page.get_text("text", clip=r).strip())
        if covered >= PDF_SCAN_TEXT_MIN:
            continue
        figures.append(r)
    figures.sort(key=lambda r: (round(r.y0), round(r.x0)))
    return figures


def _extract_pdf(path: Path, vision_llm=None) -> dict:
    """Клаузы PDF постранично (отмашка владельца 2026-08-26).

    Страница с текстовым слоем → абзацы с якорями «стр.P, абз.N»
    (нумерованный пункт даёт якорь «п.X», как у docx/md); страница
    БЕЗ текстового слоя (скан) или с крохотным слоем при картинке →
    рендер в изображение и существующий зрительный канал, дословно.
    Пустая страница (ни текста, ни графики) пропускается — нормы на
    ней нет. Зашифрованный PDF и скан без канала зрения — честные
    отказы с названной страницей.

    Смешанная страница (текстовый слой И рисунок, который слой не
    покрывает; аудит 2026-09-04, п. 2) попадает в `mixed_pages`;
    при канале зрения каждый рисунок дополнительно распознаётся
    дословно с якорями «стр.P, рис.K, абз.N» (`mixed_recognized`
    истинно), без канала — рисунки остаются нераспознанными, и запись
    об этом несут манифест и проекция (текстовый слой страницы при
    этом извлечён, отказа нет). Пустой ответ модели на рисунок
    (эмблема, схема без текста) допустим, но оставляет след:
    `mixed_empty_figures` — перечень «стр.P, рис.K» без единой
    клаузы (ревизия 2026-09-04: тихая потеря картинки-таблицы при
    ленивом ответе модели запрещена)."""
    try:
        import fitz
    except ImportError:
        raise FreezeError(
            f"{path.name}: чтение PDF требует библиотеки PyMuPDF "
            f"(fitz) — она не установлена в окружении")
    anchors = _Anchors()
    clauses: list[dict] = []
    vision_pages: list[int] = []
    mixed_pages: list[int] = []
    mixed_recognized = False
    empty_figures: list[str] = []
    doc = fitz.open(str(path))
    try:
        if doc.is_encrypted:
            raise FreezeError(
                f"{path.name}: PDF зашифрован — снимите защиту и "
                f"загрузите заново (тихая расшифровка = подмена "
                f"оригинала)")
        if doc.page_count == 0:
            raise FreezeError(f"{path.name}: в PDF нет страниц")
        tmp: Path | None = None
        for pno in range(doc.page_count):
            page, p = doc[pno], pno + 1
            paras = _pdf_text_paragraphs(page)
            text_len = sum(len(t.replace("**", "")) for t in paras)
            has_image = bool(page.get_images(full=True))
            if text_len > 0 and (text_len >= PDF_SCAN_TEXT_MIN
                                 or not has_image):
                n = 0
                for t in paras:
                    n += 1
                    anchor = _numbered_anchor(t) \
                        or f"стр.{p}, абз.{n}"
                    clauses.append({"anchor": anchors.take(anchor),
                                    "kind": "p", "text": t})
                figures = _pdf_figures(page)
                if not figures:
                    continue
                mixed_pages.append(p)
                if vision_llm is None:
                    continue
                mixed_recognized = True
                if tmp is None:
                    tmp = Path(tempfile.mkdtemp(prefix="cdpm-pdf-"))
                for k, rect in enumerate(figures, start=1):
                    png = tmp / f"page-{p}-fig-{k}.png"
                    page.get_pixmap(dpi=200, clip=rect).save(str(png))
                    got = _recognize_paragraphs(
                        png, vision_llm,
                        f"{path.name}, страница {p}, рис. {k}",
                        allow_empty=True)
                    if not got:
                        empty_figures.append(f"стр.{p}, рис.{k}")
                    for n, t in enumerate(got, start=1):
                        anchor = _numbered_anchor(t) \
                            or f"стр.{p}, рис.{k}, абз.{n}"
                        clauses.append({"anchor": anchors.take(anchor),
                                        "kind": "p", "text": t})
                continue
            if not has_image and not page.get_drawings():
                continue                    # истинно пустая страница
            if vision_llm is None:
                raise FreezeError(
                    f"{path.name}: страница {p} без текстового слоя "
                    f"(скан) требует зрительной модели, а канал "
                    f"зрения не передан")
            if tmp is None:
                tmp = Path(tempfile.mkdtemp(prefix="cdpm-pdf-"))
            png = tmp / f"page-{p}.png"
            page.get_pixmap(dpi=200).save(str(png))
            got = _recognize_paragraphs(
                png, vision_llm, f"{path.name}, страница {p}")
            vision_pages.append(p)
            for n, t in enumerate(got, start=1):
                anchor = _numbered_anchor(t) or f"стр.{p}, абз.{n}"
                clauses.append({"anchor": anchors.take(anchor),
                                "kind": "p", "text": t})
    finally:
        doc.close()
    if not clauses:
        raise FreezeError(
            f"{path.name}: из PDF не извлеклось ни одной клаузы "
            f"(документ пуст?)")
    result: dict = {"clauses": clauses, "converted": None}
    if vision_pages or mixed_recognized:
        # Распознанный текст недетерминирован — гейт заморозки повторно
        # не извлекает такие источники (признак «recognized»).
        result["recognized"] = "image"
    if vision_pages:
        result["recognized_pages"] = vision_pages
    if mixed_pages:
        result["mixed_pages"] = mixed_pages
        result["mixed_recognized"] = mixed_recognized
    if empty_figures:
        result["mixed_empty_figures"] = empty_figures
    return result


def _convert_doc(path: Path, out_docx: Path) -> None:
    """Конвертация .doc → .docx силами Microsoft Word (COM). Отказ —
    всегда с названной причиной: тихо пропущенный документ = дырка в
    покрытии, которую никто не увидит."""
    proc = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
         "-File", str(_CONVERT_SCRIPT),
         "-In", str(path), "-Out", str(out_docx)],
        capture_output=True, text=True, timeout=300)
    marker = (proc.stdout or "").strip()
    if proc.returncode == 3 or marker == "NO-WORD":
        raise FreezeError(
            f"{path.name}: формат .doc требует Microsoft Word на машине "
            f"(конвертация в .docx); Word не найден")
    if proc.returncode != 0 or not out_docx.is_file():
        raise FreezeError(
            f"{path.name}: конвертация .doc не удалась: "
            f"{marker or proc.stderr.strip()}")


def extract_file(path: Path, convert_dir: Path | None = None,
                 vision_llm=None) -> dict:
    """Клаузы одного файла + сведения о конвертации (для .doc).
    Возвращает {"clauses": [...], "converted": {...} | None} и для
    изображений — признак происхождения {"recognized": "image"}.
    `vision_llm` — зрительный канал (промпт, путь) -> ответ; без него
    изображение — честный отказ, не тихий пропуск."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise FreezeError(
            f"{path.name}: формат {ext or '<без расширения>'} не входит в "
            f"первую очередь CDPM (поддержаны: {', '.join(SUPPORTED)})")
    if not path.is_file():
        raise FreezeError(f"{path}: файл не существует")
    if ext == ".docx":
        return {"clauses": _extract_docx(path), "converted": None}
    if ext == ".xlsx":
        return {"clauses": _extract_xlsx(path), "converted": None}
    if ext == ".md":
        return {"clauses": _extract_md(path), "converted": None}
    if ext == ".json":
        return {"clauses": _extract_json(path), "converted": None}
    if ext == ".pdf":
        return _extract_pdf(path, vision_llm)
    if ext in IMAGE_SUPPORTED:
        if vision_llm is None:
            raise FreezeError(
                f"{path.name}: изображение требует зрительной модели, "
                f"а канал зрения не передан")
        return {"clauses": _extract_image(path, vision_llm),
                "converted": None, "recognized": "image"}
    # .doc: через конвертацию Word; конвертат — проекция, не оригинал.
    if convert_dir is None:
        tmp = Path(tempfile.mkdtemp(prefix="cdpm-doc-"))
    else:
        tmp = convert_dir
        tmp.mkdir(parents=True, exist_ok=True)
    out_docx = tmp / (path.stem + ".docx")
    _convert_doc(path, out_docx)
    return {"clauses": _extract_docx(out_docx),
            "converted": {"file": str(out_docx),
                          "sha256": sha256_of(out_docx)}}


def word_available() -> bool:
    """Есть ли Microsoft Word для конвертации .doc (нужно тестам, чтобы
    пропускать .doc-сценарии с названной причиной)."""
    probe = subprocess.run(
        ["powershell", "-NoProfile", "-Command",
         "try { $w = New-Object -ComObject Word.Application; $w.Quit(); "
         "'OK' } catch { 'NO' }"],
        capture_output=True, text=True, timeout=120)
    return "OK" in (probe.stdout or "")


# ---------------------------------------------------------------------------
# Проекция и комплект заморозки
# ---------------------------------------------------------------------------

def render_projection(entry: dict) -> str:
    """Каноническая md-проекция одного источника — детерминированный
    рендер машинной формы (та же дисциплина, что свод/анкета из модели:
    правка проекции мимо машинной формы не переживает гейт)."""
    out = [f"# Каноническая проекция: {entry['file']}", ""]
    out += [f"> md-файл — ПРОЕКЦИЯ оригинала для машинной обработки; "
            f"{ARBITRATION}. SHA-256 оригинала: `{entry['sha256']}`. "
            f"Жирная разметка оригинала сохранена маркерами `**…**`."]
    if entry.get("recognized") == "image":
        pages = entry.get("recognized_pages")
        if pages:
            out += ["", "> Страницы " + ", ".join(map(str, pages))
                    + " без текстового слоя РАСПОЗНАНЫ зрительной "
                      "моделью (дословная переписка); при сомнении "
                      "первичен оригинал-PDF."]
        elif not entry.get("mixed_pages"):
            out += ["", "> Текст РАСПОЗНАН зрительной моделью с "
                        "изображения (дословная переписка); при "
                        "сомнении первичен оригинал-изображение."]
    mixed = entry.get("mixed_pages")
    if mixed:
        pages = ", ".join(map(str, mixed))
        if entry.get("mixed_recognized"):
            out += ["", f"> Страницы {pages} смешанные: рисунки сверх "
                        f"текстового слоя РАСПОЗНАНЫ зрительной моделью "
                        f"(якоря «стр.P, рис.K, абз.N», дословная "
                        f"переписка); при сомнении первичен "
                        f"оригинал-PDF."]
        else:
            out += ["", f"> Страницы {pages} смешанные: рисунки сверх "
                        f"текстового слоя НЕ РАСПОЗНАНЫ (канал зрения "
                        f"не передан) — их содержимое в клаузах "
                        f"отсутствует, требуется сверка с "
                        f"оригиналом-PDF."]
    empty = entry.get("mixed_empty_figures")
    if empty:
        out += ["", f"> Рисунки {'; '.join(empty)}: зрительная модель "
                    f"текста не нашла (клауз нет) — при сомнении "
                    f"сверить с оригиналом-PDF."]
    out.append("")
    for cl in entry["clauses"]:
        out.append(f"[{cl['anchor']}] {cl['text']}")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def _projection_name(sid: str, file_name: str) -> str:
    return f"{sid}-{Path(file_name).stem}.md"


def freeze(paths: list[Path], out_dir: Path, vision_llm=None) -> dict:
    """Заморозка комплекта: клаузы + проекции + манифест в `out_dir`.
    Возвращает манифест. Порядок файлов задаёт идентификаторы S1, S2…
    `vision_llm` — зрительный канал для изображений (ПИ-1); запись
    распознанного источника несёт признак `"recognized": "image"`,
    SHA-256 в манифесте — от самого изображения."""
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "clauses").mkdir(exist_ok=True)
    (out_dir / "projections").mkdir(exist_ok=True)
    manifest = {"arbitration": ARBITRATION, "files": []}
    for i, path in enumerate(paths, start=1):
        sid = f"S{i}"
        got = extract_file(path, convert_dir=out_dir / "converted",
                           vision_llm=vision_llm)
        entry = {
            "sid": sid,
            "file": path.name,
            "sha256": sha256_of(path),
            "clauses": got["clauses"],
        }
        if got["converted"]:
            entry["converted"] = got["converted"]
        if got.get("recognized"):
            entry["recognized"] = got["recognized"]
        if got.get("recognized_pages"):
            entry["recognized_pages"] = got["recognized_pages"]
        if got.get("mixed_pages"):
            entry["mixed_pages"] = got["mixed_pages"]
            entry["mixed_recognized"] = bool(got.get("mixed_recognized"))
        if got.get("mixed_empty_figures"):
            entry["mixed_empty_figures"] = got["mixed_empty_figures"]
        (out_dir / "clauses" / f"{sid}.json").write_text(
            json.dumps(entry, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8")
        (out_dir / "projections" /
         _projection_name(sid, entry["file"])).write_text(
            render_projection(entry), encoding="utf-8")
        m_entry = {k: v for k, v in entry.items() if k != "clauses"}
        m_entry["clause_count"] = len(entry["clauses"])
        manifest["files"].append(m_entry)
    (out_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8")
    return manifest


# Поля записи источника, которые манифест и машинная форма несут
# одинаково (гейт заморозки сверяет их между собой).
SOURCE_MARK_KEYS = ("recognized", "recognized_pages", "mixed_pages",
                    "mixed_recognized", "mixed_empty_figures")


def mixed_warnings(manifest: dict) -> list[str]:
    """Предупреждения статуса прогона по манифесту заморозки (ревизия
    2026-09-04): смешанные страницы PDF без распознавания рисунков
    (канала зрения не было) и рисунки, на которых модель текста не
    нашла, — обязательная сверка с оригиналом, о которой панель
    обязана сказать, а не только манифест. Пустой список — сказать
    нечего. Детерминировано: порядок файлов манифеста."""
    unrecognized: list[str] = []
    empty: list[str] = []
    for f in manifest.get("files", []):
        sid = f.get("sid", "?")
        if f.get("mixed_pages") and not f.get("mixed_recognized"):
            unrecognized.append(
                f"{sid} стр. " + ", ".join(map(str, f["mixed_pages"])))
        if f.get("mixed_empty_figures"):
            empty.append(f"{sid} " + "; ".join(f["mixed_empty_figures"]))
    out: list[str] = []
    if unrecognized:
        out.append("смешанные страницы без распознавания: "
                   + "; ".join(unrecognized)
                   + " — сверка с оригиналом обязательна")
    if empty:
        out.append("рисунки, на которых зрительная модель текста не "
                   "нашла: " + "; ".join(empty)
                   + " — при сомнении сверить с оригиналом")
    return out


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")   # У-14: печать кириллицы
    ap = argparse.ArgumentParser(
        description="Заморозка исходных документов CDPM: клаузы + "
                    "md-проекции + манифест SHA-256")
    ap.add_argument("files", nargs="+", help="исходные документы "
                    "(doc/docx/md/xlsx/pdf/json/jpg/jpeg/png)")
    ap.add_argument("--out-dir", required=True, help="каталог комплекта")
    ap.add_argument("--cloud-images", action="store_true",
                    help="явно разрешить облачный канал зрения для "
                         "изображений (без локальной зрительной модели "
                         "иначе — честный отказ)")
    args = ap.parse_args(argv)
    out_dir = Path(args.out_dir)
    paths = [Path(f) for f in args.files]
    # Изображения — строжайшим доступным каналом (правило ПИ-1, как в
    # webui): локальная зрительная модель, иначе облако только по
    # явному флагу, иначе честный отказ.
    vision = None
    if any(p.suffix.lower() in IMAGE_SUPPORTED for p in paths):
        from cdpm import llm as llm_mod
        try:
            if llm_mod.local_vision_available():
                vision = llm_mod.vision_channel("any", requested="local")
            elif args.cloud_images:
                vision = llm_mod.vision_channel("any", requested="cloud")
            else:
                print("ОТКАЗ: распознавание изображений локальной "
                      "моделью недоступно: зрительная модель не "
                      "настроена; явное разрешение облака — флаг "
                      "--cloud-images")
                return 1
        except llm_mod.LLMCallError as e:
            print(f"ОТКАЗ: {e}")
            return 1
    try:
        manifest = freeze(paths, out_dir, vision_llm=vision)
    except FreezeError as e:
        print(f"ОТКАЗ: {e}")
        return 1
    except Exception as e:                                 # noqa: BLE001
        # Отказ канала зрения посреди заморозки — тоже честный отказ,
        # не трассировка.
        if vision is None:
            raise
        print(f"ОТКАЗ: {e}")
        return 1
    total = 0
    for f in manifest["files"]:
        total += f["clause_count"]
        conv = " (конвертирован из .doc)" if f.get("converted") else ""
        print(f"{f['sid']}: {f['file']} — клауз {f['clause_count']}, "
              f"SHA-256 {f['sha256'][:12]}…{conv}")
        print(f"  клаузы:   "
              f"{(out_dir / 'clauses' / (f['sid'] + '.json')).resolve()}")
        print(f"  проекция: "
              f"{(out_dir / 'projections' / _projection_name(f['sid'], f['file'])).resolve()}")
    print(f"итого: файлов {len(manifest['files'])}, клауз {total}")
    print(f"манифест: {(out_dir / 'manifest.json').resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
