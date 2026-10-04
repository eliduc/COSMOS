"""A04 legacy files: freeze before writers; v1 hash is not a scope proof."""
import copy

import pytest
from openpyxl import load_workbook

from cdpm import answers, anketa_xlsx as ax, family_params as fp
from cdpm import model as m, question_editions as qe
from test_cdpm_question_editions import _doc, _filled
from test_cdpm_family_params import _doc_with_home_floor, QID


@pytest.mark.parametrize("book_name", ["customer-answers.xlsx", "elsewhere/answers.xlsx"])
def test_legacy_identity_does_not_depend_on_book_name_or_location(tmp_path, book_name):
    doc, q = _doc()
    book = _filled(tmp_path, doc, q)
    target = tmp_path / book_name
    target.parent.mkdir(parents=True, exist_ok=True)
    book.rename(target)
    path = tmp_path / "model.json"
    m.dump_model(doc, path)
    old_bytes = path.read_bytes()
    old_fp = ax.question_fingerprint(q)
    loaded = m.load_model(path)
    migrated = next(x for x in loaded["questions"] if x["id"] == q["id"])
    assert migrated["edition"]["fingerprint_v1"] == old_fp
    assert migrated["edition"]["targets"] == sorted(q["rule_refs"])
    assert path.read_bytes() == old_bytes
    out = tmp_path / "new-model.json"
    m.publish_model(loaded, out)
    assert answers.read_book(target, loaded).answers == {q["id"]: "600 минут"}
    assert m.load_model(out) == loaded


@pytest.mark.parametrize("growth", ["source", "carrier"])
def test_loaded_legacy_family_is_frozen_before_automatic_scope_growth(tmp_path, growth):
    doc, _ = _doc_with_home_floor()
    fp.ensure_family_param_questions(doc)
    path = tmp_path / "model.json"
    m.dump_model(doc, path)
    loaded = m.load_model(path)
    q = next(x for x in loaded["questions"] if x["id"] == QID)
    before = copy.deepcopy(q)
    carrier = next(x for x in loaded["rules"] if x["id"] in q["rule_refs"])
    if growth == "source":
        carrier["sources"].append({"sid": "S1", "anchor": "П.9.1"})
    else:
        other = copy.deepcopy(carrier)
        other["id"] = "R-legacy-extra"
        other["blockers"] = []
        loaded["rules"].append(other)
    assert fp.ensure_family_param_questions(loaded) == [QID + "-E2"]
    assert q == before
    assert qe.problems(q) == []


def test_draft_serializer_is_still_read_only_and_migration_is_repeatable(tmp_path):
    doc, _ = _doc()
    before = copy.deepcopy(doc)
    path = tmp_path / "model.json"
    m.dump_model(doc, path)
    assert doc == before
    raw = path.read_bytes()
    first = m.load_model(path)
    assert all("edition" in q for q in first["questions"])
    assert first == m.load_model(path)
    assert path.read_bytes() == raw


def test_loading_does_not_repair_an_invalid_existing_snapshot(tmp_path):
    doc, q = _doc()
    qe.freeze(q)
    q["edition"]["fingerprint"] = "000000000000"
    path = tmp_path / "model.json"
    m.dump_model(doc, path)
    loaded = m.load_model(path)
    current = next(x for x in loaded["questions"] if x["id"] == q["id"])
    assert current == q
    assert qe.problems(current)


@pytest.mark.parametrize("bad", [None, {"id": "Q-bad"}])
def test_malformed_legacy_record_remains_available_to_schema_diagnostics(tmp_path, bad):
    doc, _ = _doc()
    doc["questions"].append(bad)
    path = tmp_path / "model.json"
    m.dump_model(doc, path)
    loaded = m.load_model(path)
    assert loaded["questions"][-1] == bad
    assert m.validate_schema(loaded)


@pytest.mark.parametrize("field", ["rule_refs", "source_refs", "note"])
def test_v1_compatibility_does_not_cover_scope_changed_before_migration(tmp_path, field):
    doc, q = _doc()
    old_fp = ax.question_fingerprint(q)
    book = _filled(tmp_path, doc, q)
    if field == "rule_refs":
        q[field].append(doc["rules"][0]["id"])
    elif field == "source_refs":
        q[field].append({"sid": "S1", "anchor": "П.9.9"})
    else:
        q["answer_format"]["note"] = "часы, не минуты"
    assert ax.question_fingerprint(q) == old_fp  # actual v1 blind spot
    path = tmp_path / "model.json"
    m.dump_model(doc, path)
    result = answers.read_book(book, m.load_model(path))
    assert q["id"] not in result.answers
    assert "контекст" in result.refused[q["id"]]["reason"]


@pytest.mark.parametrize("column", [1, 2, 3, 4])
def test_old_v1_hash_does_not_authorize_a_different_printed_context(tmp_path, column):
    doc, q = _doc()
    book = _filled(tmp_path, doc, q)
    path = tmp_path / "model.json"
    m.dump_model(doc, path)
    loaded = m.load_model(path)
    assert answers.read_book(book, loaded).answers == {q["id"]: "600 минут"}
    wb = load_workbook(book)
    ws = wb[ax.SHEET_TITLES[q["part"]]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == q["id"])
    row[column].value = str(row[column].value) + " ЧУЖОЙ КОНТЕКСТ"
    wb.save(book)
    wb.close()
    result = answers.read_book(book, loaded)
    assert q["id"] not in result.answers
    assert q["id"] in result.refused
    assert "контекст" in result.refused[q["id"]]["reason"]
