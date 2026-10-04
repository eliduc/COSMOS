"""A04/K04 composition on synthetic documents; no live model or service."""
import copy
import json

import pytest
from openpyxl import load_workbook

from cdpm import answers, family_params as fp, model as m, provenance, webui
from cdpm import anketa_xlsx as ax
from test_cdpm_edition_publication import _family_scope_doc, _complete_family_answer
from test_cdpm_family_params import QID

PARAM = "min_rest_outstation"
OTHER = "min_rest_home"


def _q(doc, qid):
    return next(q for q in doc["questions"] if q["id"] == qid)


def _apply_reading(doc, reading):
    return answers.apply_answers(doc, reading.answers, None,
                                 received_at=reading.received_at,
                                 book_sha256=reading.book_sha256)


def _book(doc, path, qid, raw):
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path)
    ws = wb[ax.SHEET_TITLES[_q(doc, qid)["part"]]]
    found = 0
    for row in ws.iter_rows(min_row=2):
        if row[0].value == qid:
            row[ax.ANSWER_COL - 1].value = raw
            found += 1
    assert found == 1
    wb.save(path)
    wb.close()
    return path


def _sync_reload(doc, run):
    (run / "extract").mkdir(parents=True, exist_ok=True)
    assert webui._sync_canon(run, doc, "edition-revision-integration") == []
    after = m.load_model(run / "extract/model.json")
    assert after == doc
    assert m.validate_schema(after) + m.validate_model(after) == []
    return after


def _pending_reply(doc, qid):
    q = _q(doc, qid)
    assert q["answer"]["status"] == m.PARTIAL_ANSWER_STATUS
    pending = answers._affected_rules(doc, q)
    assert len(pending) == 1
    patch = {k: copy.deepcopy(v) for k, v in pending[0].items()
             if k in answers._PATCH_FIELDS}
    patch.pop("blockers", None)
    patch["template_fit"]["params"][PARAM] = fp._answer_cell(
        q, PARAM, 540, "540 минут", "540 минут")
    return json.dumps({"patches": [{"question": qid, "change": "Остаток пересмотра",
                                   "rules": [patch]}]}, ensure_ascii=False)


def _finish_pending(doc, qid):
    reply = _pending_reply(doc, qid)
    report = answers.apply_worker_patches(doc, lambda prompt: reply)
    assert report["worker_rejected"] == [] and report["worker_applied"] == 1
    assert _q(doc, qid)["answer"]["status"] == "применён"


def _expanded_answer(tmp_path):
    doc, carrier = _family_scope_doc()
    carrier_id = carrier["id"]
    fp.ensure_family_param_questions(doc)
    m.publish_model(doc, tmp_path / "initial.json")
    old_book = _book(doc, tmp_path / "e1.xlsx", QID, "600 минут")
    reading = answers.read_book(old_book, doc)
    assert reading.problems == [] and reading.answers == {QID: "600 минут"}
    assert _apply_reading(doc, reading)["applied"] == 1
    assert _q(doc, QID)["answer"]["answer_version"]["book_sha256"] == reading.book_sha256
    old_question = copy.deepcopy(_q(doc, QID))
    carrier = next(r for r in doc["rules"] if r["id"] == carrier_id)
    other = copy.deepcopy(carrier)
    other["id"] = "R-integration-expanded"
    if QID in other.get("blockers", []):
        other["blockers"].remove(QID)
    other["template_fit"]["params"].pop(PARAM)
    doc["rules"].append(other)
    assert fp.ensure_family_param_questions(doc) == [QID + "-E2"]
    e2 = _q(doc, QID + "-E2")
    assert e2["edition"]["targets"] == sorted([carrier["id"], other["id"]])
    assert e2["edition"]["fingerprint"] != old_question["edition"]["fingerprint"]
    assert _q(doc, QID) == old_question
    before_rules = copy.deepcopy(doc["rules"])
    old_reading = answers.read_book(old_book, doc)
    assert e2["id"] not in old_reading.answers
    assert _apply_reading(doc, old_reading)["applied"] == 0
    assert doc["rules"] == before_rules and m.is_open_question(e2)
    assert "answer" not in e2
    # A renamed old row must not become an answer to the new scope.
    wb = load_workbook(old_book)
    ws = wb[ax.SHEET_TITLES[e2["part"]]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == QID:
            row[0].value = e2["id"]
    forged = tmp_path / "e1-renamed-e2.xlsx"
    wb.save(forged)
    wb.close()
    refused = answers.read_book(forged, doc)
    assert e2["id"] in refused.refused and e2["id"] not in refused.answers
    assert _q(doc, QID) == old_question
    # Fixed synthetic worker reply completes the second issued row.
    # Expected values below are literals, not the worker's parser output.
    _complete_family_answer(doc, e2["id"], "600 минут")
    ids = [carrier_id, other["id"]]
    for row in (r for r in doc["rules"] if r["id"] in ids):
        cell = row["template_fit"]["params"][PARAM]
        assert cell["value"] == 600
        assert cell["source"]["question"]["id"] == e2["id"]
    return doc, ids, old_question


@pytest.mark.parametrize("late_edit", ["neighbor", "same_cell"])
def test_edition_growth_revision_book_and_cas_compose(tmp_path, late_edit):
    doc, ids, old_question = _expanded_answer(tmp_path)
    rows = [r for r in doc["rules"] if r["id"] in ids]
    e2id = QID + "-E2"
    # A later neighboring edit makes whole-rule rollback unsafe.
    rows[0]["template_fit"]["params"][OTHER]["quote"] = "уточнённый отдых дома"
    opened = answers.open_revision(doc, e2id)
    rid = opened["created"]
    assert rid == e2id + "-R" and opened["restored"] is False
    assert len(opened["pending_cells"]) == 2
    baseline = copy.deepcopy(_q(doc, rid)["revision_cells"])
    doc = _sync_reload(doc, tmp_path / "run")
    assert _q(doc, QID) == old_question
    assert _q(doc, rid)["revision_cells"] == baseline
    assert m.is_open_question(_q(doc, rid))
    assert len(m.pending_revision_facts(doc)) == 2
    book = _book(doc, tmp_path / "revision.xlsx", rid, "540 минут")
    rows = [r for r in doc["rules"] if r["id"] in ids]
    target = OTHER if late_edit == "neighbor" else PARAM
    rows[1]["template_fit"]["params"][target]["value"] = 780 if target == OTHER else 660
    rows[1]["template_fit"]["params"][target]["quote"] = "позднее независимое изменение"
    before = copy.deepcopy(doc["rules"])
    reading = answers.read_book(book, doc)
    assert reading.problems == [] and reading.answers == {rid: "540 минут"}
    result = _apply_reading(doc, reading)
    if late_edit == "same_cell":
        assert result["applied"] == 0 and result["conflicted"] == 1
        assert doc["rules"] == before
        assert _q(doc, rid)["answer"]["status"] == m.REVISION_CONFLICT_STATUS
        assert _q(doc, rid)["revision_cells"] == baseline
        doc = _sync_reload(doc, tmp_path / "run")
        assert m.is_open_question(_q(doc, rid))
        assert _q(doc, rid)["revision_cells"] == baseline
        assert _apply_reading(doc, reading)["conflicted"] == 1
        # Only an explicit fresh reconsideration may accept the late baseline.
        rid2 = answers.open_revision(doc, rid)["created"]
        assert rid2 == e2id + "-R-2"
        retried = answers.apply_answers(doc, {rid2: "540 минут"}, None)
        assert retried["applied"] == 0 and retried["partial"] == 1
        _finish_pending(doc, rid2)
        active = rid2
    else:
        assert result["applied"] == 0 and result["partial"] == 1 and result.get("conflicted", 0) == 0
        _finish_pending(doc, rid)
        active = rid
    doc = _sync_reload(doc, tmp_path / "run")
    assert _q(doc, QID) == old_question
    assert m.pending_revision_facts(doc) == []
    rows = [r for r in doc["rules"] if r["id"] in ids]
    for row in rows:
        cell = row["template_fit"]["params"][PARAM]
        assert cell["value"] == 540 and "pending_revision" not in cell
        assert cell["source"]["question"]["id"] == active
        assert cell["source"]["version"]["fingerprint"] == _q(doc, active)["answer"]["answer_version"]["fingerprint"]
        source_q, source_raw = provenance.resolve_source(doc, cell["source"])
        assert source_q["id"] == active and source_raw == "540 минут"
        forged_source = copy.deepcopy(cell["source"])
        forged_source["version"]["fingerprint"] = old_question["answer"]["answer_version"]["fingerprint"]
        with pytest.raises(provenance.ProvenanceRefusal, match="версия ответа"):
            provenance.resolve_source(doc, forged_source)
    # Neither applying nor synchronizing may roll back a neighboring cell.
    for row in rows:
        prior = next(r for r in before if r["id"] == row["id"])
        assert row["template_fit"]["params"][OTHER] == prior["template_fit"]["params"][OTHER]
    snapshot = copy.deepcopy(doc)
    assert _apply_reading(doc, reading)["already_answered"] == 1
    # An import iteration advances model.meta.revision, not the answer's
    # edition, receipt, raw, source or the effective rule values.
    snapshot["meta"]["revision"] += 1
    assert doc == snapshot


@pytest.mark.parametrize("when", ["before_worker", "during_worker"])
def test_late_change_between_partial_revision_and_worker_is_not_overwritten(tmp_path, when):
    doc, ids, _ = _expanded_answer(tmp_path)
    first = next(r for r in doc["rules"] if r["id"] == ids[0])
    first["template_fit"]["params"][OTHER]["quote"] = "позднее уточнение дома"
    rid = answers.open_revision(doc, QID + "-E2")["created"]
    assert answers.apply_answers(doc, {rid: "540 минут"}, None)["partial"] == 1
    reply = _pending_reply(doc, rid)
    pending = answers._affected_rules(doc, _q(doc, rid))
    assert len(pending) == 1
    def change():
        pending[0]["template_fit"]["params"][PARAM] = {
            "value": 660, "quote": "позднее согласовано 660 минут"}
    if when == "before_worker":
        change()
    observed = {}
    def worker(prompt):
        if when == "during_worker":
            change()
        observed["rules"] = copy.deepcopy(doc["rules"])
        return reply
    baseline = copy.deepcopy(_q(doc, rid)["revision_cells"])
    receipt = copy.deepcopy(_q(doc, rid)["answer"]["answer_version"])
    result = answers.apply_worker_patches(doc, worker)
    assert result["worker_applied"] == 0 and result["worker_rejected"]
    assert doc["rules"] == observed["rules"]
    assert _q(doc, rid)["revision_cells"] == baseline
    assert _q(doc, rid)["answer"]["answer_version"] == receipt
    assert _q(doc, rid)["answer"]["status"] == m.REVISION_CONFLICT_STATUS
    assert m.is_open_question(_q(doc, rid))
    doc = _sync_reload(doc, tmp_path / "run")
    assert _q(doc, rid)["answer"]["status"] == m.REVISION_CONFLICT_STATUS
    unchanged = copy.deepcopy(doc)
    assert answers.apply_worker_patches(doc, worker)["worker_applied"] == 0
    assert doc == unchanged
    fresh = answers.open_revision(doc, rid)["created"]
    assert fresh != rid
    _complete_family_answer(doc, fresh, "540 минут")
    assert m.pending_revision_facts(doc) == []
    assert m.validate_schema(doc) + m.validate_model(doc) == []


@pytest.mark.parametrize("change", ["edited", "deleted", "conflicting"])
def test_neighbor_changed_while_worker_runs_is_preserved(tmp_path, change):
    doc, ids, _ = _expanded_answer(tmp_path)
    first = next(r for r in doc["rules"] if r["id"] == ids[0])
    first["template_fit"]["params"][OTHER]["quote"] = "позднее уточнение дома"
    rid = answers.open_revision(doc, QID + "-E2")["created"]
    assert answers.apply_answers(doc, {rid: "540 минут"}, None)["partial"] == 1
    reply = _pending_reply(doc, rid)
    if change == "conflicting":
        patch = json.loads(reply)
        patch["patches"][0]["rules"][0]["template_fit"]["params"][OTHER] = {
            "value": 840, "quote": "другая правка обработчика"}
        reply = json.dumps(patch, ensure_ascii=False)
    pending = answers._affected_rules(doc, _q(doc, rid))
    assert len(pending) == 1
    rowid = pending[0]["id"]
    neighbor = {"value": 780, "quote": "позднее дома 780 минут"}
    observed = {}
    def worker(prompt):
        if change == "deleted":
            pending[0]["template_fit"]["params"].pop(OTHER)
        else:
            pending[0]["template_fit"]["params"][OTHER] = copy.deepcopy(neighbor)
        observed["rules"] = copy.deepcopy(doc["rules"])
        return reply
    report = answers.apply_worker_patches(doc, worker)
    if change == "conflicting":
        assert report["worker_applied"] == 0 and report["worker_rejected"]
        assert doc["rules"] == observed["rules"]
        assert _q(doc, rid)["answer"]["status"] == m.REVISION_CONFLICT_STATUS
        return
    assert report["worker_applied"] == 1 and report["worker_rejected"] == []
    final = next(r for r in doc["rules"] if r["id"] == rowid)
    if change == "deleted":
        assert OTHER not in final["template_fit"]["params"]
    else:
        assert final["template_fit"]["params"][OTHER] == neighbor
    assert final["template_fit"]["params"][PARAM]["value"] == 540


@pytest.mark.parametrize("broken", ["missing_baseline", "broken_journal"])
def test_worker_cannot_invent_revision_preconditions(tmp_path, broken):
    doc, ids, _ = _expanded_answer(tmp_path)
    first = next(r for r in doc["rules"] if r["id"] == ids[0])
    first["template_fit"]["params"][OTHER]["quote"] = "позднее уточнение дома"
    rid = answers.open_revision(doc, QID + "-E2")["created"]
    assert answers.apply_answers(doc, {rid: "540 минут"}, None)["partial"] == 1
    reply = _pending_reply(doc, rid)
    q = _q(doc, rid)
    if broken == "missing_baseline":
        q.pop("revision_cells")
    else:
        effect = next(e for e in q["answer"]["effects"]
                      if e["rule"] == ids[1] and e["field"] == "template_fit")
        effect["before"]["params"][PARAM]["value"] = 999
    before = copy.deepcopy(doc["rules"])
    report = answers.apply_worker_patches(doc, lambda prompt: reply)
    assert report["worker_applied"] == 0 and report["worker_rejected"]
    assert doc["rules"] == before
    assert q["answer"]["status"] == m.REVISION_CONFLICT_STATUS
