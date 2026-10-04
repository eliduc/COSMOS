# -*- coding: utf-8 -*-
"""Браузер-интерфейс CDPM (П-УИ): загрузка → заморозка → извлечение →
артефакты, всё через HTTP-обвязку и с изолированным каталогом прогонов.

Языковая модель подменяется той же заготовленной функцией, что в тестах
извлечения; каталог прогонов уводится в tmp через CDPM_RUNS_DIR —
рабочий каталог данных тесты не трогают (та же дисциплина, что у
хранилищ наблюдений в conftest)."""
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from cdpm import webui
from test_cdpm_extract import _RULES_A, _RULES_B, fake_llm


@pytest.fixture(autouse=True)
def _s06r2_declared_layout(declared_test_layout):
    """S06-R2 (06.09.2026): облачные стадии панели требуют объявленной
    раскладки ролей — тесты этого файла объявляют явную тестовую
    (conftest.declared_test_layout), молчаливого облака нет."""


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("CDPM_RUNS_DIR", str(tmp_path / "runs"))
    # Настройки CDPM — тоже наблюдение: тесты не пишут в рабочий
    # каталог данных (та же дисциплина, что _OBSERVATION_STORES).
    monkeypatch.setenv("CDPM_SETTINGS_FILE",
                       str(tmp_path / "cdpm-settings.json"))
    monkeypatch.setattr(webui, "llm_factory", lambda scope: fake_llm)
    # Политика файлов по умолчанию «local» ведёт извлечение через
    # локальный канал (exchange_llm_factory) — тоже подменный.
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: fake_llm)
    return TestClient(webui.app)


def test_roles_and_effective_layout(client, monkeypatch):
    """Полная независимость выбора LLM (2.09): перечень ролей с
    родителями и фактическая раскладка прогона — своя / унаследована
    / умолчание."""
    roles = client.get("/api/agents/roles").json()["roles"]
    by = {r["role"]: r for r in roles}
    assert by["extractor"]["parent"] == "worker"
    assert by["auditor"]["parent"] is None
    assert by["vision"]["label"]
    run_id = _upload(client).json()["run"]
    # Носитель площадки снимается ПОСЛЕ рождения прогона: с шага 2
    # проекта «Запись о прогоне» гейт «Новый прогон» не создаёт прогон
    # при допуске, разрешающем облако, и необъявленной раскладке
    # (раскладка прогона ниже — своя, «run-explicit»).
    monkeypatch.setenv("CDPM_AGENTS_DEFAULT", "")
    resp = client.put(f"/api/runs/{run_id}/agents", json={
        "worker": {"channels": {"cloud": {"provider": "cloud",
                                          "model": "claude-sonnet-5"}}},
        "data_compiler": {"channels": {"cloud": {
            "provider": "cloud", "model": "claude-opus-5"}}},
    })
    assert resp.status_code == 200, resp.text
    eff = client.get(f"/api/runs/{run_id}/agents/effective").json()
    e = {r["role"]: r for r in eff["roles"]}
    assert e["extractor"]["effective"] == "worker"
    assert e["extractor"]["inherited"] is True
    assert e["data_compiler"]["effective"] == "data_compiler"
    assert e["data_compiler"]["inherited"] is False
    assert e["data_compiler"]["channels"]["cloud"]["model"] \
        == "claude-opus-5"
    assert e["auditor"]["effective"] is None


def _upload(client):
    # Признак допуска моделей — у комплекта (решение владельца 03.09):
    # тестовый комплект синтетический, значит допуск any по умолчанию;
    # без класса данных прогон шёл бы local-only, и просьбы scope=any
    # ниже честно отказывали бы.
    return client.post("/api/runs", files=[
        ("files", ("RULES-A.md", _RULES_A.encode("utf-8"),
                   "text/markdown")),
        ("files", ("RULES-B.md", _RULES_B.encode("utf-8"),
                   "text/markdown")),
    ], params={"data_class": "synthetic"})


def _allow_cloud(client, run_id):
    """Политика всех файлов прогона — «cloud» (стадии, смешивающие
    файлы, с облачными каналами иначе честно отказывают)."""
    files = client.get(f"/api/runs/{run_id}").json()["files"]
    cfg = {f: {"processing": "cloud", "note": ""} for f in files}
    resp = client.put(f"/api/runs/{run_id}/files-config",
                      json={"files": cfg})
    assert resp.status_code == 200, resp.text


def _extract_and_wait(client, run_id, timeout_s=30.0):
    resp = client.post(f"/api/runs/{run_id}/extract",
                       params={"customer": "Полярис (тест)",
                               "scope": "any"})
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "извлечение идёт":
            return st
        time.sleep(0.1)
    raise AssertionError("извлечение не завершилось за отведённое время")


def test_page_is_served(client):
    from cdpm import webui as webui_mod
    resp = client.get("/", follow_redirects=False)
    if (webui_mod.DIST_DIR / "index.html").is_file():
        # Собранный интерфейс закоммичен (вариант А): корень ведёт на
        # него, старый одностраничник живёт по /legacy до паритета.
        assert resp.status_code in (302, 307)
        assert resp.headers["location"] == "/ui/"
        ui = client.get("/ui/")
        assert ui.status_code == 200
        assert "CDPM" in ui.text
    else:
        assert resp.status_code == 200
        assert "Заморозить" in resp.text
    legacy = client.get("/legacy")
    assert legacy.status_code == 200
    assert "Заморозить" in legacy.text


def test_upload_mixed_pdf_without_vision_warns_in_status(client, tmp_path):
    """Ревизия 2026-09-04 (A2.3): смешанная страница PDF без канала
    зрения (локальной зрительной модели нет, облако не разрешено) —
    заморозка проходит, статус прогона несёт предупреждение о
    обязательной сверке с оригиналом, и оно переживает сброс стадий."""
    fitz = pytest.importorskip("fitz")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "3.2. Нормы экипажа по направлениям — "
                               "в таблице-картинке ниже.", fontsize=11)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 60))
    pix.clear_with(200)
    page.insert_image(fitz.Rect(72, 150, 400, 350),
                      stream=pix.tobytes("png"))
    pdf_bytes = doc.tobytes()
    doc.close()
    # Комплект без класса данных — допуск local-only: облачного зрения
    # нет по признаку (у синтетического комплекта оно включилось бы
    # само, 05.09), и заморозка обязана предупредить.
    resp = client.post("/api/runs", files=[
        ("files", ("SMESH.pdf", pdf_bytes, "application/pdf"))])
    assert resp.status_code == 200, resp.text
    got = resp.json()
    assert got["state"] == "заморожено"
    assert got["freeze_problems"] == []
    assert got["freeze_warnings"] == [
        "смешанные страницы без распознавания: S1 стр. 1 — сверка с "
        "оригиналом обязательна"]
    run_id = got["run"]
    st = client.get(f"/api/runs/{run_id}").json()
    assert st["freeze_warnings"] == got["freeze_warnings"]
    # Обычный текстовый комплект предупреждений не несёт.
    plain = _upload(client).json()
    assert "freeze_warnings" not in plain


def test_upload_freezes_and_reports_counters(client):
    resp = _upload(client)
    assert resp.status_code == 200, resp.text
    got = resp.json()
    assert got["run"] == "run-0001"
    assert got["state"] == "заморожено"
    assert got["freeze_counters"]["files"] == 2
    assert got["freeze_counters"]["originals"] == 2
    assert got["freeze_problems"] == []
    # Второй прогон получает следующий номер.
    assert _upload(client).json()["run"] == "run-0002"


def test_run_title_set_show_and_clear(client):
    """УИ-13: название прогона хранится, отдаётся статусом и списком,
    пустая строка убирает его; сверхдлинное — честный отказ."""
    _upload(client)
    r = client.put("/api/runs/run-0001/title",
                   json={"title": "Полярис — первый экзамен"})
    assert r.status_code == 200
    assert r.json()["title"] == "Полярис — первый экзамен"
    assert client.get("/api/runs/run-0001").json()["title"] \
        == "Полярис — первый экзамен"
    listed = {x["run"]: x for x in client.get("/api/runs").json()}
    assert listed["run-0001"]["title"] == "Полярис — первый экзамен"
    assert client.put("/api/runs/run-0001/title",
                      json={"title": "х" * 200}).status_code == 400
    r = client.put("/api/runs/run-0001/title", json={"title": "  "})
    assert r.json()["title"] is None
    assert client.get("/api/runs/run-0001").json().get("title") is None


def test_upload_of_unsupported_format_is_refused(client):
    # .pdf поддержан с 2026-08-26 — непокрытый пример теперь .rtf.
    resp = client.post("/api/runs", files=[
        ("files", ("scan.rtf", b"{\\rtf1", "application/rtf"))])
    assert resp.status_code == 400
    assert ".docx" in resp.json()["detail"]


def test_extract_flow_to_green_gate_and_artifacts(client):
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено", st
    assert st["extract_summary"]["rules"] == 6
    assert st["extract_summary"]["questions"] == 6
    assert st["extract_problems"] == []

    svod = client.get(f"/api/runs/{run_id}/file/свод")
    assert svod.status_code == 200
    assert "R-S1-п-1-9" in svod.text            # md отрендерен в html
    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    assert len(model["rules"]) == 6
    rej = client.get(f"/api/runs/{run_id}/file/отбраковка").json()
    assert rej["rejected"] == []
    manifest = client.get(f"/api/runs/{run_id}/file/манифест").json()
    assert len(manifest["files"]) == 2


def test_artifact_endpoint_is_a_closed_list(client):
    run_id = _upload(client).json()["run"]
    resp = client.get(f"/api/runs/{run_id}/file/..%2fmanifest")
    assert resp.status_code == 404
    resp = client.get(f"/api/runs/{run_id}/file/свод")
    assert resp.status_code == 404              # извлечения ещё не было
    assert "ещё не построен" in resp.json()["detail"]


def test_extract_validation_errors(client):
    run_id = _upload(client).json()["run"]
    assert client.post(f"/api/runs/{run_id}/extract",
                       params={"customer": " ", "scope": "any"}
                       ).status_code == 400
    assert client.post(f"/api/runs/{run_id}/extract",
                       params={"customer": "х", "scope": "куда-то"}
                       ).status_code == 400
    assert client.post("/api/runs/run-9999/extract",
                       params={"customer": "х", "scope": "any"}
                       ).status_code == 404
    assert client.post("/api/runs/run-;rm/extract",
                       params={"customer": "х", "scope": "any"}
                       ).status_code in (400, 404)


def test_extract_failure_is_reported_not_silent(client, monkeypatch):
    run_id = _upload(client).json()["run"]

    def broken(prompt):
        return "это не JSON"
    monkeypatch.setattr(webui, "llm_factory", lambda scope: broken)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: broken)
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "отказ извлечения"
    assert st["error"]


def test_exchange_flow_through_api(client, monkeypatch):
    """Аудит-цикл через HTTP: запуск с каналами по ролям, опрос живого
    состояния, артефакты обменов по закрытому перечню. Модели
    подменные: аудитор находит одно искажение и замолкает, обработчик
    принимает с патчем."""
    from cdpm import webui as webui_mod
    from test_cdpm_exchange_run import (_make_auditor, _make_worker,
                                        _REMARK_DISTORT, _PATCH_FIX)
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    _allow_cloud(client, run_id)

    auditor = _make_auditor([[dict(_REMARK_DISTORT)], [], []])
    worker = _make_worker({
        "Потеряно уточнение": {"verdict": "ПРИНЯТО",
                               "basis": "цитата: «на базе» — принято",
                               "change": "строка уточнена",
                               "patch": _PATCH_FIX}})
    monkeypatch.setattr(
        webui_mod, "exchange_llm_factory",
        lambda scope, requested: auditor if requested == "cloud"
        else worker)

    resp = client.post(f"/api/runs/{run_id}/exchange",
                       params={"auditor": "cloud", "worker": "local",
                               "max_rounds": 5})
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "аудит-цикл идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "аудит-цикл: успех", st

    ex = client.get(f"/api/runs/{run_id}/exchange").json()
    assert ex["verdict"]["verdict"] == "успех"
    assert [r["weighty"] for r in ex["rounds"]] == \
        [[0, 1, 0], [0, 0, 0], [0, 0, 0]]
    assert ex["registry"].get("закрыто") == 1
    assert ex["gate_problems"] == []

    got = client.get(
        f"/api/runs/{run_id}/exchange/rounds/1/замечания").json()
    assert len(got["remarks"]) == 1
    svod = client.get(f"/api/runs/{run_id}/exchange/svod/3")
    assert svod.status_code == 200
    assert "на базе" in svod.text
    assert client.get(
        f"/api/runs/{run_id}/exchange/rounds/1/чужое"
    ).status_code == 404


def test_iteration_archive_and_canon_sync(client, monkeypatch):
    """ИР-3: финальная модель цикла становится канонической
    (extract/model.json + рендеры); повторный запуск цикла архивирует
    завершённый в exchange-rev-NNN и переносит реестр — требование
    арбитра действует сквозь ревизии."""
    import json as _json
    from cdpm import webui as webui_mod
    from test_cdpm_exchange_run import (_make_auditor, _make_worker,
                                        _REMARK_DISTORT, _PATCH_FIX)
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    _allow_cloud(client, run_id)
    auditor = _make_auditor([[dict(_REMARK_DISTORT)], [], []])
    worker = _make_worker({
        "Потеряно уточнение": {"verdict": "ПРИНЯТО",
                               "basis": "цитата: «на базе» — принято",
                               "change": "строка уточнена",
                               "patch": _PATCH_FIX}})
    monkeypatch.setattr(
        webui_mod, "exchange_llm_factory",
        lambda scope, requested: auditor if requested == "cloud"
        else worker)
    assert client.post(f"/api/runs/{run_id}/exchange",
                       params={"auditor": "cloud", "worker": "local",
                               "max_rounds": 5}).status_code == 200
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "аудит-цикл идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "аудит-цикл: успех", st

    # Канон синхронизирован: модель extract равна финальной версии
    # цикла, свод отрендерен из неё (несёт принятый патч).
    ex = client.get(f"/api/runs/{run_id}/exchange").json()
    run_dir = webui_mod._run_dir(run_id)
    final = _json.loads(
        (run_dir / "exchange" / "versions"
         / ex["verdict"]["final_version"]).read_text(encoding="utf-8"))
    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    assert model == final
    svod = client.get(f"/api/runs/{run_id}/file/свод")
    assert "на базе" in svod.text

    # Итерация: второй цикл (чистый аудит) архивирует первый.
    auditor2 = _make_auditor([[], []])
    monkeypatch.setattr(
        webui_mod, "exchange_llm_factory",
        lambda scope, requested: auditor2 if requested == "cloud"
        else worker)
    assert client.post(f"/api/runs/{run_id}/exchange",
                       params={"auditor": "cloud", "worker": "local",
                               "max_rounds": 5}).status_code == 200
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "аудит-цикл идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "аудит-цикл: успех", st
    assert (run_dir / "exchange-rev-001" / "verdict.json").is_file()
    assert (run_dir / "exchange-rev-001" / "registry.json").is_file()
    ex2 = client.get(f"/api/runs/{run_id}/exchange").json()
    assert ex2["config"]["archived"] == "exchange-rev-001"
    assert len(ex2["rounds"]) == 2      # свежий цикл, не replay старого


def test_exchange_requires_extraction_first(client):
    run_id = _upload(client).json()["run"]
    resp = client.post(f"/api/runs/{run_id}/exchange")
    assert resp.status_code == 409
    assert "сначала извлечение" in resp.json()["detail"]


def _deadlock_via_api(client, monkeypatch) -> str:
    """Прогон до тупика через HTTP: извлечение, допуск облака, цикл, в
    котором аудитор настаивает, а обработчик отклоняет. Возвращает id
    прогона с непустым списком тупиков."""
    from cdpm import webui as webui_mod
    from test_cdpm_exchange_run import (_make_auditor, _make_worker,
                                        _REMARK_DISTORT)
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    _allow_cloud(client, run_id)
    # До аудит-цикла арбитражу нечего разбирать.
    assert client.post(f"/api/runs/{run_id}/arbiter").status_code == 409

    same = [dict(_REMARK_DISTORT)]
    auditor = _make_auditor([same, same, same, same])
    worker = _make_worker({})
    monkeypatch.setattr(
        webui_mod, "exchange_llm_factory",
        lambda scope, requested: auditor if requested == "cloud"
        else worker)
    resp = client.post(f"/api/runs/{run_id}/exchange",
                       params={"auditor": "cloud", "worker": "local",
                               "max_rounds": 6})
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "аудит-цикл идёт":
            break
        time.sleep(0.1)
    ex = client.get(f"/api/runs/{run_id}/exchange").json()
    assert ex["verdict"]["deadlocks"], ex["verdict"]
    return run_id


def _arbitrate_via_api(client, monkeypatch, run_id, worker,
                       arbiter=None) -> dict:
    """Арбитраж через HTTP с заданным обработчиком-исполнителем и
    арбитром (по умолчанию — «аудитор прав»); возвращает статус
    прогона по итогу."""
    from cdpm import webui as webui_mod
    from test_cdpm_exchange_run import _make_arbiter
    arb = arbiter if arbiter is not None else _make_arbiter("аудитор прав")
    monkeypatch.setattr(
        webui_mod, "_channel_for",
        lambda scope, run_dir, role, fb=None, klass="cloud":
        arb if role == "arbiter" else worker)
    resp = client.post(f"/api/runs/{run_id}/arbiter")
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "арбитраж идёт":
            break
        time.sleep(0.1)
    return st


def test_arbiter_flow_through_api(client, monkeypatch):
    """Арбитраж через HTTP: цикл заходит в тупик (аудитор настаивает,
    обработчик отклоняет), третья модель решает «аудитор прав»,
    обработчик обязан исполнить — гейт арбитража зелёный, сводка без
    неисполненных вердиктов."""
    from test_cdpm_exchange_run import _enforcing_worker
    run_id = _deadlock_via_api(client, monkeypatch)
    st = _arbitrate_via_api(client, monkeypatch, run_id,
                            _enforcing_worker)
    assert st["state"] == "арбитраж готов", st

    got = client.get(f"/api/runs/{run_id}/arbiter").json()
    assert got["summary"]["deadlocks"] > 0
    assert got["summary"]["by_verdict"]["аудитор прав"] == \
        got["summary"]["deadlocks"]
    assert got["summary"]["enforced"] == got["summary"]["deadlocks"]
    assert got["summary"]["unenforced"] == 0
    assert got["summary"]["unenforced_deadlocks"] == []
    assert got["gate_problems"] == []
    assert got["report"]["final_version"] == "model-arbiter.json"
    assert all(v["basis"] for v in got["report"]["verdicts"])


def test_arbiter_api_reports_unenforced_deadlocks(client, monkeypatch):
    """Отбракованное исполнение доходит до оператора (аудит A3.5):
    сводка арбитража в статусе прогона несёт unenforced (счётчик) и
    unenforced_deadlocks (тупик и причина), состояние — «красный гейт
    арбитража», гейт называет тупик."""
    from test_cdpm_arbiter_audit import _breaking_worker
    run_id = _deadlock_via_api(client, monkeypatch)
    st = _arbitrate_via_api(client, monkeypatch, run_id,
                            _breaking_worker)
    assert st["state"] == "красный гейт арбитража", st

    got = client.get(f"/api/runs/{run_id}/arbiter").json()
    s = got["summary"]
    assert s["deadlocks"] > 0 and s["enforced"] == 0
    assert s["unenforced"] == s["deadlocks"]
    assert len(s["unenforced_deadlocks"]) == s["deadlocks"]
    assert all(u["deadlock"] and "не исполнен" in u["reason"]
               for u in s["unenforced_deadlocks"])
    assert any("не исполнен" in p for p in got["gate_problems"]), got


def test_arbiter_api_reports_unresolved_deadlocks(client, monkeypatch):
    """A09: спор без однозначного решения арбитра доходит до оператора
    — сводка несёт unresolved (счётчик) и unresolved_deadlocks («тупик
    → причина»), решений по нему нет вовсе. Гейт зелёный: форма отчёта
    цела, — но спор этим НЕ закрыт: находка реестра живая и держит
    цикл в «не закрыто», и СОСТОЯНИЕ прогона это говорит словами
    «арбитраж готов: не разрешено N» (ревью A09, находка 3): зелёный
    гейт при живых спорах — не «всё решено».
    """
    import json as _json
    import re as _re2
    from test_cdpm_exchange_run import _enforcing_worker

    def _contradictory_arbiter(prompt):
        """Одному тупику — два противоречивых решения (и в повторе)."""
        ids = _re2.findall(r"Тупик \[([^\]]+)\]", prompt)
        rows = []
        for i in ids:
            rows.append({"deadlock": i, "verdict": "обработчик прав",
                         "basis": "клауза за обработчика"})
            rows.append({"deadlock": i, "verdict": "аудитор прав",
                         "basis": "клауза за аудитора"})
        return _json.dumps({"verdicts": rows}, ensure_ascii=False)

    run_id = _deadlock_via_api(client, monkeypatch)
    st = _arbitrate_via_api(client, monkeypatch, run_id,
                            _enforcing_worker,
                            arbiter=_contradictory_arbiter)

    got = client.get(f"/api/runs/{run_id}/arbiter").json()
    s = got["summary"]
    assert s["unresolved"] > 0
    # Состояние называет число неразрешённых споров — и начинается
    # прежними словами: триггеры и водители сверяют состояние началом
    # строки (state.startswith), автоподача заявок каталога после
    # арбитража продолжает срабатывать.
    assert st["state"] == f"арбитраж готов: не разрешено {s['unresolved']}"
    from cdpm import webui as webui_mod
    assert any(st["state"].startswith(p)
               for p in webui_mod._CATALOG_AFTER_ARBITER)
    assert len(s["unresolved_deadlocks"]) == s["unresolved"]
    assert all(u["deadlock"] and "противореч" in u["reason"]
               for u in s["unresolved_deadlocks"])
    assert sum(s["by_verdict"].values()) == 0
    assert got["report"]["verdicts"] == []
    assert got["gate_problems"] == []             # форма отчёта цела

    # Ключ спора не решён, поэтому живы и сам тупик, и его эхо —
    # повторы того же замечания прежних раундов (закон live_findings,
    # A07): гасит их только решение арбитра «арбитр: …».
    ex_got = client.get(f"/api/runs/{run_id}/exchange").json()
    rv = ex_got["registry_verdict"]
    assert {u["deadlock"] for u in s["unresolved_deadlocks"]} \
        <= {b["id"] for b in rv["blockers"]}
    assert any(b["state"] == "не разрешено арбитром"
               for b in rv["blockers"])
    assert rv["count"] >= s["unresolved"] > 0


def test_files_config_defaults_and_validation(client):
    """Политика файлов: умолчание — по допуску комплекта (синтетический
    комплект → «cloud», решение владельца 03.09; без признака — строжайшее
    «local»), сохранение политики и примечания, честные отказы
    валидации. Живой урок 05.09: прогон синтетического набора, созданный
    руками в панели, получал «local» и извлечение уходило в незапущенную
    локальную модель."""
    run_id = _upload(client).json()["run"]
    got = client.get(f"/api/runs/{run_id}/files-config").json()
    assert got["policies"] == ["local", "cloud-country", "cloud"]
    assert got["files"] == {
        "RULES-A.md": {"processing": "cloud", "note": ""},
        "RULES-B.md": {"processing": "cloud", "note": ""}}
    # Без признака комплекта — строжайшее умолчание.
    strict = client.post("/api/runs", files=[
        ("files", ("RULES-A.md", _RULES_A.encode("utf-8"), "text/markdown"))]
        ).json()["run"]
    assert client.get(f"/api/runs/{strict}/files-config").json()["files"] == {
        "RULES-A.md": {"processing": "local", "note": ""}}

    resp = client.put(f"/api/runs/{run_id}/files-config", json={
        "files": {"RULES-A.md": {"processing": "local",
                                 "note": "и1 — правила"}}})
    assert resp.status_code == 200, resp.text
    got = client.get(f"/api/runs/{run_id}/files-config").json()["files"]
    assert got["RULES-A.md"] == {"processing": "local",
                                 "note": "и1 — правила"}
    assert got["RULES-B.md"]["processing"] == "cloud"   # не тронут

    assert client.put(f"/api/runs/{run_id}/files-config", json={
        "files": {"чужой.md": {"processing": "local"}}}
        ).status_code == 400
    assert client.put(f"/api/runs/{run_id}/files-config", json={
        "files": {"RULES-A.md": {"processing": "куда-то"}}}
        ).status_code == 400


def test_extract_routes_by_file_policy(client, monkeypatch):
    """Маршрутизация извлечения: клаузы файла «cloud» идут облачным
    каналом, файла «local» — локальным; перекрёстная стадия
    (противоречия) — по строжайшей политике прогона (локально)."""
    run_id = _upload(client).json()["run"]
    # Синтетический комплект по умолчанию даёт «cloud» всем файлам —
    # маршрутизацию проверяем на явно разной политике двух файлов.
    client.put(f"/api/runs/{run_id}/files-config", json={
        "files": {"RULES-A.md": {"processing": "cloud", "note": ""},
                  "RULES-B.md": {"processing": "local", "note": ""}}})

    seen = {"cloud": [], "local": []}

    def recorder(channel):
        def llm(prompt):
            seen[channel].append(prompt)
            return fake_llm(prompt)
        return llm
    monkeypatch.setattr(webui, "llm_factory",
                        lambda scope: recorder("cloud"))
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: recorder("local"))

    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено", st
    assert all("RULES-A.md" in p for p in seen["cloud"]
               if "Документ" in p)
    assert not any("RULES-B.md" in p for p in seen["cloud"])
    assert any("RULES-B.md" in p for p in seen["local"])
    # Перекрёстная стадия противоречий — только локальным каналом.
    assert not any("противореч" in p.lower() for p in seen["cloud"])
    assert any("противореч" in p.lower() for p in seen["local"])


def test_exchange_uses_strictest_class_of_file_policies(client,
                                                        monkeypatch):
    """Цикл смешивает содержимое всех файлов: при политике «local»
    хотя бы одного файла каналы ОБЕИХ ролей строятся локальным классом,
    даже если запрошено облако — данные не уходят с машины."""
    from test_cdpm_exchange_run import _make_auditor
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"

    clean_auditor = _make_auditor([[], []])
    asked: list[str] = []

    def recording_factory(scope, requested):
        asked.append(requested)
        return clean_auditor
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        recording_factory)
    # Синтетический комплект по умолчанию даёт всем файлам «cloud»;
    # строжайший класс проверяется явной политикой «local» у одного файла.
    assert client.put(f"/api/runs/{run_id}/files-config", json={
        "files": {"RULES-B.md": {"processing": "local"}}}).status_code == 200
    resp = client.post(f"/api/runs/{run_id}/exchange",
                       params={"auditor": "cloud", "worker": "cloud"})
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "аудит-цикл идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "аудит-цикл: успех", st
    assert st["exchange"]["class"] == "local"
    assert set(asked) == {"local"}      # облако не строилось вовсе


def _xlsx_bytes(sheets: dict) -> bytes:
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_data_line_flow_through_api(client, monkeypatch):
    """Линия данных 3.3 через HTTP: загрузка и заморозка таблиц,
    политика файлов данных, компиляция проверок подменной моделью,
    отчёт с зелёным гейтом."""
    import json as _json
    import re as _re
    run_id = _upload(client).json()["run"]

    roster = _xlsx_bytes({"Ростер": [
        ["Табельный", "Дата"],
        ["EMP-1", "05.03.2026"],
        ["EMP-2", "07.03.2026"]]})
    quals = _xlsx_bytes({"Допуски": [
        ["Табельный", "Действителен до"],
        ["EMP-1", "31.12.2026"],
        ["EMP-2", "06.03.2026"]]})
    xl = "application/vnd.openxmlformats-officedocument" \
         ".spreadsheetml.sheet"
    resp = client.post(f"/api/runs/{run_id}/data", files=[
        ("files", ("ROSTER.xlsx", roster, xl)),
        ("files", ("QUAL.xlsx", quals, xl))])
    assert resp.status_code == 200, resp.text
    got = resp.json()
    assert got["state"] == "данные заморожены"
    assert got["counters"] == {"files": 2, "sheets": 2, "rows": 4}
    assert got["problems"] == []
    # Чужой формат — честный отказ.
    assert client.post(f"/api/runs/{run_id}/data", files=[
        ("files", ("scan.pdf", b"%PDF-", "application/pdf"))]
        ).status_code == 400
    # Файлы данных видны политике файлов; умолчание — по допуску
    # комплекта (синтетический → «cloud», 05.09), не «local» вслепую.
    cfg = client.get(f"/api/runs/{run_id}/files-config").json()["files"]
    assert cfg["ROSTER.xlsx"]["processing"] == "cloud"

    # Проверки без свода невозможны — сначала извлечение правил.
    assert client.post(f"/api/runs/{run_id}/data-run"
                       ).status_code == 409
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"

    def compiler(prompt):
        ids = _re.findall(r"^(R-\S+) \|", prompt, _re.M)
        assert "D1/Допуски" in prompt      # схема набора в промпте
        return _json.dumps({"checks": [], "skipped": [
            {"rule": rid, "reason": "нет данных для проверки"}
            for rid in ids]}, ensure_ascii=False)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: compiler)

    resp = client.post(f"/api/runs/{run_id}/data-run")
    assert resp.status_code == 200, resp.text
    # Строжайший класс политик файлов: у синтетического комплекта все
    # файлы по умолчанию «cloud» (05.09) — им и идёт линия данных.
    assert resp.json()["class"] == "cloud"
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}/data").json()
        if got["state"] != "проверки идут":
            break
        time.sleep(0.1)
    assert got["state"] == "отчёт готов", got
    assert got["gate_problems"] == []
    assert got["summary"]["rows"] == 4
    assert got["summary"]["checks_run"] >= 2      # встроенные (дубль)
    assert got["summary"]["checks_skipped"] >= 1  # строки свода честно

    md = client.get(f"/api/runs/{run_id}/file/отчёт-данных")
    assert md.status_code == 200
    assert "Отчёт о качестве операционных данных" in md.text
    model = client.get(
        f"/api/runs/{run_id}/file/отчёт-данных-модель").json()
    assert model["counters"]["rows_total"] == 4


def test_data_review_flow_through_api(client, monkeypatch):
    """Сверка спорного (ДЛ-7) через HTTP: отчёт с rule-находкой,
    сверка подменной моделью, вердикты в отчёте, зелёный гейт,
    артефакт сводки."""
    import json as _json
    import re as _re
    run_id = _upload(client).json()["run"]
    roster = _xlsx_bytes({"Ростер": [
        ["Табельный", "Дата"],
        ["EMP-1", "05.03.2026"],
        ["EMP-2", ""]]})           # пустая дата → находка проверки
    xl = "application/vnd.openxmlformats-officedocument" \
         ".spreadsheetml.sheet"
    assert client.post(f"/api/runs/{run_id}/data", files=[
        ("files", ("ROSTER.xlsx", roster, xl))]).status_code == 200
    # Сверка до отчёта — честный отказ.
    assert client.post(f"/api/runs/{run_id}/data-review"
                       ).status_code == 409
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    # Все правила подменного извлечения блокированы вопросами — для
    # rule-находки нужна неблокированная строка: дописываем её в
    # модель (данные-конвейер читает model.json как есть).
    model_path = webui._run_dir(run_id) / "extract" / "model.json"
    model = _json.loads(model_path.read_text(encoding="utf-8"))
    model["rules"].append({
        "id": "R-TEST-DATA", "kind": "HARD", "zone": "ROSTERING",
        "statement": "Дата строки ростера обязательна", "sources": []})
    model_path.write_text(_json.dumps(model, ensure_ascii=False),
                          encoding="utf-8")

    def compiler(prompt):
        ids = _re.findall(r"^(R-\S+) \|", prompt, _re.M)
        assert "R-TEST-DATA" in ids
        checks = [{"rule": "R-TEST-DATA",
                   "template": "обязательное-поле",
                   "params": {"sheet": "D1/Ростер", "column": "Дата"},
                   "statement": "Дата строки ростера заполнена",
                   "unit": "строка", "severity": "КРИТИЧЕСКОЕ"}]
        return _json.dumps({"checks": checks, "skipped": [
            {"rule": rid, "reason": "нет данных"}
            for rid in ids if rid != "R-TEST-DATA"]},
            ensure_ascii=False)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: compiler)
    assert client.post(f"/api/runs/{run_id}/data-run"
                       ).status_code == 200
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}/data").json()
        if got["state"] != "проверки идут":
            break
        time.sleep(0.1)
    assert got["state"] == "отчёт готов", got
    assert got["summary"]["findings"] >= 1

    def reviewer(prompt):
        fids = _re.findall(r"^(F-\d+) \|", prompt, _re.M)
        assert fids, "в промпте сверщика нет находок"
        return _json.dumps({"reviews": [
            {"finding": fid, "verdict": "законное исключение",
             "basis": "дата уточняется отдельным реестром"}
            for fid in fids]}, ensure_ascii=False)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: reviewer)
    resp = client.post(f"/api/runs/{run_id}/data-review")
    assert resp.status_code == 200, resp.text
    # Строжайший класс политик файлов: у синтетического комплекта все
    # файлы по умолчанию «cloud» (05.09) — им и идёт линия данных.
    assert resp.json()["class"] == "cloud"
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}/data").json()
        if got["state"] != "сверка идёт":
            break
        time.sleep(0.1)
    assert got["state"] == "спорное сверено", got.get("review_problems")
    assert got["review_problems"] == []
    assert got["review_summary"]["reviewed"] >= 1
    assert got["review_summary"]["by_verdict"] == {
        "законное исключение": got["review_summary"]["reviewed"]}

    summary = client.get(
        f"/api/runs/{run_id}/file/сводка-сверки").json()
    assert summary["reviewed"] == got["review_summary"]["reviewed"]
    md = client.get(f"/api/runs/{run_id}/file/отчёт-данных")
    assert "Вердикт сверки" in md.text
    assert "законное исключение" in md.text
    # Повторная сверка без несверённых находок — «кандидатов 0».
    resp = client.post(f"/api/runs/{run_id}/data-review")
    assert resp.status_code == 200
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}/data").json()
        if got["state"] != "сверка идёт":
            break
        time.sleep(0.1)
    assert got["state"] == "спорное сверено", got
    assert got["review_summary"]["candidates"] == 0


def test_answers_flow_through_api(client, tmp_path):
    """Интерактивный цикл (ИР-1) через HTTP: заполненная книга анкеты
    → применение → ревизия свода, сводка, зелёный гейт в статусе."""
    from openpyxl import load_workbook
    from cdpm import anketa_xlsx as ax
    run_id = _upload(client).json()["run"]
    # Ответы до извлечения — честный отказ.
    resp = client.post(
        f"/api/runs/{run_id}/answers",
        files={"file": ("anketa.xlsx", b"PK", "application/vnd.ms")})
    assert resp.status_code == 409
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    rule = next(r for r in model["rules"] if r.get("conditional"))
    qid = rule["conditional"]["question"]
    key = sorted(rule["conditional"]["branches"])[0]

    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(model, book)
    wb = load_workbook(str(book))
    for ws in wb.worksheets:
        if ws.title in (ax.SUMMARY_TITLE, ax.ANSWERED_TITLE):
            continue
        for row in ws.iter_rows(min_row=2):
            if row[0].value == qid:
                row[5].value = key
    wb.save(str(book))

    resp = client.post(
        f"/api/runs/{run_id}/answers",
        files={"file": ("anketa.xlsx", book.read_bytes(),
                        "application/vnd.ms")})
    assert resp.status_code == 200, resp.text
    got = resp.json()
    assert got["revision"] == 2 and got["received"] == 1
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "применение ответов идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "ответы применены", st
    assert st["answers_problems"] == []
    assert st["answers_summary"]["applied"] == 1
    assert st["answers_summary"]["revision"] == 2

    new_model = client.get(f"/api/runs/{run_id}/file/модель").json()
    assert new_model["meta"]["revision"] == 2
    import hashlib
    from datetime import datetime
    received = next(q for q in new_model["questions"] if q["id"] == qid)["answer"]
    assert received["answer_version"]["book_sha256"] == hashlib.sha256(book.read_bytes()).hexdigest()
    assert datetime.fromisoformat(received["answer_version"]["received_at"]).utcoffset() is not None
    new_rule = next(r for r in new_model["rules"]
                    if r["id"] == rule["id"])
    assert "conditional" not in new_rule
    summary = client.get(
        f"/api/runs/{run_id}/file/сводка-ответов").json()
    assert summary["book_sha256"]
    anketa = client.get(f"/api/runs/{run_id}/file/анкета")
    assert "Принятые ответы" in anketa.text
    # Пустая книга (без единого ответа) — честный отказ.
    empty = tmp_path / "empty.xlsx"
    ax.render_anketa_xlsx(new_model, empty)
    resp = client.post(
        f"/api/runs/{run_id}/answers",
        files={"file": ("anketa.xlsx", empty.read_bytes(),
                        "application/vnd.ms")})
    assert resp.status_code == 400
    assert "нет ни одного ответа" in resp.text


def _slow_llm(reached, release):
    def llm(prompt):
        reached.set()
        release.wait(10)
        return fake_llm(prompt)
    return llm


def test_stage_control_pause_and_resume(client, monkeypatch):
    """ПС-1: пауза срабатывает на контрольной точке между чанками
    (честное состояние «на паузе: …»), продолжение доводит стадию до
    конца."""
    import threading as th
    run_id = _upload(client).json()["run"]
    reached, release = th.Event(), th.Event()
    monkeypatch.setattr(webui, "llm_factory",
                        lambda scope: _slow_llm(reached, release))
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested:
                        _slow_llm(reached, release))
    resp = client.post(f"/api/runs/{run_id}/extract",
                       params={"customer": "т", "scope": "any"})
    assert resp.status_code == 200, resp.text
    assert reached.wait(10)
    assert client.post(f"/api/runs/{run_id}/control",
                       params={"command": "pause"}).status_code == 200
    release.set()
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] == "на паузе: извлечение":
            break
        time.sleep(0.05)
    assert st["state"] == "на паузе: извлечение", st
    assert client.post(f"/api/runs/{run_id}/control",
                       params={"command": "resume"}).status_code == 200
    # Продолжение доводит ТУ ЖЕ нить до конца; повторное «Извлечь» на
    # паузе — 409 (пауза = занятость линии, вторая нить не рождается;
    # прежде тест молча запускал вторую нить поверх первой).
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] not in ("извлечение идёт", "на паузе: извлечение"):
            break
        time.sleep(0.1)
    assert st["state"] == "извлечено", st


def test_stage_control_stop_is_not_a_failure(client, monkeypatch):
    """ПС-1: остановка — «остановлено: …» без ошибки (это не отказ);
    повторный запуск стадии чистит команду и доходит до конца."""
    import threading as th
    run_id = _upload(client).json()["run"]
    reached, release = th.Event(), th.Event()
    monkeypatch.setattr(webui, "llm_factory",
                        lambda scope: _slow_llm(reached, release))
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested:
                        _slow_llm(reached, release))
    assert client.post(f"/api/runs/{run_id}/extract",
                       params={"customer": "т", "scope": "any"}
                       ).status_code == 200
    assert reached.wait(10)
    assert client.post(f"/api/runs/{run_id}/control",
                       params={"command": "stop"}).status_code == 200
    release.set()
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] == "остановлено: извлечение":
            break
        time.sleep(0.05)
    assert st["state"] == "остановлено: извлечение", st
    assert not st.get("error")
    # Недопустимая команда — честный отказ.
    assert client.post(f"/api/runs/{run_id}/control",
                       params={"command": "взорви"}).status_code == 400
    # Повторный запуск стартует заново и доходит до конца.
    monkeypatch.setattr(webui, "llm_factory", lambda scope: fake_llm)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: fake_llm)
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено", st


def test_orphaned_stage_is_stopped_on_server_restart(client,
                                                     monkeypatch):
    """Гибель процесса посреди стадии (случилось 2026-08-19 на живой
    сверке run-0002): стадии живут в фоновых потоках, после смерти
    сервера статус остаётся «… идёт» и повторный запуск отвечает 409.
    Старт сервера (lifespan) переводит осиротевшие состояния обеих
    линий в «остановлено: <стадия>», error не трогает, и стадия
    принимается заново."""
    import json as _json
    import re as _re
    run_id = _upload(client).json()["run"]
    roster = _xlsx_bytes({"Ростер": [
        ["Табельный", "Дата"],
        ["EMP-1", "05.03.2026"]]})
    xl = "application/vnd.openxmlformats-officedocument" \
         ".spreadsheetml.sheet"
    assert client.post(f"/api/runs/{run_id}/data", files=[
        ("files", ("ROSTER.xlsx", roster, xl))]).status_code == 200
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"

    def compiler(prompt):
        ids = _re.findall(r"^(R-\S+) \|", prompt, _re.M)
        return _json.dumps({"checks": [], "skipped": [
            {"rule": rid, "reason": "нет данных"} for rid in ids]},
            ensure_ascii=False)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: compiler)
    assert client.post(f"/api/runs/{run_id}/data-run"
                       ).status_code == 200
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}/data").json()
        if got["state"] != "проверки идут":
            break
        time.sleep(0.1)
    assert got["state"] == "отчёт готов", got

    # Смерть сервера посреди стадий: статус «… идёт» осиротел (потока
    # больше нет); error от давнего отказа восстановление не трогает.
    run_dir = webui._run_dir(run_id)
    crashed = webui._read_status(run_dir)
    crashed.update(state="арбитраж идёт", data_state="сверка идёт",
                   error="давний отказ (восстановлению не трогать)")
    webui._write_status(run_dir, crashed)

    # Без восстановления обе линии отвечали бы 409 «прогон занят».
    assert client.post(f"/api/runs/{run_id}/data-review"
                       ).status_code == 409

    # Новый сервер: контекстный менеджер TestClient проигрывает
    # lifespan-событие старта — как настоящий перезапуск.
    with TestClient(webui.app) as restarted:
        st = restarted.get(f"/api/runs/{run_id}").json()
        assert st["state"] == "остановлено: арбитраж"
        assert st["data_state"] == "остановлено: сверка спорного"
        assert st["error"] == "давний отказ (восстановлению не трогать)"
        # Стадия принимается заново и доходит до конца (кандидатов
        # спорного в чистом отчёте нет — сверка честно пустая).
        resp = restarted.post(f"/api/runs/{run_id}/data-review")
        assert resp.status_code == 200, resp.text
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            got = restarted.get(f"/api/runs/{run_id}/data").json()
            if got["state"] != "сверка идёт":
                break
            time.sleep(0.1)
        assert got["state"] == "спорное сверено", got


def test_runs_listing(client):
    _upload(client)
    got = client.get("/api/runs").json()
    assert [r["run"] for r in got] == ["run-0001"]
    assert got[0]["files"] == ["RULES-A.md", "RULES-B.md"]
    assert got[0]["state"] == "заморожено"


def test_mode2_flow_through_api(client, monkeypatch, tmp_path):
    """Режим 2 через HTTP: загрузка чужой формализации, сверка
    подменной моделью, вопросы части 4 в анкете, аудиторский отчёт,
    зелёный гейт."""
    import json as _json
    run_id = _upload(client).json()["run"]
    # Сверка до загрузки файла — честный отказ.
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    assert client.post(f"/api/runs/{run_id}/mode2").status_code == 409

    third = ("1.1. Отдых на базе перед сменой — не менее 14 часов "
             "непрерывно.\n\n2.2. Полёты в дни затмений запрещены "
             "внутренним стандартом перевозчика.\n")
    resp = client.post(
        f"/api/runs/{run_id}/third-party",
        files={"file": ("THIRD.md", third.encode("utf-8"),
                        "text/markdown")})
    assert resp.status_code == 200, resp.text
    assert resp.json()["claims"] == 2
    # Чужой формат — честный отказ.
    assert client.post(
        f"/api/runs/{run_id}/third-party",
        files={"file": ("scan.pdf", b"%PDF-", "application/pdf")}
    ).status_code == 400

    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    rid = model["rules"][0]["id"]

    def checker(prompt):
        assert "УТВЕРЖДЕНИЯ ТРЕТЬЕЙ СТОРОНЫ" in prompt
        return _json.dumps({"checks": [
            {"anchor": "п.1.1", "verdict": "расходится", "rule": rid,
             "note": "у нас 10 часов, у третьей стороны 14"},
            {"anchor": "п.2.2", "verdict": "в источниках отсутствует",
             "rule": rid, "note": "затмений в источниках нет"},
        ]}, ensure_ascii=False)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: checker)
    resp = client.post(f"/api/runs/{run_id}/mode2")
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "сверка формализации идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "формализация сверена", st
    assert st["mode2_problems"] == []
    assert st["mode2_summary"]["part4_added"] == 2
    assert st["mode2_summary"]["by_verdict"] == {
        "расходится": 1, "в источниках отсутствует": 1}

    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    assert model["meta"]["mode"] == 2
    q4 = [q for q in model["questions"] if q["part"] == "4"]
    assert len(q4) == 2
    anketa = client.get(f"/api/runs/{run_id}/file/анкета")
    assert "Часть 4" in anketa.text
    md = client.get(
        f"/api/runs/{run_id}/file/отчёт-сверки-формализации")
    assert "Аудит формализации третьей стороны" in md.text


def test_full_cycle_through_api(client, monkeypatch):
    """СКВ-1: сквозной процесс одной кнопкой — цикл с тупиком →
    арбитраж («обработчик прав») → новая ревизия → чистый цикл →
    останов «чисто по Б/К». Все три роли подменные."""
    import json
    import re as _re
    from cdpm import webui as webui_mod
    from test_cdpm_exchange_run import (_make_auditor, _make_worker,
                                        _REMARK_DISTORT)
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    _allow_cloud(client, run_id)

    same = [dict(_REMARK_DISTORT)]
    auditor = _make_auditor([same] * 12)
    worker = _make_worker({})          # всегда отклоняет — тупик

    def arbiter(prompt):
        ids = _re.findall(r"Тупик \[([^\]]+)\]", prompt)
        return json.dumps({"verdicts": [
            {"deadlock": i, "verdict": "обработчик прав",
             "basis": "по клаузе источника"} for i in ids]},
            ensure_ascii=False)

    def route(prompt):
        if "Тупик [" in prompt:
            return arbiter(prompt)
        if "Клаузы документа" in prompt:
            return auditor(prompt)
        return worker(prompt)

    monkeypatch.setattr(webui_mod, "exchange_llm_factory",
                        lambda scope, requested: route)

    resp = client.post(f"/api/runs/{run_id}/full-cycle",
                       params={"auditor": "cloud", "worker": "cloud",
                               "max_rounds": 4, "max_megacycles": 2})
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "сквозной цикл идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "сквозной цикл: чисто по Б/К (кругов: 2)", st
    # Арбитраж случился и закрыл тупик в пользу обработчика.
    assert st["arbiter_summary"]["by_verdict"].get(
        "обработчик прав") == 1
    # Первая ревизия ушла в архив, реестр перенесён сквозь ревизии.
    runs_root = webui_mod.runs_root()
    assert (runs_root / run_id / "exchange-rev-001").is_dir()
    ex = client.get(f"/api/runs/{run_id}/exchange").json()
    assert ex["verdict"]["verdict"] == "успех"


# ---------------------------------------------------------------------------
# С-1 (анализ Лайта-v3): чистоту по Б/К доказывает последний ПОЛНЫЙ
# раунд, чистый инкрементальный хвост её не опровергает
# ---------------------------------------------------------------------------

def test_c1_bk_clean_tail_logic():
    def rnd(r, b, k, mode):
        return {"round": r, "weighty": [b, k, 0], "E": 0, "mode": mode}

    # Живой случай круга 3 Лайта: полный чистый р4, инкрементальный
    # чистый р5 — раньше давало «лимит кругов», теперь чисто.
    hist = [rnd(3, 0, 1, "полный"), rnd(4, 0, 0, "полный"),
            rnd(5, 0, 0, "инкрементальный")]
    assert webui._bk_clean_tail(hist) == 1

    # Последний полный грязен по К — не чисто.
    assert webui._bk_clean_tail(
        [rnd(1, 0, 0, "полный"), rnd(2, 0, 1, "полный"),
         rnd(3, 0, 0, "инкрементальный")]) is None

    # Инкрементальный ПОСЛЕ чистого полного принёс К — не чисто.
    assert webui._bk_clean_tail(
        [rnd(1, 0, 0, "полный"),
         rnd(2, 0, 1, "инкрементальный")]) is None

    # Ни одного полного раунда — чистота недоказуема.
    assert webui._bk_clean_tail(
        [rnd(1, 0, 0, "инкрементальный")]) is None
    assert webui._bk_clean_tail([]) is None

    # Раунд без mode считается полным (совместимость со старыми
    # verdict.json).
    assert webui._bk_clean_tail(
        [{"round": 1, "weighty": [0, 0, 2], "E": 0}]) == 0


def test_b2_bk_clean_tail_rejects_patched_tail(tmp_path):
    """В2-5: раунд, принявший правки, менял свод ПОСЛЕ своего аудита —
    хвост с патчами не финализируется как «чисто по Б/К»."""
    def rnd(r, b, k, mode, patched=False):
        return {"round": r, "weighty": [b, k, 0], "E": 0,
                "mode": mode, "patched": patched}

    # Полный раунд с принятым патчем С-замечания — финала нет.
    assert webui._bk_clean_tail(
        [rnd(1, 0, 0, "полный", patched=True)]) is None
    # Чистый полный + инкрементальный, патчивший свод, — финала нет.
    assert webui._bk_clean_tail(
        [rnd(1, 0, 0, "полный"),
         rnd(2, 0, 0, "инкрементальный", patched=True)]) is None
    # Чистый полный + чистый непатчивший хвост — финал есть.
    assert webui._bk_clean_tail(
        [rnd(1, 0, 0, "полный"),
         rnd(2, 0, 0, "инкрементальный")]) == 0


def test_b2_live_findings_read_registry(tmp_path):
    """В2-5: бэклог и страховка финала читают реестр — живыми
    считаются «открыто»/«отклонено»/«тупик», закрытые и решённые
    арбитром не всплывают."""
    import json as _json
    ex_dir = tmp_path / "run" / "exchange"
    ex_dir.mkdir(parents=True)
    findings = [
        {"id": "F1", "key": "k1", "state": "открыто",
         "severity": "СЕРЬЁЗНОЕ"},
        {"id": "F2", "key": "k2", "state": "закрыто",
         "severity": "СЕРЬЁЗНОЕ"},
        {"id": "F3", "key": "k3", "state": "тупик",
         "severity": "БЛОКИРУЮЩЕЕ"},
        {"id": "F4", "key": "k4", "state": "арбитр: обработчик прав",
         "severity": "КРИТИЧЕСКОЕ"},
        {"id": "F5", "key": "k5", "state": "отклонено",
         "severity": "СЕРЬЁЗНОЕ"},
        # Эхо решённого спора: живое состояние, но ключ k4 закрыт
        # арбитром — не всплывает (пока спор зрел, каждый повтор
        # открывал отдельную находку).
        {"id": "F6", "key": "k4", "state": "открыто",
         "severity": "КРИТИЧЕСКОЕ"},
    ]
    (ex_dir / "registry.json").write_text(
        _json.dumps({"findings": findings}, ensure_ascii=False),
        encoding="utf-8")
    live = webui._live_findings(tmp_path / "run")
    assert sorted(f["id"] for f in live) == ["F1", "F3", "F5"]
    # Пустой каталог без реестра — пустой список, не падение.
    assert webui._live_findings(tmp_path / "нет") == []


def test_exchange_anketa_version_endpoint(client, monkeypatch):
    """Анкета версии обмена отдаётся рядом со сводом; для прогонов,
    сыгранных до этой точки, честно рендерится из модели версии."""
    from cdpm import webui as webui_mod
    from test_cdpm_exchange_run import (_make_auditor, _make_worker,
                                        _REMARK_DISTORT, _PATCH_FIX)
    import json as _json

    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    _allow_cloud(client, run_id)
    auditor = _make_auditor([[dict(_REMARK_DISTORT)], [], []])
    worker = _make_worker({
        "Потеряно уточнение": {"verdict": "ПРИНЯТО",
                               "basis": "цитата: «на базе» — принято",
                               "change": "строка уточнена",
                               "patch": _PATCH_FIX}})

    def route(prompt):
        if "Клаузы документа" in prompt:
            return auditor(prompt)
        return worker(prompt)
    monkeypatch.setattr(webui_mod, "exchange_llm_factory",
                        lambda scope, requested: route)
    client.post(f"/api/runs/{run_id}/exchange",
                params={"auditor": "cloud", "worker": "cloud",
                        "max_rounds": 4})
    import time as _t
    deadline = _t.monotonic() + 60
    while _t.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}").json()
        if got["state"] != "аудит-цикл идёт":
            break
        _t.sleep(0.1)
    r = client.get(f"/api/runs/{run_id}/exchange/anketa/1")
    assert r.status_code == 200
    assert "Анкета" in r.text
    # Прогон прежнего кода: md-файла нет — рендер из модели версии.
    import pathlib
    vdir = pathlib.Path(webui_mod.runs_root()) / run_id / "exchange" \
        / "versions"
    (vdir / "anketa-v01.md").unlink()
    r = client.get(f"/api/runs/{run_id}/exchange/anketa/1")
    assert r.status_code == 200
    assert "Анкета" in r.text
    assert client.get(
        f"/api/runs/{run_id}/exchange/anketa/99").status_code == 404


def test_export_endpoint_serves_solver_formats(client):
    """ЭКС-1: скачивание Rave/CSV/JSON из канонической модели;
    Meridian — 501 до документации; до извлечения — 409."""
    run_id = _upload(client).json()["run"]
    assert client.get(f"/api/runs/{run_id}/export/rave"
                      ).status_code == 409
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    r = client.get(f"/api/runs/{run_id}/export/csv")
    # Образец блокирован вопросами → готовность «а» → честный отказ.
    if st.get("readiness", {}).get("state") == "а":
        assert r.status_code == 409
        assert "недостаточно" in r.json()["detail"]
    else:
        assert r.status_code == 200
    # МЕР-1: meridian реализован — он выгружает КНИГУ ДАННЫХ; без
    # неё честный 409 (прежде был 501 «ждёт документацию»).
    assert client.get(f"/api/runs/{run_id}/export/meridian"
                      ).status_code == 409
    assert client.get(f"/api/runs/{run_id}/export/чужой"
                      ).status_code == 404


def test_reset_returns_run_to_frozen_keeping_files_and_agents(client):
    """Сброс (решение владельца 2026-08-21): состояние — «заморожено»,
    файлы/заморозка/модели блоков/название сохранены, результаты
    стадий удалены; повторное извлечение работает."""
    import pathlib
    from cdpm import webui as webui_mod
    run_id = _upload(client).json()["run"]
    client.put(f"/api/runs/{run_id}/title", json={"title": "Проба"})
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    # Конфигурация ролей задаётся ПОСЛЕ извлечения: с настроенной
    # ролью канал шёл бы живым провайдером, а не подменной моделью.
    cfg = {"worker": {"channels": {
        "cloud": {"provider": "openai", "model": "m-1"}}}}
    client.put(f"/api/runs/{run_id}/agents", json=cfg)

    r = client.post(f"/api/runs/{run_id}/reset")
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["state"] == "заморожено"
    assert got["title"] == "Проба"
    assert got["files"] == ["RULES-A.md", "RULES-B.md"]
    assert "extract_summary" not in got
    run_dir = pathlib.Path(webui_mod.runs_root()) / run_id
    assert not (run_dir / "extract").exists()
    assert not (run_dir / "exchange").exists()
    assert (run_dir / "freeze" / "manifest.json").is_file()
    assert client.get(f"/api/runs/{run_id}/agents").json() == cfg
    # Повторное извлечение после сброса — штатно (роли предварительно
    # очищаются: настроенная роль повела бы канал живым провайдером,
    # а не подменной моделью теста).
    client.put(f"/api/runs/{run_id}/agents", json={})
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"


def test_meridian_export_needs_data_book(client):
    """МЕР-1: формат meridian выгружает КНИГУ ДАННЫХ (не правила);
    без загруженной книги — честный 409 с объяснением природы
    формата."""
    run_id = _upload(client).json()["run"]
    r = client.get(f"/api/runs/{run_id}/export/meridian",
                   params={"preset_id": "x"})
    assert r.status_code == 409
    assert "данн" in r.json()["detail"]


def test_clean_data_export_requires_report(client):
    """ЭКС-3: без отчёта проверок «очищенных» данных не существует."""
    run_id = _upload(client).json()["run"]
    r = client.get(f"/api/runs/{run_id}/export-clean-data")
    assert r.status_code == 409


def test_reset_leaves_no_exchange_ghosts(client, monkeypatch):
    """Живой сброс владельца 2026-08-21: ignore_errors молча оставлял
    каталоги раундов — панель показывала «Обменов 4» с пустыми
    метриками. После сброса живое состояние цикла обязано быть пустым."""
    from cdpm import webui as webui_mod
    from test_cdpm_exchange_run import (_make_auditor, _make_worker,
                                        _REMARK_DISTORT, _PATCH_FIX)
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    _allow_cloud(client, run_id)
    auditor = _make_auditor([[dict(_REMARK_DISTORT)], [], []])
    worker = _make_worker({
        "Потеряно уточнение": {"verdict": "ПРИНЯТО",
                               "basis": "цитата: «на базе» — принято",
                               "change": "строка уточнена",
                               "patch": _PATCH_FIX}})

    def route(prompt):
        if "Клаузы документа" in prompt:
            return auditor(prompt)
        return worker(prompt)
    monkeypatch.setattr(webui_mod, "exchange_llm_factory",
                        lambda scope, requested: route)
    client.post(f"/api/runs/{run_id}/exchange",
                params={"auditor": "cloud", "worker": "cloud",
                        "max_rounds": 4})
    import time as _t
    deadline = _t.monotonic() + 60
    while _t.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}").json()
        if got["state"] != "аудит-цикл идёт":
            break
        _t.sleep(0.1)
    ex = client.get(f"/api/runs/{run_id}/exchange").json()
    assert len(ex["rounds"]) >= 1

    assert client.post(f"/api/runs/{run_id}/reset").status_code == 200
    ex = client.get(f"/api/runs/{run_id}/exchange").json()
    assert ex["rounds"] == []
    assert ex["verdict"] is None
    assert ex["registry"] in ({}, None)


# ── Удаление файла операционных данных (живой случай 2026-08-23) ──

_XL_MIME = ("application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet")


def _upload_two_books(client, run_id):
    """Две книги данных; по алфавиту имён QUAL → D1, ROSTER → D2."""
    roster = _xlsx_bytes({"Ростер": [
        ["Табельный", "Дата"],
        ["EMP-1", "05.03.2026"],
        ["EMP-2", "07.03.2026"]]})
    quals = _xlsx_bytes({"Допуски": [
        ["Табельный", "Действителен до"],
        ["EMP-1", "31.12.2026"]]})
    resp = client.post(f"/api/runs/{run_id}/data", files=[
        ("files", ("ROSTER.xlsx", roster, _XL_MIME)),
        ("files", ("QUAL.xlsx", quals, _XL_MIME))])
    assert resp.status_code == 200, resp.text
    assert resp.json()["counters"] == {"files": 2, "sheets": 2,
                                       "rows": 3}


def _run_checks_and_wait(client, run_id, monkeypatch, timeout_s=30.0):
    import json as _json
    import re as _re

    def compiler(prompt):
        ids = _re.findall(r"^(R-\S+) \|", prompt, _re.M)
        return _json.dumps({"checks": [], "skipped": [
            {"rule": rid, "reason": "нет данных для проверки"}
            for rid in ids]}, ensure_ascii=False)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: compiler)
    resp = client.post(f"/api/runs/{run_id}/data-run")
    assert resp.status_code == 200, resp.text
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}/data").json()
        if got["state"] != "проверки идут":
            return got
        time.sleep(0.1)
    raise AssertionError("проверки данных не завершились вовремя")


def test_delete_data_file_cascades_through_all_four_places(
        client, monkeypatch):
    """Живой случай 2026-08-23 (эталоны run-0036/run-0038): замена
    книги данных требовала ручной правки ЧЕТЫРЁХ мест, а файл без
    записи в конфигурации получал строжайшее умолчание «local» и
    уводил перегон в локальный канал. DELETE обязан пройти каскадом:
    файл в data-src/, запись files-config.json, data_files в
    status.json, SID в dataset.json + manifest.json (пере-заморозка);
    отчёт проверок и «очищенные данные» снимаются — данные им больше
    не соответствуют; строжайший класс политики пересчитывается."""
    import json as _json
    import pathlib
    run_id = _upload(client).json()["run"]
    _upload_two_books(client, run_id)
    run_dir = pathlib.Path(webui.runs_root()) / run_id
    # Политика: удаляемый файл — «local», остающийся — «cloud».
    # Пока «local» в прогоне, строжайший класс — «local» (ровно так
    # перегон эталона падал в локальный канал).
    files = client.get(f"/api/runs/{run_id}").json()["files"]
    cfg = {f: {"processing": "cloud", "note": ""} for f in files}
    cfg["ROSTER.xlsx"] = {"processing": "cloud",
                          "note": "остаётся"}
    cfg["QUAL.xlsx"] = {"processing": "local", "note": "уйдёт"}
    assert client.put(f"/api/runs/{run_id}/files-config",
                      json={"files": cfg}).status_code == 200
    assert webui._strictest_class(run_dir) == "local"
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    got = _run_checks_and_wait(client, run_id, monkeypatch)
    assert got["state"] == "отчёт готов", got
    assert got["summary"]["rows"] == 3
    assert (run_dir / "data-report" / "data-report.json").is_file()
    assert client.get(f"/api/runs/{run_id}/export-clean-data"
                      ).status_code == 200
    ds = _json.loads((run_dir / "data-freeze" / "dataset.json")
                     .read_text(encoding="utf-8"))
    assert [s["file"] for s in ds["sources"]] == ["QUAL.xlsx",
                                                  "ROSTER.xlsx"]

    r = client.delete(f"/api/runs/{run_id}/data/QUAL.xlsx")
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["deleted"] == "QUAL.xlsx"
    assert got["removed"] == {"file": True, "files_config": True}
    assert got["state"] == "данные заморожены"
    assert got["files"] == ["ROSTER.xlsx"]
    assert got["counters"] == {"files": 1, "sheets": 1, "rows": 2}
    assert got["problems"] == []

    # 1. файл в data-src/
    assert not (run_dir / "data-src" / "QUAL.xlsx").exists()
    assert (run_dir / "data-src" / "ROSTER.xlsx").is_file()
    # 2. запись в files-config.json — на диске и через API; политика
    #    оставшегося файла цела, строжайший класс стал «cloud».
    raw = _json.loads((run_dir / "files-config.json")
                      .read_text(encoding="utf-8"))["files"]
    assert "QUAL.xlsx" not in raw
    assert raw["ROSTER.xlsx"] == {"processing": "cloud",
                                  "note": "остаётся"}
    api_cfg = client.get(f"/api/runs/{run_id}/files-config"
                         ).json()["files"]
    assert "QUAL.xlsx" not in api_cfg
    assert api_cfg["ROSTER.xlsx"]["processing"] == "cloud"
    assert webui._strictest_class(run_dir) == "cloud"
    # 3. status.json: data_files, счётчики заморозки, снятые выходы.
    status = _json.loads((run_dir / "status.json")
                         .read_text(encoding="utf-8"))
    assert status["data_files"] == ["ROSTER.xlsx"]
    assert status["data_freeze_counters"]["files"] == 1
    assert status["data_state"] == "данные заморожены"
    for key in ("data_summary", "data_problems", "data_counters",
                "data_class", "data_review_summary"):
        assert key not in status, key
    # 4. dataset.json + manifest.json: один источник, SID пересчитан.
    ds = _json.loads((run_dir / "data-freeze" / "dataset.json")
                     .read_text(encoding="utf-8"))
    assert [(s["sid"], s["file"]) for s in ds["sources"]] \
        == [("D1", "ROSTER.xlsx")]
    assert sorted(ds["tables"]) == ["D1"]
    mf = _json.loads((run_dir / "data-freeze" / "manifest.json")
                     .read_text(encoding="utf-8"))
    assert [(f["sid"], f["file"]) for f in mf["files"]] \
        == [("D1", "ROSTER.xlsx")]
    assert mf["sheets"] == {"D1": ["Ростер"]}
    # Оригинал сверен с манифестом по SHA-256 (гейт заморозки).
    assert status["data_freeze_problems"] == []
    assert status["data_freeze_counters"]["originals"] == 1
    # Инвалидация: отчёта нет, очищенные данные честно отказывают,
    # панель данных без сводки.
    assert not (run_dir / "data-report").exists()
    assert client.get(f"/api/runs/{run_id}/export-clean-data"
                      ).status_code == 409
    dl = client.get(f"/api/runs/{run_id}/data").json()
    assert dl["files"] == ["ROSTER.xlsx"]
    assert dl["summary"] is None and dl["class"] is None
    # Повторное удаление того же файла — 404, чужое имя — 404.
    assert client.delete(f"/api/runs/{run_id}/data/QUAL.xlsx"
                         ).status_code == 404
    assert client.delete(f"/api/runs/{run_id}/data/NOPE.xlsx"
                         ).status_code == 404
    # Имя с путём — отказ (наружу из data-src/ не выйти): маршрут
    # не принимает «/» в имени (404), а обработчик — имя с каталогом
    # (400); status.json цел.
    from fastapi import HTTPException
    assert client.delete(f"/api/runs/{run_id}/data/..%2Fstatus.json"
                         ).status_code in (400, 404)
    with pytest.raises(HTTPException) as ei:
        webui.delete_data_file(run_id, "../status.json")
    assert ei.value.status_code == 400
    with pytest.raises(HTTPException) as ei:
        webui.delete_data_file(run_id, "..")
    assert ei.value.status_code == 400
    assert (run_dir / "status.json").is_file()
    # Проверки по оставшейся книге идут штатно.
    got = _run_checks_and_wait(client, run_id, monkeypatch)
    assert got["state"] == "отчёт готов", got
    assert got["summary"]["rows"] == 2


def test_delete_last_data_file_leaves_no_freeze(client):
    """Последний файл данных: заморозки не остаётся вовсе (набор с
    нулём источников проходил бы гейт и пускал проверки «ни по
    чему»), состояние линии данных — как до первой загрузки."""
    import json as _json
    import pathlib
    run_id = _upload(client).json()["run"]
    _upload_two_books(client, run_id)
    run_dir = pathlib.Path(webui.runs_root()) / run_id
    assert client.delete(f"/api/runs/{run_id}/data/ROSTER.xlsx"
                         ).status_code == 200
    r = client.delete(f"/api/runs/{run_id}/data/QUAL.xlsx")
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["files"] == [] and got["state"] is None
    assert got["counters"] == {"files": 0, "sheets": 0, "rows": 0}
    assert not (run_dir / "data-freeze").exists()
    assert list((run_dir / "data-src").iterdir()) == []
    status = _json.loads((run_dir / "status.json")
                         .read_text(encoding="utf-8"))
    assert status["data_files"] == []
    assert "data_state" not in status
    assert "data_freeze_counters" not in status
    # Проверки данных без набора — честный отказ, как до загрузки.
    _extract_and_wait(client, run_id)
    assert client.post(f"/api/runs/{run_id}/data-run"
                       ).status_code == 409
    # Новая загрузка после полного удаления — штатно.
    _upload_two_books(client, run_id)
    assert client.get(f"/api/runs/{run_id}/data").json()["files"] \
        == ["QUAL.xlsx", "ROSTER.xlsx"]


def test_delete_data_file_refused_while_checks_run(client, monkeypatch):
    """Пока идут проверки данных — удаление файла отвечает 409 и не
    трогает ни одно из четырёх мест; после завершения стадии
    удаление проходит."""
    import json as _json
    import pathlib
    import threading as th
    run_id = _upload(client).json()["run"]
    _upload_two_books(client, run_id)
    run_dir = pathlib.Path(webui.runs_root()) / run_id
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    reached, release = th.Event(), th.Event()
    # Подменный свод целиком блокирован вопросами анкеты — модель
    # компилятор не зовёт; медленной делается сама компиляция.
    real_compile = webui.data_checks_mod.compile_checks

    def slow_compile(*args, **kwargs):
        reached.set()
        release.wait(10)
        return real_compile(*args, **kwargs)
    monkeypatch.setattr(webui.data_checks_mod, "compile_checks",
                        slow_compile)
    assert client.post(f"/api/runs/{run_id}/data-run"
                       ).status_code == 200
    assert reached.wait(10)
    assert client.get(f"/api/runs/{run_id}/data").json()["state"] \
        == "проверки идут"
    r = client.delete(f"/api/runs/{run_id}/data/QUAL.xlsx")
    assert r.status_code == 409, r.text
    assert "проверки идут" in r.json()["detail"]
    # Загрузка новых данных в бегущую стадию — тот же отказ.
    assert client.post(f"/api/runs/{run_id}/data", files=[
        ("files", ("NEW.xlsx", _xlsx_bytes({"Л": [["a"], ["1"]]}),
                   _XL_MIME))]).status_code == 409
    assert not (run_dir / "data-src" / "NEW.xlsx").exists()
    assert (run_dir / "data-src" / "QUAL.xlsx").is_file()
    ds = _json.loads((run_dir / "data-freeze" / "dataset.json")
                     .read_text(encoding="utf-8"))
    assert [s["file"] for s in ds["sources"]] == ["QUAL.xlsx",
                                                  "ROSTER.xlsx"]
    assert client.get(f"/api/runs/{run_id}/data").json()["files"] \
        == ["QUAL.xlsx", "ROSTER.xlsx"]
    release.set()
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        got = client.get(f"/api/runs/{run_id}/data").json()
        if got["state"] != "проверки идут":
            break
        time.sleep(0.1)
    assert got["state"] == "отчёт готов", got
    # Стадия завершена — удаление проходит, отчёт снимается.
    assert client.delete(f"/api/runs/{run_id}/data/QUAL.xlsx"
                         ).status_code == 200
    assert client.get(f"/api/runs/{run_id}/data").json()["summary"] \
        is None


def test_upload_data_after_report_drops_stale_report(client,
                                                     monkeypatch):
    """Догрузка книги данных после отчёта проверок: отчёт строился по
    прежнему составу и больше ему не соответствует — снимается тем же
    каскадом, что при удалении."""
    import pathlib
    run_id = _upload(client).json()["run"]
    _upload_two_books(client, run_id)
    run_dir = pathlib.Path(webui.runs_root()) / run_id
    _extract_and_wait(client, run_id)
    got = _run_checks_and_wait(client, run_id, monkeypatch)
    assert got["state"] == "отчёт готов", got
    r = client.post(f"/api/runs/{run_id}/data", files=[
        ("files", ("NEW.xlsx", _xlsx_bytes({"Л": [["a"], ["1"]]}),
                   _XL_MIME))])
    assert r.status_code == 200, r.text
    assert r.json()["counters"] == {"files": 3, "sheets": 3, "rows": 4}
    assert not (run_dir / "data-report").exists()
    dl = client.get(f"/api/runs/{run_id}/data").json()
    assert dl["state"] == "данные заморожены"
    assert dl["summary"] is None
    assert client.get(f"/api/runs/{run_id}/export-clean-data"
                      ).status_code == 409


def test_проекция_карты_из_книги_данных(tmp_path, monkeypatch):
    """Пакет трёх 30.08, п.1: детерминированная проекция data-freeze
    → формат NetMap (станции/рёбра/сотрудники/geo); без книги —
    честный отказ 409."""
    import json as _j

    from fastapi.testclient import TestClient

    from cdpm import webui as w
    monkeypatch.setattr(w, "_run_dir", lambda rid: tmp_path)
    c = TestClient(w.app)

    r = c.get("/api/runs/run-x/map-projection")
    assert r.status_code == 409
    assert "книга данных не загружена" in r.json()["detail"]

    (tmp_path / "data-freeze").mkdir()
    tables = {"D1": {
        "Рейсы": {"columns": ["Номер рейса", "Вылет", "Прилёт"],
                  "rows": [
                      {"Номер рейса": "PL1", "Вылет": "QA1",
                       "Прилёт": "QB2"},
                      {"Номер рейса": "PL2", "Вылет": "QB2",
                       "Прилёт": "QA1"},
                      {"Номер рейса": "PL3", "Вылет": "QA1",
                       "Прилёт": "QC3"}]},
        "Экипажи": {"columns": ["Табельный номер", "Фамилия И.О.",
                                "База"],
                    "rows": [
                        {"Табельный номер": "E-2",
                         "Фамилия И.О.": "Второв В.В.",
                         "База": "QB2"},
                        {"Табельный номер": "E-1",
                         "Фамилия И.О.": "Первов П.П.",
                         "База": "QA1"}]},
    }}
    (tmp_path / "data-freeze" / "dataset.json").write_text(
        _j.dumps({"tables": tables}, ensure_ascii=False),
        encoding="utf-8")
    r = c.get("/api/runs/run-x/map-projection")
    assert r.status_code == 200, r.text[:200]
    d = r.json()
    assert [n["code"] for n in d["nodes"]] == ["QA1", "QB2", "QC3"]
    assert {n["code"] for n in d["nodes"] if n.get("base")} \
        == {"QA1", "QB2"}
    assert d["arcs"] == [{"a": "QA1", "b": "QB2", "n": 2},
                        {"a": "QA1", "b": "QC3", "n": 1}]
    # Сотрудники отсортированы (base, id).
    assert [e["id"] for e in d["employees"]] == ["E-1", "E-2"]
    # Вымышленные коды не полны координатами — «мир» не включится.
    assert d["geo"]["complete"] is False

    # Общий движок и подложка отдаются панелью тем же файлом.
    assert c.get("/net-map.js").status_code == 200
    assert "NetMap" in c.get("/net-map.js").text[:2000]
    assert c.get("/geo/world-110m.json").status_code == 200


def test_переизвлечение_инвалидирует_обмен(tmp_path, monkeypatch):
    """Класс прод-демо run-0001 (30.08): регресс-токены старого свода
    судили новый после переизвлечения; извлечение при бегущем цикле
    стартовало параллельную стадию. Лечение: 409 при бегущей стадии;
    переизвлечение архивирует exchange ЦЕЛИКОМ (токены не едут)."""
    import json as _j

    from fastapi.testclient import TestClient

    from cdpm import webui as w
    monkeypatch.setattr(w, "_run_dir", lambda rid: tmp_path)
    c = TestClient(w.app)

    # Бегущий цикл — честный 409.
    (tmp_path / "freeze").mkdir()
    (tmp_path / "freeze" / "manifest.json").write_text(
        '{"files": []}', encoding="utf-8")
    (tmp_path / "status.json").write_text(_j.dumps(
        {"state": "сквозной цикл идёт", "customer": "Т"},
        ensure_ascii=False), encoding="utf-8")
    r = c.post("/api/runs/run-x/extract")
    assert r.status_code == 409
    assert "запрещено" in r.json()["detail"]

    # Переизвлечение: exchange с токенами уезжает в rev целиком.
    ex = tmp_path / "exchange"
    ex.mkdir()
    (ex / "tokens.json").write_text("[]", encoding="utf-8")
    (ex / "registry.json").write_text("{}", encoding="utf-8")
    (tmp_path / "status.json").write_text(_j.dumps(
        {"state": "извлечено", "customer": "Т"},
        ensure_ascii=False), encoding="utf-8")
    # Фоновая работа извлечения не нужна тесту гонки состояния —
    # глушится до старта потока.
    monkeypatch.setattr(w.threading, "Thread",
                        lambda *a, **k: type("T", (), {
                            "start": lambda self: None})())
    r = c.post("/api/runs/run-x/extract")
    assert r.status_code == 200, r.text[:200]
    ex_now = tmp_path / "exchange"
    assert (not ex_now.exists()) or not any(ex_now.iterdir()), \
        "exchange нового свода обязан быть пустым (родится циклом)"
    rev = tmp_path / "exchange-rev-001"
    assert (rev / "tokens.json").is_file(), \
        "токены старого свода — в архиве, не в новом цикле"


def test_channel_delivery_receipt_endpoint(client):
    """Кнопка «Передать в COSMOS» (05.09): квитанция последней доставки
    читается отдельной точкой; без доставок — sent=false и состояние
    канала словами (в тестах канал выключен autouse-фикстурой)."""
    import json as _j
    import hashlib
    import zipfile
    run_id = _upload(client).json()["run"]
    got = client.get(f"/api/runs/{run_id}/channel").json()
    assert got["sent"] is False
    assert got["configured"] is False
    assert "канал выключен" in got["detail"]
    ch = Path(webui.runs_root()) / run_id / "channel"
    ch.mkdir()
    package = ch / "task-package.zip"
    # Artifact identity only, not a gate-5-valid task-package fixture.
    with zipfile.ZipFile(package, "w") as archive:
        archive.writestr("fixture.txt", "channel delivery projection")
    package_sha = hashlib.sha256(package.read_bytes()).hexdigest()
    (ch / "delivery.json").write_text(_j.dumps({
        "sent": True, "sha": package_sha, "attempts": 1,
        "verdict": "OK", "detail": "пакет принят",
        "mismatch": [], "already_accepted": False,
        "mg": {"pairing": {"state": "feasible_soft"}},
        "at": "2026-09-05 21:28:40"}, ensure_ascii=False),
        encoding="utf-8")
    got = client.get(f"/api/runs/{run_id}/channel").json()
    assert got["sent"] is True
    assert got["sha"] == package_sha
    assert got["verdict"] == "OK" and got["at"] == "2026-09-05 21:28:40"
    assert got["run"] == run_id
    assert client.get("/api/runs/run-9999/channel").status_code == 404


def test_readiness_counts_all_questions_of_current_anketa(client):
    """Метрика «Вопросов» — все вопросы действующей анкеты (в том числе
    рождённые линией данных), а не только извлечённые (замечание
    владельца на репетиции 05.09)."""
    import json as _j
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    model = _j.loads((Path(webui.runs_root()) / run_id / "extract"
                      / "model.json").read_text(encoding="utf-8"))
    rd = st["readiness"]
    assert rd["vmp_questions_total"] == len(model.get("questions", []))
    assert rd["vmp_questions_total"] >= rd["vmp_open_questions"]
    # Вопрос, добавленный в каноническую модель позже извлечения
    # (как делает линия данных), тоже входит в счёт.
    model.setdefault("questions", []).append({
        "id": "Q-3b-история-нет", "part": "3b",
        "question": "Истории сотрудников в данных нет — как поступить?",
        "status": "open"})
    (Path(webui.runs_root()) / run_id / "extract" / "model.json") \
        .write_text(_j.dumps(model, ensure_ascii=False), encoding="utf-8")
    rd2 = client.get(f"/api/runs/{run_id}").json()["readiness"]
    assert rd2["vmp_questions_total"] == rd["vmp_questions_total"] + 1


def test_settings_carry_catalog_service_address(client, monkeypatch):
    """Шаг 5 каталога: адрес службы каталога шаблонов интерфейс берёт
    настройкой, а не из кода. Переменной нет — `null`, и панель говорит
    словами «служба каталога не настроена», вместо мёртвой ссылки.

    Рядом — отметка каталога, которым работает панель: по ней строится
    ссылка на карточку версии, и она же стоит в паспорте прогона.
    """
    monkeypatch.delenv("CREW_CATALOG_URL", raising=False)
    got = client.get("/api/settings").json()
    assert got["catalog_url"] is None
    assert "catalog" in got, sorted(got)
    stamp = got["catalog"]
    assert stamp is None or set(stamp) == {"version", "fingerprint"}

    # Пробелы по краям адреса — частая беда переменных среды: адрес
    # либо есть, либо его нет, «строка из пробелов» адресом не бывает.
    monkeypatch.setenv("CREW_CATALOG_URL", "  ")
    assert client.get("/api/settings").json()["catalog_url"] is None
    # Адрес — из зарезервированной зоны .invalid (RFC 2606): ссылка только
    # собирается, в сеть не ходит, и адрес рабочего места владельца
    # (127.0.0.1:8105) в тесте не живёт литералом (сторож литералов среды,
    # волна 9).
    monkeypatch.setenv("CREW_CATALOG_URL", "http://catalog.invalid")
    assert client.get("/api/settings").json()["catalog_url"] \
        == "http://catalog.invalid"
