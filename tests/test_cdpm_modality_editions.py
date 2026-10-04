"""Published modality questions retain immutable scope and real answer history."""
import copy

import pytest
from openpyxl import load_workbook

from cdpm import answers, anketa_xlsx as ax, model as m, question_editions as qe, render
from test_cdpm_traits import _mini, _fork_question


def _doc(shared=False, answer=None, published=True):
    doc = _mini()
    first, second = doc["rules"]
    q = _fork_question("Q-2-MOD", first["id"], {"а": "SOFT", "б": "HARD"})
    first["conditional"] = {"question": q["id"], "branches": {"а": "SOFT", "б": "HARD"}}
    first["blockers"] = [q["id"]]
    if shared:
        q["rule_refs"].append(second["id"])
        second["conditional"] = {"question": q["id"], "branches": {"а": "HARD", "б": "OUT"}}
        second["blockers"] = [q["id"]]
    if answer is not None:
        q["answer"] = copy.deepcopy(answer)
    doc["questions"] = [q]
    if published:
        qe.freeze(q)
    return doc, q, first, second


def test_published_single_modality_keeps_edition_and_records_code_outcome():
    doc, q, first, _ = _doc()
    old = copy.deepcopy(q)
    m.derive_rule_traits(doc, first, "test")
    assert q in doc["questions"]
    assert {k: v for k, v in q.items() if k != "code_resolutions"} == old
    assert qe.problems(q) == []
    assert not m.is_open_question(q)
    assert "answer" not in q
    assert q["code_resolutions"][0]["rule"] == first["id"]
    assert m.validate_schema(doc) == []
    assert m.derive_rule_traits(doc, first, "again") == []


def test_partial_resolution_keeps_scope_and_answer_applies_only_live_fork():
    doc, q, first, second = _doc(shared=True)
    original = copy.deepcopy(q)
    m.derive_rule_traits(doc, first, "test")
    assert q["rule_refs"] == original["rule_refs"]
    assert q["edition"] == original["edition"] and qe.problems(q) == []
    assert m.is_open_question(q)
    before_first = copy.deepcopy(first)
    result = answers.apply_answers(doc, {q["id"]: "б"})
    assert result["applied"] == 1
    assert first == before_first
    assert second["kind"] == "OUT" and not second.get("blockers")


@pytest.mark.parametrize("status", ["не разобран", "ожидает обработчика", "применён"])
def test_code_outcome_preserves_preexisting_real_answer(status):
    answer = {"raw": "мой исходный ответ", "status": status, "revision": 2}
    doc, q, first, _ = _doc(answer=answer)
    m.derive_rule_traits(doc, first, "test")
    assert q in doc["questions"] and q["answer"] == answer
    assert q["code_resolutions"][0]["rule"] == first["id"]
    assert not m.is_open_question(q)
    assert qe.problems(q) == []


def test_code_resolved_question_refuses_stale_book_answer():
    doc, q, first, _ = _doc()
    m.derive_rule_traits(doc, first, "test")
    assert q in doc["questions"]
    old = copy.deepcopy(q)
    result = answers.apply_answers(doc, {q["id"]: "а"})
    assert result["applied"] == 0
    assert q == old
    assert first["kind"] == "HARD"


def test_code_resolved_pending_worker_does_not_run():
    doc, q, first, _ = _doc(answer={"raw": "с оговоркой", "status": "ожидает обработчика", "revision": 2})
    m.derive_rule_traits(doc, first, "test")
    old = copy.deepcopy(doc)
    def forbidden(prompt):
        pytest.fail("code-resolved question must not invoke the worker")
    result = answers.apply_worker_patches(doc, forbidden)
    assert result["worker_applied"] == 0
    assert result["worker_rejected"]
    assert doc == old
    assert answers.remaining_awaiting(doc) == 0


def test_partial_caveat_cannot_reblock_code_resolved_rule():
    doc, q, first, second = _doc(shared=True)
    m.derive_rule_traits(doc, first, "test")
    before_first = copy.deepcopy(first)
    answers.apply_answers(doc, {q["id"]: "а, но только для дальних рейсов"})
    assert q["answer"]["status"] == "ожидает обработчика"
    assert first == before_first
    assert {r["id"] for r in answers._affected_rules(doc, q)} == {second["id"]}


def test_code_outcome_visible_in_markdown_and_saved_workbook(tmp_path):
    doc, q, first, _ = _doc()
    m.derive_rule_traits(doc, first, "test")
    md = render.render_anketa(doc)
    assert q["id"] in md and "снят кодом" in md.lower()
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [row for sheet in wb for row in sheet.iter_rows(values_only=True) if row[0] == q["id"]]
        assert len(rows) == 1
        assert "снят кодом" in str(rows[0]).lower()
    finally:
        wb.close()
    from cdpm.gates.gate_anketa_xlsx import check_workbook
    assert check_workbook(doc, path)[0] == []
    wb = load_workbook(path)
    wb[ax.ANSWERED_TITLE].cell(2, 4).value = "применён"
    wb.save(path)
    wb.close()
    assert check_workbook(doc, path)[0]


def test_unpublished_draft_modality_still_can_disappear():
    doc, q, first, _ = _doc(published=False)
    m.derive_rule_traits(doc, first, "test")
    assert q not in doc["questions"]


def test_all_targets_close_only_after_each_code_decision():
    doc, q, first, second = _doc(shared=True)
    second["conditional"]["branches"] = {"а": "SOFT", "б": "HARD"}
    m.derive_rule_traits(doc, first, "first")
    assert m.is_open_question(q)
    m.derive_rule_traits(doc, second, "second")
    assert not m.is_open_question(q)
    assert len(q["code_resolutions"]) == 2 and "answer" not in q
    assert qe.problems(q) == []
    assert dict((r[0], r[1]) for r in ax.summary_rows(doc) if len(r) == 2)["Принятых ответов"] == 0


def test_new_edition_does_not_inherit_old_code_decisions():
    doc, q, first, _ = _doc()
    m.derive_rule_traits(doc, first, "test")
    new = qe.expand(doc, q, [], content={"text": q["text"] + " Новое условие."})
    assert "code_resolutions" not in new and m.is_open_question(new)
    assert not m.is_open_question(q)


def test_legacy_partial_question_round_trips_actual_new_book(tmp_path):
    doc, q, first, _ = _doc(shared=True, published=False)
    qe.freeze(q, legacy=True)
    m.derive_rule_traits(doc, first, "test")
    path = tmp_path / "new-book.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path)
    ws = wb[ax.SHEET_TITLES[q["part"]]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == q["id"])
    row[ax.ANSWER_COL - 1].value = "б"
    wb.save(path)
    wb.close()
    got = answers.read_book(path, doc)
    assert got.answers == {q["id"]: "б"} and not got.refused


@pytest.mark.parametrize("bad", [None, [{"rule": "R-S1-a1"}]])
def test_malformed_code_decision_is_open_and_refused(bad):
    doc, q, _, _ = _doc()
    q["code_resolutions"] = bad
    assert m.is_open_question(q)
    assert m.validate_schema(doc)
    assert qe.application_problems(doc, q)


def test_reactivated_target_invalidates_code_decision_before_mutation():
    doc, q, first, _ = _doc()
    m.derive_rule_traits(doc, first, "test")
    first["conditional"] = {"question": q["id"], "branches": {"а": "SOFT", "б": "HARD"}}
    first["blockers"] = [q["id"]]
    before = copy.deepcopy(doc)
    with pytest.raises(qe.EditionRefusal, match="живой строкой"):
        m.derive_rule_traits(doc, first, "bad reuse")
    assert doc == before
    assert any("живой строкой" in p for p in m.validate_model(doc))


def test_changed_family_makes_old_code_resolution_non_authoritative():
    doc, q, first, _ = _doc()
    m.derive_rule_traits(doc, first, "test")
    first["template_fit"]["family"] = "calendar_windows"
    first["kind"] = "SOFT"
    assert q in m.open_questions(doc)
    assert q["id"] not in {row[0] for row in ax.answered_rows(doc)}
    assert any(q["id"] == row[0] for _, rows in ax.sheet_plan(doc) for row in rows)
    assert "не подтверждено" in render.render_anketa(doc)
