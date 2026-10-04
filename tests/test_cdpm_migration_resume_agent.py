"""Interrupted on-disk migration publication; no service, solver or LLM.

Temporary real catalogues and legacy models use the existing validated
fixtures. Only the filesystem write is interrupted; migration, validation,
publication, archive parsing and execution refusals are the real judges.
"""
import copy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cdpm import model as m
from test_cdpm_landing_migration import legacy
from test_cdpm_landing_v2_reader import receive
from test_k01_family_manifest import manifest
from test_k01_semantic_transition import clean_bridge
from webapp import task_package as receiver


@pytest.mark.parametrize("failure", [OSError, KeyboardInterrupt])
def test_interrupted_publication_keeps_legacy_source_and_retry_cannot_execute_archive(
        tmp_path, monkeypatch, failure):
    from cdpm import webui

    doc, _rule, cat = legacy(tmp_path, monkeypatch)
    assert m.validate_schema(doc) == m.validate_model(doc) == []
    raw = json.dumps(doc, ensure_ascii=False).encode("utf-8")
    run = tmp_path / "run"
    path = run / "extract" / "model.json"
    path.parent.mkdir(parents=True)
    path.write_bytes(raw)

    # A healthy control proves this is a legal, confirmable migration,
    # not a refusal caused by bad input or a made-up catalogue record.
    proposal = m.load_model(path)
    assert proposal["meta"]["landing_migrations"][-1]["items"][0]["state"] == "confirmed"
    assert m.validate_schema(proposal) == m.validate_model(proposal) == []
    healthy = tmp_path / "healthy.json"
    m.publish_model(copy.deepcopy(proposal), healthy)
    expected = healthy.read_bytes()
    assert m.load_model(healthy) == proposal
    assert path.read_bytes() == raw

    declared = manifest(cat)
    declared["landing_compat"] = 1
    archived = receive(tmp_path, doc, declared, landing_compat=1)
    archive = tmp_path / "landing.zip"
    archive_bytes = archive.read_bytes()
    assert archived["landing_projection"]["archive_only"] is True

    # This public handler really loads and publishes the same model path
    # before creating a missing questionnaire. Only run-directory lookup
    # is redirected; no HTTP server or network request is involved.
    monkeypatch.setattr(webui, "_run_dir", lambda _run_id: run)
    original_write = Path.write_text
    interrupted_paths = []

    def partial_write_then_interrupt(target, text, *args, **kwargs):
        if target.parent == path.parent:
            interrupted_paths.append(target)
            original_write(target, text[:len(text) // 3], *args, **kwargs)
            raise failure("injected interruption after a real partial model write")
        return original_write(target, text, *args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(Path, "write_text", partial_write_then_interrupt)
        with pytest.raises(failure, match="real partial model write"):
            webui.run_file("migration-resume", "анкета-xlsx")
    assert len(interrupted_paths) == 1
    assert path.read_bytes() == raw
    assert set(path.parent.iterdir()) == {path}
    assert not (path.parent / "anketa.xlsx").exists()
    assert m.load_model(path) == proposal

    # Retry the actual consumer of publication, not just pure proposal.
    response = webui.run_file("migration-resume", "анкета-xlsx")
    assert Path(response.path).is_file()
    book = load_workbook(response.path, read_only=True)
    try:
        from cdpm.anketa_xlsx import SUMMARY_TITLE
        assert SUMMARY_TITLE in book.sheetnames
        assert list(book[SUMMARY_TITLE].values)
    finally:
        book.close()
    assert path.read_bytes() == expected
    restored = m.load_model(path)
    assert m.validate_schema(restored) == m.validate_model(restored) == []
    assert restored == proposal
    webui.run_file("migration-resume", "анкета-xlsx")
    assert path.read_bytes() == expected

    # A successful NEW model publication does not rewrite an old received
    # package or turn archive compatibility into execution permission.
    assert archive.read_bytes() == archive_bytes
    reread = receiver.load(archive, run_gate=False, landing_compat=1)
    assert reread["model"] == doc
    assert reread["landing_projection"]["archive_only"] is True
    reread.pop("landing_projection")
    with pytest.raises(receiver.TaskPackageRefused) as error:
        receiver.to_payload(reread)
    assert error.value.code == "landing-version"
    assert archive.read_bytes() == archive_bytes


@pytest.mark.parametrize("failure", [PermissionError, KeyboardInterrupt])
def test_failed_replace_keeps_source_and_callers_model_then_retry_succeeds(
        tmp_path, monkeypatch, failure):
    from cdpm import file_replace

    doc, _rule, _cat = legacy(tmp_path, monkeypatch)
    raw = json.dumps(doc, ensure_ascii=False).encode("utf-8")
    directory = tmp_path / "publication"
    directory.mkdir()
    path = directory / "model.json"
    path.write_bytes(raw)
    proposal = m.load_model(path)
    assert m.validate_schema(proposal) == m.validate_model(proposal) == []
    before = copy.deepcopy(proposal)
    healthy = tmp_path / "healthy.json"
    m.publish_model(copy.deepcopy(proposal), healthy)
    expected = healthy.read_bytes()
    attempts = []

    def fail_replacement(stage, target):
        attempts.append((stage, target))
        assert target == path and stage.parent == path.parent and stage != path
        assert stage.read_bytes() == expected
        assert target.read_bytes() == raw
        raise failure("injected failure before replacement")

    with monkeypatch.context() as patch:
        patch.setattr(file_replace, "replace_with_retry", fail_replacement)
        with pytest.raises(failure, match="before replacement"):
            m.publish_model(proposal, path)
    assert len(attempts) == 1
    assert path.read_bytes() == raw
    assert set(directory.iterdir()) == {path}
    assert proposal == before == m.load_model(path)

    m.publish_model(proposal, path)
    assert path.read_bytes() == expected
    assert m.load_model(path) == before == proposal
    assert set(directory.iterdir()) == {path}


def test_cleanup_failure_does_not_hide_interruption_or_touch_source(tmp_path, monkeypatch):
    from cdpm import file_replace

    doc, _rule, _cat = legacy(tmp_path, monkeypatch)
    directory = tmp_path / "publication"
    directory.mkdir()
    path = directory / "model.json"
    raw = json.dumps(doc, ensure_ascii=False).encode("utf-8")
    path.write_bytes(raw)
    proposal = m.load_model(path)
    attempts = []

    def interrupt_replace(stage, target):
        attempts.append(stage)
        assert target == path and stage.is_file()
        raise KeyboardInterrupt("original publication interruption")

    original_unlink = Path.unlink

    def deny_staging_cleanup(target, *args, **kwargs):
        if target in attempts:
            raise PermissionError("staging cleanup denied")
        return original_unlink(target, *args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(file_replace, "replace_with_retry", interrupt_replace)
        patch.setattr(Path, "unlink", deny_staging_cleanup)
        with pytest.raises(KeyboardInterrupt, match="original publication interruption"):
            m.publish_model(proposal, path)
    assert len(attempts) == 1 and attempts[0].is_file()
    assert path.read_bytes() == raw
    assert m.load_model(path) == proposal
    # A failed cleanup leaves only a non-authoritative staging file;
    # retry neither reads it nor mistakes it for the canonical model.
    staged_bytes = attempts[0].read_bytes()
    m.publish_model(proposal, path)
    assert m.load_model(path) == proposal
    assert attempts[0].read_bytes() == staged_bytes
