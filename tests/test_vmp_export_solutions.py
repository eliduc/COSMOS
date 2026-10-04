# -*- coding: utf-8 -*-
"""Выгрузки решений пакета в известных форматах (vmp/export_solutions.py,
слово владельца 06.09.2026) — стенд на синтетическом мини-мире: три
аэропорта, семь рейсов на два дня, три цепочки (одна из двух смен с
перегоном и восстановлением смен единой моделью, одна закреплённая с
периодами единой модели цены, одна с границами duty_fns), ростер на
этих цепочках. Ни одного числа и имени авиакомпании.

Известные ответы: (1) все форматы детерминированы (два вызова — одни
байты, created — из журнала выдачи, не «сейчас»); (2) Meridian-пейринг
проходит зеркало гейта BA (П1…П7) и сам гейт BA подпроцессом, когда он
есть на площадке; на мутациях документа зеркало ОТКАЗЫВАЕТ там же, где
гейт; (3) шкала времени — метки как в наборе CDPM по умолчанию, перевод в
UTC только по utc=1 с переводом planSince/planUntil и legId; (4) круг
«выгрузка → импорт COSMOS» замыкается на алфавите source (Р / D/H);
(5) ростер: роли мест — из назначений (не из индекса slot), разбиение
смен сравнивается множествами участков (duty_id — ярлык), сквозной
ростер и разорванная цепочка — отказ словами, номера цепочек не
пересекаются и при метках вне формы «каталог:<k>»; (6) iCal: события
без Z при неизвестной либо shifted шкале, TZID при base-local, Z только
при utc; DTSTAMP всегда в форме UTC (RFC 5545) с заметкой о шкале
журнала; UID различает прогоны; (7) персональные данные («Фамилия
И.О.») не выгружаются; (8) шкала `shifted` пакета названа своим словом,
не «не объявлена»; (9) байты XLSX не зависят от площадки (create_system
и права записей zip фиксированы); (10) указатель доступности отказывает
там же, где отказал бы файл (выдача без журнала)."""
from __future__ import annotations

import copy
import io
import json
import os
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))

from vmp import export_solutions as xs  # noqa: E402

SHA = "ab" * 32
MODEL_SHA = "cd" * 32
CAT_SHA = "ef" * 32
#: Уровней вверх от корня репозитория, на которых ищется клон линии
#: BA: своё дерево git уводит корень на три уровня вглубь.
_BA_SEARCH_DEPTH = 6


def _ba_gate() -> Path | None:
    """Гейт линии BA: переменная BA_MERIDIAN_GATE, иначе клон линии
    BA РЯДОМ с репозиторием.

    Зашитый абсолютный путь машины разработчика делал вид, что гейт
    есть везде: на чистом раннере тест уходил в skip по ЧУЖОМУ пути,
    а на машине разработчика был зелёным только потому, что клон
    линии BA лежал ровно там. Ищем по-соседски и вверх по дереву —
    так работает и из рабочего дерева git."""
    named = os.environ.get("BA_MERIDIAN_GATE", "").strip()
    if named:
        path = Path(named)
        return path if path.is_file() else None
    for parent in list(HERE.parents)[:_BA_SEARCH_DEPTH]:
        candidate = parent / "BA" / "tools" / "ba_meridian_gate.py"
        if candidate.is_file():
            return candidate
    return None


GATE_PATH = _ba_gate()
#: Слова пропуска — одни на оба стенда гейта.
NO_GATE = ("гейт линии BA не найден: ни переменной BA_MERIDIAN_GATE, "
           "ни клона линии BA рядом с репозиторием")

F1 = "PL101|2026-09-04T08:00:00"    # PLY→SVG — работа цепочки A
F2 = "PL102|2026-09-04T11:00:00"    # SVG→GRT — работа цепочки B, перегон цепочки A
F3 = "PL103|2026-09-05T07:00:00"    # GRT→PLY — работа цепочки A
F4 = "PL201|2026-09-04T06:00:00"    # PLY→SVG — цепочка B (закреплена)
F5 = "PL202|2026-09-04T14:00:00"    # GRT→PLY — цепочка B
F6 = "PL301|2026-09-05T10:00:00"    # PLY→SVG — цепочка C
F7 = "PL302|2026-09-05T12:30:00"    # SVG→PLY — цепочка C

_FLIGHTS = [
    (F1, "PLY", "2026-09-04T08:00:00", "SVG", "2026-09-04T09:30:00"),
    (F2, "SVG", "2026-09-04T11:00:00", "GRT", "2026-09-04T12:30:00"),
    (F3, "GRT", "2026-09-05T07:00:00", "PLY", "2026-09-05T09:00:00"),
    (F4, "PLY", "2026-09-04T06:00:00", "SVG", "2026-09-04T07:30:00"),
    (F5, "GRT", "2026-09-04T14:00:00", "PLY", "2026-09-04T15:30:00"),
    (F6, "PLY", "2026-09-05T10:00:00", "SVG", "2026-09-05T11:00:00"),
    (F7, "SVG", "2026-09-05T12:30:00", "PLY", "2026-09-05T13:30:00"),
]


def flights() -> list[dict]:
    return [{"flightNumber": fn, "flightLegId": fn, "departureAirport": dep,
             "departureUTCDateTime": t0, "arrivalAirport": arr,
             "arrivalUTCDateTime": t1, "aircraftType": "T-1", "transportOnly": False}
            for fn, dep, t0, arr, t1 in _FLIGHTS]


def payload() -> dict:
    # Спрос мест: indexInFlight 0 — FirstOfficer, 1 — Captain (нарочно НЕ в
    # порядке seats каталога: роль места берётся из назначений/спроса, не
    # из индекса slot).
    fas = []
    for f in flights():
        fas.append({"id": f["flightNumber"] + "#0", "flight": f["flightNumber"],
                    "indexInFlight": 0, "requiredSkill": "FirstOfficer"})
        fas.append({"id": f["flightNumber"] + "#1", "flight": f["flightNumber"],
                    "indexInFlight": 1, "requiredSkill": "Captain"})
    return {"airports": [{"code": c, "name": c} for c in ("PLY", "SVG", "GRT")],
            "employees": [{"id": "H0%d" % i, "homeAirport": "PLY",
                           "skills": ["Captain" if i % 2 else "FirstOfficer"]}
                          for i in range(1, 6)],
            "flights": flights(), "flightAssignments": fas,
            "_seatRoles": ["Captain", "FirstOfficer"],
            "_coverage": {"obligations": 7, "transport": 0},
            "window_horizon": {"start_inclusive": "2026-09-04T00:00",
                               "end_exclusive": "2026-09-06T00:00"}}


def catalog() -> dict:
    a = {"fns": [F1, F2, F3], "active_fns": [F1, F3], "seats": ["Captain", "FirstOfficer"],
         "seatCodes": [["T-1_Captain"], ["T-1_FirstOfficer"]], "pinned": False,
         "sourceUnit": "pairing-builder-l3:0", "kind": "chain", "minRestAfterMin": 720,
         "cost": 5, "price": {"known_subtotal": 5, "full_cost": None,
                              "cost_certainty": "partial", "nights": 1, "crew": 2}}
    b = {"fns": [F4, F2, F5], "active_fns": [F4, F2, F5], "duty_fns": [[F4, F2, F5]],
         "seats": ["Captain", "FirstOfficer"], "pinned": True,
         "sourceUnit": "pairing-builder-l3:1", "kind": "chain", "minRestAfterMin": 720,
         "cost": 0, "price": {"known_subtotal": 0, "full_cost": 0, "cost_certainty": "full",
                              "nights": 0, "crew": 2,
                              "accounting_scope": {"periods": [
                                  {"start": "2026-09-04T05:00:00", "end": "2026-09-04T16:30:00",
                                   "rest_at": "PLY", "fns": [F4, F2, F5]}]}}}
    c = {"fns": [F6, F7], "active_fns": [F6, F7], "duty_fns": [[F6, F7]],
         "seats": ["Captain", "FirstOfficer"], "pinned": False,
         "sourceUnit": "=SUM(1)", "kind": "chain", "minRestAfterMin": 600, "cost": 0}
    return {"label": "package-abcdef123456", "ip_status": "OPTIMAL", "ip_objective": 5,
            "ip_objective_scope": "known_subtotal", "known_subtotal": 5, "full_cost": None,
            "coverage": {"legs_total": 7, "uncovered": 0, "verdict": "OK", "complete": True},
            "proven_impossible_legs": [], "proven_impossible_proofs": [],
            "config_sig": {"rest_base": 720, "rest_off": 600},
            "price_certainty": {"status": "partial"}, "selected": [a, b, c]}


def manifest() -> dict:
    return {"spec_version": "1.18", "source_run": "run-T", "customer": "стенд",
            "horizon": {"start": "2026-09-04T00:00", "end_exclusive": "2026-09-06T00:00",
                        "tz": "Europe/Moscow"},
            "time": {"time_scale": "base-local", "base_tz_offset_min": 0,
                     "base_time_zone": "Europe/Moscow"},
            "kit": {"data_class": "synthetic", "llm_scope": "local-only"}}


def rules() -> dict:
    return {"layover_tiers": [[1440, 720, 600]]}


def pairing_delivery(certified=True) -> dict:
    return {"verdict": "пейринг построен", "certified": certified,
            "certificate_verdict": "допустимо" if certified else "недопустимо",
            "findings": 0 if certified else 2,
            "result": {"feasibility": "confirmed" if certified else "refuted",
                       "cost": "partial", "optimality": "catalog_optimum"},
            "mode": "feasibility-soft", "objective": None, "coverage_verdict": "OK",
            "chains": 3, "coverage": {"complete": True}, "catalog_sha256": CAT_SHA,
            "cost_certainty": "partial",
            "schedule": {"solver_run_id": "run-p-1"},
            "trace": {"cdpm_run": SHA, "solver_run_id": "run-p-1", "task": "pairing",
                      "package_fingerprint": "f" * 64, "spec_version": "1.18"},
            "journal": [{"step": "Ш-0", "status": "пройден", "at": "2026-09-06T16:41:17"},
                        {"step": "Ш-4", "status": "сертифицирован",
                         "at": "2026-09-06T16:45:23"}]}


def certificate() -> dict:
    return {"customer": "стенд", "certificate": {"certified": True, "findings": 0,
                                                 "solution": "стенд"}}


def pairing_sources(**over) -> xs.PairingSources:
    base = dict(delivery=pairing_delivery(), catalog=catalog(), payload=payload(),
                manifest=manifest(), rules=rules(), certificate=certificate(),
                model_sha256=MODEL_SHA, package_sha=SHA, attempt=None)
    base.update(over)
    return xs.PairingSources(**base)


# ───────────────────────────────────────────────── ростер на цепочках

def _asg(emp, fn, role, pid, did, base="PLY", deadhead=None, date=True):
    f = {x[0]: x for x in _FLIGHTS}[fn]
    row = {"employee": emp, "flight": fn, "dep_time": f[2], "arr_time": f[4],
           "dep_airport": f[1], "arr_airport": f[3], "role": role,
           "pairing_id": pid, "duty_id": did, "base": base}
    if date:
        row["date"] = f[2][:10]
    if deadhead is not None:
        row["deadhead"] = deadhead
    return row


def roster_solution() -> dict:
    a = [_asg("H01", F1, "Captain", "каталог:0:slot:0", "H01/D1"),
         _asg("H01", F2, "Captain", "каталог:0:slot:0", "H01/D1", deadhead="да"),
         _asg("H01", F3, "Captain", "каталог:0:slot:0", "H01/D2"),
         _asg("H02", F1, "FirstOfficer", "каталог:0:slot:1", "H02/D1"),
         _asg("H02", F2, "FirstOfficer", "каталог:0:slot:1", "H02/D1", deadhead="да"),
         _asg("H02", F3, "FirstOfficer", "каталог:0:slot:1", "H02/D2"),
         _asg("H03", F4, "Captain", "каталог:1:slot:0", "H03/D1", date=False),
         _asg("H03", F2, "Captain", "каталог:1:slot:0", "H03/D1"),
         _asg("H03", F5, "Captain", "каталог:1:slot:0", "H03/D1"),
         _asg("H04", F4, "FirstOfficer", "каталог:1:slot:1", "H04/D1"),
         _asg("H04", F2, "FirstOfficer", "каталог:1:slot:1", "H04/D1"),
         _asg("H04", F5, "FirstOfficer", "каталог:1:slot:1", "H04/D1", deadhead="нет"),
         _asg("H03", F6, "Captain", "каталог:2:slot:0", "H03/D2"),
         _asg("H03", F7, "Captain", "каталог:2:slot:0", "H03/D2"),
         _asg("H04", F6, "FirstOfficer", "каталог:2:slot:1", "H04/D2"),
         _asg("H04", F7, "FirstOfficer", "каталог:2:slot:1", "H04/D2")]
    duties = [{"id": "H01/D1", "employee": "H01", "start": "2026-09-04T07:00:00",
               "end": "2026-09-04T13:00:00", "base": "PLY", "last_arr_airport": "GRT"},
              {"id": "H01/D2", "employee": "H01", "start": "2026-09-05T06:00:00",
               "end": "2026-09-05T09:30:00", "base": "PLY", "last_arr_airport": "PLY"},
              {"id": "H02/D1", "employee": "H02", "start": "2026-09-04T07:00:00",
               "end": "2026-09-04T13:00:00", "base": "PLY", "last_arr_airport": "GRT"},
              {"id": "H02/D2", "employee": "H02", "start": "2026-09-05T06:00:00",
               "end": "2026-09-05T09:30:00", "base": "PLY", "last_arr_airport": "PLY"},
              {"id": "H03/D1", "employee": "H03", "start": "2026-09-04T05:00:00",
               "end": "2026-09-04T16:30:00", "base": "PLY", "last_arr_airport": "PLY"},
              {"id": "H04/D1", "employee": "H04", "start": "2026-09-04T05:00:00",
               "end": "2026-09-04T16:30:00", "base": "PLY", "last_arr_airport": "PLY"}]
    pairings = [{"employee": "H01", "id": "каталог:0:slot:0", "start": "2026-09-04T08:00:00",
                 "end": "2026-09-05T09:00:00"},
                {"employee": "H02", "id": "каталог:0:slot:1", "start": "2026-09-04T08:00:00",
                 "end": "2026-09-05T09:00:00"}]
    crews = {"columns": ["Табельный номер", "Фамилия И.О.", "База приписки", "Должность"],
             "rows": [{"Табельный номер": "H0%d" % i, "Фамилия И.О.": "Пилот %d" % i,
                       "База приписки": "PLY", "Должность": "КВС" if i % 2 else "ВП"}
                      for i in range(1, 6)]}
    return {"solution": "кнопка стенда", "assignments": a, "duties": duties,
            "pairings": pairings,
            "window_horizon": {"start_inclusive": "2026-09-04T00:00",
                               "end_exclusive": "2026-09-06T00:00"},
            "window_employees": [{"employee": "H0%d" % i, "base": "PLY"} for i in range(1, 6)],
            "reference_tables": {"D1": {"Экипажи": crews}},
            "thresholds": {"applied": [{"threshold": "min_rest_home_min", "value": 720},
                                       {"threshold": "min_rest_outstation_min", "value": 600}]},
            "runInfo": {"score": "0hard/-2soft"}}


def solved(unfilled: bool = False) -> dict:
    fas = []
    slots = {F1: ("H02", "H01"), F2: ("H04", "H03"), F3: ("H02", "H01"),
             F4: ("H04", "H03"), F5: ("H04", "H03"), F6: ("H04", "H03"), F7: ("H04", "H03")}
    chain_of = {F1: 0, F2: 1, F3: 0, F4: 1, F5: 1, F6: 2, F7: 2}
    for fn, (fo, cpt) in slots.items():
        k = chain_of[fn]
        fas.append({"id": fn + "#0", "flight": fn, "indexInFlight": 0,
                    "requiredSkill": "FirstOfficer", "employee": fo,
                    "pinned": False, "chainId": "каталог:%d:slot:1" % k, "deadhead": False})
        fas.append({"id": fn + "#1", "flight": fn, "indexInFlight": 1,
                    "requiredSkill": "Captain",
                    "employee": None if (unfilled and k == 2) else cpt,
                    "pinned": fn == F4, "chainId": "каталог:%d:slot:0" % k, "deadhead": False})
    # перегонные места цепочки A на рейсе цепочки B
    fas.append({"id": F2 + "#2", "flight": F2, "indexInFlight": 2, "requiredSkill": "FirstOfficer",
                "employee": "H02", "pinned": False, "chainId": "каталог:0:slot:1", "deadhead": True})
    fas.append({"id": F2 + "#3", "flight": F2, "indexInFlight": 3, "requiredSkill": "Captain",
                "employee": "H01", "pinned": False, "chainId": "каталог:0:slot:0", "deadhead": True})
    return {"flights": flights(), "employees": payload()["employees"],
            "flightAssignments": fas, "score": "0hard/-2soft"}


def roster_delivery(certified=True) -> dict:
    return {"verdict": "допустимо", "certified": certified,
            "certificate_verdict": "допустимо" if certified else "недопустимо",
            "findings": 0, "result": {"feasibility": "confirmed", "cost": "partial",
                                      "optimality": "feasible_found"},
            "mode": "feasibility-soft", "objective": None,
            "pairing_run_id": "run-p-1", "pairing_chains": 6, "pairing_catalog_sha256": CAT_SHA,
            "pairing_selection": {"pairing_run_id": "run-p-1", "attempt_id": "A1",
                                  "catalog_sha256": CAT_SHA},
            "solver_hard": 0, "counted_share": {"counted": 2, "of": 2},
            "window_coverage": {"expected": 4, "evaluated": 4, "unknown": 0, "checks": []},
            "cost_certainty": "partial",
            "schedule": {"solver_run_id": "run-r-1", "score": "0hard/-2soft"},
            "trace": {"cdpm_run": SHA, "solver_run_id": "run-r-1", "task": "rostering",
                      "package_fingerprint": "f" * 64, "spec_version": "1.18"},
            "journal": [{"step": "Ш-0", "status": "пройден", "at": "2026-09-06T16:50:00"},
                        {"step": "Ш-4г", "status": "гейт зелёный", "at": "2026-09-06T16:55:10"}]}


def roster_sources(**over) -> xs.RosterSources:
    base = dict(delivery=roster_delivery(), solution=roster_solution(), solved=solved(),
                payload=solved(), manifest=manifest(), certificate=certificate(),
                model_sha256=MODEL_SHA, package_sha=SHA, attempt=None, path="on-pairing",
                pairing_catalog=catalog(), pairing_catalog_note=None)
    base.update(over)
    return xs.RosterSources(**base)


def _xlsx_sheets(data: bytes) -> dict[str, list[list]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


# ══════════════════════════════════════════════════════════════ пейринг

def test_pairing_csv_rows_activity_and_formula_guard():
    text = xs.pairing_csv(pairing_sources())
    assert not text.startswith("\ufeff")            # BOM добавляет ручка
    assert "\r\n" in text
    lines = text.rstrip("\r\n").split("\r\n")
    assert lines[0] == ",".join(xs.PAIRING_CSV_HEADER)
    assert len(lines) - 1 == 8                        # Σ участков цепочек (F2 дважды)
    rows = {(l.split(",")[1], l.split(",")[6]): l.split(",") for l in lines[1:]}
    assert rows[("1", F2)][13] == "deadhead" and rows[("2", F2)][13] == "work"
    assert rows[("1", F1)][8] == "PLY" and rows[("1", F1)][10] == "SVG"   # участки соединены
    assert rows[("3", F6)][17] == "'=SUM(1)"                           # защита от формул
    assert rows[("2", F4)][16] == "true"                               # закреплена


def test_pairing_xlsx_sheets_and_determinism():
    src = pairing_sources()
    data = xs.pairing_xlsx(src)
    assert data == xs.pairing_xlsx(src)
    sheets = _xlsx_sheets(data)
    assert list(sheets) == ["Паспорт", "Цепочки", "Участки", "Смены", "Покрытие", "Показатели"]
    assert len(sheets["Цепочки"]) - 1 == 3 and len(sheets["Участки"]) - 1 == 8
    assert sheets["Участки"][0] == xs.PAIRING_XLSX_LEGS_HEADER
    passport = {r[0]: r[1] for r in sheets["Паспорт"][1:]}
    assert passport["package.sha256"] == SHA and passport["created"] == "2026-09-06T16:45:23"
    assert passport["time.zone"] == "Europe/Moscow"
    # строка-формула осталась строкой
    assert any(r[17] == "=SUM(1)" for r in sheets["Участки"][1:])
    kpi = {r[0]: r[1] for r in sheets["Показатели"][1:]}
    assert kpi["Сертифицировано"] is True and kpi["ip_status"] == "OPTIMAL"


def test_pairing_json_chains_connected_and_sorted():
    doc = json.loads(xs.pairing_json(pairing_sources()))
    assert doc["kind"] == "cosmos-pairing-export/1"
    assert list(doc) == sorted(doc)
    assert "journal" not in doc["delivery"] and doc["delivery"]["verdict"] == "пейринг построен"
    a = doc["chains"][0]
    assert a["id"] == "COSMOS-abcdef123456-0001"
    assert [l["activity"] for l in a["legs"]] == ["work", "deadhead", "work"]
    assert a["legs"][1]["dep"] == "SVG" and a["legs"][1]["arr_time"] == "2026-09-04T12:30:00"
    assert len(a["duties"]) == 2 and a["duties"][0]["flight_min"] == 90     # без перегона
    assert a["duties"][0]["rest_after_min"] == 1110 and a["duties"][1]["rest_after_min"] == 720
    assert "восстановлены единой моделью" in a["notes"][0]
    b = doc["chains"][1]
    assert b["duties"][0]["start"] == "2026-09-04T05:00:00" and b["pinned"] is True
    assert b["duties"][0]["flight_min"] == 270 and b["duties"][0]["work_min"] == 690
    assert doc["coverage"]["covered"] == 7 and doc["certificate"]["certified"] is True


def test_meridian_pairing_passes_mirror_and_keeps_invariants():
    src = pairing_sources()
    doc = xs.meridian_pairing_doc(src)
    cat, pay = xs.meridian_gate_inputs(src, doc)
    assert xs.meridian_check(doc, "pairing", catalog=cat, payload=pay) == []
    assert doc["type"] == "pairing" and doc["division"] == "FD" and doc["format"] == 1
    legs = {l["flightNumber"]: l for l in doc["legs"]}
    # умолчание — метки как есть (как в наборе CDPM), legId по дате листа
    assert legs["PL101"]["start"] == "2026-09-04T08:00:00"
    assert legs["PL101"]["legId"] == "PL101$04.09.2026$1" and legs["PL101"]["legNumber"] == 1
    assert legs["PL101"]["carrier"] == "PL" and legs["PL101"]["fleet"] == "T-1"
    duties = {d["dutyId"]: d for d in doc["duties"]}
    d1, d2 = duties["COSMOS-abcdef123456-0001$1"], duties["COSMOS-abcdef123456-0001$2"]
    assert legs["PL102"]["dutyCrewTypes"] == {d1["dutyId"]: "П",
                                             "COSMOS-abcdef123456-0002$1": "Р"}
    assert legs["PL101"]["dutyCrewTypes"] == {d1["dutyId"]: "Р"}
    assert d1["flightTimeCredit"] == 90 and d1["restOverride"] == 1110 and d2["restOverride"] == 720
    assert d1["complements"] == [{"employeeId": "", "complementRank": "Captain",
                                  "requirements": [], "locked": False},
                                 {"employeeId": "", "complementRank": "FirstOfficer",
                                  "requirements": [], "locked": False}]
    pairs = {p["pairingId"]: p for p in doc["pairings"]}
    assert pairs["COSMOS-abcdef123456-0002"]["locked"] is True
    assert pairs["COSMOS-abcdef123456-0001"]["restOverride"] == 720
    assert pairs["COSMOS-abcdef123456-0001"]["flightMinutesCredit"] == 90 + 120
    assert all(isinstance(m["metricValue"], str) for m in doc["metrics"])
    m = {x["metricCode"]: x["metricValue"] for x in doc["metrics"]}
    assert m["LEGS_TOTAL"] == "7" and m["LEGS_OBLIGATIONS"] == "7" and m["LEGS_COVERED"] == "7"
    assert m["CERTIFIED"] == "1" and m["ROLE_ALPHABET"] == "Р/П" and m["TIME_SCALE"] == "base-local"
    assert doc["planSince"] <= min(l["start"] for l in doc["legs"])
    assert max(l["end"] for l in doc["legs"]) <= doc["planUntil"]
    assert all(l["legId"].split("$")[1] == datetime.fromisoformat(l["start"]).strftime("%d.%m.%Y")
               for l in doc["legs"])
    assert doc["created"] == "2026-09-06T16:45:23"
    assert "перевод в UTC не выполнен — как в наборе CDPM" in doc["description"]
    assert "местное время машины, пояс не записан" in doc["description"]
    assert doc["datasetId"] == SHA and doc["presetId"] == MODEL_SHA
    assert doc["groundActivities"] == [] and "employees" not in doc


def test_meridian_pairing_utc_mode_shifts_labels_plan_and_legids():
    src = pairing_sources()
    doc = xs.meridian_pairing_doc(src, xs.Options(utc=True))
    legs = {l["flightNumber"]: l for l in doc["legs"]}
    assert legs["PL101"]["start"] == "2026-09-04T05:00:00"        # Europe/Moscow → UTC
    assert legs["PL101"]["legId"] == "PL101$04.09.2026$1"
    assert doc["planSince"] <= min(l["start"] for l in doc["legs"])
    assert max(l["end"] for l in doc["legs"]) <= doc["planUntil"]
    assert all(l["legId"].split("$")[1] == datetime.fromisoformat(l["start"]).strftime("%d.%m.%Y")
               for l in doc["legs"])
    d = {x["dutyId"]: x for x in doc["duties"]}["COSMOS-abcdef123456-0001$1"]
    assert d["start"] == "2026-09-04T05:00:00"
    cat, pay = xs.meridian_gate_inputs(src, doc)
    assert xs.meridian_check(doc, "pairing", catalog=cat, payload=pay) == []
    assert "переведено из base-local" in doc["description"]
    # ночной рейс до 03:00 местного — другая дата legId в UTC-режиме
    fl = payload()
    fl["flights"].append({"flightNumber": "PL900|2026-09-05T01:00:00",
                          "flightLegId": "PL900|2026-09-05T01:00:00",
                          "departureAirport": "PLY", "departureUTCDateTime": "2026-09-05T01:00:00",
                          "arrivalAirport": "SVG", "arrivalUTCDateTime": "2026-09-05T02:30:00",
                          "aircraftType": "T-1", "transportOnly": True})
    src2 = pairing_sources(payload=fl)
    ids_default = {l["flightNumber"]: l["legId"] for l in xs.meridian_pairing_doc(src2)["legs"]}
    ids_utc = {l["flightNumber"]: l["legId"]
               for l in xs.meridian_pairing_doc(src2, xs.Options(utc=True))["legs"]}
    assert ids_default["PL900"] == "PL900$05.09.2026$1"
    assert ids_utc["PL900"] == "PL900$04.09.2026$1"


def test_meridian_legids_match_cdpm_dataset_writer():
    """Набор для Meridian из той же книги пишет cdpm/export_meridian.py
    (чтение разрешено, правка — нет): legId решения обязаны совпасть с
    legId набора по тем же листам."""
    from cdpm import export_meridian as cem
    legs_rows = [{"Номер рейса": fn.split("|")[0], "Дата": t0[:10],
                  "Вылет": dep, "Время вылета": t0[11:16],
                  "Прилёт": arr, "Время прилёта": t1[11:16], "Тип ВС": "T-1"}
                 for fn, dep, t0, arr, t1 in _FLIGHTS]
    emp_rows = [{"Табельный": "H01", "База": "PLY", "Дата рождения": "1990-01-01",
                 "Пол": "M", "Гражданство": "XX", "Лимит налёта": 80, "Лимит работы": 160,
                 "Тип ВС": "T-1", "Эскадрилья": "1"}]
    apt_rows = [{"Код": c, "Страна": "XX", "Часовой пояс": "Europe/Moscow",
                 "Категория": "A", "Провозка": "NONE", "МВЛ": "нет", "СНГ": "нет"}
                for c in ("PLY", "SVG", "GRT")]
    tables = {"D1": {"Рейсы": {"columns": list(legs_rows[0]), "rows": legs_rows},
                     "Экипажи": {"columns": list(emp_rows[0]), "rows": emp_rows},
                     "Аэропорты": {"columns": list(apt_rows[0]), "rows": apt_rows}}}
    dataset = cem.build_dataset(tables, preset_id="p-1")
    ours = xs.meridian_pairing_doc(pairing_sources())
    assert {l["legId"] for l in dataset["legs"]} == {l["legId"] for l in ours["legs"]}
    assert {l["legId"]: l["start"] for l in dataset["legs"]} == \
        {l["legId"]: l["start"] for l in ours["legs"]}
    assert (dataset["planSince"], dataset["planUntil"]) == (ours["planSince"], ours["planUntil"])


def test_meridian_pairing_gate_ba_subprocess(tmp_path):
    """Гейт линии BA подпроцессом на файлах meridian_gate_inputs. Файла
    гейта на площадке нет — skip с причиной и путём (BA_MERIDIAN_GATE);
    зеркало meridian_check прогоняется всегда (тест выше): оно повторяет
    П1…П7 гейта пункт в пункт (та же семантика, тот же алфавит Р/П) и
    служит эквивалентом там, где каталога BA нет. Гейт судит только
    type=pairing (на roster падает assert-ом) — ростер судится зеркалом."""
    if GATE_PATH is None:
        pytest.skip(NO_GATE + " — судить нечего")
    gate = GATE_PATH
    src = pairing_sources()
    doc = xs.meridian_pairing_doc(src)
    cat, pay = xs.meridian_gate_inputs(src, doc)
    for name, obj in (("sol.json", doc), ("cat.json", cat), ("pay.json", pay)):
        (tmp_path / name).write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
    r = subprocess.run([sys.executable, str(gate), str(tmp_path / "sol.json"),
                        str(tmp_path / "cat.json"), str(tmp_path / "pay.json")],
                       capture_output=True, text=True, encoding="utf-8",
                       env={**os.environ, "PYTHONUTF8": "1"})
    assert r.returncode == 0, r.stdout + r.stderr
    assert "ГЕЙТ MERIDIAN PASS" in r.stdout


def _mutations() -> list[tuple[str, callable]]:
    def m_rank(d): d["duties"][0]["complements"][0]["complementRank"] = ""
    def m_role(d): d["legs"][0]["dutyCrewTypes"] = {k: "X" for k in d["legs"][0]["dutyCrewTypes"]}
    def m_window(d): d["duties"][0]["start"] = "2026-09-04T09:00:00"
    def m_credit(d): d["pairings"][0]["flightMinutesCredit"] += 1
    def m_emp(d): d["duties"][0]["complements"][0]["employeeId"] = "H01"
    def m_covered(d):
        for m in d["metrics"]:
            if m["metricCode"] == "LEGS_COVERED":
                m["metricValue"] = "0"
    def m_dup(d): d["legs"][1]["legId"] = d["legs"][0]["legId"]
    def m_numbers(d):
        for du in d["duties"]:
            if du["pairingId"] == "COSMOS-abcdef123456-0001":
                du["numberOfDuty"] += 1
    return [("пустой complementRank", m_rank), ("роль вне алфавита", m_role),
            ("участок вне окна смены", m_window), ("сумма кредитов ≠", m_credit),
            ("employeeId у пейринга", m_emp), ("LEGS_COVERED ≠ каталогу", m_covered),
            ("дубль legId", m_dup), ("несплошные numberOfDuty", m_numbers)]


@pytest.mark.parametrize("name,mutate", _mutations(), ids=[n for n, _ in _mutations()])
def test_meridian_mirror_refuses_where_gate_refuses(name, mutate, tmp_path):
    src = pairing_sources()
    base = xs.meridian_pairing_doc(src)
    cat, pay = xs.meridian_gate_inputs(src, base)
    doc = copy.deepcopy(base)
    mutate(doc)
    fails = xs.meridian_check(doc, "pairing", catalog=cat, payload=pay)
    assert fails, name
    if GATE_PATH is None:
        pytest.skip(NO_GATE + " — мутация «%s» проверена зеркалом"
                    % name)
    gate = GATE_PATH
    for fname, obj in (("sol.json", doc), ("cat.json", cat), ("pay.json", pay)):
        (tmp_path / fname).write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
    r = subprocess.run([sys.executable, str(gate), str(tmp_path / "sol.json"),
                        str(tmp_path / "cat.json"), str(tmp_path / "pay.json")],
                       capture_output=True, text=True, encoding="utf-8",
                       env={**os.environ, "PYTHONUTF8": "1"})
    assert r.returncode != 0, "гейт BA пропустил мутацию «%s»: %s" % (name, r.stdout)


def test_meridian_pairing_refusals_in_words():
    # роль мест вне канона → division не определим
    pl = payload()
    pl["_seatRoles"] = ["Navigator"]
    with pytest.raises(xs.ExportRefused, match="division не определим"):
        xs.meridian_pairing_doc(pairing_sources(payload=pl))
    # объявление division в манифесте первично
    man = manifest()
    man["division"] = "CC"
    assert xs.meridian_pairing_doc(pairing_sources(payload=pl, manifest=man))["division"] == "CC"
    # перевозчик не определим
    pl2 = payload()
    for f in pl2["flights"]:
        if f["flightNumber"] == F1:
            f["flightNumber"] = "1|2026-09-04T08:00:00"
    cat = catalog()
    cat["selected"][0]["fns"][0] = "1|2026-09-04T08:00:00"
    cat["selected"][0]["active_fns"][0] = "1|2026-09-04T08:00:00"
    with pytest.raises(xs.ExportRefused, match="перевозчик не определим"):
        xs.meridian_pairing_doc(pairing_sources(payload=pl2, catalog=cat))
    # несертифицированное — отказ без флага, с флагом — пометка
    src = pairing_sources(delivery=pairing_delivery(certified=False))
    with pytest.raises(xs.ExportRefused, match="allow_uncertified"):
        xs.meridian_pairing_doc(src)
    doc = xs.meridian_pairing_doc(src, xs.Options(allow_uncertified=True))
    assert "НЕ СЕРТИФИЦИРОВАНО" in doc["description"]
    assert {m["metricCode"]: m["metricValue"] for m in doc["metrics"]}["CERTIFIED"] == "0"
    # certified: null — не установлено = не сертифицировано
    d = pairing_delivery()
    d["certified"] = None
    cert = xs.certification(d)
    assert cert == dict(cert, certified=False, established=False)
    # выдача-остановка
    d2 = pairing_delivery()
    d2["verdict"] = "остановка: решатель отверг пакет"
    with pytest.raises(xs.ExportRefused, match="остановка"):
        xs.meridian_pairing_doc(pairing_sources(delivery=d2))
    # пустые обязательства
    cat3 = catalog()
    cat3["coverage"]["legs_total"] = 0
    with pytest.raises(xs.ExportRefused, match="нет рейсов-обязательств"):
        xs.meridian_pairing_doc(pairing_sources(catalog=cat3))
    # алфавит вне словаря
    with pytest.raises(xs.ExportBadRequest, match="алфавит ролей"):
        xs.meridian_pairing_doc(pairing_sources(), xs.Options(roles="x"))


def test_meridian_roundtrip_through_cosmos_importer(tmp_path):
    """Круг «выгрузка → импорт COSMOS»: документ на алфавите source
    (константы писателя и импортёра, Р / D/H) читается
    _meridian_to_payload.convert; цепочки и перегоны восстанавливаются
    один в один. Зеркало гейта на том же документе честно отказывает П5
    (гейт BA знает только Р/П) — два алфавита, один документ."""
    import _meridian_to_payload as m2p
    src = pairing_sources()
    doc = xs.meridian_pairing_doc(src, xs.Options(roles="source"))
    roles = {v for l in doc["legs"] for v in l["dutyCrewTypes"].values()}
    assert roles == {"Р", m2p.DEADHEAD_DUTY_ROLE}
    cat, pay = xs.meridian_gate_inputs(src, doc)
    assert any("[П5]" in f for f in xs.meridian_check(doc, "pairing", catalog=cat, payload=pay))
    assert xs.meridian_check(doc, "pairing", catalog=cat, payload=pay,
                             roles=xs.role_alphabets()["source"]) == []
    feed = dict(doc, employees=[], airports=[{"IATA": c} for c in ("PLY", "SVG", "GRT")])
    pl, check, errors, units = m2p.convert(feed, None, "стенд", "check.json", None, None, True)
    assert not errors, errors
    src_of = {f["flightNumber"]: f["sourceLegId"] for f in pl["flights"]}
    leg_of_fn = {l["flightNumber"]: l["legId"] for l in doc["legs"]}
    got = {(frozenset(src_of[x] for x in u["fns"]), frozenset(src_of[x] for x in u["active_fns"]))
           for u in units}
    want = set()
    for ch in xs.pairing_chains(src):
        want.add((frozenset(leg_of_fn[l["flight_no"]] for l in ch["legs"]),
                  frozenset(leg_of_fn[l["flight_no"]] for l in ch["legs"] if l["activity"] == "work")))
    assert got == want


def test_created_stamp_scales_and_missing_journal():
    d = pairing_delivery()
    assert xs.created_stamp(d) == ("2026-09-06T16:45:23", "machine-local")
    d["journal"][-1]["at_utc"] = "2026-09-06T13:45:23"
    assert xs.created_stamp(d) == ("2026-09-06T13:45:23", "utc")
    d2 = pairing_delivery()
    d2["journal"][-1]["at"] = "2026-09-06T16:45:23+03:00"
    assert xs.created_stamp(d2) == ("2026-09-06T13:45:23", "utc")
    d3 = pairing_delivery()
    d3["journal"] = []
    with pytest.raises(xs.ExportRefused, match="нет журнала"):
        xs.created_stamp(d3)
    d3.pop("journal")
    with pytest.raises(xs.ExportRefused, match="нет журнала"):
        xs.pairing_json(pairing_sources(delivery=d3))


def test_ids_order_and_preset_parity_between_solutions():
    p = xs.meridian_pairing_doc(pairing_sources())
    r = xs.meridian_roster_doc(roster_sources())
    assert p["presetId"] == r["presetId"] == MODEL_SHA and p["datasetId"] == r["datasetId"] == SHA
    doc = xs.meridian_pairing_doc(pairing_sources(), xs.Options(dataset_id="DS-1", preset_id="PR-1"))
    assert (doc["datasetId"], doc["presetId"]) == ("DS-1", "PR-1")
    assert "задан явно параметром" in doc["description"]
    man = manifest()
    man["source"] = {"kind": "meridian", "meridian": {"datasetId": "DS-M", "presetId": "PR-M"}}
    doc2 = xs.meridian_pairing_doc(pairing_sources(manifest=man))
    assert (doc2["datasetId"], doc2["presetId"]) == ("DS-M", "PR-M")
    # пакет из источника Meridian — алфавит source по умолчанию
    assert {m["metricCode"]: m["metricValue"] for m in doc2["metrics"]}["ROLE_ALPHABET"] == \
        "%s/%s" % xs.role_alphabets()["source"]


def test_legs_total_counts_transport_and_obligations_separately():
    pl = payload()
    pl["flights"].append({"flightNumber": "XX900|2026-09-05T20:00:00",
                          "flightLegId": "XX900|2026-09-05T20:00:00",
                          "departureAirport": "GRT", "departureUTCDateTime": "2026-09-05T20:00:00",
                          "arrivalAirport": "SVG", "arrivalUTCDateTime": "2026-09-05T21:00:00",
                          "aircraftType": "T-9", "transportOnly": True})
    src = pairing_sources(payload=pl)
    doc = xs.meridian_pairing_doc(src)
    m = {x["metricCode"]: x["metricValue"] for x in doc["metrics"]}
    assert m["LEGS_TOTAL"] == "8" and m["LEGS_OBLIGATIONS"] == "7" and m["LEGS_COVERED"] == "7"
    legs = {l["flightNumber"]: l for l in doc["legs"]}
    assert legs["XX900"]["deadhead"] is True and legs["XX900"]["dutyCrewTypes"] == {}
    cat, pay = xs.meridian_gate_inputs(src, doc)
    assert xs.meridian_check(doc, "pairing", catalog=cat, payload=pay) == []


def test_augmented_crew_counts_as_extended_and_demand_gap_refuses():
    cat = catalog()
    cat["selected"][1]["seats"] = ["Captain", "FirstOfficer", "FirstOfficer"]
    doc = xs.meridian_pairing_doc(pairing_sources(catalog=cat))
    assert {m["metricCode"]: m["metricValue"] for m in doc["metrics"]}["DUTIES_EXTENDED_CREW"] == "1"
    d = {x["dutyId"]: x for x in doc["duties"]}["COSMOS-abcdef123456-0002$1"]
    assert [c["complementRank"] for c in d["complements"]] == ["Captain", "FirstOfficer", "FirstOfficer"]
    cat2 = catalog()
    cat2["selected"][1]["seats"] = ["FirstOfficer"]
    with pytest.raises(xs.ExportRefused, match="не покрывают спрос"):
        xs.meridian_pairing_doc(pairing_sources(catalog=cat2))
    cat3 = catalog()
    cat3["selected"][2].pop("seats")
    doc3 = xs.meridian_pairing_doc(pairing_sources(catalog=cat3))
    d3 = {x["dutyId"]: x for x in doc3["duties"]}["COSMOS-abcdef123456-0003$1"]
    assert [c["complementRank"] for c in d3["complements"]] == ["Captain", "FirstOfficer"]
    assert "состав по спросу участков" in doc3["description"]


def test_division_roles_cover_the_cosmos_role_canon():
    from webapp import task_package
    assert set().union(*xs.DIVISION_ROLES.values()) == set(task_package.ROLE_ORDER)


# ══════════════════════════════════════════════════════════════ ростер

def test_roster_csv_xlsx_json_normalization_and_no_personal_data():
    src = roster_sources()
    text = xs.roster_csv(src)
    lines = text.rstrip("\r\n").split("\r\n")
    assert lines[0] == ",".join(xs.ROSTER_CSV_HEADER)
    assert len(lines) - 1 == 16
    by = {(l.split(",")[0], l.split(",")[4]): l.split(",") for l in lines[1:]}
    assert by[("H01", F2)][10] == "deadhead" and by[("H04", F5)][10] == "work"
    assert by[("H03", F2)][10] == "work"
    assert by[("H03", F4)][2] == "2026-09-04"                 # date восстановлен
    assert "Пилот" not in text and "КВС" not in text           # персональных данных нет
    data = xs.roster_xlsx(src)
    sheets = _xlsx_sheets(data)
    assert list(sheets) == ["Паспорт", "Назначения", "Смены", "Цепочки", "Доска", "Штат",
                            "Показатели"]
    assert sheets["Назначения"][0] == xs.ROSTER_XLSX_ASSIGN_HEADER
    staff = {r[0] for r in sheets["Штат"][1:]}
    assert "H05" in staff and len(sheets["Штат"][0]) == 6
    board = {r[0]: r for r in sheets["Доска"][1:]}
    assert board["H01"][3] == "PL101 *PL102" and (board["H05"][3] or "") == ""
    assert board["H03"][3] == "PL201 PL102 PL202"
    assert sheets["Доска"][0][3:] == ["04.09", "05.09"]
    flat = json.dumps(sheets, ensure_ascii=False)
    assert "Пилот " not in flat and "КВС" not in flat
    assert not any("Фамилия" in str(h) or "Должность" in str(h)
                   for sh in sheets.values() for h in sh[0])
    doc = json.loads(xs.roster_json(src))
    assert doc["kind"] == "cosmos-roster-export/1" and list(doc) == sorted(doc)
    emp = {e["id"]: e for e in doc["employees"]}
    assert emp["H05"]["legs"] == 0 and emp["H01"]["deadheads"] == 1 and "name" not in emp["H01"]
    a = [x for x in doc["assignments"] if x["employee"] == "H01" and x["flight_key"] == F2][0]
    assert a["deadhead"] is True and a["date"] == "2026-09-04"
    assert "Пилот" not in json.dumps(doc, ensure_ascii=False)


def test_roster_ics_events_filters_uid_scales_and_folding():
    src = roster_sources()
    ics = xs.roster_ics(src)
    assert ics.count("BEGIN:VEVENT") == 16
    def event_z(s):
        """Строки DTSTART/DTEND, записанные как UTC («…Z»)."""
        return [l for l in s.split("\r\n")
                if l.startswith(("DTSTART", "DTEND")) and l.endswith("Z")]
    # DTSTAMP — штамп выдачи (не «сейчас») в форме UTC по RFC 5545 §3.8.7.2:
    # при журнале без пояса — без перевода, с заметкой словами
    assert "DTSTAMP:20260906T165510Z" in ics
    assert "X-COSMOS-DTSTAMP-NOTE:" in ics and "без перевода" in ics.replace("\r\n ", "")
    assert "DTSTART;TZID=Europe/Moscow:20260904T080000" in ics    # base-local → TZID
    assert "BEGIN:VTIMEZONE" in ics and "TZID:Europe/Moscow" in ics
    assert not event_z(ics)                                          # события — не UTC
    assert all(len(l.encode("utf-8")) <= 75 for l in ics.split("\r\n"))
    assert "X-COSMOS-CERTIFIED:true" in ics and "X-COSMOS-CREATED-SCALE:machine-local" in ics
    # журнал с at_utc — штамп UTC по праву, заметки нет
    d_utc = roster_delivery()
    d_utc["journal"][-1]["at_utc"] = "2026-09-06T13:55:10"
    ics_utc = xs.roster_ics(roster_sources(delivery=d_utc))
    assert "DTSTAMP:20260906T135510Z" in ics_utc and "X-COSMOS-DTSTAMP-NOTE" not in ics_utc
    assert "X-COSMOS-CREATED-SCALE:utc" in ics_utc
    assert "\\;" in ics                                             # экранирование
    one = xs.roster_ics(src, xs.Options(employee="H01"))
    assert one.count("BEGIN:VEVENT") == 3 and "SUMMARY:DH PL102 SVG→GRT (Captain)" in one
    with pytest.raises(xs.ExportNotFound, match="не найден"):
        xs.roster_ics(src, xs.Options(employee="H99"))
    # UID: равен между вызовами, не зависит от порядка строк, различает прогоны
    def uids(s):
        return sorted(l for l in s.split("\r\n") if l.startswith("UID:"))
    sol = roster_solution()
    sol["assignments"].reverse()
    assert uids(xs.roster_ics(roster_sources(solution=sol))) == uids(ics)
    d = roster_delivery()
    d["schedule"]["solver_run_id"] = "run-r-2"
    assert uids(xs.roster_ics(roster_sources(delivery=d))) != uids(ics)
    # utc=1 → Z; неизвестная шкала → плавающее время без Z и TZID
    utc = xs.roster_ics(src, xs.Options(utc=True))
    assert "DTSTART:20260904T050000Z" in utc and "TZID=" not in utc
    man = manifest()
    man.pop("time")
    unk = xs.roster_ics(roster_sources(manifest=man))
    assert "DTSTART:20260904T080000\r\n" in unk and not event_z(unk) and "TZID" not in unk
    assert "X-COSMOS-TIME-SCALE:unknown" in unk
    man2 = manifest()
    man2["time"] = {"time_scale": "utc"}
    assert "DTSTART:20260904T080000Z" in xs.roster_ics(roster_sources(manifest=man2))


def test_meridian_roster_draft_form_mirror_and_roles_from_assignments():
    src = roster_sources()
    doc = xs.meridian_roster_doc(src)
    rows = xs.roster_rows(src)
    assert doc["type"] == "roster" and doc["description"].startswith(xs.ROSTER_FORM_NOTE)
    assert xs.meridian_check(doc, "roster", rows=rows) == []
    duties = {d["dutyId"]: d for d in doc["duties"]}
    d1 = duties["COSMOS-abcdef123456-0001$1"]
    # роль места — из назначений: slot 0 ↔ Captain, хотя indexInFlight 0 — FirstOfficer
    assert [(c["employeeId"], c["complementRank"]) for c in d1["complements"]] == \
        [("H01", "Captain"), ("H02", "FirstOfficer")]
    assert d1["flightTimeCredit"] == 90 and d1["start"] == "2026-09-04T07:00:00"
    legs = {l["flightNumber"]: l for l in doc["legs"]}
    assert legs["PL102"]["dutyCrewTypes"] == {d1["dutyId"]: "П",
                                             "COSMOS-abcdef123456-0002$1": "Р"}
    d2 = duties["COSMOS-abcdef123456-0002$1"]
    assert [c["locked"] for c in d2["complements"]] == [True, False]       # закрепление H03 на PL201
    assert d2["start"] == "2026-09-04T05:00:00" and d2["flightTimeCredit"] == 270
    m = {x["metricCode"]: x["metricValue"] for x in doc["metrics"]}
    assert m["SEATS_TOTAL"] == "8" and m["SEATS_FILLED"] == "8" and m["SEATS_UNFILLED"] == "0"
    assert m["FORM_STATUS"] == "draft-unconfirmed" and m["EMPLOYEES_USED"] == "4"
    assert m["CHAINS_DIVERGED_FROM_CATALOG"] == "0"
    pairs = {p["pairingId"]: p for p in doc["pairings"]}
    assert pairs["COSMOS-abcdef123456-0002"]["locked"] is True
    assert pairs["COSMOS-abcdef123456-0002"]["restOverride"] == 720
    assert pairs["COSMOS-abcdef123456-0003"]["restOverride"] == 600
    # разные ярлыки duty_id у людей одной цепочки (H01/D1 и H02/D1) — не отказ


def test_meridian_roster_refusals_unfilled_and_divergence():
    # сквозной ростер (без слоя цепочек)
    sol = roster_solution()
    for a in sol["assignments"]:
        a.pop("pairing_id")
    with pytest.raises(xs.ExportRefused, match="без слоя цепочек"):
        xs.meridian_roster_doc(roster_sources(solution=sol, path="cold"))
    # цепочка разделена между сотрудниками по участкам
    sol2 = roster_solution()
    sol2["assignments"] = [a for a in sol2["assignments"]
                           if not (a["employee"] == "H02" and a["flight"] == F3)]
    with pytest.raises(xs.ExportRefused, match="разделена между сотрудниками"):
        xs.meridian_roster_doc(roster_sources(solution=sol2))
    # одинаковые множества участков, разное разбиение на смены
    sol3 = roster_solution()
    for a in sol3["assignments"]:
        if a["employee"] == "H02" and a["flight"] == F3:
            a["duty_id"] = "H02/D1"
    with pytest.raises(xs.ExportRefused, match="разбиты по-разному"):
        xs.meridian_roster_doc(roster_sources(solution=sol3))
    # место без человека — не выкидывается: employeeId "" и SEATS_UNFILLED
    sol4 = roster_solution()
    sol4["assignments"] = [a for a in sol4["assignments"]
                           if not (a["employee"] == "H03" and a["flight"] in (F6, F7))]
    src4 = roster_sources(solution=sol4, solved=solved(unfilled=True), payload=solved(unfilled=True))
    doc4 = xs.meridian_roster_doc(src4)
    d = {x["dutyId"]: x for x in doc4["duties"]}["COSMOS-abcdef123456-0003$1"]
    assert [(c["employeeId"], c["complementRank"]) for c in d["complements"]] == \
        [("", "Captain"), ("H04", "FirstOfficer")]
    m = {x["metricCode"]: x["metricValue"] for x in doc4["metrics"]}
    assert m["SEATS_UNFILLED"] == "1" and m["SEATS_FILLED"] == "7" and m["SEATS_TOTAL"] == "8"
    assert xs.meridian_check(doc4, "roster", rows=xs.roster_rows(src4)) == []
    # решатель добавил перегон сверх каталога — истина по назначениям, метрика
    sol5 = roster_solution()
    for emp, pid, did in (("H03", "каталог:1:slot:0", "H03/D1"), ("H04", "каталог:1:slot:1", "H04/D1")):
        sol5["assignments"].append(_asg(emp, F1, "Captain" if emp == "H03" else "FirstOfficer",
                                        pid, did, deadhead=True))
    doc5 = xs.meridian_roster_doc(roster_sources(solution=sol5))
    m5 = {x["metricCode"]: x["metricValue"] for x in doc5["metrics"]}
    assert m5["CHAINS_DIVERGED_FROM_CATALOG"] == "1"
    assert "отличается от каталога" in doc5["description"]
    # без каталога пейринга — идентификаторы по метке, отдых по порогам
    doc6 = xs.meridian_roster_doc(roster_sources(pairing_catalog=None,
                                                 pairing_catalog_note="каталога нет"))
    pairs = {p["pairingId"]: p for p in doc6["pairings"]}
    assert set(pairs) == {"COSMOS-%s-%04d" % (SHA[:12], i) for i in (1, 2, 3)}
    assert pairs["COSMOS-%s-0002" % SHA[:12]]["restOverride"] == 720 and "порог решателя" in doc6["description"]
    # несертифицированное — только с флагом
    with pytest.raises(xs.ExportRefused, match="allow_uncertified"):
        xs.meridian_roster_doc(roster_sources(delivery=roster_delivery(certified=False)))


def test_meridian_roster_no_rest_override_when_chain_continues_in_same_duty():
    """Сотрудник продолжает следующей цепочкой в той же смене решателя —
    переопределение отдыха после цепочки не пишется (0), иначе файл
    утверждал бы, что решение нарушает собственный отдых; зеркало Р3
    держит невязку."""
    sol = roster_solution()
    for a in sol["assignments"]:
        if a["employee"] in ("H03", "H04") and a["flight"] in (F6, F7):
            a["flight"] = {F6: "PL301|2026-09-04T19:00:00", F7: "PL302|2026-09-04T21:00:00"}[a["flight"]]
            a["dep_time"] = {"PL301|2026-09-04T19:00:00": "2026-09-04T19:00:00",
                             "PL302|2026-09-04T21:00:00": "2026-09-04T21:00:00"}[a["flight"]]
            a["arr_time"] = {"PL301|2026-09-04T19:00:00": "2026-09-04T20:00:00",
                             "PL302|2026-09-04T21:00:00": "2026-09-04T22:00:00"}[a["flight"]]
            a["date"] = "2026-09-04"
            a["duty_id"] = a["employee"] + "/D1"
    pl = solved()
    for f in pl["flights"]:
        if f["flightNumber"] in (F6, F7):
            new = {F6: ("PL301|2026-09-04T19:00:00", "2026-09-04T19:00:00", "2026-09-04T20:00:00"),
                   F7: ("PL302|2026-09-04T21:00:00", "2026-09-04T21:00:00", "2026-09-04T22:00:00")}[f["flightNumber"]]
            f["flightNumber"], f["flightLegId"] = new[0], new[0]
            f["departureUTCDateTime"], f["arrivalUTCDateTime"] = new[1], new[2]
    for fa in pl["flightAssignments"]:
        fa["flight"] = {F6: "PL301|2026-09-04T19:00:00", F7: "PL302|2026-09-04T21:00:00"}.get(fa["flight"], fa["flight"])
    cat = catalog()
    cat["selected"][2]["fns"] = cat["selected"][2]["active_fns"] = \
        ["PL301|2026-09-04T19:00:00", "PL302|2026-09-04T21:00:00"]
    cat["selected"][2]["duty_fns"] = [list(cat["selected"][2]["fns"])]
    src = roster_sources(solution=sol, solved=pl, payload=pl, pairing_catalog=cat)
    doc = xs.meridian_roster_doc(src)
    pairs = {p["pairingId"]: p for p in doc["pairings"]}
    assert pairs["COSMOS-abcdef123456-0002"]["restOverride"] == 0
    assert "продолжает следующей цепочкой" in doc["description"]
    assert xs.meridian_check(doc, "roster", rows=xs.roster_rows(src)) == []


# ═══════════════════════════════════════════ имена, детерминированность

def test_file_names_ascii_and_utf8_form():
    a, u = xs.file_name("roster", "ics", SHA, None, True, "H01")
    assert a == u == "cosmos-roster-%s-H01.ics" % SHA[:12]
    a, u = xs.file_name("roster", "ics", SHA, "20260906T164117-842d52", False, "Иванов-1")
    assert a.isascii() and a.endswith("-uncertified.ics") and "Иванов-1" in u
    assert "filename*=UTF-8''" in xs.content_disposition(a, u)
    a, u = xs.file_name("pairing", "meridian", SHA, None, True)
    assert a == "cosmos-pairing-%s-meridian.json" % SHA[:12]


def test_every_format_is_deterministic_and_created_is_not_now():
    now = datetime.now().strftime("%Y-%m-%dT%H:%M")
    for src, fmts in ((pairing_sources(), xs.FORMATS["pairing"]),
                      (roster_sources(), xs.FORMATS["roster"])):
        for fmt in fmts:
            one = xs.export(src, fmt)
            two = xs.export(src, fmt)
            assert one[0] == two[0], fmt
            assert one[4]["created"] == ("2026-09-06T16:45:23" if isinstance(src, xs.PairingSources)
                                         else "2026-09-06T16:55:10")
            assert not one[4]["created"].startswith(now)
    csv_bytes = xs.export(pairing_sources(), "csv")[0]
    assert csv_bytes.startswith("\ufeff".encode("utf-8"))
    with pytest.raises(xs.ExportNotFound, match="перечень"):
        xs.export(pairing_sources(), "ics")


def test_availability_index_words():
    av = xs.availability(pairing_sources(delivery=pairing_delivery(certified=False)), "pairing")
    assert av["formats"]["csv"]["ok"] is True and av["certified"] is False
    assert av["formats"]["meridian"]["ok"] is False and av["formats"]["meridian"]["needs_confirm"]
    sol = roster_solution()
    for a in sol["assignments"]:
        a.pop("pairing_id")
    av2 = xs.availability(roster_sources(solution=sol, path="cold"), "roster")
    assert av2["formats"]["ics"]["ok"] and "без слоя цепочек" in av2["formats"]["meridian"]["reason"]
    assert av2["employees"] == ["H01", "H02", "H03", "H04", "H05"]
    # выдача без журнала — штамп created не установлен: файл ответил бы 409,
    # и указатель говорит ту же причину по каждому формату, а не рисует
    # живую ссылку (та же базовая проверка, что у export())
    d = pairing_delivery()
    d.pop("journal")
    av3 = xs.availability(pairing_sources(delivery=d), "pairing")
    assert all(v["ok"] is False and "нет журнала" in v["reason"]
               for v in av3["formats"].values())
    assert "needs_confirm" not in av3["formats"]["meridian"]
    d2 = roster_delivery()
    d2["journal"][-1].pop("at")
    av4 = xs.availability(roster_sources(delivery=d2), "roster")
    assert all("момент выдачи не установлен" in v["reason"] for v in av4["formats"].values())
    with pytest.raises(xs.ExportRefused, match="момент выдачи не установлен"):
        xs.export(roster_sources(delivery=d2), "ics")


def test_time_scale_shifted_is_declared_not_unknown():
    """CDPM пишет time_scale «shifted» при ненулевом base_tz_offset_min:
    шкала объявлена, смысл сдвига контрактом не закреплён — перевод в UTC
    не выполняется, и причина названа этим словом (описание, метрика,
    заголовок), а не «не объявлена пакетом»."""
    man = manifest()
    man["time"] = {"time_scale": "shifted", "base_tz_offset_min": 180,
                   "base_time_zone": "Europe/Moscow"}
    sc = xs.time_scale(man)
    assert (sc.kind, sc.offset_min, sc.zone, sc.convertible) == \
        ("shifted", 180, "Europe/Moscow", False)
    assert sc.header() == "shifted;offset=180;zone=Europe/Moscow"
    assert xs.time_scale({"time": {"time_scale": "shifted",
                                   "base_tz_offset_min": 180}}).header() == "shifted;offset=180"
    assert "не объявлена" not in sc.words() and "сдвиг 180 мин" in sc.words()
    note = xs.scale_note(sc, True)
    assert "shifted (сдвиг 180 мин)" in note and "не выполнен" in note and "не объявлена" not in note
    assert "shifted, сдвиг 180 мин" in xs.scale_note(sc, False)
    src = pairing_sources(manifest=man)
    doc = xs.meridian_pairing_doc(src, xs.Options(utc=True))
    legs = {l["flightNumber"]: l for l in doc["legs"]}
    assert legs["PL101"]["start"] == "2026-09-04T08:00:00"           # без перевода
    assert legs["PL101"]["legId"] == "PL101$04.09.2026$1"
    assert "шкала меток: shifted (сдвиг 180 мин)" in doc["description"]
    assert {m["metricCode"]: m["metricValue"] for m in doc["metrics"]}["TIME_SCALE"] == "shifted"
    cat, pay = xs.meridian_gate_inputs(src, doc)
    assert xs.meridian_check(doc, "pairing", catalog=cat, payload=pay) == []
    assert xs.export(src, "csv", xs.Options(utc=True))[4]["time_scale"] == \
        "shifted;offset=180;zone=Europe/Moscow"
    assert xs.export(src, "json")[4]["time_scale"] == "shifted;offset=180;zone=Europe/Moscow"
    assert json.loads(xs.pairing_json(src))["passport"]["time"] == \
        {"scale": "shifted", "zone": "Europe/Moscow", "offset_min": 180,
         "labels": "как в артефактах (шкала пакета)"}
    ics = xs.roster_ics(roster_sources(manifest=man), xs.Options(utc=True))
    flat = ics.replace("\r\n ", "")
    assert "X-COSMOS-TIME-SCALE:shifted;offset=180;zone=Europe/Moscow;note=" in flat
    assert "DTSTART:20260904T080000\r\n" in ics and "TZID" not in ics
    # прежние шкалы не тронуты
    assert xs.time_scale(manifest()).header() == "base-local;zone=Europe/Moscow"
    assert xs.time_scale({"time": {"time_scale": "utc"}}).header() == "utc"
    assert xs.time_scale({}).header() == "unknown"


def test_meridian_roster_chain_numbers_unique_with_mixed_labels():
    """Индексная метка «каталог:k» — номер k+1 (как у выгрузки пейринга);
    метка иной формы — с первого номера, свободного и от каталога, и от
    индексных меток. Порядковый номер в списке давал бы при разрежённых
    индексах («каталог:0», «каталог:2», «прочее») дубль 0003 —
    pairingId и dutyId обязаны быть различны."""
    sol = roster_solution()
    for a in sol["assignments"]:
        if a["pairing_id"].startswith("каталог:1:"):
            a["pairing_id"] = "прочее"
    src = roster_sources(solution=sol)
    doc = xs.meridian_roster_doc(src)
    ids = [p["pairingId"] for p in doc["pairings"]]
    assert len(ids) == len(set(ids)) == 3
    assert set(ids) == {"COSMOS-abcdef123456-%04d" % n for n in (1, 3, 4)}
    duty_ids = [d["dutyId"] for d in doc["duties"]]
    assert len(duty_ids) == len(set(duty_ids))
    assert "пронумерованы после каталожных, с номера 0004" in doc["description"]
    assert xs.meridian_check(doc, "roster", rows=xs.roster_rows(src)) == []
    # все метки индексные — пометки нет, номера прежние
    base = xs.meridian_roster_doc(roster_sources())
    assert "пронумерованы после каталожных" not in base["description"]
    assert {p["pairingId"] for p in base["pairings"]} == \
        {"COSMOS-abcdef123456-%04d" % n for n in (1, 2, 3)}


def test_xlsx_zip_entries_do_not_depend_on_platform():
    """Записи zip-контейнера XLSX: система создания и права фиксированы
    (стандартная библиотека ставит create_system 0 на Windows и 3 на
    прочих — один файл с Local и с legion отличался бы на байт в каждой
    записи центрального каталога), даты — created выдачи."""
    for data, stamp in ((xs.pairing_xlsx(pairing_sources()), (2026, 9, 6, 16, 45, 23)),
                        (xs.roster_xlsx(roster_sources()), (2026, 9, 6, 16, 55, 10))):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            core = z.read("docProps/core.xml").decode("utf-8")
        assert infos
        assert all(i.create_system == 3 for i in infos)
        assert all(i.external_attr == 0o600 << 16 for i in infos)
        # у zip секунды хранятся с шагом 2 (формат DOS) — 23 → 22
        assert all(i.date_time == stamp[:5] + (stamp[5] // 2 * 2,) for i in infos)
        # свойства книги — created выдачи, не «сейчас»: save_workbook openpyxl
        # перебивал dcterms:modified текущим моментом (два вызова по разные
        # стороны секунды — разные байты)
        want = "%04d-%02d-%02dT%02d:%02d:%02dZ" % stamp
        assert ("<dcterms:modified" in core and (">%s</dcterms:modified>" % want) in core), core
        assert (">%s</dcterms:created>" % want) in core
