# -*- coding: utf-8 -*-
"""Excel-издание анкеты CDPM: полнота книги и гейт содержимого.

Книга из образца «Полярис» несёт лист «Сводка» и лист на каждую
непустую часть анкеты, строку на каждый вопрос со всеми колонками и
трассировкой; гейт зелёный на честной книге и красный после порчи
ячейки идентификатора в файле (правка книги мимо модели ловится по
содержимому, не по байтам — байты xlsx нестабильны из-за меток
времени zip-архива). Сеть и языковые модели не нужны."""
from openpyxl import load_workbook

from cdpm import anketa_xlsx as ax
from cdpm import model as m
from cdpm.gates import gate_anketa_xlsx
from cdpm.sample import sample_model


def _book(tmp_path, doc=None):
    """Книга из образца во временном каталоге."""
    doc = doc or sample_model()
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    return doc, path


def test_book_has_summary_and_sheet_per_nonempty_part(tmp_path):
    _, path = _book(tmp_path)
    wb = load_workbook(str(path))
    assert wb.sheetnames == [
        "Сводка", "Часть 1—противоречия", "Часть 2—неоднозначности",
        "Часть 3а—постановка пейринга",
        "Часть 3б—постановка ростеринга"]
    # Режим 1: части 4 в модели нет — и листа её нет.
    assert ax.SHEET_TITLES["4"] not in wb.sheetnames


def test_part4_sheet_appears_with_part4_question(tmp_path):
    doc = sample_model()
    doc["meta"]["mode"] = 2
    doc["questions"].append({
        "id": "Q-4-TP-77", "part": "4",
        "title": "Несоответствие лимита лётного времени",
        "text": "Файл третьей стороны задаёт 90 часов, источник — "
                "80 часов. Какое значение верно?",
        "answer_format": {"type": "choice", "options": [
            {"key": "а", "text": "80 часов (источник)"},
            {"key": "б", "text": "90 часов (файл третьей стороны)"},
        ]},
        "rule_refs": ["R-S2-П-7-1"],
        "source_refs": [{"sid": "S2", "anchor": "П.7.1"}],
        "third_party_refs": [{"file": "TP-RULES.json",
                              "anchor": "r-77"}],
    })
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(str(path))
    assert "Часть 4—несоответствия" in wb.sheetnames
    assert wb["Часть 4—несоответствия"].cell(row=2, column=1).value \
        == "Q-4-TP-77"


def test_question_row_carries_all_columns_and_empty_answer(tmp_path):
    _, path = _book(tmp_path)
    wb = load_workbook(str(path))
    ws = wb["Часть 1—противоречия"]
    assert ws.cell(row=2, column=1).value == "Q-1-REST"
    q_cell = ws.cell(row=2, column=2).value
    assert "Минимальный отдых на базе: 10 или 12 часов" in q_cell
    assert "Какое требование действует?" in q_cell
    fmt = ws.cell(row=2, column=3).value
    assert "закрытый перечень — выбрать один вариант:" in fmt
    assert "а) 12 часов (колдоговор сильнее)" in fmt
    assert ws.cell(row=2, column=4).value == "R-S1-П-1-4, R-S2-П-5-1"
    # Источник — файл + якорь, разворачивается через doc["sources"].
    assert ws.cell(row=2, column=5).value == \
        "POLARIS-OM-A.docx, П.1.4; POLARIS-KD.docx, П.5.1"
    assert ws.cell(row=2, column=6).value is None
    # Лишних строк нет: под последним вопросом части — пусто.
    assert ws.cell(row=3, column=1).value is None


def test_header_style_freeze_and_autofilter(tmp_path):
    _, path = _book(tmp_path)
    wb = load_workbook(str(path))
    ws = wb["Часть 3а—постановка пейринга"]
    assert tuple(c.value for c in ws[1]) == ax.HEADERS
    assert all(c.font.bold for c in ws[1])
    assert all(c.fill.fill_type == "solid" for c in ws[1])
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:F3"   # шапка + два вопроса 3а


def test_summary_customer_counts_and_instruction(tmp_path):
    doc, path = _book(tmp_path)
    wb = load_workbook(str(path))
    ws = wb["Сводка"]
    got = {ws.cell(row=i, column=1).value:
           ws.cell(row=i, column=2).value
           for i in range(1, ws.max_row + 1)}
    assert got["Авиакомпания"] == doc["meta"]["customer"]
    assert got["Часть 1—противоречия"] == 1
    assert got["Часть 3а—постановка пейринга"] == 2
    assert got["Всего вопросов"] == len(doc["questions"])
    # Инструкция называет колонку ответа и правило буквы варианта.
    texts = "\n".join(str(k) for k in got if k)
    assert "ОТВЕТ АВИАКОМПАНИИ" in texts
    assert "букву выбранного варианта" in texts


def test_gate_green_on_honest_book(tmp_path, capsys):
    doc, path = _book(tmp_path)
    model_path = tmp_path / "model.json"
    m.dump_model(doc, model_path)
    rc = gate_anketa_xlsx.run(str(path), str(model_path))
    out = capsys.readouterr().out
    assert rc == 0
    # У-05: гейт печатает, сколько объектов рассмотрел.
    assert "рассмотрено: листов 5, вопросов 5, ссылок 12" in out
    assert "самотест пройден" in out
    assert "ВЕРДИКТ: книга анкеты соответствует модели" in out


def test_gate_red_on_id_tampered_in_file(tmp_path, capsys):
    doc, path = _book(tmp_path)
    model_path = tmp_path / "model.json"
    m.dump_model(doc, model_path)
    wb = load_workbook(str(path))
    wb["Часть 2—неоднозначности"].cell(row=2, column=1).value = \
        "Q-ЧУЖОЙ"
    wb.save(str(path))
    rc = gate_anketa_xlsx.run(str(path), str(model_path))
    out = capsys.readouterr().out
    assert rc == 1
    assert "НАРУШЕНИЕ" in out
    assert "Q-ЧУЖОЙ" in out
    assert "ВЕРДИКТ: КРАСНЫЙ" in out


def test_cli_writes_book(tmp_path, capsys):
    doc = sample_model()
    model_path = tmp_path / "model.json"
    m.dump_model(doc, model_path)
    out_path = tmp_path / "kniga" / "anketa.xlsx"
    rc = ax.main([str(model_path), "--out", str(out_path)])
    assert rc == 0
    out = capsys.readouterr().out
    assert out_path.is_file()
    # Полные пути в выводе — рабочее соглашение владельца (У-10).
    assert str(out_path.resolve()) in out
    wb = load_workbook(str(out_path))
    assert wb.sheetnames[0] == "Сводка"
