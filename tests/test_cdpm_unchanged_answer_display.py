"""An unchanged result must be visible without a worker-authored rule note."""
import copy
import json

import pytest
from openpyxl import load_workbook

from cdpm import answers, anketa_xlsx as ax, model as m, render
from cdpm.gates import gate_anketa_xlsx
from test_cdpm_a02_worker_row_results import _pending, _result_reply


BASIS = "Второе правило без изменений."
NOTICE = "R-H2 — признано неизменным по ответу: «" + BASIS + "»"


def _completed(*, unchanged, note=False):
    doc, q = _pending()
    q["answer"]["raw"] = "Первое правило подтверждено. " + BASIS
    patches = []
    for row in (doc["rules"][0], doc["rules"][-1]):
        if unchanged and row["id"] == "R-H2":
            patches.append({"id": row["id"], "unchanged": {"basis": BASIS}})
        else:
            patch = copy.deepcopy(row)
            patch.pop("blockers")
            if note:
                patch["note"] = "Заметка обработчика сохранена"
            patches.append(patch)
    assert answers.apply_worker_patches(doc, _result_reply(q, patches)) == {
        "worker_applied": 1, "worker_rejected": []}
    assert q["answer"]["status"] == "применён"
    assert not any(q["id"] in r.get("blockers", []) for r in doc["rules"])
    assert m.validate_schema(doc) + m.validate_model(doc) == []
    return doc, q


@pytest.mark.parametrize("surface", ["anketa", "accepted_xlsx"])
@pytest.mark.parametrize("note", [False, True])
def test_unchanged_result_names_row_and_basis_independently_of_rule_note(surface, note):
    changed, _ = _completed(unchanged=False, note=note)
    unchanged, _ = _completed(unchanged=True, note=note)
    before = copy.deepcopy(unchanged)
    def shown(doc):
        return render.render_anketa(doc) if surface == "anketa" else json.dumps(
            ax.answered_rows(doc), ensure_ascii=False)
    assert NOTICE in shown(unchanged)
    assert "признано неизменным по ответу" not in shown(changed)
    assert unchanged == before
    if note:
        assert "Заметка обработчика сохранена" in render.render_svod(unchanged)


def test_actual_book_delivers_unchanged_basis_without_changing_answer_status(tmp_path):
    doc, q = _completed(unchanged=True)
    before = copy.deepcopy(doc)
    path = tmp_path / "unchanged.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path)
    try:
        row = next(r for r in wb[ax.ANSWERED_TITLE].iter_rows(min_row=2, values_only=True)
                   if r[0] == q["id"])
        assert NOTICE in row[2] and row[3] == "применён" and row[4] == 2
    finally:
        wb.close()
    assert gate_anketa_xlsx.check_workbook(doc, path)[0] == []
    assert doc == before


@pytest.mark.parametrize("surface", ["anketa", "svod", "xlsx"])
def test_unchanged_basis_also_reaches_partial_answer_surfaces(surface):
    doc, q = _pending()
    q["answer"]["raw"] = "Первое пока уточняется. " + BASIS
    result = answers.apply_worker_patches(doc, _result_reply(q, [
        {"id": "R-H2", "unchanged": {"basis": BASIS}}]))
    assert result["worker_partial"] == 1 and result["worker_rejected"] == []
    before = copy.deepcopy(doc)
    text = (render.render_anketa(doc) if surface == "anketa" else
            render.render_svod(doc) if surface == "svod" else
            ax.published_question_row(doc, q, ax._files_by_sid(doc))[1])
    assert NOTICE in text
    assert "Обработанные строки: R-H2." in text
    assert "Остались без результата обработчика: R-H." in text
    assert doc == before
