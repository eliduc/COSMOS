# -*- coding: utf-8 -*-
"""Замки чтения источников при заморозке (аудит 2026-09-04, пункт 2,
план A2.1–A2.3): три воспроизводящих документа, каждый — дефект,
который старый код проходил молча.

- docx: абзац → таблица → абзац; раздельные обходы абзацев и таблиц
  уносили таблицу в конец списка клауз, порядок оригинала терялся;
- xlsx: книга, сохранённая без пересчёта формул (формула без
  сохранённого значения), теряла значение ячейки молча — теперь
  честный отказ с адресом ячейки, а пустые ячейки отказа не дают;
- pdf: страница с текстовым слоем И картинкой (картинка-таблица)
  целиком шла текстовым путём — теперь страница помечена `mixed_pages`,
  рисунок при канале зрения распознаётся с якорем «стр.P, рис.K».
Зрительная модель везде подменная — живых вызовов тесты не делают."""
import json
import re
import zipfile
from pathlib import Path

import pytest

from cdpm import freeze

fitz = pytest.importorskip("fitz")


# ---------------------------------------------------------------------------
# A2.1 docx: порядок тела документа
# ---------------------------------------------------------------------------

def _docx_para_table_para(path):
    import docx
    d = docx.Document()
    d.add_paragraph("1.1. Нормы экипажа по направлениям — в таблице ниже.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Север"
    t.cell(0, 1).text = "2 пилота"
    d.add_paragraph("Исключение к таблице выше: зимой 3 пилота.")
    d.save(str(path))


def test_docx_table_keeps_its_place_between_paragraphs(tmp_path):
    p = tmp_path / "PORYADOK.docx"
    _docx_para_table_para(p)
    anchors = [c["anchor"] for c in freeze.extract_file(p)["clauses"]]
    # Порядок клауз = порядок тела документа: таблица между абзацами
    # («абз.2» — второй непустой абзац, счёт абзацев прежний).
    assert anchors == ["п.1.1",
                       "табл.1, строка 1, кол 1",
                       "табл.1, строка 1, кол 2",
                       "абз.2"]


def test_docx_anchors_unchanged_by_body_walk(tmp_path):
    """Якоря прежние: «абз.N» считает непустые абзацы, «табл.T» —
    таблицы; таблицы в счёт абзацев не входят (комплекты H0/H0-К и
    существующие тесты адресуют клаузы по этим якорям)."""
    import docx
    p = tmp_path / "YAKORYA.docx"
    d = docx.Document()
    d.add_paragraph("Преамбула.")
    d.add_paragraph("")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "ячейка"
    d.add_paragraph("Второй абзац.")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "ячейка 2"
    d.save(str(p))
    anchors = [c["anchor"] for c in freeze.extract_file(p)["clauses"]]
    assert anchors == ["абз.1", "табл.1, строка 1, кол 1", "абз.2",
                       "табл.2, строка 1, кол 1"]


# ---------------------------------------------------------------------------
# A2.2 xlsx: формула без сохранённого значения
# ---------------------------------------------------------------------------

def _xlsx_uncalculated(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Нормы"
    ws["A1"] = "Норма отдыха, ч"
    ws["B1"] = "=10+2"          # openpyxl сохраняет формулу БЕЗ значения
    ws["A3"] = "Пустая строка 2 — не дефект"
    wb.save(str(path))


def test_xlsx_formula_without_cached_value_is_refused(tmp_path):
    p = tmp_path / "BEZ-PERESCHETA.xlsx"
    _xlsx_uncalculated(p)
    with pytest.raises(freeze.FreezeError) as e:
        freeze.extract_file(p)
    msg = str(e.value)
    assert "без пересчёта формул" in msg
    assert "лист «Нормы», ячейка B1" in msg
    assert "BEZ-PERESCHETA.xlsx" in msg


def test_xlsx_formula_with_cached_value_and_empty_cells_pass(tmp_path):
    """Отказ — только за формулу без значения: та же книга с
    записанным значением формулы (как после пересчёта в Excel)
    замораживается, значение попадает в клаузу; пустые ячейки и
    пустые строки отказа не дают."""
    src = tmp_path / "src.xlsx"
    _xlsx_uncalculated(src)
    dst = tmp_path / "PERESCHITANA.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                assert "<f>10+2</f>" in xml
                xml = re.sub(r"<f>10\+2</f>(<v></v>)?",
                             "<f>10+2</f><v>12</v>", xml)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    clauses = freeze.extract_file(dst)["clauses"]
    by_anchor = {c["anchor"]: c["text"] for c in clauses}
    assert by_anchor["лист «Нормы», строка 1"] == "Норма отдыха, ч | 12"
    assert "лист «Нормы», строка 3" in by_anchor
    assert "лист «Нормы», строка 2" not in by_anchor


def test_xlsx_string_formula_with_empty_cached_value_passes(tmp_path):
    """Ревизия 2026-09-04: формула с сохранённым результатом-пустой
    строкой (Excel пишет <c t="str"><f>IF(...)</f><v></v></c> —
    типичный приём журналов и таблиц норм) — книга ПЕРЕСЧИТАНА,
    отказа нет; строка с пустым результатом выходит как пустая
    ячейка. Непересчитанной считается только ячейка без
    типизированного кэша."""
    src = tmp_path / "src.xlsx"
    _xlsx_uncalculated(src)
    dst = tmp_path / "PUSTAYA-STROKA.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                new, n = re.subn(
                    r'<c r="B1"[^>]*>\s*<f>10\+2</f>(?:<v\s*/>|<v></v>)?\s*</c>',
                    '<c r="B1" t="str"><f>IF(1=1,"","x")</f><v></v></c>',
                    xml)
                assert n == 1, xml
                data = new.encode("utf-8")
            zout.writestr(item, data)
    clauses = freeze.extract_file(dst)["clauses"]       # отказа нет
    by_anchor = {c["anchor"]: c["text"] for c in clauses}
    assert by_anchor["лист «Нормы», строка 1"] == "Норма отдыха, ч"
    assert "лист «Нормы», строка 3" in by_anchor
    # Отрицательный контроль остаётся: та же книга без типизированного
    # кэша (исходник openpyxl) по-прежнему отказ.
    with pytest.raises(freeze.FreezeError):
        freeze.extract_file(src)


# ---------------------------------------------------------------------------
# A2.3 pdf: смешанная страница (текстовый слой + рисунок)
# ---------------------------------------------------------------------------

_FIG_PARAS = ["Направление | Норма экипажа", "Север | 2 пилота"]


def _fake_vision(reply=json.dumps({"paragraphs": _FIG_PARAS},
                                  ensure_ascii=False)):
    def fn(prompt, image_path):
        fn.calls += 1
        fn.paths.append(str(image_path))
        return reply
    fn.calls = 0
    fn.paths = []
    return fn


def _image_bytes():
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 60))
    pix.clear_with(200)
    return pix.tobytes("png")


def _mixed_pdf(path, text_over_image=False):
    """Страница: текстовый слой длиннее PDF_SCAN_TEXT_MIN + картинка.
    `text_over_image` — текстовый слой поверх картинки (скан после OCR):
    картинка покрыта, страница не смешанная."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "3.2. Нормы экипажа по направлениям — "
                               "в таблице-картинке ниже.", fontsize=11)
    img = fitz.Rect(72, 150, 400, 350)
    page.insert_image(img, stream=_image_bytes())
    if text_over_image:
        page.insert_text((80, 200), "Направление Норма экипажа: Север "
                                    "два пилота (распознанный слой)",
                         fontsize=11)
    doc.save(str(path))
    doc.close()


def test_pdf_mixed_page_figure_goes_to_vision_with_figure_anchor(tmp_path):
    p = tmp_path / "SMESHANNAYA.pdf"
    _mixed_pdf(p)
    fake = _fake_vision()
    got = freeze.extract_file(p, vision_llm=fake)
    assert fake.calls == 1                      # рисунок ушёл зрению
    assert got["mixed_pages"] == [1]
    assert got["mixed_recognized"] is True
    assert got["recognized"] == "image"
    assert "recognized_pages" not in got        # скан-страниц нет
    anchors = [c["anchor"] for c in got["clauses"]]
    assert anchors == ["п.3.2", "стр.1, рис.1, абз.1", "стр.1, рис.1, абз.2"]
    assert got["clauses"][1]["text"] == _FIG_PARAS[0]


def test_pdf_mixed_page_without_vision_marked_not_refused(tmp_path):
    p = tmp_path / "SMESHANNAYA2.pdf"
    _mixed_pdf(p)
    got = freeze.extract_file(p)
    assert got["mixed_pages"] == [1]
    assert got["mixed_recognized"] is False
    assert "recognized" not in got              # ничего не распознано
    assert [c["anchor"] for c in got["clauses"]] == ["п.3.2"]


def test_pdf_image_covered_by_text_layer_is_not_mixed(tmp_path):
    """Картинка с текстовым слоем поверх (скан после OCR) покрыта —
    зрению не уходит, страница не смешанная: иначе каждый OCR-скан
    распознавался бы дважды."""
    p = tmp_path / "OCR.pdf"
    _mixed_pdf(p, text_over_image=True)
    fake = _fake_vision()
    got = freeze.extract_file(p, vision_llm=fake)
    assert fake.calls == 0
    assert "mixed_pages" not in got


def test_pdf_figure_without_text_is_not_a_refusal(tmp_path):
    """Эмблема или схема без текста: пустой список абзацев от модели —
    допустимый ответ для рисунка, не провал файла, но со следом
    (ревизия 2026-09-04): `mixed_empty_figures` в результате, записи
    манифеста и машинной формы, строка проекции, предупреждение
    статуса — ленивый ответ модели на картинку-таблицу не теряется
    молча."""
    from cdpm.gates import gate_freeze
    p = tmp_path / "EMBLEMA.pdf"
    _mixed_pdf(p)
    fake = _fake_vision(json.dumps({"paragraphs": []}))
    got = freeze.extract_file(p, vision_llm=fake)
    assert fake.calls == 1
    assert got["mixed_pages"] == [1]
    assert got["mixed_empty_figures"] == ["стр.1, рис.1"]
    assert [c["anchor"] for c in got["clauses"]] == ["п.3.2"]
    # След доходит до комплекта: манифест, машинная форма, проекция.
    out = tmp_path / "kit"
    manifest = freeze.freeze([p], out,
                             vision_llm=_fake_vision(
                                 json.dumps({"paragraphs": []})))
    entry = manifest["files"][0]
    assert entry["mixed_empty_figures"] == ["стр.1, рис.1"]
    machine = json.loads((out / "clauses" / "S1.json").read_text(
        encoding="utf-8"))
    assert machine["mixed_empty_figures"] == ["стр.1, рис.1"]
    proj = (out / "projections" / f"S1-{p.stem}.md").read_text(
        encoding="utf-8")
    assert ("Рисунки стр.1, рис.1: зрительная модель текста не нашла"
            in proj)
    assert gate_freeze.check_kit(out, [p])[0] == []
    assert freeze.mixed_warnings(manifest) == [
        "рисунки, на которых зрительная модель текста не нашла: "
        "S1 стр.1, рис.1 — при сомнении сверить с оригиналом"]
    # Положительный контроль: рисунок с текстом следа не оставляет.
    got2 = freeze.extract_file(p, vision_llm=_fake_vision())
    assert "mixed_empty_figures" not in got2


def test_pdf_source_marks_tampered_in_manifest_fail_gate(tmp_path):
    """Гейт сверяет признаки источника (recognized, mixed_pages,
    mixed_empty_figures…) между манифестом и машинной формой: стёртый
    в манифесте след — правка мимо заморозки, красный гейт."""
    from cdpm.gates import gate_freeze
    p = tmp_path / "EMBLEMA2.pdf"
    _mixed_pdf(p)
    out = tmp_path / "kit"
    freeze.freeze([p], out, vision_llm=_fake_vision(
        json.dumps({"paragraphs": []})))
    assert gate_freeze.check_kit(out, [p])[0] == []
    m_path = out / "manifest.json"
    manifest = json.loads(m_path.read_text(encoding="utf-8"))
    del manifest["files"][0]["mixed_empty_figures"]
    m_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2)
                      + "\n", encoding="utf-8")
    problems = gate_freeze.check_kit(out, [p])[0]
    assert any("mixed_empty_figures" in x and "не совпадает с манифестом"
               in x for x in problems), problems


def _vector_table_pdf(path, rows=2, cols=4, diagram=False):
    """Текстовая страница с векторной таблицей: ячейки — заливка с
    рамкой (draw_rect), значения короткие («1», «180»), как в экспорте
    Word/LibreOffice; `diagram` — плюс схема из многих штрихов без
    текста (одна связная область ≥ PDF_FIGURE_MIN_ITEMS элементов)."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "3.2. Нормы экипажа по направлениям — "
                               "в таблице ниже.", fontsize=11)
    values = [["Напр.", "Норма", "Зима", "Лето"],
              ["Север", "2", "3", "180"]]
    for r in range(rows):
        for c in range(cols):
            cell = fitz.Rect(72 + c * 80, 150 + r * 30,
                             72 + (c + 1) * 80, 150 + (r + 1) * 30)
            page.draw_rect(cell, color=(0, 0, 0), fill=(0.9, 0.9, 0.9),
                           width=0.5)
            page.insert_text((cell.x0 + 5, cell.y0 + 20),
                             values[r % 2][c % 4], fontsize=10)
    if diagram:
        sh = page.new_shape()
        for i in range(10):
            sh.draw_line((72 + i * 20, 400), (72 + i * 20 + 15, 480))
        sh.draw_circle((300, 440), 30)
        sh.finish(color=(0, 0, 1), width=1)
        sh.commit()
    doc.save(str(path))
    doc.close()


def test_pdf_vector_table_is_not_a_figure(tmp_path):
    """Отрицательный контроль (ревизия 2026-09-04): векторная таблица
    из заливок и рамок ячеек с короткими значениями — не рисунок;
    страница не смешанная, зрение не зовётся, клаузы — только
    текстовый слой."""
    p = tmp_path / "TABLICA.pdf"
    _vector_table_pdf(p)
    doc = fitz.open(str(p))
    try:
        assert len(doc[0].get_drawings()) == 8      # по ячейке
        assert freeze._pdf_figures(doc[0]) == []
    finally:
        doc.close()
    fake = _fake_vision()
    got = freeze.extract_file(p, vision_llm=fake)
    assert fake.calls == 0
    assert "mixed_pages" not in got
    assert "recognized" not in got
    assert got["clauses"][0]["anchor"] == "п.3.2"


def test_pdf_vector_diagram_is_a_figure(tmp_path):
    """Положительный контроль: схема из многих штрихов без текста на
    той же странице — рисунок, страница смешанная, зрение зовётся
    один раз (таблица рядом рисунком не считается)."""
    p = tmp_path / "SHEMA.pdf"
    _vector_table_pdf(p, diagram=True)
    fake = _fake_vision()
    got = freeze.extract_file(p, vision_llm=fake)
    assert fake.calls == 1
    assert got["mixed_pages"] == [1]
    assert got["mixed_recognized"] is True
    assert [c["anchor"] for c in got["clauses"]][-2:] == \
        ["стр.1, рис.1, абз.1", "стр.1, рис.1, абз.2"]


_REPORT_DIR = Path(__file__).resolve().parent.parent / "final-report"


@pytest.mark.parametrize("name,expected_mixed", [
    # Word-экспорт с таблицами: смешанные — только страницы с
    # растровыми диаграммами (замечание ревизии: было 10 из 14).
    ("FINAL-REPORT-RU.pdf", [4, 5, 6, 8, 9, 10]),
    # 15 страниц таблиц: было 209 ложных рисунков, теперь ни одного.
    ("PROFESSIONAL-REPORT-RU.pdf", []),
])
def test_pdf_real_reports_have_no_false_figures(name, expected_mixed):
    p = _REPORT_DIR / name
    if not p.is_file():
        pytest.skip(f"{p} отсутствует")
    got = freeze.extract_file(p)                      # без зрения
    assert got.get("mixed_pages", []) == expected_mixed
    if expected_mixed:
        assert got["mixed_recognized"] is False


def test_pdf_mixed_in_freeze_kit_manifest_projection_gate(tmp_path):
    from cdpm.gates import gate_freeze
    src = tmp_path / "src"
    src.mkdir()
    p = src / "PRAVILA-SMESH.pdf"
    _mixed_pdf(p)
    out = tmp_path / "kit"
    manifest = freeze.freeze([p], out, vision_llm=_fake_vision())
    entry = manifest["files"][0]
    assert entry["mixed_pages"] == [1]
    assert entry["mixed_recognized"] is True
    assert entry["recognized"] == "image"
    assert entry["clause_count"] == 3
    proj = (out / "projections" / f"S1-{p.stem}.md").read_text(
        encoding="utf-8")
    assert "Страницы 1 смешанные" in proj
    assert "[стр.1, рис.1, абз.1]" in proj
    problems, counters = gate_freeze.check_kit(out, [p])
    assert problems == []
    # Без канала зрения: пометка в манифесте и проекции, гейт зелёный
    # (повторное извлечение детерминировано).
    out2 = tmp_path / "kit2"
    manifest2 = freeze.freeze([p], out2)
    entry2 = manifest2["files"][0]
    assert entry2["mixed_pages"] == [1]
    assert entry2["mixed_recognized"] is False
    assert "recognized" not in entry2
    proj2 = (out2 / "projections" / f"S1-{p.stem}.md").read_text(
        encoding="utf-8")
    assert "НЕ РАСПОЗНАНЫ" in proj2
    assert gate_freeze.check_kit(out2, [p])[0] == []
    # Предупреждение статуса прогона — по манифесту (панель обязана
    # сказать о сверке, не только манифест).
    assert freeze.mixed_warnings(manifest2) == [
        "смешанные страницы без распознавания: S1 стр. 1 — сверка с "
        "оригиналом обязательна"]
    assert freeze.mixed_warnings(manifest) == []      # распознано
