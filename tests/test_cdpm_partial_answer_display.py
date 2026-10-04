"""Partial processing is visible without changing the answer or its effects."""
import copy
import json

import pytest
from openpyxl import load_workbook

from cdpm import answers, anketa_xlsx as ax, model as m, question_editions, render
from cdpm.gates import gate_answers, gate_anketa_xlsx
from test_cdpm_a02_worker_row_results import _pending, _worker


def _partial(*, published=False, extra_blocker=False):
    doc, q = _pending()
    if extra_blocker:
        other = copy.deepcopy(q)
        other.update(id="Q-OTHER", title="Другая проверка первой строки",
                     text="Уточните отдельное основание первого правила",
                     rule_refs=["R-H"])
        other.pop("answer")
        doc["questions"].append(other)
        doc["rules"][0]["blockers"].append(other["id"])
    if published:
        question_editions.freeze(q)
    assert m.validate_schema(doc) + m.validate_model(doc) == []
    result = answers.apply_worker_patches(doc, _worker(q, [doc["rules"][0]]))
    assert result["worker_rejected"] == [] and result["worker_partial"] == 1
    return doc, q


def _surface(doc, q, surface):
    if surface == "anketa":
        return render.render_anketa(doc)
    if surface == "svod":
        return render.render_svod(doc)
    return ax.published_question_row(doc, q, ax._files_by_sid(doc))[1]


@pytest.mark.parametrize("surface", ["anketa", "svod", "xlsx"])
@pytest.mark.parametrize("published", [False, True])
def test_partial_surfaces_show_saved_answer_and_distinct_row_outcomes(surface, published):
    doc, q = _partial(published=published)
    before = copy.deepcopy(doc)
    fingerprint = ax.question_fingerprint(q)
    text = _surface(doc, q, surface)
    assert "Ответ применён частично" in text
    assert q["answer"]["raw"] in text
    assert "Обработанные строки: R-H." in text
    assert "Остались без результата обработчика: R-H2." in text
    assert "Повторно заполнять поле ответа не нужно" in text
    assert "пересмотр" in text
    assert doc == before and ax.question_fingerprint(q) == fingerprint


def test_partial_anketa_does_not_invite_a_duplicate_answer():
    doc, q = _partial()
    text = render.render_anketa(doc)
    block = text.split("ВОПРОС " + q["id"] + ":", 1)[1].split("\n### ", 1)[0]
    assert "\nОтвет:\n" not in block


@pytest.mark.parametrize("surface", ["anketa", "svod", "xlsx"])
def test_full_result_has_no_partial_processing_noise(surface):
    doc, q = _pending()
    result = answers.apply_worker_patches(doc, _worker(q, [doc["rules"][0], doc["rules"][-1]]))
    assert result == {"worker_applied": 1, "worker_rejected": []}
    before = copy.deepcopy(doc)
    text = _surface(doc, q, surface)
    for marker in ("Ответ применён частично", "Обработанные строки:",
                   "Остались без результата обработчика:",
                   "Повторно заполнять поле ответа не нужно"):
        assert marker not in text
    assert doc == before


@pytest.mark.parametrize("published", [False, True])
def test_another_question_blocker_does_not_mark_processed_row_pending(published):
    doc, q = _partial(published=published, extra_blocker=True)
    text = _surface(doc, q, "xlsx")
    assert "Обработанные строки: R-H." in text
    assert "Остались без результата обработчика: R-H2." in text
    assert doc["rules"][0]["blockers"] == ["Q-OTHER"]


def test_gate_names_pending_rows_without_marking_processed_rows_pending(tmp_path):
    doc, q = _partial()
    model, summary = tmp_path / "model.json", tmp_path / "summary.json"
    model.write_text(json.dumps(doc, ensure_ascii=False), encoding="utf-8")
    summary.write_text(json.dumps({"revision": 2}), encoding="utf-8")
    issues, counts = gate_answers.check(model, summary)
    issue = next(p for p in issues if p.startswith(q["id"] + ": ответ применён частично"))
    assert issue.endswith("без результата обработчика: R-H2")
    # The fixture already contains two applied goal answers; this partial
    # answer must not become a third applied answer.
    assert counts["partial"] == 1 and counts["applied"] == 2


def test_partial_book_contains_progress_but_answer_cell_is_not_fabricated(tmp_path):
    doc, q = _partial(published=True)
    before = copy.deepcopy(doc)
    path = tmp_path / "partial.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path)
    try:
        sheet = wb[ax.SHEET_TITLES[q["part"]]]
        row = next(row for row in sheet.iter_rows(min_row=2, values_only=True)
                   if row[0] == q["id"])
        assert "Обработанные строки: R-H." in row[1]
        assert "Остались без результата обработчика: R-H2." in row[1]
        assert row[ax.ANSWER_COL - 1] is None
        assert row[ax.FINGERPRINT_COL - 1] == ax.question_fingerprint(q)
    finally:
        wb.close()
    assert gate_anketa_xlsx.check_workbook(doc, path)[0] == []
    assert doc == before


def test_second_worker_pass_removes_partial_notice_and_keeps_first_result():
    doc, q = _partial(published=True)
    first = copy.deepcopy(doc["rules"][0])
    assert "Обработанные строки: R-H." in _surface(doc, q, "xlsx")
    result = answers.apply_worker_patches(doc, _worker(q, [doc["rules"][-1]]))
    assert result == {"worker_applied": 1, "worker_rejected": []}
    assert doc["rules"][0] == first
    assert not m.is_open_question(q, doc=doc)
    assert "Ответ применён частично" not in render.render_anketa(doc)


def test_unanswered_question_keeps_answer_slot_and_original_xlsx_row():
    doc, q = _pending()
    q.pop("answer")
    text = render.render_anketa(doc)
    assert "\nОтвет:\n" in text
    assert "Ответ применён частично" not in text
    assert ax.published_question_row(doc, q, ax._files_by_sid(doc))[1] == q["title"] + "\n" + q["text"]


def test_progress_is_cumulative_over_two_partial_passes():
    doc, q = _pending()
    third = copy.deepcopy(doc["rules"][-1])
    third.update(id="R-H3", title="Третье правило")
    doc["rules"].append(third)
    q["rule_refs"].append(third["id"])
    for rid in ("R-H", "R-H2"):
        row = next(r for r in doc["rules"] if r["id"] == rid)
        result = answers.apply_worker_patches(doc, _worker(q, [row]))
        assert result["worker_partial"] == 1 and result["worker_rejected"] == []
    text = _surface(doc, q, "xlsx")
    assert "Обработанные строки: R-H, R-H2." in text
    assert "Остались без результата обработчика: R-H3." in text


def test_progress_does_not_expand_published_scope_to_outside_blocked_row():
    doc, q = _partial(published=True)
    outside = copy.deepcopy(doc["rules"][-1])
    outside.update(id="R-OUTSIDE", title="Чужая строка")
    doc["rules"].append(outside)
    text = _surface(doc, q, "xlsx")
    assert "R-OUTSIDE" not in text
    assert "Остались без результата обработчика: R-H2." in text


@pytest.mark.parametrize("surface", ["anketa", "svod", "xlsx"])
def test_saved_answer_is_not_silently_truncated(surface):
    doc, q = _partial()
    q["answer"]["raw"] = "Длинный ответ | первая строка\n" + "Содержание. " * 100 + "КОНЕЦ ОТВЕТА"
    before = copy.deepcopy(doc)
    assert q["answer"]["raw"] in _surface(doc, q, surface)
    assert doc == before
