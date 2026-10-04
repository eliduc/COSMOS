# -*- coding: utf-8 -*-
"""K04-барьер: пересмотр как явное состояние (контракт П0 §2.4 (а)/(б),
ранняя часть волны 17а, 12.09.2026).

Закон. Пока ячейка посадки несёт признак открытого пересмотра
(`pending_revision` = id вопроса-пересмотра), прежнее значение решением
авиакомпании НЕ является: (а) готовность затронутого режима ниже «в»
независимо от присутствия старого значения, матрица называет
семью/параметр/вопрос словами, выпуск отказывает кодом
`family-incomplete` со словами «пересмотр открыт», гейт производителя
краснеет ТЕМ ЖЕ расчётом по `model.json` пакета; (б) синхронизация
(вопросы полноты параметров семьи, канон после цикла, ручка вопросов
условий, извлечение) НЕ снимает открытый пересмотр кодом: единственный
кодовый исход — «вопрос больше не требуется», когда исчезли ВСЕ
носители и эффекты параметра; иначе пересмотр закрывает только человек
ответом.

Ожидания заданы независимо от реализации: доля фактов (ячейка, семья,
параметр, вопрос) выписана в тестах поимённо, причинное условие дефекта
(«свод считает параметр заданным, значит вопрос лишний») проверяется
отдельным утверждением — барьер держит пересмотр открытым ИМЕННО при
нём. Набор синтетический (`tmp_path`, самодельные своды), каталог
шаблонов — настоящий, канал переводчика SMT — подменный (автофикстура
`tests/conftest.py::_smt_gate_translator`): подменяется обращение к
языковой модели, отказы рубежей настоящие.

Не входит в раннюю часть (волна 17б): сверка-и-замена по отпечатку
прежнего значения (§2.4 (в)), статус «пересмотр: конфликт» (§2.4 (г)),
resolver `_reapply_values` по точной версии ответа (§2.4 (е)),
миграция старых ячеек (§2.5).
"""
import copy
import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest

from cdpm import answers as answers_mod
from cdpm import export_task_package as tp
from cdpm import family_params as fp
from cdpm import model as m
from cdpm import webui
from cdpm.gates import gate_task_package as gate
from cdpm.sample import sample_model

REV = m.registry_fingerprint()

#: Семья каталога с двумя параметрами и правило-носитель — предмет
#: закона; имена берутся из каталога, а не из набора данных.
FAMILY = "rest_floors"
PARAM = "min_rest_outstation"
OTHER_PARAM = "min_rest_home"
ROOT_Q = f"Q-3b-параметр-{FAMILY}-{PARAM}"
REV_Q = ROOT_Q + "-R"
NOTE = ("длительность в минутах — например «600 минут» или «10:00»; "
        "либо «не применяется»")

#: Ответ на пересмотр, ПРИМЕНЁННЫЙ обработчиком. Нужен там, где пробе
#: требуется ЗАКОННЫЙ пакет-опора: по закону происхождения A03-мин
#: (принят 12.09.2026) ячейка, помеченная ответом, выпускается только
#: если ДЕЙСТВУЮЩИЙ ответ цепочки применён. Свод, у которого пересмотр
#: ждёт обработчика, законным пакетом больше не является — и это
#: усиление, а не поломка: прежде такой свод выпускался, и прежнее
#: значение ехало в пакет без действующего решения авиакомпании.
#: Барьер K04 ловил этот случай ТОЛЬКО с признаком на ячейке; закон
#: происхождения ловит его и без признака. Для проб самого барьера K04
#: опора не нужна — там пакет и обязан отказать.
APPLIED_REVISION = {"raw": "10:00 (600 минут)", "status": "применён",
                    "revision": 3}


# ── своды ────────────────────────────────────────────────────────────

def _landed_rule(rid: str, params: dict, anchor: str = "п.2") -> dict:
    return {"id": rid, "title": "Полы отдыха",
            "statement": "отдых дома не менее 12:00, вне базы — по ответу",
            "zone": "ROSTERING", "kind": "HARD",
            "sources": [{"sid": "S1", "anchor": anchor}],
            "template_fit": {"category": "A", "registry_rev": REV,
                             "family": FAMILY, "params": params}}


def _retired_rule(rid: str = "R-F", anchor: str = "п.2") -> dict:
    """Строка, отменённая ответом авиакомпании: посадки у неё больше
    нет — носителем параметра семьи она не является."""
    return {"id": rid, "title": "Отменённое требование",
            "statement": "требование отменено ответом авиакомпании",
            "zone": "ROSTERING", "kind": "OUT",
            "note": "отменено ответом авиакомпании",
            "sources": [{"sid": "S1", "anchor": anchor}]}


def _questions(rule_ids: list[str], revision_answer: dict | None) -> list:
    """Корневой вопрос параметра (ответ пересмотрен) и вопрос-пересмотр."""
    rq = {"id": REV_Q, "part": "3b",
          "title": f"Пересмотр: недостающий параметр шаблона: {FAMILY}",
          "text": "Прежний ответ принят и применён — дайте новый ответ.",
          "answer_format": {"type": "value", "note": NOTE},
          "rule_refs": list(rule_ids),
          "source_refs": [{"sid": "S1", "anchor": "п.2"}],
          "revision_of": ROOT_Q}
    if revision_answer is not None:
        rq["answer"] = revision_answer
    return [
        {"id": ROOT_Q, "part": "3b",
         "title": f"Недостающий параметр шаблона: {FAMILY} — {PARAM}",
         "text": f"У шаблона «{FAMILY}» есть параметр «{PARAM}» — укажите "
                 f"значение либо ответьте «не применяется».",
         "answer_format": {"type": "value", "note": NOTE},
         "rule_refs": list(rule_ids),
         "source_refs": [{"sid": "S1", "anchor": "п.2"}],
         "answer": {"raw": "10:00 (600 минут)", "status": "пересмотрен",
                    "revision": 2, "revised_by": REV_Q}},
        rq,
    ]


def _doc(*, pending: bool, rules: list[dict] | None = None,
         revision_answer: dict | None = None) -> dict:
    """Свод уровня «в» с носителем семьи; параметр `PARAM` внесён
    ответом. `pending=True` — ячейка под открытым пересмотром.

    Ответ на пересмотр по умолчанию «ожидает обработчика»: вопрос при
    этом НЕ открыт (`model.is_open_question` — ложь) и строки не
    блокированы, поэтому готовность держит ровно признак ячейки, а не
    открытый вопрос — так факт (а) судится отдельно от прочих оснований.
    """
    cell = {"value": 600, "quote": "10:00 (600 минут)", "answer": ROOT_Q}
    if pending:
        cell["pending_revision"] = REV_Q
    if rules is None:
        rules = [_landed_rule("R-F", {
            OTHER_PARAM: {"value": 720, "quote": "дома не менее 12:00"},
            PARAM: cell})]
    if revision_answer is None:
        revision_answer = {"raw": "уточним после совещания",
                           "status": "ожидает обработчика", "revision": 3}
    return {
        "meta": {"customer": "Т", "mode": 1, "cdpm_version": "0.1",
                 "revision": 3},
        "sources": [{"sid": "S1", "file": "a.md", "sha256": "0" * 64}],
        "rules": rules + [
            {"id": "R-S", "title": "Ночёвки",
             "statement": "штраф за ночь вне базы: вес 10",
             "zone": "PAIRING", "kind": "SOFT",
             "sources": [{"sid": "S1", "anchor": "п.1"}]},
        ],
        "questions": _questions([r["id"] for r in rules], revision_answer) + [
            {"id": "Q-3a-GOAL", "part": "3a", "title": "Целевая",
             "text": "целевая пейринга", "answer_format": {"type": "registry"},
             "rule_refs": ["R-S"],
             "source_refs": [{"sid": "S1", "anchor": "п.1"}],
             "answer": {"raw": "таблица", "status": "применён",
                        "revision": 2, "goal_decline": False}},
            {"id": "Q-3b-GOAL", "part": "3b", "title": "Целевая",
             "text": "целевая ростеринга",
             "answer_format": {"type": "registry"}, "rule_refs": ["R-S"],
             "source_refs": [{"sid": "S1", "anchor": "п.1"}],
             "answer": {"raw": "любое допустимое решение",
                        "status": "применён", "revision": 2,
                        "goal_decline": True}},
        ],
        "objectives": [
            {"id": "O-1", "title": "Критерий оптимизации пейринга",
             "statement": "Минимум человеко-дней командировок; цена "
                          "рабочего дня (смены) — 100 условных единиц",
             "zone": "PAIRING", "direction": "minimize",
             "sources": [{"sid": "S1", "anchor": "п.1"}]},
        ],
    }


def _open_revision_question(rule_ids: list[str]) -> dict:
    """Открытый вопрос-пересмотр (ответа ещё нет) — состояние сразу
    после `answers.open_revision`."""
    return {"id": REV_Q, "part": "3b", "title": "Пересмотр параметра",
            "text": "Дайте новый ответ о параметре семьи.",
            "answer_format": {"type": "value", "note": NOTE},
            "rule_refs": list(rule_ids),
            "source_refs": [{"sid": "S1", "anchor": "п.2"}],
            "revision_of": ROOT_Q}


def _with_open_revision(doc: dict, rule_ids: list[str]) -> dict:
    """Тот же свод, но вопрос-пересмотр в нём ОТКРЫТ."""
    doc["questions"] = [q for q in doc["questions"] if q["id"] != REV_Q]         + [_open_revision_question(rule_ids)]
    return doc


def _run(tmp_path: Path, doc: dict, name: str = "run-9401") -> Path:
    run = tmp_path / name
    (run / "extract").mkdir(parents=True, exist_ok=True)
    (run / "extract" / "model.json").write_text(
        json.dumps(doc, ensure_ascii=False), encoding="utf-8")
    (run / "extract" / "svod.md").write_text("# свод\n", encoding="utf-8")
    return run


def _q(doc: dict, qid: str) -> dict:
    return next(x for x in doc["questions"] if x["id"] == qid)


def _cell(doc: dict, rid: str = "R-F", param: str = PARAM) -> dict:
    rule = next(r for r in doc["rules"] if r["id"] == rid)
    return rule["template_fit"]["params"][param]


# ── (1) барьер: ячейка под пересмотром держит режим ниже «в» ─────────

def test_pending_cell_holds_mode_and_release_and_gate_agree(tmp_path):
    """Ячейка со СТАРЫМ значением под открытым пересмотром: готовность
    ниже «в», ячейка матрицы своей задачи не готова, выпуск отказывает
    кодом `family-incomplete` словами «пересмотр открыт», гейт
    производителя краснеет тем же расчётом. Ожидание независимо: факт
    выписан поимённо (строка, семья, параметр, вопрос)."""
    doc = _doc(pending=True)
    assert m.validate_schema(doc) == [] and m.validate_model(doc) == []
    # Вопрос-пересмотр НЕ открыт и строки не блокированы — готовность
    # держит ровно признак ячейки.
    assert not m.is_open_question(_q(doc, REV_Q))
    assert not any(r.get("blockers") for r in doc["rules"])

    facts = m.pending_revision_facts(doc)
    assert [(f["rule"], f["family"], f["param"], f["question"], f["problem"])
            for f in facts] == [("R-F", FAMILY, PARAM, REV_Q, None)]

    ready = m.readiness(doc)
    assert ready["state"] != "в" and ready["rostering_ready"] is False
    assert ready["rostering_open"][m.PENDING_REVISION_TAG] == 1
    for word in (FAMILY, PARAM, REV_Q, "R-F"):
        assert word in ready["label"], word

    matrix = m.readiness_matrix(doc)
    assert matrix["rostering"]["feasibility-hard"]["ready"] is False
    assert m.PENDING_REVISION_TAG in \
        matrix["rostering"]["feasibility-hard"]["reasons"]
    assert tp.allowed_modes(doc)[0]["rostering"] == []

    run = _run(tmp_path, doc)
    with pytest.raises(tp.TaskPackageError) as e:
        tp.build_package(run)
    assert e.value.code == tp.FAMILY_INCOMPLETE
    # Перечень `items` — грамматики K01 («<семья>/<параметр>»): один
    # барьер — один читатель перечня; строки названы словами отказа.
    assert e.value.items == [f"{FAMILY}/{PARAM}"]
    for word in ("пересмотр открыт", FAMILY, PARAM, REV_Q, "R-F"):
        assert word in str(e.value), word

    # Записи барьера для интеграции с K01 (частичная передача):
    # те же поля, что у `readiness_barriers`.
    assert m.pending_revision_barriers(doc) == [{
        "code": tp.FAMILY_INCOMPLETE, "family": FAMILY, "param": PARAM,
        "item": f"{FAMILY}/{PARAM}", "question": REV_Q, "rules": ["R-F"],
        "zones": ["ROSTERING"], "hard": True,
        "words": m.pending_revision_barriers(doc)[0]["words"]}]
    words = m.pending_revision_barriers(doc)[0]["words"]
    for word in ("пересмотр открыт", FAMILY, PARAM, REV_Q, "R-F"):
        assert word in words, word


def test_ready_without_pending_is_the_positive_control(tmp_path):
    """Положительный контроль: тот же свод БЕЗ признака пересмотра —
    готовность «в», пакет выпускается, гейт зелёный. Исправление не
    подменяет пропуск дефекта запретом выпуска вообще."""
    doc = _doc(pending=False, revision_answer=APPLIED_REVISION)
    assert m.validate_schema(doc) == [] and m.validate_model(doc) == []
    assert m.pending_revision_facts(doc) == []
    ready = m.readiness(doc)
    assert ready["state"] == "в" and ready["rostering_open"] == {}
    assert m.PENDING_REVISION_TAG not in ready["label"]
    assert m.readiness_matrix(doc)["rostering"]["feasibility-hard"]["ready"]
    assert tp.allowed_modes(doc)[0]["rostering"] == [
        "feasibility-soft", "feasibility-hard"]

    run = _run(tmp_path, doc)
    pkg = tmp_path / "p.zip"
    pkg.write_bytes(tp.export_zip(run))
    assert gate.check(pkg)[0] == []


# ── (3) синхронизация не снимает открытый пересмотр ──────────────────

def _doc_after_conflicting_revision():
    """Настоящий причинный путь дефекта K04: ответ применён, посадку
    ПОЗЖЕ правит кто-то ещё, пересмотр открывается — откат по журналу
    невозможен, старое значение остаётся в ячейке. Ровно эта цепочка в
    находке аудита закрывалась синхронизацией прежним значением."""
    doc = sample_model()
    rule = next(x for x in doc["rules"] if x["id"] == "R-S1-П-2-3")
    rule["template_fit"] = {"category": "A", "registry_rev": REV,
                            "family": FAMILY,
                            "params": {OTHER_PARAM: {
                                "value": 720, "quote": "ранее 08:00"}}}
    assert fp.ensure_family_param_questions(doc) == [ROOT_Q]
    summary = answers_mod.apply_answers(
        doc, {ROOT_Q: "Отдых вне базы — 10:00 (600 минут)."})
    assert summary["applied"] == 1
    assert rule["template_fit"]["params"][PARAM]["value"] == 600
    # Чужая поздняя правка посадки — журнал следствий перекрыт.
    rule["template_fit"]["params"][OTHER_PARAM]["quote"] = "иная цитата клаузы"
    report = answers_mod.open_revision(doc, ROOT_Q)
    assert report["created"] == REV_Q and report["restored"] is False
    assert report["conflicts"] and report["pending_cells"] == [
        f"R-S1-П-2-3.{PARAM}"]
    assert rule["template_fit"]["params"][PARAM]["pending_revision"] == REV_Q
    assert m.validate_schema(doc) == [] and m.validate_model(doc) == []
    return doc, rule


def test_sync_does_not_close_open_revision_while_carrier_lives():
    """Причинное условие дефекта налицо: свод считает параметр
    ЗАДАННЫМ (старым значением) — по прежнему закону вопрос стал бы
    «лишним» и был бы снят кодом. Барьер держит пересмотр открытым:
    вопрос открыт, признак на месте, носитель заблокирован."""
    doc, rule = _doc_after_conflicting_revision()
    coverage = fp.family_coverage(doc)[FAMILY]
    assert PARAM in coverage["given"] and coverage["missing"] == []

    refusal = m.code_close_refusal(doc, _q(doc, REV_Q))
    assert refusal and "R-S1-П-2-3" in refusal
    assert m.CODE_CLOSE_OUTCOME in refusal

    assert fp.ensure_family_param_questions(doc) == []
    q = _q(doc, REV_Q)
    assert m.is_open_question(q) and "answer" not in q
    assert rule["template_fit"]["params"][PARAM]["pending_revision"] == REV_Q
    assert REV_Q in rule["blockers"]
    assert m.validate_model(doc) == []


def test_sync_canon_and_conditions_handler_keep_revision_open(tmp_path,
                                                              monkeypatch):
    """Те же слова на двух других путях синхронизации: канон итерации
    (`webui._sync_canon`) и ручка вопросов условий расчёта
    (`POST /api/runs/{id}/conditions-questions`). Проверяется файл
    прогона, а не только объект в памяти."""
    from fastapi.testclient import TestClient

    doc, rule = _doc_after_conflicting_revision()
    monkeypatch.setenv("CDPM_RUNS_DIR", str(tmp_path / "runs"))
    monkeypatch.setenv("CDPM_SETTINGS_FILE", str(tmp_path / "settings.json"))
    run = _run(tmp_path / "runs", doc, "run-9402")

    assert webui._sync_canon(run, doc, "канон итерации") == []
    on_disk = json.loads((run / "extract" / "model.json")
                         .read_text(encoding="utf-8"))
    assert m.is_open_question(_q(on_disk, REV_Q))
    assert _cell(on_disk, "R-S1-П-2-3")["pending_revision"] == REV_Q

    resp = TestClient(webui.app).post(
        "/api/runs/run-9402/conditions-questions")
    assert resp.status_code == 200, resp.text
    after = json.loads((run / "extract" / "model.json")
                       .read_text(encoding="utf-8"))
    assert m.is_open_question(_q(after, REV_Q))
    assert _cell(after, "R-S1-П-2-3")["pending_revision"] == REV_Q
    assert m.pending_revision_facts(after)[0]["problem"] is None


# ── (4)/(5) исход «вопрос больше не требуется» ───────────────────────

def test_all_carriers_gone_closes_by_code_with_named_outcome():
    """Все носители параметра правомерно исчезли (строка отменена,
    посадки нет): код снимает пересмотр ЕДИНСТВЕННЫМ исходом —
    «вопрос больше не требуется» — и называет его словами в записи
    вопроса; признака в ячейках не остаётся, готовность зелёная."""
    doc = _with_open_revision(_doc(pending=False, rules=[_retired_rule()]),
                              ["R-F"])
    assert m.validate_schema(doc) == [] and m.validate_model(doc) == []
    assert m.revision_carriers(doc, _q(doc, REV_Q)) == []
    assert m.code_close_refusal(doc, _q(doc, REV_Q)) is None

    assert fp.ensure_family_param_questions(doc) == []
    ans = _q(doc, REV_Q)["answer"]
    assert ans["status"] == "применён" and ans["closed_by"] == "code"
    assert ans["close_reason"] == m.CODE_CLOSE_OUTCOME
    assert m.pending_revision_facts(doc) == []
    assert m.readiness(doc)["state"] == "в"
    assert m.validate_model(doc) == [] and m.validate_schema(doc) == []


def test_partly_vanished_carrier_keeps_revision_open():
    """Исчез ОДИН носитель из двух — исход «вопрос больше не требуется»
    незаконен: второй носитель жив, пересмотр остаётся открытым."""
    rules = [_retired_rule("R-F"),
             _landed_rule("R-F2", {OTHER_PARAM: {
                 "value": 720, "quote": "дома не менее 12:00"}}, "п.3")]
    doc = _with_open_revision(_doc(pending=False, rules=rules),
                              ["R-F", "R-F2"])
    assert m.validate_schema(doc) == [] and m.validate_model(doc) == []
    assert m.revision_carriers(doc, _q(doc, REV_Q)) == ["R-F2"]
    refusal = m.code_close_refusal(doc, _q(doc, REV_Q))
    assert refusal and "R-F2" in refusal

    assert fp.ensure_family_param_questions(doc) == []
    q = _q(doc, REV_Q)
    assert m.is_open_question(q) and "answer" not in q
    assert REV_Q in next(r for r in doc["rules"]
                         if r["id"] == "R-F2")["blockers"]


def test_carrier_relanded_on_another_family_is_a_lawful_vanish():
    """Носитель ПЕРЕСАЖЕН на другую семью: параметра прежней семьи в
    своде больше нет нигде — носитель исчез правомерно, и исход
    «вопрос больше не требуется» законен. Прежде строка считалась
    носителем по одному факту посадки (семья не сверялась): пересмотр
    оставался открытым навсегда, строка — заблокированной, а ответ
    авиакомпании значением попадал в тупик «значение применить некуда»."""
    doc = _with_open_revision(_doc(pending=True), ["R-F"])
    rule = next(r for r in doc["rules"] if r["id"] == "R-F")
    rule["template_fit"] = {
        "category": "A", "registry_rev": REV, "family": "night_series",
        "params": {"max_consecutive_night_duties": {"value": 3,
                                                    "quote": "не более 3"},
                   "night_duty_start": {"value": 1320, "quote": "с 22:00"},
                   "night_duty_end": {"value": 300, "quote": "до 05:00"}}}
    rule["blockers"] = [REV_Q]
    assert m.validate_schema(doc) == [] and m.validate_model(doc) == []
    assert FAMILY not in fp.family_coverage(doc)   # параметра негде жить

    assert m.revision_carriers(doc, _q(doc, REV_Q)) == []
    assert m.code_close_refusal(doc, _q(doc, REV_Q)) is None
    assert fp.ensure_family_param_questions(doc) == []
    answer = _q(doc, REV_Q)["answer"]
    assert answer["closed_by"] == "code"
    assert answer["close_reason"] == m.CODE_CLOSE_OUTCOME
    assert "blockers" not in rule
    assert m.pending_revision_facts(doc) == []
    assert m.readiness(doc)["state"] == "в"
    assert m.validate_model(doc) == []


def test_relanded_carrier_of_another_question_is_still_a_carrier():
    """Обратная сторона того же закона: строка, пересаженная на ДРУГУЮ
    семью, носителем ЭТОЙ семьи быть перестала, но носителем СВОЕЙ
    новой семьи остаётся — сверка идёт по семье вопроса, а не по факту
    посадки вообще."""
    doc = _with_open_revision(_doc(pending=True), ["R-F"])
    other_root = "Q-3b-параметр-night_series-max_consecutive_night_duties"
    other_rev = other_root + "-R"
    rule = next(r for r in doc["rules"] if r["id"] == "R-F")
    rule["template_fit"]["family"] = "night_series"
    rule["template_fit"]["params"] = {
        "max_consecutive_night_duties": {"value": 3, "quote": "не более 3"},
        "night_duty_start": {"value": 1320, "quote": "с 22:00"},
        "night_duty_end": {"value": 300, "quote": "до 05:00"}}
    q = {"id": other_rev, "part": "3b", "title": "Пересмотр серии ночей",
         "text": "текст вопроса-пересмотра параметра другой семьи",
         "answer_format": {"type": "value", "note": "целое число"},
         "rule_refs": ["R-F"],
         "source_refs": [{"sid": "S1", "anchor": "п.2"}],
         "revision_of": other_root}
    assert m.revision_carriers(doc, q) == ["R-F"]
    assert m.code_close_refusal(doc, q)


def test_code_writer_refuses_with_words_not_silently():
    """Писатель кодового снятия не обходится: попытка снять открытый
    пересмотр при живом носителе — отказ словами, не тихая запись."""
    doc, _rule = _doc_after_conflicting_revision()
    q = _q(doc, REV_Q)
    with pytest.raises(m.CodeCloseRefused) as e:
        m.close_question_by_code(doc, q, raw="снят кодом", note="лишний")
    assert m.CODE_CLOSE_OUTCOME in str(e.value)
    assert "answer" not in q


# ── (6) ответ человека закрывает пересмотр обычным путём ────────────

def test_human_answer_closes_revision_and_clears_pending():
    """Человек ответил на пересмотр: ответ применяется обычным путём
    (без пометки кодового исхода), новое значение садится в ячейку,
    признак пересмотра снят, готовность больше им не держится."""
    doc, rule = _doc_after_conflicting_revision()
    summary = answers_mod.apply_answers(doc, {REV_Q: "9 часов"})
    assert summary["applied"] == 1
    cell = rule["template_fit"]["params"][PARAM]
    # Композиция K04 × A03-мин: ячейку пишет ОДИН писатель, и с
    # 12.09.2026 он кладёт в неё ещё и точное происхождение (`source`:
    # издание вопроса, отпечаток версии ответа, адресат следствия).
    # Прежнее ожидание точного равенства трём полям описывало форму
    # ячейки ДО принятого A03-мин; предмет этой пробы — что новое
    # значение село, признак снят и кодового исхода нет — не менялся.
    assert cell["value"] == 540
    assert cell["quote"] == "9 часов"
    assert cell["answer"] == ROOT_Q
    assert cell["source"]["kind"] == "answer"
    assert cell["source"]["question"]["id"] == REV_Q
    ans = _q(doc, REV_Q)["answer"]
    assert ans["status"] == "применён"
    assert "closed_by" not in ans and "close_reason" not in ans
    assert m.pending_revision_facts(doc) == []
    # Признак снят самим внесением нового значения в ту же ячейку —
    # общему дожиму в конце прохода снимать уже нечего.
    assert summary["settled_revisions"] == []
    assert m.validate_model(doc) == []
    # Синхронизация после ответа ничего не переоткрывает и не снимает.
    assert fp.ensure_family_param_questions(doc) == []
    assert rule["template_fit"]["params"][PARAM]["value"] == 540


def test_declined_revision_removes_the_withdrawn_value():
    """«Не применяется» на пересмотр — решение авиакомпании: прежнее
    значение того же вопроса не остаётся жить рядом с ним (иначе старое
    значение выдавалось бы за новое решение)."""
    doc, rule = _doc_after_conflicting_revision()
    summary = answers_mod.apply_answers(doc, {REV_Q: "не применяется"})
    assert summary["applied"] == 1
    assert PARAM not in rule["template_fit"].get("params", {})
    # Соседний параметр семьи остаётся — посадка без параметров не
    # остаётся, и отказ применяется решением (граница — тест ниже).
    assert OTHER_PARAM in rule["template_fit"]["params"]
    assert m.pending_revision_facts(doc) == []
    assert _q(doc, REV_Q)["answer"]["status"] == "применён"
    assert m.validate_model(doc) == []


def _only_param_doc() -> tuple[dict, str, str]:
    """Семья с ЕДИНСТВЕННЫМ параметром (reserve_floors): ячейка внесена
    ответом и стоит под открытым пересмотром. После слияния с K01 это
    обычный случай (нечитаемая ячейка клаузы → параметр недостаёт →
    вопрос → ответ вносит значение → пересмотр)."""
    root = "Q-3b-параметр-reserve_floors-reserve_crews"
    rid = root + "-R"
    fmt = {"type": "value", "note": "целое число; либо «не применяется»"}
    doc = {
        "meta": {"customer": "Т", "mode": 1, "cdpm_version": "0.1",
                 "revision": 3},
        "sources": [{"sid": "S1", "file": "a.md", "sha256": "0" * 64}],
        "rules": [{"id": "R-R", "title": "Резерв",
                   "statement": "резервных экипажей не менее 2",
                   "zone": "ROSTERING", "kind": "HARD",
                   "blockers": [rid],
                   "sources": [{"sid": "S1", "anchor": "п.9"}],
                   "template_fit": {
                       "category": "A", "registry_rev": REV,
                       "family": "reserve_floors",
                       "params": {"reserve_crews": {
                           "value": 2, "quote": "2", "answer": root,
                           "pending_revision": rid}}}}],
        "questions": [
            {"id": root, "part": "3b", "title": "Параметр",
             "text": "текст вопроса о параметре семьи",
             "answer_format": dict(fmt), "rule_refs": ["R-R"],
             "source_refs": [{"sid": "S1", "anchor": "п.9"}],
             "answer": {"raw": "2", "status": "пересмотрен",
                        "revision": 2, "revised_by": rid}},
            {"id": rid, "part": "3b", "title": "Пересмотр параметра",
             "text": "текст вопроса-пересмотра параметра семьи",
             "answer_format": dict(fmt), "rule_refs": ["R-R"],
             "source_refs": [{"sid": "S1", "anchor": "п.9"}],
             "revision_of": root}],
        "objectives": [],
    }
    return doc, root, rid


def _filled_book(tmp_path: Path, doc: dict, filled: dict) -> bytes:
    """Настоящая книга анкеты прогона с заполненной колонкой ответа."""
    from openpyxl import load_workbook
    from cdpm import anketa_xlsx as ax
    book = tmp_path / "anketa-otvety.xlsx"
    ax.render_anketa_xlsx(doc, book)
    wb = load_workbook(str(book))
    for ws in wb.worksheets:
        if ws.title in (ax.SUMMARY_TITLE, ax.ANSWERED_TITLE):
            continue
        for row in ws.iter_rows(min_row=2):
            if row[0].value in filled:
                row[5].value = filled[row[0].value]
    wb.save(str(book))
    return book.read_bytes()


def _wait_stage(run: Path, timeout: float = 60.0) -> dict:
    import time
    end = time.time() + timeout
    while time.time() < end:
        status = json.loads((run / "status.json").read_text(encoding="utf-8"))
        if "идёт" not in str(status.get("state") or ""):
            return status
        time.sleep(0.1)
    raise AssertionError("стадия применения ответов не завершилась")


def test_decline_of_the_only_family_param_is_not_a_decision(tmp_path,
                                                            monkeypatch):
    """Отказ «не применяется» на ЕДИНСТВЕННЫЙ параметр семьи — НЕ
    решение о значении: снять такую ячейку значило бы оставить посадку
    без параметров (контракт «params обязателен»), и живой приём книги
    уронил бы стадию целиком — слова авиакомпании пропали бы, а
    пересмотр остался бы открытым навсегда.

    Закон: такой отказ записывается «не разобран» словами («посадка на
    семью не стоит — нужна перепосадка строки, не ответ»); книга
    принята, свод валиден, прежнее значение остаётся ПОД признаком
    (решением оно не является — барьер держит готовность). Путь —
    настоящий: `webui._upload_answers_bytes` с книгой анкеты; канал
    языковой модели падает на любом обращении."""
    doc, _root, rid = _only_param_doc()
    assert m.validate_schema(doc) == [] and m.validate_model(doc) == []
    monkeypatch.setenv("CDPM_RUNS_DIR", str(tmp_path / "runs"))
    monkeypatch.setenv("CDPM_SETTINGS_FILE", str(tmp_path / "settings.json"))
    run = _run(tmp_path / "runs", doc, "run-9404")
    (run / "status.json").write_text(json.dumps({"state": "готов"}),
                                     encoding="utf-8")

    def _no_llm(*args, **kwargs):
        def _channel(prompt):
            raise AssertionError("языковая модель звалась на пути ответов")
        return _channel
    monkeypatch.setattr(webui, "_channel_for", _no_llm)

    webui._upload_answers_bytes(
        "run-9404", "anketa-otvety.xlsx",
        _filled_book(tmp_path, doc, {rid: "не применяется"}))
    status = _wait_stage(run)
    assert status.get("state") == "ответы применены", status

    after = json.loads((run / "extract" / "model.json")
                       .read_text(encoding="utf-8"))
    answer = _q(after, rid)["answer"]
    assert answer["status"] == "не разобран"
    assert answer["raw"] == "не применяется"          # слова сохранены
    for word in ("единственный", "перепосадка строки, не ответ"):
        assert word in answer["note"], word
    # Прежнее значение стоит ПОД признаком: решением оно не является.
    cell = after["rules"][0]["template_fit"]["params"]["reserve_crews"]
    assert cell["value"] == 2 and cell["pending_revision"] == rid
    facts = m.pending_revision_facts(after)
    assert [f["problem"] for f in facts] == [None]
    assert m.readiness(after)["state"] != "в"
    assert m.validate_schema(after) == [] and m.validate_model(after) == []
    # Синхронизация вопрос не снимает: носитель жив.
    assert fp.ensure_family_param_questions(after) == []
    assert m.is_open_question(_q(after, rid))


def test_reapply_does_not_resurrect_value_under_open_revision():
    """Самолечение посадки не возвращает прежнее значение, пока
    пересмотр открыт: иначе прежний ответ стал бы решением нового
    вопроса, а синхронизация следом сняла бы пересмотр как лишний."""
    doc, rule = _doc_after_conflicting_revision()
    rule["template_fit"]["params"].pop(PARAM)      # заплата переиздала посадку
    assert fp.ensure_family_param_questions(doc) == []
    assert PARAM not in rule["template_fit"]["params"]
    assert m.is_open_question(_q(doc, REV_Q))


# ── (7) подложный признак ────────────────────────────────────────────

@pytest.mark.parametrize("case,qid,fragment", [
    ("вопроса нет", "Q-3b-параметр-нет-такого-вопроса-R", "в анкете нет"),
    ("не пересмотр", ROOT_Q, "не является пересмотром"),
])
def test_forged_pending_reference_is_named_not_green(case, qid, fragment,
                                                     tmp_path):
    """Признак указывает не на неразрешённый вопрос-пересмотр: барьер
    стоит, беда названа словами — валидатор модели, готовность, отказ
    выпуска и гейт говорят о ней, а не молчат зелёным."""
    doc = _doc(pending=True)
    _cell(doc)["pending_revision"] = qid
    problems = m.validate_model(doc)
    assert len(problems) == 1 and fragment in problems[0], (case, problems)
    fact = m.pending_revision_facts(doc)[0]
    assert fact["problem"] and fragment in fact["problem"]
    assert m.readiness(doc)["state"] != "в"
    run = _run(tmp_path, doc)
    with pytest.raises(tp.TaskPackageError) as e:
        tp.build_package(run)
    assert e.value.code == tp.FAMILY_INCOMPLETE and fragment in str(e.value)


@pytest.mark.parametrize("bad", [True, 5, 0, "", None, {}, []])
def test_forged_pending_tag_of_wrong_form_is_named_not_green(bad, tmp_path):
    """Признак НЕ по форме (`true`, `5`, `0`, `""`, `null`, `{}`, `[]`):
    признак ЕСТЬ, годного значения нет — это беда словами, а не
    отсутствие признака. Прежде такой признак молча пропускался: фактов
    ноль, готовность «в», пакет ВЫПУСКАЛСЯ (выпуск схему не проверяет) и
    выдавался по ссылке скачивания. Барьер обязан стоять на всех
    читателях."""
    doc = _doc(pending=False)
    _cell(doc)["pending_revision"] = bad
    assert m.validate_schema(doc), "схема обязана отвергать признак не-строку"

    facts = m.pending_revision_facts(doc)
    assert len(facts) == 1 and facts[0]["problem"]
    assert "не по форме" in facts[0]["problem"]
    assert m.validate_model(doc) == [facts[0]["words"]]
    ready = m.readiness(doc)
    assert ready["state"] != "в" and ready["rostering_ready"] is False

    run = _run(tmp_path, doc)
    with pytest.raises(tp.TaskPackageError) as e:
        tp.build_package(run)
    assert e.value.code == tp.FAMILY_INCOMPLETE
    assert "не по форме" in str(e.value)

    # Гейт производителя по подложному пакету — тем же расчётом.
    clean = _run(tmp_path, _doc(pending=False,
                                revision_answer=APPLIED_REVISION),
                 "run-9405")
    pkg = tmp_path / "forged-form.zip"
    pkg.write_bytes(_repack(tp.export_zip(clean), doc))
    problems, _counters = gate.check(pkg)
    assert [p for p in problems if tp.FAMILY_INCOMPLETE in p]


def test_tag_pointing_at_a_foreign_revision_is_named_not_green():
    """Признак указывает на открытый пересмотр ЧУЖОГО ответа (другого
    параметра той же семьи): своего пересмотра у ячейки нет — снятие
    чужого пересмотра сняло бы барьер с непересмотренного значения."""
    other_root = f"Q-3b-параметр-{FAMILY}-{OTHER_PARAM}"
    other_rev = other_root + "-R"
    doc = _doc(pending=True)
    doc["questions"] += [
        {"id": other_root, "part": "3b", "title": "Параметр (дома)",
         "text": "текст вопроса о ДРУГОМ параметре семьи",
         "answer_format": {"type": "value", "note": NOTE},
         "rule_refs": ["R-F"],
         "source_refs": [{"sid": "S1", "anchor": "п.2"}],
         "answer": {"raw": "12:00", "status": "пересмотрен",
                    "revision": 2, "revised_by": other_rev}},
        {"id": other_rev, "part": "3b", "title": "Пересмотр (дома)",
         "text": "текст пересмотра ДРУГОГО параметра семьи",
         "answer_format": {"type": "value", "note": NOTE},
         "rule_refs": ["R-F"],
         "source_refs": [{"sid": "S1", "anchor": "п.2"}],
         "revision_of": other_root}]
    _cell(doc)["pending_revision"] = other_rev
    fact = m.pending_revision_facts(doc)[0]
    assert fact["problem"] and "ЧУЖОЙ ответ" in fact["problem"]
    assert m.validate_model(doc) == [fact["words"]]
    assert m.readiness(doc)["state"] != "в"


def test_pending_left_after_applied_answer_is_named_not_green():
    """Ответ на пересмотр применён, а признак в ячейке остался —
    осиротевший признак: барьер стоит и беда названа (снятие признака
    обязано идти одним из двух законных путей)."""
    doc = _doc(pending=True,
               revision_answer={"raw": "9 часов", "status": "применён",
                                "revision": 3})
    problems = m.validate_model(doc)
    assert len(problems) == 1 and "применён, а признак" in problems[0]
    assert m.readiness(doc)["state"] != "в"
    # Законный путь снятия возвращает свод в зелёное состояние.
    assert m.settle_applied_revisions(doc) == ["R-F"]
    assert m.pending_revision_facts(doc) == []
    assert m.validate_model(doc) == [] and m.readiness(doc)["state"] == "в"


# ── (8) порядок и независимость ─────────────────────────────────────

def test_verdict_independent_of_order_and_of_question_generator():
    """Вердикт не зависит от порядка строк и вопросов и от того,
    отработал ли построитель вопросов: перестановка даёт те же факты,
    независимая семья и лишняя строка ничего не меняют."""
    doc = _doc(pending=True)
    base = m.pending_revision_facts(doc)
    assert base and base[0]["problem"] is None

    shuffled = copy.deepcopy(doc)
    shuffled["rules"] = list(reversed(shuffled["rules"]))
    shuffled["questions"] = list(reversed(shuffled["questions"]))
    assert m.pending_revision_facts(shuffled) == base
    assert m.readiness(shuffled)["state"] == m.readiness(doc)["state"]

    # Построитель вопросов не участвует в вердикте: до и после его
    # прохода факты одни и те же.
    after_generator = copy.deepcopy(doc)
    fp.ensure_family_param_questions(after_generator)
    assert m.pending_revision_facts(after_generator) == base

    # Независимая семья рядом вердикта не меняет.
    with_neighbour = copy.deepcopy(doc)
    with_neighbour["rules"].append(_landed_rule("R-N", {
        "max_consecutive_night_duties": {"value": 3, "quote": "не более 3"},
        "night_duty_start": {"value": 1320, "quote": "с 22:00"},
        "night_duty_end": {"value": 300, "quote": "до 05:00"}}, "п.4"))
    with_neighbour["rules"][-1]["template_fit"]["family"] = "night_series"
    assert m.pending_revision_facts(with_neighbour) == base


# ── (9) перезапуск и идемпотентность ────────────────────────────────

def test_two_syncs_give_byte_identical_document():
    """Повтор синхронизации байт в байт: открытый пересмотр не
    снимается ни первым проходом, ни вторым, и документ не «дышит»."""
    doc, _rule = _doc_after_conflicting_revision()
    fp.ensure_family_param_questions(doc)
    first = json.dumps(doc, ensure_ascii=False, sort_keys=True)
    fp.ensure_family_param_questions(doc)
    second = json.dumps(doc, ensure_ascii=False, sort_keys=True)
    assert first == second
    assert m.is_open_question(_q(doc, REV_Q))

    # Тот же закон у законного кодового исхода: повтор не меняет байты.
    gone = _with_open_revision(_doc(pending=False, rules=[_retired_rule()]),
                               ["R-F"])
    fp.ensure_family_param_questions(gone)
    one = json.dumps(gone, ensure_ascii=False, sort_keys=True)
    fp.ensure_family_param_questions(gone)
    assert json.dumps(gone, ensure_ascii=False, sort_keys=True) == one


# ── (10) настоящие демонстрационные наборы ──────────────────────────

def _kit_models() -> list[tuple[str, dict]]:
    root = Path(__file__).resolve().parents[1] / "final-report" / "DEMO-H0-KITS"
    out: list[tuple[str, dict]] = []
    for kit in ("H0", "H0-B", "H0-C"):
        for path in sorted((root / kit / "cdpm-run").glob("*/model.json")):
            out.append((f"{kit}/{path.parent.name}",
                        json.loads(path.read_text(encoding="utf-8"))))
    return out


def test_real_demo_kits_get_no_false_refusal():
    """Настоящие демонстрационные наборы (своды прогонов и своды
    пакетов) барьером не задеваются: признака пересмотра в них нет,
    выпуск не отказывает, слов о пересмотре в подписи готовности нет."""
    models = _kit_models()
    assert len(models) >= 3, "своды демонстрационных наборов не найдены"
    for name, doc in models:
        assert m.pending_revision_facts(doc) == [], name
        assert tp.pending_revision_problem(doc) is None, name
        assert m.PENDING_REVISION_TAG not in m.readiness(doc)["label"], name
        assert not any("пересмотр" in p and "ячейк" in p
                       for p in m.validate_model(doc)), name


def test_real_demo_kit_packages_pass_the_producer_gate_check():
    """Своды ПАКЕТОВ тех же наборов: проверка гейта по K04 молчит —
    ранняя часть не краснит замороженные пакеты задним числом."""
    root = Path(__file__).resolve().parents[1] / "final-report" / "DEMO-H0-KITS"
    seen = 0
    for kit in ("H0", "H0-B", "H0-C"):
        for pkg in sorted((root / kit / "cdpm-run").glob("*/task-package.zip")):
            with zipfile.ZipFile(pkg) as z:
                doc = json.loads(z.read("model.json").decode("utf-8"))
            assert m.pending_revision_facts(doc) == [], pkg.parent.name
            seen += 1
    assert seen >= 3, "пакеты демонстрационных наборов не найдены"


# ── (11) обход сборщика ─────────────────────────────────────────────

def _repack(body: bytes, doc: dict) -> bytes:
    """Подложный пакет: model.json подменён, контрольная сумма в
    манифесте пересчитана — обход сборщика, а не порча файла."""
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        files = {n: z.read(n) for n in z.namelist()}
    model = json.dumps(doc, ensure_ascii=False, indent=1).encode("utf-8")
    files["model.json"] = model
    manifest = json.loads(files["manifest.json"].decode("utf-8"))
    for entry in manifest["files"]:
        if entry["name"] == "model.json":
            entry["sha256"] = hashlib.sha256(model).hexdigest()
            entry["bytes"] = len(model)
    files["manifest.json"] = json.dumps(
        manifest, ensure_ascii=False, indent=1).encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name in sorted(files):
            z.writestr(name, files[name])
    return out.getvalue()


def test_bypass_package_with_pending_cell_is_red_at_the_producer_gate(
        tmp_path):
    """Пакет, собранный в обход рубежа (подменённый `model.json` с
    честной контрольной суммой), краснеет у гейта производителя ТЕМ ЖЕ
    расчётом и теми же словами — и прямой вызов сборщика отказывает."""
    clean = _doc(pending=False, revision_answer=APPLIED_REVISION)
    run = _run(tmp_path, clean)
    body = tp.export_zip(run)

    forged = _doc(pending=True)
    pkg = tmp_path / "forged.zip"
    pkg.write_bytes(_repack(body, forged))
    problems, _counters = gate.check(pkg)
    k04 = [p for p in problems if tp.FAMILY_INCOMPLETE in p]
    assert len(k04) == 1, problems
    for word in (FAMILY, PARAM, REV_Q, "R-F"):
        assert word in k04[0], word

    # Прямой вызов сборщика на том же своде — отказ с тем же кодом.
    run2 = _run(tmp_path, forged, "run-9403")
    with pytest.raises(tp.TaskPackageError) as e:
        tp.build_package(run2)
    assert e.value.code == tp.FAMILY_INCOMPLETE


# ── (12) писатель кодового снятия не гасит ЧУЖОЙ признак пересмотра ──
#
# Находка Б-1 неавторского ревью общей головы 17а (13.09.2026). Закон
# §2.4 (б) знает ровно ДВА пути снятия признака `pending_revision`:
# ответ человека на пересмотр и кодовый исход «вопрос больше не
# требуется» у САМОГО пересмотра. Единый писатель снимал признак
# безусловно — всякое кодовое снятие ЛЮБОГО вопроса гасило признак,
# который на этот вопрос указывает. Так заводился третий путь: снятие
# ПОСТОРОННЕГО вопроса (линия данных, история), у которого ни ответа
# авиакомпании нет, ни носители никуда не делись.
#
# Происхождение названо честно: дефект предсуществующий — ветвь
# «посадка C» ходила через того же писателя ещё в ПРИНЯТОЙ единице K04
# (3322ca1e:cdpm/model.py:4965); композиция 12.09.2026 перевела на
# писателя ВТОРУЮ ветвь («правило судится решением») и удвоила
# поверхность. Пробы держат обе половины шва: писателя самого и боевой
# путь линии данных, который его зовёт.

NO_DATA_Q = "Q-3b-нет-данных-R-C"


def _doc_with_foreign_pending_mark() -> dict:
    """Свод, где признак `pending_revision` ячейки указывает на вопрос
    ЛИНИИ ДАННЫХ (`Q-3b-нет-данных-…`), а не на пересмотр: признак
    подделан либо испорчен переносом. Барьер K04 обязан стоять и
    называть беду словами («вопрос … не является пересмотром»), а
    снятие постороннего вопроса кодом обязано его не трогать."""
    return {
        "meta": {"customer": "Т", "mode": 1, "cdpm_version": "0.1",
                 "revision": 4},
        "sources": [{"sid": "S1", "file": "a.md", "sha256": "0" * 64}],
        "rules": [
            {"id": "R-C", "title": "Без шаблона",
             "statement": "механики такого требования в каталоге нет",
             "zone": "ROSTERING", "kind": "HARD", "blockers": [NO_DATA_Q],
             "sources": [{"sid": "S1", "anchor": "п.1"}],
             "template_fit": {"category": "C", "registry_rev": REV,
                              "gap": "нет механики"}},
            _landed_rule("R-A", {
                OTHER_PARAM: {"value": 720, "quote": "дома не менее 12:00"},
                PARAM: {"value": 600, "quote": "10:00 (600 минут)",
                        "answer": ROOT_Q, "pending_revision": NO_DATA_Q}}),
        ],
        "questions": [
            {"id": NO_DATA_Q, "part": "3b",
             "title": "Данных для правила R-C нет в книге",
             "text": "Проверка правила не прогоняется. Считать ли до "
                     "получения данных?",
             "answer_format": {"type": "choice",
                               "options": [{"key": "а", "text": "да"},
                                           {"key": "б", "text": "нет"}]},
             "rule_refs": ["R-C"],
             "source_refs": [{"sid": "S1", "anchor": "п.1"}]},
        ],
    }


def _no_data_report(rid: str = "R-C") -> dict:
    return {"checks": [{"id": f"C-{rid}", "status": "не скомпилирована",
                        "reason": "нет листа «Смены»",
                        "basis": {"rule": rid}}],
            "findings": []}


def test_code_writer_leaves_a_foreign_pending_mark_alone():
    """Половина шва «писатель»: кодовое снятие вопроса, который
    пересмотром НЕ является, чужой признак `pending_revision` не
    трогает. Прежде писатель снимал признак безусловно — и посторонний
    вопрос уносил с собой барьер чужой ячейки."""
    doc = _doc_with_foreign_pending_mark()
    q = _q(doc, NO_DATA_Q)
    assert not m.is_revision_question(q)          # причинное условие
    assert m.code_close_refusal(doc, q) is None   # отказ тут не сторож
    m.close_question_by_code(doc, q, key="неприменимо",
                             raw="снят кодом: у правила нет шаблона "
                                 "(посадка C)",
                             note="пробел покрытия")
    assert _cell(doc, "R-A")["pending_revision"] == NO_DATA_Q
    assert not m.is_open_question(q)              # вопрос всё же снят
    # Any code closure identifies its author (CDPM-409). Only a revision
    # carries the special disappearance outcome; a foreign CAS stays intact.
    assert q["answer"]["closed_by"] == "code"
    assert "close_reason" not in q["answer"]
    assert "answer_version" not in q["answer"]
    # Preserve existing key order and append the explicit author marker.
    assert list(q["answer"]) == ["raw", "status", "key", "revision", "note", "closed_by"]


def test_data_line_withdrawal_does_not_take_the_barrier_with_it():
    """Половина шва «боевой путь»: линия данных снимает кодом свой
    вопрос «нет данных» у строки без шаблона (посадка C) — барьер K04
    ЧУЖОЙ ячейки обязан остаться, и валидатор свода обязан назвать
    беду. Прежде после этого снятия фактов барьера становилось 0,
    `validate_model` замолкал, а отказ выпуска исчезал."""
    doc = _doc_with_foreign_pending_mark()
    before = m.pending_revision_facts(doc)
    assert [f["cell"] for f in before] == [f"R-A.{PARAM}"]
    assert "не является пересмотром" in before[0]["problem"]
    assert len(m.pending_revision_problems(doc)) == 1
    assert tp.pending_revision_problem(doc) is not None

    report = _no_data_report()
    assert m.ensure_missing_data_questions(doc, report) == []
    assert report["withdrawn_no_data_questions"] == [NO_DATA_Q]
    assert not doc["rules"][0].get("blockers")    # свой вопрос снят

    assert _cell(doc, "R-A")["pending_revision"] == NO_DATA_Q
    after = m.pending_revision_facts(doc)
    assert [f["cell"] for f in after] == [f["cell"] for f in before]
    assert [f["problem"] for f in after] == [f["problem"] for f in before]
    assert len(m.pending_revision_problems(doc)) == 1
    words, code, items = tp.pending_revision_problem(doc)
    assert code == tp.FAMILY_INCOMPLETE
    assert items == [f"{FAMILY}/{PARAM}"]
    assert "пересмотр открыт" in words
    assert m.readiness(doc)["state"] != "в"


def test_code_writer_still_clears_the_mark_of_its_own_revision():
    """Цена правки: законная дорога писателя осталась. Кодовое снятие
    ПЕРЕСМОТРА снимает признак СВОЕЙ ячейки — иначе осиротевший признак
    держал бы готовность вечно.

    Проба контрактная, не поточная: боевые зовы писателя стоят под
    `is_open_question`, а сюда доходит только разрешённый пересмотр
    (`code_close_refusal` возвращает `None` лишь у разрешённого либо
    у оставшегося без носителей, а у последнего снимать уже нечего —
    всякая ячейка с признаком делает строку носителем). Предмет пробы —
    контракт самого писателя, и он же ловит удаление зова целиком."""
    doc = _doc(pending=True, revision_answer=APPLIED_REVISION)
    q = _q(doc, REV_Q)
    assert _cell(doc)["pending_revision"] == REV_Q
    assert m.code_close_refusal(doc, q) is None
    m.close_question_by_code(doc, q, raw="снят кодом", note="лишний")
    assert "pending_revision" not in _cell(doc)
    assert q["answer"]["closed_by"] == "code"
    assert m.pending_revision_facts(doc) == []


# ── (13) самолечение: рубеж K04 и пустой `params` (С-1, С-2 ревью) ───
#
# Находки С-1 и С-2 того же ревью. Свод, на котором рубеж K04 в
# `family_params._reapply_values` ДОСТИЖИМ: корень ПРИМЕНЁН (цепочка на
# пересмотр не переведена, `revised_by` нет), а вопрос-пересмотр того
# же параметра ОТКРЫТ. В штатном потоке `answers.open_revision` ставит
# корню «пересмотрен» и `revised_by`, и проход останавливает закон
# происхождения A03-мин РАНЬШЕ рубежа — оттого мутация «удалить рубеж»
# не краснила ни одной пробы. Здесь отказывает ровно рубеж.


def _doc_open_revision_but_applied_root() -> dict:
    """Носитель БЕЗ ячейки параметра («посадку переиздала заплата»),
    корень применён, пересмотр открыт."""
    doc = _doc(pending=False)
    rule = next(r for r in doc["rules"] if r["id"] == "R-F")
    rule["template_fit"].pop("params")              # ячеек нет вовсе
    _q(doc, ROOT_Q)["answer"] = {"raw": "10:00 (600 минут)",
                                 "status": "применён", "revision": 2}
    return _with_open_revision(doc, ["R-F"])


def _fit(doc: dict, rid: str = "R-F") -> dict:
    return next(r for r in doc["rules"] if r["id"] == rid)["template_fit"]


def test_self_healing_refuses_while_the_revision_of_the_param_is_open():
    """С-2: рубеж K04 вторым замком — на своде, который ДО него
    доходит. Прежний ответ в посадку не возвращается, причина названа
    словами, параметр остаётся недостающим. Без рубежа старый ответ
    оказался бы решением НОВОГО вопроса."""
    doc = _doc_open_revision_but_applied_root()
    rev, root = _q(doc, REV_Q), _q(doc, ROOT_Q)
    assert m.is_revision_question(rev) and m.is_open_question(rev)
    assert root["answer"]["status"] == "применён"
    assert "revised_by" not in root["answer"]     # причинное условие
    notes: list[str] = []
    assert fp._reapply_values(doc, fp.family_coverage(doc), notes) == []
    assert notes == [f"{FAMILY}.{PARAM}: пересмотр параметра открыт — "
                     f"значение в посадку не возвращено, пока нет нового "
                     f"решения авиакомпании"]
    assert PARAM not in _fit(doc).get("params", {})


def test_refused_self_healing_does_not_touch_the_bytes_of_the_model():
    """С-1: проход, который ОТКАЗАЛ, свод не трогает ни одной буквой.
    Прежде рубеж K04 отказывал уже ПОСЛЕ `_writable_cells`, а тот у
    посадки без `params` словарь ЗАВОДИТ: в своде оставался пустой
    `params: {}`, который схема краснит («minProperties»: 1), а
    валидатор модели зовёт «params посадки не непустой объект».

    Обе ветви отказа — рубеж K04 и предсуществующий близнец «значение
    не разобрано» — лечатся одной правкой и проверяются вместе."""
    def _unparsable(d):
        """Пересмотра НЕТ вовсе, но сырой ответ корня значением не
        разбирается — проход обязан отказать так же тихо."""
        d["questions"] = [q for q in d["questions"] if q["id"] != REV_Q]
        _q(d, ROOT_Q)["answer"]["raw"] = "уточним после совещания"

    for name, damage in (("рубеж K04", lambda d: None),
                         ("значение не разобрано", _unparsable)):
        doc = _doc_open_revision_but_applied_root()
        damage(doc)
        before = json.dumps(doc, ensure_ascii=False)
        assert m.validate_schema(doc) == [], name
        notes: list[str] = []
        assert fp._reapply_values(doc, fp.family_coverage(doc),
                                  notes) == [], name
        assert json.dumps(doc, ensure_ascii=False) == before, name
        assert "params" not in _fit(doc), name
        assert m.validate_schema(doc) == [], name


def test_refused_self_healing_keeps_the_model_clean_on_the_live_call():
    """То же на БОЕВОМ зове (`ensure_family_param_questions`): свод
    после отказа схеме не противоречит, и валидатор модели не получает
    НОВОЙ беды про пустые `params` — прежние беды свода (семья без
    ячеек) остаются, их лечит ответ авиакомпании."""
    doc = _doc_open_revision_but_applied_root()
    was = m.validate_model(doc)
    notes: list[str] = []
    fp.ensure_family_param_questions(doc, notes)
    assert m.validate_schema(doc) == []
    assert [p for p in m.validate_model(doc) if p not in was] == []
    assert not any("не непустой объект" in p for p in m.validate_model(doc))
    assert notes and "пересмотр параметра открыт" in notes[0]


def test_self_healing_never_rewrites_crooked_landing_cells():
    """С-1 стороны барьеров, сохранённый перестановкой: ячейки посадки
    кривой формы самолечение не переписывает и не заводит взамен —
    чужую форму судит валидатор свода, а не заплата.

    Проба стоит на НОВОМ читателе формы (`_cells_writable`): он спрошен
    ДО рубежа и ничего не пишет, а `_writable_cells` зовётся уже на
    записи. Сними читателя — проход дойдёт до записи и упадёт на
    списке (`cells[param] = …`), то есть замок исчезнет."""
    doc = _doc_open_revision_but_applied_root()
    doc["questions"] = [q for q in doc["questions"] if q["id"] != REV_Q]
    _fit(doc)["params"] = []                       # форма не по схеме
    assert m.validate_schema(doc) != []             # её судит валидатор
    before = json.dumps(doc, ensure_ascii=False)
    notes: list[str] = []
    assert fp._reapply_values(doc, fp.family_coverage(doc), notes) == []
    assert notes == []            # молчит: беду формы называет валидатор
    assert json.dumps(doc, ensure_ascii=False) == before
    assert _fit(doc)["params"] == []


def test_self_healing_tells_an_absent_key_from_a_present_null():
    """С-1, вторая половина (находка GPT-166 13.09.2026): `params:
    null` — это КРИВАЯ ФОРМА, а не отсутствие ключа, хотя
    `fit.get("params")` отдаёт `None` в обоих случаях.

    На этом совпадении читатель формы ЛГАЛ: он звал `null`
    записываемым, `_writable_cells` возвращал на нём `None`
    (`setdefault` у СУЩЕСТВУЮЩЕГО ключа отдаёт прежнее значение), и
    самолечение падало на `cells[param] = …` с `TypeError`. Живой путь
    `apply_param_answer` спасала его собственная проверка `cells is
    None`; у `_reapply_values` её нет — он доверяет читателю формы, и
    докстрока прохода прямо на это ссылается.

    Проба доходит ДО записи: действующий ответ цепочки есть и
    разбирается (600), открытого пересмотра нет — ни происхождение
    A03-мин, ни рубеж K04 её раньше не отсекают. Проверка на `[]`
    (соседняя проба) этого случая НЕ различает: там читатель отказывал
    и без правки."""
    doc = _doc_open_revision_but_applied_root()
    doc["questions"] = [q for q in doc["questions"] if q["id"] != REV_Q]
    _fit(doc)["params"] = None              # ключ ЕСТЬ, значение null
    assert m.validate_schema(doc) != []      # форму судит валидатор
    before = json.dumps(doc, ensure_ascii=False)

    notes: list[str] = []
    assert fp._reapply_values(doc, fp.family_coverage(doc), notes) == []

    assert json.dumps(doc, ensure_ascii=False) == before
    assert "params" in _fit(doc)             # ключ не исчез
    assert _fit(doc)["params"] is None       # и null НЕ подменён на {}


def test_the_live_path_names_a_crooked_landing_in_words():
    """ВТОРАЯ ПОЛОВИНА закона С-1 — та, на которую соседняя проба
    ссылается СЛОВАМИ и которую до сих пор не судила НИ ОДНА проба из
    2296 (находка неавторского ревью 13.09.2026: мутация «снять рубеж
    `cells is None` в `apply_param_answer`» проходила зелёной на всей
    маске).

    Живой путь зовёт `cdpm/answers.py` БЕЗ `try/except`: без рубежа
    человек получил бы `TypeError` наружу из обработчика ответа анкеты —
    пятисотку вместо слов — на первом же ответе по параметру у
    повреждённой посадки.

    Кривая форма — ЛЮБАЯ не-словарная, и каждая судится отдельно:
    проверка на одной форме доказывает одну форму."""
    ФОРМЫ = (None, [], "нет", 0)
    for форма in ФОРМЫ:
        doc = _doc_open_revision_but_applied_root()
        doc["questions"] = [q for q in doc["questions"] if q["id"] != REV_Q]
        _fit(doc)["params"] = форма
        было = json.dumps(doc["rules"], ensure_ascii=False)

        got = fp.apply_param_answer(doc, _q(doc, ROOT_Q),
                                    "10:00 (600 минут)", 3)

        assert got["applied"] is False, (форма, got)
        assert got["touched"] == [], (форма, got)
        ответ = _q(doc, ROOT_Q)["answer"]
        assert ответ["status"] == "не разобран", (форма, ответ)
        assert "записать некуда" in ответ["note"], (форма, ответ)
        # свод НЕ ТРОНУТ, и кривая форма сохранена как была
        assert json.dumps(doc["rules"], ensure_ascii=False) == было, форма

    # ПОЛОЖИТЕЛЬНЫЙ контроль: ключа нет — ответ применяется штатно,
    # иначе рубеж отказывал бы всем подряд и проба была бы беззубой.
    doc = _doc_open_revision_but_applied_root()
    doc["questions"] = [q for q in doc["questions"] if q["id"] != REV_Q]
    assert "params" not in _fit(doc)
    got = fp.apply_param_answer(doc, _q(doc, ROOT_Q), "10:00 (600 минут)", 3)
    assert got["applied"] is True, got
    assert _fit(doc)["params"][PARAM]["value"] == 600


def test_self_healing_keeps_both_lawful_ways_of_writing_cells():
    """Обе ЗАКОННЫЕ дороги живы: ключа нет — запись его заводит; ключ
    есть словарём — пишем В НЕГО, не подменяя и не теряя чужих ячеек.
    Без этой пробы лечение кривой формы могло забрать с собой и
    здоровые случаи, а зелень соседних проб этого бы не показала."""
    def _healthy():
        d = _doc_open_revision_but_applied_root()
        d["questions"] = [q for q in d["questions"] if q["id"] != REV_Q]
        return d

    doc = _healthy()                        # (а) ключа нет вовсе
    assert "params" not in _fit(doc)
    assert fp._reapply_values(doc, fp.family_coverage(doc), []) == ["R-F"]
    assert _fit(doc)["params"][PARAM]["value"] == 600

    doc = _healthy()                        # (б) словарь с ЧУЖОЙ ячейкой
    _fit(doc)["params"] = {"иной_параметр": {"value": 1}}
    assert fp._reapply_values(doc, fp.family_coverage(doc), []) == ["R-F"]
    cells = _fit(doc)["params"]
    assert cells[PARAM]["value"] == 600
    assert cells["иной_параметр"] == {"value": 1}   # чужое не тронуто


def test_the_shape_reader_and_the_writer_agree_on_every_shape():
    """Договор двух половин: `_cells_writable` — РОВНО предикат
    «`_writable_cells` вернёт словарь». На этом договоре стоит
    докстрока самолечения («возврата None здесь быть не может»), и до
    правки она была неверна ровно на `params: null`. Проверяем договор
    на всех формах сразу, а не на одной: одна форма доказывает одну
    форму."""
    def rule(**kw):
        return {"template_fit": dict({"category": "A", "family": FAMILY},
                                     **kw)}

    shapes = [("ключа нет", rule(), True),
              ("пустой словарь", rule(params={}), True),
              ("словарь с ячейкой", rule(params={"p": {"value": 1}}), True),
              ("null", rule(params=None), False),
              ("список", rule(params=[]), False),
              ("строка", rule(params="нет"), False),
              ("число", rule(params=0), False)]
    for name, r, writable in shapes:
        assert fp._cells_writable(r) is writable, name
        cells = fp._writable_cells(r)
        if writable:
            assert isinstance(cells, dict), name
        else:
            assert cells is None, name

    probe = rule()                  # у ЧИТАТЕЛЯ побочного эффекта нет
    assert fp._cells_writable(probe) is True
    assert "params" not in probe["template_fit"]
