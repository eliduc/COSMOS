"""Accepted be2 residuals: cosmetic excerpts, normative E2/E3, actual book."""
import copy

import pytest
from openpyxl import load_workbook

from cdpm import answers, anketa_xlsx as ax, exchange, family_params as fp
from cdpm import history_questions as hq, model as m, question_editions as qe, render
from test_cdpm_family_params import _doc_with_home_floor, QID
from test_cdpm_history_editions import _published as _history
from test_cdpm_history_1_17_2 import _book
from test_cdpm_question_editions import _doc


def _family():
    doc, carrier = _doc_with_home_floor()
    carrier["sources"][0]["quote"] = "Отдых — не менее 10 часов"
    fp.ensure_family_param_questions(doc)
    q = next(x for x in doc["questions"] if x["id"] == QID)
    qe.freeze(q)
    assert answers.apply_answers(doc, {q["id"]: "600 минут"})["applied"] == 1
    return doc, carrier, q


@pytest.mark.parametrize("writer", ["family", "history"])
@pytest.mark.parametrize("quote", [
    "Отдых — не менее 10 часов ", "Отдых — не  менее 10 часов",
    "Отдых —не менее 10 часов", "не менее 10 часов",
])
def test_cosmetic_excerpt_of_same_frozen_clause_does_not_reissue(writer, quote):
    if writer == "family":
        doc, carrier, q = _family()
        run = lambda: fp.ensure_family_param_questions(doc)
    else:
        doc, q = _history()
        carrier = next(x for x in doc["rules"] if x["id"] == "R-H-REST")
        assert answers.apply_answers(doc, {q["id"]: "б"})["applied"] == 1
        run = lambda: hq.ensure_history_questions(doc, _book())
    old = copy.deepcopy(q)
    assert run() == []  # healthy no-change control
    carrier["sources"][0]["quote"] = quote
    assert run() == []
    assert run() == []
    assert q == old
    assert not any(x.get("edition_of") == q["id"] for x in doc["questions"])


@pytest.mark.parametrize("source", [
    {"sid": "S1", "anchor": "П.4.1"},
    {"sid": "S2", "anchor": "П.2.3"},
    {"sid": "S2", "anchor": "П.4.1"},
])
def test_new_normative_address_still_reissues(source):
    doc, carrier, q = _family()
    old = copy.deepcopy(q)
    carrier["sources"].append(source)
    assert fp.ensure_family_param_questions(doc) == [QID + "-E2"]
    new = next(x for x in doc["questions"] if x["id"] == QID + "-E2")
    assert m.is_open_question(new)
    assert source in new["edition"]["sources"]
    assert new["id"] in carrier["blockers"]
    assert q == old


@pytest.mark.parametrize("answer_e2", [False, True])
def test_genuine_new_clause_needs_answer_even_after_unanswered_e2(tmp_path, answer_e2):
    doc, carrier, q = _family()
    old = copy.deepcopy(q)
    carrier["sources"].append({"sid": "S1", "anchor": "П.9.1"})
    assert fp.ensure_family_param_questions(doc) == [QID + "-E2"]
    if answer_e2:
        assert answers.apply_answers(doc, {QID + "-E2": "660 минут"})["applied"] == 1
    carrier["sources"].append({"sid": "S1", "anchor": "П.9.2"})
    assert fp.ensure_family_param_questions(doc) == [QID + "-E3"]
    new = next(x for x in doc["questions"] if x["id"] == QID + "-E3")
    assert m.is_open_question(new)
    assert new["id"] in carrier["blockers"]
    assert fp.ensure_family_param_questions(doc) == []
    assert m.is_open_question(new)
    assert q == old
    # The historical value remains, but cannot close the new normative question.
    assert carrier["template_fit"]["params"]["min_rest_outstation"]["value"] \
        == (660 if answer_e2 else 600)
    assert new["id"] in render.render_anketa(doc)
    cell = _saved_source_cell(tmp_path, doc, new)
    assert "ответ «600 минут» применён к изданию 1" in cell
    assert "Добавлены нормативные источники:" in cell
    assert "П.9.2" in cell.split("Добавлены нормативные источники:", 1)[1]
    assert answers.apply_answers(doc, {new["id"]: "720 минут"})["applied"] == 1
    assert not m.is_open_question(new)
    assert new["id"] not in carrier["blockers"]
    assert carrier["template_fit"]["params"]["min_rest_outstation"]["value"] == 720


def _saved_source_cell(tmp_path, doc, q):
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return next(row[4] for sheet in wb for row in sheet.iter_rows(values_only=True)
                    if row[0] == q["id"])
    finally:
        wb.close()


def test_each_normative_edition_names_only_its_added_source(tmp_path):
    doc, q = _doc()
    qe.freeze(q)
    for anchor in ("П.9.1", "П.9.2"):
        previous = q
        q = qe.expand(doc, q, [], content={"source_refs": q["source_refs"] + [
            {"sid": "S1", "anchor": anchor}]})
        text = render.render_anketa(doc)
        line = next((s for s in text.splitlines()
                     if s.startswith("Добавлены нормативные источники:")), "")
        assert anchor in line
        for source in previous["source_refs"]:
            assert source["anchor"] not in line
        cell = _saved_source_cell(tmp_path, doc, q)
        assert "Добавлены нормативные источники:" in cell
        delta = cell.split("Добавлены нормативные источники:", 1)[1]
        assert anchor in delta
        assert all(s["anchor"] not in delta for s in previous["source_refs"])


@pytest.mark.parametrize("text_only_reissue", [False, True])
def test_no_source_delta_for_first_or_text_only_edition(tmp_path, text_only_reissue):
    doc, q = _doc()
    qe.freeze(q)
    if text_only_reissue:
        q = qe.expand(doc, q, [], content={"text": q["text"] + " Уточнение."})
    assert "Добавлены нормативные источники:" not in render.render_anketa(doc)
    assert "Добавлены нормативные источники:" not in _saved_source_cell(tmp_path, doc, q)


def test_explicit_normative_excerpt_patch_is_not_deduplicated_as_code_witness(tmp_path):
    doc, q = _doc()
    qe.freeze(q)
    old = copy.deepcopy(q)
    source = dict(q["source_refs"][0], quote="Кроме случаев продления отдыха")
    assert exchange._apply_question_patch(doc, {
        "question_id": q["id"], "add_source_refs": [source]}, None) == []
    new = next(x for x in doc["questions"] if x.get("edition_of") == q["id"])
    assert q == old
    assert qe.fingerprint(new) != qe.fingerprint(q)
    assert source in new["edition"]["sources"]
    assert source["quote"] in render.render_anketa(doc)
    assert source["quote"] in _saved_source_cell(tmp_path, doc, new)


def test_code_source_union_preserves_first_witness_and_distinct_addresses():
    first = {"sid": "S1", "anchor": "П.1", "quote": "первый отрывок"}
    second = dict(first, quote="другой отрывок")
    other = dict(first, anchor="П.2")
    refs = qe.normative_rule_sources([{"sources": [first, second, other]}])
    assert refs == [first, other]
    refs[0]["quote"] = "не менять носитель"
    assert first["quote"] == "первый отрывок"


def test_saved_book_gate_requires_the_new_source_explanation(tmp_path):
    from cdpm.gates.gate_anketa_xlsx import check_workbook
    doc, q = _doc()
    qe.freeze(q)
    q = qe.expand(doc, q, [], content={"source_refs": q["source_refs"] + [
        {"sid": "S1", "anchor": "П.9.1"}]})
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    assert check_workbook(doc, path)[0] == []
    wb = load_workbook(path)
    try:
        ws = wb[ax.SHEET_TITLES[q["part"]]]
        row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == q["id"])
        assert "Добавлены нормативные источники:" in row[4].value
        row[4].value = row[4].value.split("; Добавлены нормативные источники:", 1)[0]
        wb.save(path)
    finally:
        wb.close()
    assert check_workbook(doc, path)[0]
