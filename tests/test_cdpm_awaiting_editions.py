"""Pending work counts current editions, not immutable historical answers."""
import copy
import json

import pytest

from cdpm import answers, model as m, question_editions as qe
from cdpm import anketa_xlsx as ax, render
from openpyxl import load_workbook
from test_cdpm_goal_editions import _goal
from test_cdpm_modality_editions import _doc as _modality
from test_cdpm_objectives import _base_doc


def _pending(goal=False):
    if goal:
        doc, q = _goal(soft=False)
    else:
        doc = _base_doc()
        q = {"id": "Q-2-TEXT", "part": "2", "title": "Правило",
             "text": "Уточните правило", "answer_format": {"type": "registry"},
             "rule_refs": ["R-H"], "source_refs": [{"sid": "S1", "anchor": "п.1"}]}
        doc["questions"] = [q]
        qe.freeze(q)
    answers.apply_answers(doc, {q["id"]: "любое допустимое решение"})
    assert q["answer"]["status"] == "ожидает обработчика"
    return doc, q


def _next(doc, q):
    return qe.expand(doc, q, [], content={"text": q["text"] + " Новые условия."})


@pytest.mark.parametrize("goal", [False, True])
def test_superseded_answer_is_history_not_remaining_worker_work(goal):
    doc, q = _pending(goal)
    current = _next(doc, q)
    before = copy.deepcopy(doc)
    def forbidden(prompt):
        pytest.fail("historical answer must not invoke a worker")
    assert answers.apply_worker_patches(doc, forbidden)["worker_applied"] == 0
    assert answers.remaining_awaiting(doc) == 0
    assert m.open_questions(doc) == [current]
    assert doc == before


@pytest.mark.parametrize("goal", [False, True])
def test_only_latest_pending_answer_counts_through_three_editions(goal):
    doc, q = _pending(goal)
    second = _next(doc, q)
    answers.apply_answers(doc, {second["id"]: "любое допустимое решение"})
    third = _next(doc, second)
    answers.apply_answers(doc, {third["id"]: "любое допустимое решение"})
    doc["questions"].reverse()
    before = copy.deepcopy(doc)
    assert answers.remaining_awaiting(doc) == 1
    assert doc == before


@pytest.mark.parametrize("damage", ["draft_child", "changed_child", "fork", "revision_gap"])
def test_invalid_successor_does_not_hide_unresolved_work(damage):
    doc, q = _pending()
    current = _next(doc, q)
    if damage == "draft_child":
        current.pop("edition")
    elif damage == "changed_child":
        current["text"] += " Подмена."
    elif damage == "fork":
        other = copy.deepcopy(current)
        other.pop("edition")
        other["id"] += "-FORK"
        qe.freeze(other, revision=2)
        doc["questions"].append(other)
    else:
        current.pop("edition")
        qe.freeze(current, revision=3)
    before = copy.deepcopy(doc)
    assert answers.remaining_awaiting(doc) == 1
    assert doc == before


def test_id_suffix_alone_cannot_retire_an_answer():
    doc, q = _pending()
    other = copy.deepcopy(q)
    other.pop("edition")
    other["id"] += "-E2"
    qe.freeze(other)
    doc["questions"].append(other)
    assert answers.remaining_awaiting(doc) == 2


@pytest.mark.parametrize("damage", ["reactivated", "family_changed", "edition_changed"])
def test_invalid_code_resolution_cannot_make_pending_count_zero(damage):
    doc, q, first, _ = _modality(answer={"raw": "с оговоркой",
        "status": "ожидает обработчика", "revision": 2})
    m.derive_rule_traits(doc, first, "test")
    assert answers.remaining_awaiting(doc) == 0
    if damage == "reactivated":
        first["blockers"] = [q["id"]]
    elif damage == "family_changed":
        first["template_fit"]["family"] = "unknown-family"
    else:
        q["text"] += " Подмена."
    before = copy.deepcopy(doc)
    assert answers.remaining_awaiting(doc) == 1
    assert doc == before


def test_open_questions_answer_api_uses_current_edition():
    doc, q = _pending()
    q.pop("answer")
    current = _next(doc, q)
    assert answers.open_questions(doc) == m.open_questions(doc) == [current]


def test_current_rejected_work_is_not_hidden_by_application_failure():
    doc, q = _pending()
    foreign = copy.deepcopy(doc["rules"][0])
    foreign.update(id="R-FOREIGN", blockers=[q["id"]])
    doc["rules"].append(foreign)
    assert qe.application_problems(doc, q)
    assert answers.remaining_awaiting(doc) == 1


@pytest.mark.parametrize("goal", [False, True])
def test_repeated_worker_rounds_do_not_report_retired_answer_as_failure(goal):
    doc, q = _pending(goal)
    _next(doc, q)
    before = copy.deepcopy(doc)
    def forbidden(prompt):
        pytest.fail("retired work must never be sent to a worker")
    for _ in range(2):
        result = answers.apply_worker_patches(doc, forbidden)
        assert result == {"worker_applied": 0, "worker_rejected": []}
    assert doc == before


@pytest.mark.parametrize("goal", [False, True])
def test_mid_call_reissue_refuses_once_then_retires_without_losing_answer(goal):
    doc, q = _pending(goal)
    old = copy.deepcopy(q)
    calls = []
    def worker(prompt):
        calls.append(prompt)
        _next(doc, q)
        return json.dumps({"patches": [], "decline": True, "objectives": [], "weights": []})
    first = answers.apply_worker_patches(doc, worker)
    assert first["worker_applied"] == 0
    assert len(first["worker_rejected"]) == 1
    assert "переиздан" in first["worker_rejected"][0]["reasons"][0]
    assert len(calls) == 1 and q == old
    second = answers.apply_worker_patches(doc, worker)
    assert second == {"worker_applied": 0, "worker_rejected": []}
    assert len(calls) == 1 and q == old
    assert answers.remaining_awaiting(doc) == 0


def test_pending_predecessor_trace_is_explicit_in_saved_book_and_markdown(tmp_path):
    doc, q = _pending(goal=True)
    current = _next(doc, q)
    before = copy.deepcopy(doc)
    notes = qe.prior_answer_notes(doc, current)
    assert len(notes) == 1
    assert q["answer"]["raw"] in notes[0] and "не применён" in notes[0]
    assert "новый ответ" in notes[0]
    md = render.render_anketa(doc)
    block = md.split(f"### ВОПРОС {current['id']}:", 1)[1]
    block = block.split("\n### ", 1)[0].split("\n## ", 1)[0]
    assert notes[0] in block  # specifically this question, not the history table
    path = tmp_path / "pending-trace.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path)
    try:
        ws = wb[ax.SHEET_TITLES[current["part"]]]
        row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == current["id"])
        assert notes[0] in str(row[4].value)
    finally:
        wb.close()
    assert doc == before


@pytest.mark.parametrize("finish_latest", [False, True])
def test_retired_answer_history_never_looks_like_live_pending_work(tmp_path, finish_latest):
    doc, q = _pending(goal=True)
    current = _next(doc, q)
    old = copy.deepcopy(q)
    if finish_latest:
        answers.apply_answers(doc, {current["id"]: "любое допустимое решение"})
        result = answers.apply_worker_patches(doc, lambda prompt: json.dumps({
            "decline": True, "objectives": [], "weights": []}))
        assert result == {"worker_applied": 1, "worker_rejected": []}
    before = copy.deepcopy(doc)
    status = "не применён — вопрос переиздан"
    row = next(r for r in ax.answered_rows(doc) if r[0] == q["id"])
    assert row[3] == status and row[2] == old["answer"]["raw"]
    summary = ax.summary_rows(doc)
    assert ("Прежних неприменённых ответов", 1) in summary
    assert ("Принятых ответов", int(finish_latest)) in summary
    md_row = next(line for line in render.render_anketa(doc).splitlines()
                  if line.startswith(f"| {q['id']} — "))
    assert status in md_row and "ожидает обработчика" not in md_row
    path = tmp_path / "history-status.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path)
    try:
        saved = next(r for r in wb[ax.ANSWERED_TITLE].iter_rows(min_row=2)
                     if r[0].value == q["id"])
        assert saved[3].value == status and saved[2].value == old["answer"]["raw"]
    finally:
        wb.close()
    assert q == old and answers.remaining_awaiting(doc) == 0
    assert doc == before


def test_current_pending_answer_still_displays_awaiting_status():
    doc, q = _pending(goal=True)
    row = next(r for r in ax.answered_rows(doc) if r[0] == q["id"])
    assert row[3] == "ожидает обработчика"
    assert not any(r and r[0] == "Прежних неприменённых ответов" for r in ax.summary_rows(doc))
    assert answers.remaining_awaiting(doc) == 1


@pytest.mark.parametrize("column", [1, 2, 3, 4])
def test_legacy_book_context_tolerates_edge_whitespace_only(tmp_path, column):
    from test_cdpm_question_editions import _doc, _filled
    doc, q = _doc()
    path = _filled(tmp_path, doc, q)
    qe.migrate_legacy_document(doc)
    wb = load_workbook(path)
    try:
        ws = wb[ax.SHEET_TITLES[q["part"]]]
        row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == q["id"])
        original = str(row[column].value or "")
        row[column].value = " \t" + original + "\u00a0 "
        wb.save(path)
        assert answers.read_book(path, doc).answers == {q["id"]: "600 минут"}
        row[column].value = original + " ДРУГОЙ СМЫСЛ"
        wb.save(path)
        rejected = answers.read_book(path, doc)
        assert q["id"] not in rejected.answers
        assert "контекст" in rejected.refused[q["id"]]["reason"]
    finally:
        wb.close()
