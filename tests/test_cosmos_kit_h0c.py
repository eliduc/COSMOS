# -*- coding: utf-8 -*-
"""ИДАК H0-C — набор H0 плюс правило без семьи в каталоге (поручение
владельца 09.09.2026, вариант C). Проверяется сборка фазы 1: 17 записей
документами (H0 плюс VN-026 после VN-060), якоря H0 со сдвигом строк листа
ниже вставленной, эталон с VN-026 категории C и словами пробела, книга
ответов только об истории, ожидания фазы 1 называют заявку и пробел;
фаза 2 без семьи в каталоге отказывает словами; отпечатки источников
ставит НОСИТЕЛЬ базового сборщика (bk.stamp_sources), своей формулы у
вариации нет, и обе проекции эталона получают их по факту сборки.

Равенство пересобранных docx/xlsx/jpg с ВЫЛОЖЕННЫМИ файлами комплекта не
судится: в zip docx/xlsx попадает время записи, а байты jpg зависят от
шрифтов машины (находка CI № 81, 11.09.2026) — сравнивается то, что
сборщик записал сам."""
import hashlib
import importlib.util
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

KIT = ROOT / "open-bench" / "polyaris-h0c"
H0 = ROOT / "open-bench" / "polyaris-h0"


def _builder():
    spec = importlib.util.spec_from_file_location(
        "build_kit_h0c", KIT / "build_kit_h0c.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _sha(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_h0c_kit_is_h0_plus_rule_without_family(tmp_path):
    b = _builder()
    res = b.build(tmp_path)
    form = json.loads((tmp_path / "etalon" / "VMP-H0C-form.json").read_text(encoding="utf-8"))
    ids = [r["id"] for r in form["rostering"]["hard"]]
    assert ids.index(b.NEW_RULE) == ids.index(b.AFTER_RULE) + 1
    row = next(r for r in form["rostering"]["hard"] if r["id"] == b.NEW_RULE)
    assert row["template"] is None and row["gap"] == b.GAP_WORDS
    # параметров у записи нет: имя даст черновик заявки и запись каталога
    assert "params" not in row and str(b.NEW_LIMIT) in row["statement"]
    assert row["machine"]["template_fit"]["category"] == "C"
    vf = __import__("cdpm.vmp_form", fromlist=["x"])
    h0_form = json.loads((H0 / "etalon" / "VMP-H0-form.json").read_text(encoding="utf-8"))
    assert res["rows"] == sum(len(rows) for _p, _s, rows in vf.form_sections(h0_form)) + 1
    # якоря: VN-026 — строка листа сразу за VN-060, строки ниже опустились
    amap = json.loads((tmp_path / "etalon" / "anchors-map.json").read_text(encoding="utf-8"))
    h0_map = json.loads((H0 / "etalon" / "anchors-map.json").read_text(encoding="utf-8"))
    assert amap["VN-060"] == h0_map["VN-060"]
    assert amap[b.NEW_RULE][0]["anchor"].endswith("строка 5")
    assert amap["VN-050"][0]["anchor"].endswith("строка 7") \
        and h0_map["VN-050"][0]["anchor"].endswith("строка 6")
    assert amap["VN-010"] == h0_map["VN-010"]
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "docs" / b.XLSX_NAME)
    texts = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value]
    line = next(t for t in texts if t.startswith(b.NEW_RULE))
    assert b.NEW_STATEMENT in line and "Шаблона нет" in line \
        and "Параметры" not in line


def test_h0c_etalon_phase1_marks_coverage_gap(tmp_path):
    b = _builder()
    b.build(tmp_path)
    et = json.loads((tmp_path / "etalon" / "VMP-H0C.json").read_text(encoding="utf-8"))
    h0 = json.loads((H0 / "etalon" / "VMP-H0.json").read_text(encoding="utf-8"))
    ids = [r["id"] for r in et["rules"]]
    assert len(ids) == len(h0["rules"]) + 1
    assert ids.index(b.NEW_RULE) == ids.index(b.AFTER_RULE) + 1
    rule = next(r for r in et["rules"] if r["id"] == b.NEW_RULE)
    assert rule["kind"] == "HARD" and rule["zone"] == "ROSTERING"
    fit = rule["template_fit"]
    assert fit["category"] == "C" and fit["gap"] == b.GAP_WORDS
    assert fit["registry_rev"] == h0["meta"].get("registry_rev", fit["registry_rev"]) \
        and len(fit["registry_rev"]) == 64
    assert "family" not in fit and "params" not in fit
    assert rule["sources"][0]["anchor"].endswith("строка 5")
    # якоря остальных правил листа — по карте комплекта (VN-050 опустился)
    vn050 = next(r for r in et["rules"] if r["id"] == "VN-050")
    assert vn050["sources"][0]["anchor"].endswith("строка 7")
    vn060 = next(r for r in et["rules"] if r["id"] == "VN-060")
    assert vn060["sources"][0]["anchor"].endswith("строка 4")
    assert et["questions"] == [] and len(et["objectives"]) == 2


def test_h0c_manifest_phase1_expects_request_and_stop(tmp_path):
    b = _builder()
    b.build(tmp_path)
    m = json.loads((tmp_path / "kit-manifest.json").read_text(encoding="utf-8"))
    assert m["phase"] == 1
    e = m["expect"]
    # Каталог 0.5.0: вопросы об истории — по виду чтения семей эталона
    # фазы 1 (семьи VN-026 в каталоге нет — её строка блоком «логика»).
    et = json.loads((tmp_path / "etalon" / "VMP-H0C.json").read_text(encoding="utf-8"))
    history_q = b.bk.history_question_ids(et)
    assert history_q and all(q.startswith(b.bk.history_question_prefix()) for q in history_q)
    assert e["questions"] == history_q
    assert e["coverage_gap_rules"] == [b.NEW_RULE]
    assert e["catalog_requests"]["count"] == 1 and e["package_ready"] is False
    assert set(m["expect_phases"]) == {"1", "2"}
    assert m["expect_phases"]["2"]["rule_landing"] == {b.NEW_RULE: "A"}
    book = json.loads((tmp_path / "etalon" / "ANSWER-BOOK-MACHINE.json").read_text(encoding="utf-8"))
    assert list(book["exact"]) == []
    assert book["prefix"][b.bk.history_question_prefix()] == b.bk.HISTORY_CONTINUE_ZERO
    # дефект линии данных CDPM («нет данных» у посадки C) исправлен
    # 09.09.2026 (1316023b) — обхода «а» по префиксу больше нет
    assert book["prefix"] == {b.bk.history_question_prefix(): b.bk.HISTORY_CONTINUE_ZERO}
    assert "исправлен" in book["_comment"]
    assert m["catalog"] == json.loads((H0 / "kit-manifest.json").read_text(encoding="utf-8"))["catalog"]


def test_сборщик_h0c_ставит_отпечатки_источников_по_факту_сборки(tmp_path):
    """Отпечатки источников обеих проекций эталона (форма и ВМП) — sha
    ТОЛЬКО ЧТО собранных файлов комплекта, и ставит их носитель базового
    сборщика: своей формулы отпечатка у вариации больше нет (11.09.2026 —
    одна формула на все сборщики формы H0)."""
    b = _builder()
    assert not hasattr(b, "_sha"), "у вариации своя формула отпечатка"
    b.build(tmp_path)
    form = json.loads((tmp_path / "etalon" / "VMP-H0C-form.json")
                      .read_text(encoding="utf-8"))
    et = json.loads((tmp_path / "etalon" / "VMP-H0C.json")
                    .read_text(encoding="utf-8"))
    assert [s["file"] for s in form["sources"]] == [b.DOCX_NAME, b.XLSX_NAME,
                                                    b.JPG_NAME]
    assert et["sources"] == form["sources"]
    for s in form["sources"]:
        assert s["sha256"] == _sha(tmp_path / "docs" / s["file"]), s["file"]


def test_сборщик_h0c_отказывает_словами_когда_файла_источника_нет(tmp_path,
                                                                  monkeypatch):
    """Файла из `sources` на диске нет — отказ словами носителя, а не
    молчаливый пропуск и не голое падение разбора."""
    b = _builder()
    real = b.bk.build_jpg

    def _без_файла(path, items):
        lines = real(path, items)
        Path(path).unlink()
        return lines

    monkeypatch.setattr(b.bk, "build_jpg", _без_файла)
    with pytest.raises(SystemExit) as e:
        b.build(tmp_path)
    assert "источник формы не найден" in str(e.value)
    assert not (tmp_path / "kit-manifest.json").is_file()


def test_сборщик_h0c_отказывает_когда_проекции_эталона_о_разных_источниках(
        tmp_path, monkeypatch):
    """ВМП эталона о других файлах, чем форма, — отказ словами (закон
    базового сборщика): молча выдать эталон о чужих файлах нельзя."""
    b = _builder()
    srcs = [{"sid": "S1", "file": "a.docx", "sha256": "0" * 64}]
    b.check_sources(srcs, {"sources": [dict(srcs[0])]})        # молчит
    for et in ({"sources": [{"sid": "S1", "file": "ЧУЖОЙ.docx"}]}, {}):
        with pytest.raises(SystemExit) as e:
            b.check_sources(srcs, et)
        assert "О РАЗНЫХ ИСТОЧНИКАХ" in str(e.value)
    real = b.derive_etalon
    monkeypatch.setattr(b, "derive_etalon", lambda *a, **kw: dict(
        real(*a, **kw), sources=[]))
    with pytest.raises(SystemExit) as e:
        b.build(tmp_path)
    assert "О РАЗНЫХ ИСТОЧНИКАХ" in str(e.value)
    assert not (tmp_path / "kit-manifest.json").is_file()


def test_h0c_phase2_without_family_refuses(tmp_path):
    b = _builder()
    with pytest.raises(SystemExit) as e:
        b.build(tmp_path, phase=2, family="нет_такой_семьи")
    assert "семьи" in str(e.value)


def test_h0c_builder_is_deterministic(tmp_path):
    b = _builder()
    b.build(tmp_path / "a")
    b.build(tmp_path / "b")
    # Манифест не сравнивается (пути эталона вне репозитория абсолютны);
    # отпечатки docx/xlsx в `sources` несут метки времени файлов —
    # сравнивается содержание: записи, правила, целевые, якоря, книга.
    def _read(p):
        return json.loads(p.read_text(encoding="utf-8"))
    for name in ("etalon/anchors-map.json", "etalon/ANSWER-BOOK-MACHINE.json"):
        assert _read(tmp_path / "a" / name) == _read(tmp_path / "b" / name), name
    fa, fb = _read(tmp_path / "a/etalon/VMP-H0C-form.json"), _read(tmp_path / "b/etalon/VMP-H0C-form.json")
    fa.pop("sources"), fb.pop("sources")
    assert fa == fb
    ea, eb = _read(tmp_path / "a/etalon/VMP-H0C.json"), _read(tmp_path / "b/etalon/VMP-H0C.json")
    assert ea["rules"] == eb["rules"] and ea["objectives"] == eb["objectives"]


TWO = ["max_duty_series_length", "post_series_days_off"]     # черновик CDPM 09.09
ROLES = {"max_duty_series_length": "limit", "post_series_days_off": "days_off"}


def test_h0c_bind_roles_one_param_is_limit_and_two_need_mapping():
    b = _builder()
    assert b.bind_roles("f", ["max_consecutive_duties"], None) == \
        {"max_consecutive_duties": "limit"}
    with pytest.raises(SystemExit) as e:
        b.bind_roles("f", TWO, None)
    assert "--param" in str(e.value)
    with pytest.raises(SystemExit) as e:
        b.bind_roles("f", TWO, {**ROLES, "post_series_days_off": "rest"})
    assert "ролей" in str(e.value)
    with pytest.raises(SystemExit) as e:
        b.bind_roles("f", TWO, {n: "days_off" for n in TWO})
    assert "дважды" in str(e.value)
    with pytest.raises(SystemExit) as e:
        b.bind_roles("f", TWO, {"max_duty_series_length": "limit", "nope": "days_off"})
    assert "нет" in str(e.value)
    with pytest.raises(SystemExit) as e:
        b.bind_roles("f", [], None)
    assert "нет параметров" in str(e.value)
    # порядок — как у семьи каталога, а не как у ключей
    assert list(b.bind_roles("f", TWO, dict(reversed(list(ROLES.items()))))) == TWO


def test_h0c_phase2_etalon_with_two_params(tmp_path, monkeypatch):
    b = _builder()
    monkeypatch.setattr(b, "catalog_family", lambda fam, mapping=None: (
        b.bind_roles(fam, TWO, mapping), "0.3.0", "f" * 64))
    res = b.build(tmp_path, phase=2, family="duty_series_limit", params=ROLES)
    assert res["params"] == ROLES
    et = json.loads((tmp_path / "etalon" / "VMP-H0C.json").read_text(encoding="utf-8"))
    rule = next(r for r in et["rules"] if r["id"] == b.NEW_RULE)
    fit = rule["template_fit"]
    assert fit["category"] == "A" and fit["family"] == "duty_series_limit"
    assert fit["registry_rev"] == "f" * 64 and "gap" not in fit
    assert fit["params"]["max_duty_series_length"] == {"value": 3, "quote": b.NEW_QUOTE}
    assert fit["params"]["post_series_days_off"]["value"] == 1
    for p in fit["params"].values():            # цитаты — из формулировки
        assert p["quote"] in b.NEW_STATEMENT
    m = json.loads((tmp_path / "kit-manifest.json").read_text(encoding="utf-8"))
    assert m["phase"] == 2 and m["catalog"] == {"version": "0.3.0", "fingerprint": "f" * 64}
    assert m["expect"]["family"] == {
        "id": "duty_series_limit", "roles": ROLES,
        "params": {"max_duty_series_length": 3, "post_series_days_off": 1}}
    assert m["expect"]["family_complete"] == {"duty_series_limit": TWO}
    assert m["expect"]["rule_landing"] == {b.NEW_RULE: "A"}
    assert m["expect"]["cosmos_numbers"]["rostering"] == "14/0"
    # без связки ролей при двух параметрах сборка отказывает словами
    with pytest.raises(SystemExit) as e:
        b.build(tmp_path, phase=2, family="duty_series_limit")
    assert "--param" in str(e.value)


def test_h0c_числа_сверки_фазы_2_справочные_и_по_прогонам_09_09(tmp_path):
    """Числа COSMOS фазы 2 — 11/0 и 14/0 (фактические прогоны Test и Prod
    09.09.2026; прежние 9/0 и 13/0 устарели ещё с волны 7), и при них
    стоит пометка «справочно»: число строк сверки гуляет по чтению
    модели, критерий «без ошибок» — вердикты, сертификаты и НОЛЬ
    нарушений (слово владельца). Числа задаёт сборщик — манифест
    комплекта обязан совпасть с ним."""
    b = _builder()
    got = json.loads((KIT / "kit-manifest.json").read_text(encoding="utf-8"))
    for numbers in (b.EXPECT_PHASE_2["cosmos_numbers"],
                    got["expect"]["cosmos_numbers"],
                    got["expect_phases"]["2"]["cosmos_numbers"]):
        assert numbers["pairing"] == "11/0", numbers
        assert numbers["rostering"] == "14/0", numbers
        assert "справочно" in numbers["_note"], numbers
        assert "0 нарушений" in numbers["_note"], numbers
    # фаза 1 чисел COSMOS не ждёт: пакета там нет
    assert "cosmos_numbers" not in b.EXPECT_PHASE_1


def test_h0c_main_parses_param_keys(monkeypatch):
    b = _builder()
    seen = {}
    monkeypatch.setattr(b, "build", lambda **kw: seen.update(kw))
    assert b.main(["--phase", "2", "--family", "duty_series_limit",
                   "--param", "max_duty_series_length=limit",
                   "--param", "post_series_days_off=days_off"]) == 0
    assert seen == {"phase": 2, "family": "duty_series_limit", "params": ROLES}
    with pytest.raises(SystemExit):
        b.main(["--phase", "2", "--family", "x", "--param", "без_знака"])


def test_h0c_phase2_reads_the_catalog_of_the_setting(tmp_path, monkeypatch):
    """Фаза 2 читает каталог НАСТРОЙКИ продукта (CREW_TEMPLATE_CATALOG), а
    не только каталог дерева: сухой прогон по временному каталогу стенда
    CR (кандидат записи семьи) обязан видеть семью, которой в дереве нет,
    и класть в манифест версию с отпечатком ЭТОГО каталога."""
    import catalog_copies as cc
    import template_catalog as tc
    folder = cc.copy_catalog(tmp_path, families={"duty_series_limit": {
        "status": "есть", "zone": "ROSTERING", "kind": "HARD",
        "violation_unit": None,       # единица нарушения — только у SOFT
        "params": [{"name": n, "unit": "count", "qualifier": None,
                    "relation": None, "note": "проба копии каталога"}
                   for n in TWO]}})
    monkeypatch.setenv("CREW_TEMPLATE_CATALOG", str(folder))
    b = _builder()
    assert str(b.bk.catalog_dir()) == str(folder)
    res = b.build(tmp_path / "kit", phase=2, family="duty_series_limit",
                  params=ROLES)
    assert res["params"] == ROLES
    cat = tc.load(folder)
    m = json.loads((tmp_path / "kit" / "kit-manifest.json").read_text(encoding="utf-8"))
    assert m["catalog"] == {"version": str(cat.version()),
                            "fingerprint": str(cat.fingerprint())}
    et = json.loads((tmp_path / "kit" / "etalon" / "VMP-H0C.json").read_text(encoding="utf-8"))
    fit = next(r for r in et["rules"] if r["id"] == b.NEW_RULE)["template_fit"]
    assert fit["category"] == "A" and fit["family"] == "duty_series_limit"
    assert set(fit["params"]) == set(TWO)
    assert fit["registry_rev"] == str(cat.fingerprint())


def test_h0c_phase2_refusal_leaves_the_kit_untouched(tmp_path):
    """Отказ фазы 2 (семьи «есть» в каталоге нет) — ДО записи документов:
    сухой прогон 09.09.2026 показал, что неудачный запуск переписывал
    docx/xlsx комплекта (те же байты содержания, новые отметки времени в
    zip) и оставлял комплект грязным."""
    b = _builder()
    kit = tmp_path / "kit"
    b.build(kit)                                   # фаза 1 — комплект есть
    docs = sorted(p for p in (kit / "docs").iterdir() if p.is_file())
    before = {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in docs}
    with pytest.raises(SystemExit) as e:
        b.build(kit, phase=2, family="no_such_family_h0c")
    assert "no_such_family_h0c" in str(e.value)
    after = {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in docs}
    assert after == before, "отказ фазы 2 переписал документы комплекта"
