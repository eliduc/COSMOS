"""Offline source boundary: the reader receives answers, never model values."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import socket
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def raw(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True) + "\n").encode()


def sha(body):
    return hashlib.sha256(body).hexdigest()


@pytest.fixture(autouse=True)
def offline(monkeypatch, tmp_path):
    def forbidden(*args, **kwargs):
        raise AssertionError("network is outside the offline evidence profile")
    monkeypatch.setattr(socket.socket, "connect", forbidden)
    monkeypatch.setattr(socket, "create_connection", forbidden)
    monkeypatch.setenv("CREW_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("COSMOS_SKIP_CONFIG_MIGRATION", "1")


@pytest.fixture
def answered_run(tmp_path):
    from cdpm import anketa_xlsx, answer_versions, question_editions
    from openpyxl import load_workbook

    run = tmp_path / "run-1"
    (run / "freeze/clauses").mkdir(parents=True)
    (run / "freeze/clauses/S1.json").write_bytes(raw({"clauses": [
        {"anchor": "VN023", "text": "Минимальный отдых дома — 720 минут."}]}))
    q = {"id": "Q-rest-outstation", "part": "3b", "title": "Отдых вне базы",
         "text": "Каков минимальный отдых между сменами вне базы?",
         "answer_format": {"type": "value", "note": "минуты"},
         "rule_refs": ["R1"], "source_refs": [{"sid": "S1", "anchor": "VN023"}]}
    question_editions.freeze(q)
    doc = {"meta": {"customer": "Synthetic evidence", "mode": 1,
                    "cdpm_version": "0.1", "revision": 1},
           "sources": [{"sid": "S1", "file": "rules.txt", "sha256": "0" * 64}],
           "rules": [], "objectives": [], "questions": [q]}
    d = run / "answers/rev-002"
    d.mkdir(parents=True)
    book = d / "anketa-otvety.xlsx"
    anketa_xlsx.render_anketa_xlsx(doc, book)
    wb = load_workbook(book)
    ws = wb[anketa_xlsx.SHEET_TITLES["3b"]]
    ws.cell(2, anketa_xlsx.ANSWER_COL, "600 минут")
    wb.save(book)
    wb.close()
    q["answer"] = {"raw": "600 минут", "status": "применён", "revision": 2}
    answer_versions.stamp(q["answer"], None, answer_versions.context(
        "2026-09-20T10:00:00+00:00", sha(book.read_bytes())))
    doc["meta"]["revision"] = 2
    summary = {"revision": 2, "applied": 1, "worker_applied": 0,
               "refused": [], "worker_rejected": [],
               "book_sha256": sha(book.read_bytes())}
    (d / "answers-summary.json").write_bytes(raw(summary))
    (run / "extract").mkdir()
    (run / "extract/model.json").write_bytes(raw(doc))
    return run, doc


def test_blind_reader_receives_question_and_answer_not_only_freeze(answered_run):
    """Causal baseline fails: 88c71 only reads the original 720 clause."""
    from vmp import source_check_run as scr
    run, _ = answered_run
    original = (run / "freeze/clauses/S1.json").read_bytes()
    rows = scr.load_clauses(run)
    answer = [r for r in rows if r["sid"].startswith("ANSWER-")]
    assert len(answer) == 1
    assert "минимальный отдых между сменами вне базы" in answer[0]["text"]
    assert "600 минут" in answer[0]["text"]
    assert (run / "freeze/clauses/S1.json").read_bytes() == original


def package_payloads(run, doc):
    from cdpm import source_answer_evidence as ae
    model = raw(doc)
    return {"model.json": model,
            "freeze/clauses/S1.json": (run / "freeze/clauses/S1.json").read_bytes(),
            **ae.export_files(run, model, run.name)}


def zip_body(payloads, *, source_run="run-1", manifest_edit=None):
    manifest = {"source_run": source_run, "revision": 2,
                "files": [{"name": k, "sha256": sha(v), "bytes": len(v)}
                          for k, v in sorted(payloads.items())]}
    if manifest_edit:
        manifest_edit(manifest)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("manifest.json", raw(manifest))
        for name, body in payloads.items():
            z.writestr(name, body)
    return out.getvalue()


def import_sources(body, target):
    from cdpm import source_answer_evidence as ae
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        files = ae.package_sources(z, target.name)
    for name, data in files.items():
        p = target / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(data)
    return files


def persist(run, doc):
    (run / "extract/model.json").write_bytes(raw(doc))


def test_package_roundtrip_binds_exact_sources(answered_run, tmp_path):
    from cdpm import source_answer_evidence as ae
    from vmp import source_check_run as scr
    from webapp.source_check_compile import clauses_prompt
    run, doc = answered_run
    payloads = package_payloads(run, doc)
    dest = tmp_path / "imported/run-1"
    imported = import_sources(zip_body(payloads), dest)
    assert imported["freeze/clauses/S1.json"] == (run / "freeze/clauses/S1.json").read_bytes()
    rows = scr.load_clauses(dest)
    prompt = clauses_prompt(rows)
    assert "720 минут" in prompt and "600 минут" in prompt
    assert "Каков минимальный отдых между сменами вне базы?" in prompt
    entry = next(r["answer_source"] for r in rows if "answer_source" in r)
    assert entry["answer"]["book_sha256"] == doc["questions"][0]["answer"]["answer_version"]["book_sha256"]
    assert entry["question"]["fingerprint"] == doc["questions"][0]["edition"]["fingerprint"]
    assert ae.package_sources  # no alternate importer used


def test_reader_does_not_use_parsed_model_value_or_effect(answered_run):
    from vmp import source_check_run as scr
    run, doc = answered_run
    doc["questions"][0]["answer"].update(parsed={"value": 999999}, effects=[{"after": 999999}])
    persist(run, doc)
    text = "\n".join(r["text"] for r in scr.load_clauses(run))
    assert "600 минут" in text and "999999" not in text


@pytest.mark.parametrize("name", ["anketa-otvety.xlsx", "answers-summary.json"])
def test_missing_original_source_refuses(answered_run, name):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    (run / "answers/rev-002" / name).unlink()
    with pytest.raises(ae.EvidenceError, match="отсутствует"):
        ae.export_files(run, raw(doc), run.name)


@pytest.mark.parametrize("change", ["book", "revision", "applied", "refused", "worker_rejected"])
def test_receipt_conflicts_refuse(answered_run, change):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    p = run / "answers/rev-002/answers-summary.json"
    summary = json.loads(p.read_bytes())
    if change == "book":
        summary["book_sha256"] = "f" * 64
    elif change == "revision":
        summary["revision"] = 3
    elif change == "applied":
        summary["applied"] = 0
    else:
        summary[change] = [{"question": "Q-rest-outstation", "reason": "rejected"}]
    p.write_bytes(raw(summary))
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


@pytest.mark.parametrize("field,value", [
    ("raw", "601 минут"), ("revision", True), ("revision", 1),
    ("revision", 3), ("status", "применён частично")])
def test_current_answer_must_be_exact_applied_receipt(answered_run, field, value):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    doc["questions"][0]["answer"][field] = value
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


@pytest.mark.parametrize("field,value", [("fingerprint", "0" * 12), ("book_sha256", None)])
def test_invalid_version_identity_refuses(answered_run, field, value):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    doc["questions"][0]["answer"]["answer_version"][field] = value
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


def test_changed_question_scope_refuses(answered_run):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    doc["questions"][0]["text"] = "Каков отдых дома?"
    with pytest.raises(ValueError):
        ae.export_files(run, raw(doc), run.name)


@pytest.mark.parametrize("status", ["не разобран", "ожидает обработчика", "пересмотрен"])
def test_rejected_or_pending_answer_is_not_a_norm(answered_run, status):
    from cdpm import source_answer_evidence as ae
    from vmp import source_check_run as scr
    run, doc = answered_run
    doc["questions"][0]["answer"]["status"] = status
    persist(run, doc)
    assert ae.export_files(run, raw(doc), run.name) == {}
    assert "600 минут" not in "\n".join(r["text"] for r in scr.load_clauses(run))


def test_rejected_answer_still_used_as_cell_provenance_refuses(answered_run):
    from cdpm import source_answer_evidence as ae, provenance
    run, doc = answered_run
    q = doc["questions"][0]
    cell = {"value": 600, "source": provenance.answer_source(q, "p", q["answer"]["raw"])}
    doc["rules"] = [{"id": "R1", "template_fit": {"params": {"p": cell}}}]
    q["answer"]["status"] = "не разобран"
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


def mutate_book(run, doc, mutate):
    from cdpm import anketa_xlsx
    from openpyxl import load_workbook
    path = run / "answers/rev-002/anketa-otvety.xlsx"
    wb = load_workbook(path)
    mutate(wb, wb[anketa_xlsx.SHEET_TITLES["3b"]])
    wb.save(path)
    wb.close()
    # Re-pin the container to reach the semantic guard, not just byte corruption.
    digest = sha(path.read_bytes())
    doc["questions"][0]["answer"]["answer_version"]["book_sha256"] = digest
    p = run / "answers/rev-002/answers-summary.json"
    s = json.loads(p.read_bytes())
    s["book_sha256"] = digest
    p.write_bytes(raw(s))
    persist(run, doc)


@pytest.mark.parametrize("column,value", [
    (2, "Другой предмет: отдых дома"), (3, "секунды"), (4, "R-other"),
    (5, "другой документ"), (6, "601 минут"), (7, "bad-fingerprint")])
def test_bound_book_row_semantic_mutations_refuse(answered_run, column, value):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    mutate_book(run, doc, lambda _wb, ws: ws.cell(2, column, value))
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


def test_stale_book_revision_refuses_even_with_new_container_hash(answered_run):
    from cdpm import source_answer_evidence as ae, anketa_xlsx
    run, doc = answered_run
    def change(wb, _ws):
        for row in wb[anketa_xlsx.SUMMARY_TITLE]:
            if row[0].value == anketa_xlsx.REVISION_LABEL:
                row[1].value = 0
    mutate_book(run, doc, change)
    with pytest.raises(ae.EvidenceError, match="ревизии"):
        ae.export_files(run, raw(doc), run.name)


def test_known_rejected_feedback_preserved_but_not_made_a_norm(answered_run):
    from cdpm import source_answer_evidence as ae
    from vmp import source_check_run as scr
    from webapp.source_check_compile import clauses_prompt
    run, doc = answered_run
    tail = "\nНеприменённый ответ сохранён: «777777 минут».\nОтвет не разобран: неоднозначно"
    mutate_book(run, doc, lambda _wb, ws: ws.cell(2, 2, ws.cell(2, 2).value + tail))
    sources = ae.export_files(run, raw(doc), run.name)
    index = json.loads(sources[ae.EVIDENCE])
    assert index["answers"][0]["rows"][0]["historical_feedback"] == tail
    prompt = clauses_prompt(scr.load_clauses(run))
    assert "600 минут" in prompt and "777777" not in prompt


def test_arbitrary_added_question_clause_not_hidden_as_feedback(answered_run):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    mutate_book(run, doc, lambda _wb, ws: ws.cell(2, 2, ws.cell(2, 2).value + "\nДополнительно отдых не нужен."))
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


@pytest.mark.parametrize("attack", ["all-evidence", "index", "book", "summary", "model", "run", "model-revision"])
def test_package_loss_corruption_and_wrong_identity_refuse(answered_run, tmp_path, attack):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    files = package_payloads(run, doc)
    if attack == "all-evidence":
        files = {k: v for k, v in files.items() if not k.startswith(("answers/", "source-check/"))}
    elif attack == "index":
        files[ae.EVIDENCE] = files[ae.EVIDENCE].replace(b'"revision":2', b'"revision":99')
    elif attack in ("book", "summary"):
        name = "answers/rev-002/" + ("anketa-otvety.xlsx" if attack == "book" else "answers-summary.json")
        files.pop(name)
    elif attack == "model":
        files.pop("model.json")
    elif attack == "model-revision":
        doc["meta"]["revision"] = 3
        files["model.json"] = raw(doc)
    body = zip_body(files, source_run="other-run" if attack == "run" else "run-1")
    with pytest.raises((ae.EvidenceError, KeyError)):
        import_sources(body, tmp_path / "target/run-1")


def test_manifest_hash_corruption_refuses(answered_run, tmp_path):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    files = package_payloads(run, doc)
    body = zip_body(files, manifest_edit=lambda m: m["files"][0].update(sha256="0" * 64))
    with pytest.raises(ae.EvidenceError, match="SHA"):
        import_sources(body, tmp_path / "target/run-1")


def test_imported_evidence_missing_after_extract_does_not_become_legacy(answered_run, tmp_path):
    from cdpm import source_answer_evidence as ae
    from vmp import source_check_run as scr
    run, doc = answered_run
    dest = tmp_path / "target/run-1"
    import_sources(zip_body(package_payloads(run, doc)), dest)
    (dest / ae.EVIDENCE).unlink()
    with pytest.raises((ae.EvidenceError, OSError)):
        scr.load_clauses(dest)


def test_no_answer_legacy_stays_healthy(tmp_path):
    from vmp import source_check_run as scr
    dest = tmp_path / "legacy/run-1"
    files = {"freeze/clauses/S1.json": raw({"clauses": [{"anchor": "a", "text": "Исходная норма"}]})}
    import_sources(zip_body(files), dest)
    assert scr.load_clauses(dest) == [{"sid": "S1", "anchor": "a", "text": "Исходная норма"}]


def test_actual_importer_never_falls_back_on_missing_answer_evidence(answered_run, tmp_path, monkeypatch):
    from fastapi import HTTPException
    from webapp import app, cosmos_runs
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    files = package_payloads(run, doc)
    files.pop(ae.EVIDENCE)
    pkg = tmp_path / "broken.zip"
    pkg.write_bytes(zip_body(files))
    monkeypatch.setenv("CREW_CDPM_RUNS_DIR", str(run.parent))
    monkeypatch.setattr(cosmos_runs, "workdir_of", lambda *_: tmp_path / "cosmos")
    with pytest.raises(HTTPException) as exc:
        app._cdpm_run_dir(run.name, pkg, {"sha": sha(pkg.read_bytes())})
    assert exc.value.status_code == 409
    assert not (tmp_path / "cosmos").exists()


def test_actual_importer_roundtrip_and_post_import_change(answered_run, tmp_path, monkeypatch):
    from fastapi import HTTPException
    from webapp import app, cosmos_runs
    from vmp import source_check_run as scr
    run, doc = answered_run
    pkg = tmp_path / "good.zip"
    pkg.write_bytes(zip_body(package_payloads(run, doc)))
    monkeypatch.setattr(cosmos_runs, "workdir_of", lambda *_: tmp_path / "cosmos")
    cur = {"sha": sha(pkg.read_bytes())}
    dest = app._cdpm_run_dir(run.name, pkg, cur)
    assert any("600 минут" in r["text"] for r in scr.load_clauses(dest))
    (dest / "answers/rev-002/anketa-otvety.xlsx").write_bytes(b"corrupt")
    with pytest.raises(HTTPException):
        app._cdpm_run_dir(run.name, pkg, cur)


@pytest.mark.parametrize("current", [{"sha": "0" * 64}, {"source": {"sha": "0" * 64}}])
def test_importer_current_package_sha_cannot_name_another_zip(answered_run, tmp_path, monkeypatch, current):
    from fastapi import HTTPException
    from webapp import app, cosmos_runs
    run, doc = answered_run
    pkg = tmp_path / "good.zip"
    pkg.write_bytes(zip_body(package_payloads(run, doc)))
    monkeypatch.setattr(cosmos_runs, "workdir_of", lambda *_: tmp_path / "cosmos")
    with pytest.raises(HTTPException) as exc:
        app._cdpm_run_dir(run.name, pkg, current)
    assert exc.value.status_code == 409
    assert "SHA текущего пакета" in str(exc.value.detail)
    assert not (tmp_path / "cosmos").exists()


def test_serializer_calls_answer_boundary_before_downstream_execution(answered_run, monkeypatch):
    from types import SimpleNamespace
    from cdpm import export_task_package as tp
    run, doc = answered_run
    (run / "answers/rev-002/anketa-otvety.xlsx").unlink()
    p = SimpleNamespace(run_dir=run, spec_version="1.18", canon_bytes=raw(doc),
                        model_bytes=raw(doc), doc=doc, coverage_summary=None, ms_report=None)
    monkeypatch.setattr(tp, "package_run_name", lambda _: run.name)
    def unexpected(*a, **kw):
        pytest.fail("downstream execution must not precede source evidence validation")
    monkeypatch.setattr(tp.export_solver, "build_objective_json", unexpected)
    with pytest.raises(tp.TaskPackageError) as exc:
        tp._serialize_native_content(p)
    assert exc.value.code == "source-answer-evidence"


def test_run_full_passes_authenticated_answer_to_actual_compiler_boundary(answered_run, monkeypatch):
    from vmp import source_check_run as scr
    from webapp import source_check_compile as scc
    run, _ = answered_run
    captured = []
    def capture(clauses, *args, **kwargs):
        captured.append(scc.clauses_prompt(clauses))
        # This only witnesses prompt delivery; it is deliberately NOT a
        # fabricated successful rule check or a claim about an LLM's output.
        return {"specs": [], "problems": ["capture-only; no semantic verdict"],
                "backend": "offline", "model": "capture-only", "channel": "offline"}
    monkeypatch.setattr(scc, "compile_checks", capture)
    got = scr.run_full(run, [], requested="cloud", model="capture-only", channel="api")
    assert len(captured) == 1 and "600 минут" in captured[0]
    primary = json.loads((run / "source-check/checks.json").read_bytes())
    assert primary["problems"] == ["capture-only; no semantic verdict"]
    assert got["counts"] == {}


def test_manifest_revision_is_not_inferred_from_model(answered_run, tmp_path):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    body = zip_body(package_payloads(run, doc), manifest_edit=lambda m: m.update(revision=99))
    with pytest.raises(ae.EvidenceError, match="ревизия манифеста"):
        import_sources(body, tmp_path / "target/run-1")


def test_code_closure_is_not_an_airline_answer(answered_run):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    doc["questions"][0]["answer"] = {"closed_by": "code", "status": "применён", "raw": "resolved by code"}
    assert ae.export_files(run, raw(doc), run.name) == {}


def test_duplicate_question_id_refuses(answered_run):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    doc["questions"].append(copy.deepcopy(doc["questions"][0]))
    with pytest.raises(ae.EvidenceError, match="неоднозначен"):
        ae.export_files(run, raw(doc), run.name)


def test_all_evidence_removed_with_real_answer_cell_cannot_become_legacy(answered_run, tmp_path):
    from cdpm import source_answer_evidence as ae, provenance
    run, doc = answered_run
    q = doc["questions"][0]
    doc["rules"] = [{"id": "R1", "template_fit": {"family": "rest_floors", "params": {
        "min_rest_outstation": {"value": 600, "unit": "minute",
                               "source": provenance.answer_source(q, "min_rest_outstation", "600 минут")}}}}]
    files = package_payloads(run, doc)
    files = {k: v for k, v in files.items() if not k.startswith(("answers/", "source-check/"))}
    with pytest.raises(ae.EvidenceError):
        import_sources(zip_body(files), tmp_path / "target/run-1")


def test_answer_cell_with_missing_question_cannot_hide_all_evidence(answered_run):
    from cdpm import source_answer_evidence as ae, provenance
    run, doc = answered_run
    q = doc["questions"][0]
    doc["rules"] = [{"id": "R1", "template_fit": {"params": {
        "p": {"value": 600, "source": provenance.answer_source(q, "p", "600 минут")}}}}]
    doc["questions"] = []
    with pytest.raises(ae.EvidenceError, match="в анкете"):
        ae.export_files(run, raw(doc), run.name)


def test_conflicting_duplicate_answer_row_refuses(answered_run):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    def duplicate(_wb, ws):
        ws.append([cell.value for cell in ws[2]])
        ws.cell(3, 6, "601 минут")
    mutate_book(run, doc, duplicate)
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


def test_identical_duplicate_row_keeps_every_source_location(answered_run):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    mutate_book(run, doc, lambda _wb, ws: ws.append([cell.value for cell in ws[2]]))
    files = ae.export_files(run, raw(doc), run.name)
    assert [r["row"] for r in json.loads(files[ae.EVIDENCE])["answers"][0]["rows"]] == [2, 3]


def test_index_only_metadata_cannot_override_book_source(answered_run, tmp_path):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    files = package_payloads(run, doc)
    index = json.loads(files[ae.EVIDENCE])
    index["answers"][0]["rows"][0]["raw"] = "601 минут"
    files[ae.EVIDENCE] = raw(index)
    with pytest.raises(ae.EvidenceError, match="текущему пакету"):
        import_sources(zip_body(files), tmp_path / "target/run-1")


@pytest.mark.parametrize("path", ["freeze/clauses/../foreign.json", "freeze/clauses/nested/foreign.json",
                                 "freeze/clauses/S1.json:foreign.json"])
def test_unsafe_clause_path_refuses_before_extraction(answered_run, tmp_path, path):
    from cdpm import source_answer_evidence as ae
    run, doc = answered_run
    files = package_payloads(run, doc)
    files[path] = raw({"clauses": []})
    target = tmp_path / "target/run-1"
    with pytest.raises(ae.EvidenceError, match="небезопасное имя"):
        import_sources(zip_body(files), target)
    assert not target.exists()
