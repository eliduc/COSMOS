# -*- coding: utf-8 -*-
"""Аудит 04.09, пункт A3 (ответы интегратора): замки приёмника ответов.

A3.1 — ответ привязан к ИЗДАНИЮ вопроса: книга несёт ревизию анкеты
(«Сводка») и отпечаток каждого вопроса (скрытая колонка); книга
ревизии N на своде ревизии N+1, переизданный вопрос и книга старого
издания без отпечатков отклоняются — ответ ложится «не разобран» с
причиной, вопрос задаётся повторно, молчаливого применения нет.
A3.2 — два разных ответа на один вопрос: ни один не принят; замечания
чтения краснят гейт ответов и видны в статусе прогона.
A3.3 — ключ с оговоркой («а, но только для дальних рейсов») машинно
не применяется: развилка и блокировки целы, ответ ждёт обработчика,
который получает «ключ (оговорка)».

Сеть и языковые модели не нужны (каналы подменные)."""
import copy
import json
import time

from openpyxl import load_workbook

from cdpm import answers as ans
from cdpm import anketa_xlsx as ax
from cdpm import model as m
from cdpm import render
from cdpm.gates import gate_anketa_xlsx, gate_answers
from cdpm.sample import sample_model
import pytest
from test_cdpm_webui import _extract_and_wait, _upload, client  # noqa: F401


@pytest.fixture(autouse=True)
def _s06r2_declared_layout(declared_test_layout):
    """S06-R2 (06.09.2026): облачные стадии панели требуют объявленной
    раскладки ролей — тесты этого файла объявляют явную тестовую
    (conftest.declared_test_layout), молчаливого облака нет."""

_XL_MIME = ("application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet")


def _fill(book, filled: dict[str, str]) -> None:
    """Заполнить колонку «ОТВЕТ АВИАКОМПАНИИ» по заголовку, как это
    делает авиакомпания (и машинная книга регрессии)."""
    wb = load_workbook(str(book))
    for ws in wb.worksheets:
        if ws.title in (ax.SUMMARY_TITLE, ax.ANSWERED_TITLE):
            continue
        col = next(j for j, c in enumerate(ws[1])
                   if str(c.value or "").strip() == ax.ANSWER_HEADER)
        for row in ws.iter_rows(min_row=2):
            if row[0].value in filled:
                row[col].value = filled[row[0].value]
    wb.save(str(book))


def _book(tmp_path, doc, filled: dict[str, str]):
    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, book)
    _fill(book, filled)
    return book


def _rule(doc, rid="R-S1-П-2-3"):
    return next(r for r in doc["rules"] if r["id"] == rid)


def _q(doc, qid="Q-2-MODAL"):
    return next(x for x in doc["questions"] if x["id"] == qid)


# ------------------------------------------------------------ A3.1


def test_book_carries_revision_and_hidden_fingerprint(tmp_path):
    doc = sample_model()
    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, book)
    wb = load_workbook(str(book))
    summary = {ws_row[0]: ws_row[1] for ws_row in
               wb[ax.SUMMARY_TITLE].iter_rows(values_only=True)
               if ws_row and ws_row[0]}
    assert summary[ax.REVISION_LABEL] == 1
    ws = wb[ax.SHEET_TITLES["2"]]
    assert ws.cell(row=1, column=ax.FINGERPRINT_COL).value \
        == ax.FINGERPRINT_HEADER
    assert ws.cell(row=1, column=ax.ANSWER_COL).value == ax.ANSWER_HEADER
    fp = ws.cell(row=2, column=ax.FINGERPRINT_COL).value
    assert fp == ax.question_fingerprint(_q(doc)) and len(fp) == 12
    assert ws.column_dimensions["G"].hidden
    # Отпечаток — чистая функция издания вопроса: та же формулировка
    # даёт тот же отпечаток, изменённый вариант — другой.
    assert ax.question_fingerprint(copy.deepcopy(_q(doc))) == fp
    changed = copy.deepcopy(_q(doc))
    changed["answer_format"]["options"][0]["text"] += " (кроме чартеров)"
    assert ax.question_fingerprint(changed) != fp
    # Гейт книги зелёный и на заполненной книге: ответ — колонка
    # авиакомпании, отпечаток сверяется как содержимое модели.
    _fill(book, {"Q-2-MODAL": "б"})
    assert gate_anketa_xlsx.check_workbook(doc, book)[0] == []


def test_book_of_previous_revision_is_refused(tmp_path):
    """Книга ревизии 1 на своде ревизии 2: ответ не применяется, а
    ложится «не разобран» с причиной, называющей ревизии; развилка
    строки остаётся."""
    doc = sample_model()
    book = _book(tmp_path, doc, {"Q-2-MODAL": "б"})
    doc["meta"]["revision"] = 2                  # свод ушёл вперёд
    got = ans.read_book(book, doc)
    assert got.answers == {}
    assert "Q-2-MODAL" in got.refused
    assert "ревизии 1" in got.refused["Q-2-MODAL"]["reason"] \
        and "ревизии 2" in got.refused["Q-2-MODAL"]["reason"]
    assert any("скачайте книгу" in p for p in got.problems)
    # Прежняя форма чтения: принятых нет, причина — в замечаниях.
    raw_map, problems = ans.read_answer_book(book, doc)
    assert raw_map == {} and problems == got.problems

    summary = ans.apply_answers(doc, got.answers, refused=got.refused)
    assert summary["received"] == 1 and summary["applied"] == 0
    assert summary["unparsed"] == 1
    assert summary["refused"] == [{"question": "Q-2-MODAL",
                                   "reason": got.refused["Q-2-MODAL"]
                                   ["reason"]}]
    q = _q(doc)
    assert q["answer"]["status"] == "не разобран"
    assert q["answer"]["raw"] == "б"
    assert "ревизии 1" in q["answer"]["note"]
    rule = _rule(doc)
    assert rule.get("conditional") and "Q-2-MODAL" in rule["blockers"]
    assert m.is_open_question(q)          # вопрос задаётся повторно
    assert m.validate_schema(doc) + m.validate_model(doc) == []
    # Та же книга на своде своей ревизии принимается — замок точный.
    fresh = sample_model()
    got = ans.read_book(book, fresh)
    assert got.answers == {"Q-2-MODAL": "б"} and got.refused == {}
    assert got.problems == []


def test_reissued_question_is_refused_by_fingerprint(tmp_path):
    """Тот же id, та же ревизия, но перечень вариантов изменён после
    издания книги: ключ «б» к новому перечню не применяется."""
    doc = sample_model()
    book = _book(tmp_path, doc, {"Q-2-MODAL": "б"})
    q = _q(doc)
    q["answer_format"]["options"][1]["text"] = "жёстко только зимой"
    got = ans.read_book(book, doc)
    assert got.answers == {}
    reason = got.refused["Q-2-MODAL"]["reason"]
    assert "переиздан" in reason and "Q-2-MODAL" in reason
    assert any("Q-2-MODAL" in p and "не принят" in p
               for p in got.problems)
    ans.apply_answers(doc, got.answers, refused=got.refused)
    assert q["answer"]["status"] == "не разобран"
    assert "переиздан" in q["answer"]["note"]
    assert _rule(doc).get("conditional")
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_legacy_book_without_fingerprints_is_refused(tmp_path):
    """Книга старого издания (без строки ревизии и колонки
    отпечатков) отклоняется целиком: без отпечатка нет доказательства,
    что ответ дан на этот вопрос."""
    doc = sample_model()
    book = _book(tmp_path, doc, {"Q-2-MODAL": "б", "Q-3a-DAY": "а"})
    wb = load_workbook(str(book))
    for ws in wb.worksheets:
        if ws.title in (ax.SUMMARY_TITLE, ax.ANSWERED_TITLE):
            continue
        ws.delete_cols(ax.FINGERPRINT_COL)
    ws = wb[ax.SUMMARY_TITLE]
    for i in range(1, ws.max_row + 1):
        if ws.cell(row=i, column=1).value == ax.REVISION_LABEL:
            ws.delete_rows(i)
            break
    wb.save(str(book))
    got = ans.read_book(book, doc)
    assert got.answers == {}
    assert set(got.refused) == {"Q-2-MODAL", "Q-3a-DAY"}
    assert all("старого издания" in r["reason"]
               for r in got.refused.values())
    assert any("старого издания" in p for p in got.problems)

    # Ревизия на месте, но колонка отпечатков удалена — отказ
    # построчно с причиной.
    (tmp_path / "b2").mkdir()
    book2 = _book(tmp_path / "b2", doc, {"Q-2-MODAL": "б"})
    wb = load_workbook(str(book2))
    wb[ax.SHEET_TITLES["2"]].delete_cols(ax.FINGERPRINT_COL)
    wb.save(str(book2))
    got = ans.read_book(book2, doc)
    assert got.answers == {}
    assert "нет отпечатка" in got.refused["Q-2-MODAL"]["reason"]


def test_refusal_does_not_overwrite_accepted_answer(tmp_path):
    """Устаревшая книга с ответом на уже принятый вопрос: принятый
    ответ не затирается отказом."""
    doc = sample_model()
    book = _book(tmp_path, doc, {"Q-2-MODAL": "б"})
    ans.apply_answers(doc, {"Q-2-MODAL": "а"})          # ревизия 2
    got = ans.read_book(book, doc)
    assert "Q-2-MODAL" in got.refused
    summary = ans.apply_answers(doc, got.answers, refused=got.refused)
    assert summary["already_answered"] == 1 and summary["unparsed"] == 0
    assert _q(doc)["answer"]["status"] == "применён"
    assert _q(doc)["answer"]["key"] == "а"


# ------------------------------------------------------------ A3.2


def test_conflicting_duplicate_answers_are_refused(tmp_path):
    """Два разных ответа на один вопрос (на разных листах книги): ни
    один не принят, статус «не разобран» с причиной; одинаковые дубли
    сливаются."""
    doc = sample_model()
    book = _book(tmp_path, doc, {"Q-2-MODAL": "б"})
    wb = load_workbook(str(book))
    src = wb[ax.SHEET_TITLES["2"]]
    dup = wb[ax.SHEET_TITLES["1"]]
    dup.append([c.value for c in src[2]])
    dup.cell(row=dup.max_row, column=ax.ANSWER_COL).value = "а"
    wb.save(str(book))
    got = ans.read_book(book, doc)
    assert got.answers == {}
    reason = got.refused["Q-2-MODAL"]["reason"]
    assert "два разных ответа" in reason and "«б»" in reason \
        and "«а»" in reason
    assert not any("взят первый" in p for p in got.problems)
    ans.apply_answers(doc, got.answers, refused=got.refused)
    assert _q(doc)["answer"]["status"] == "не разобран"
    assert "два разных ответа" in _q(doc)["answer"]["note"]
    assert _rule(doc).get("conditional")

    # Одинаковый дубль — один ответ, без замечаний.
    dup.cell(row=dup.max_row, column=ax.ANSWER_COL).value = "б"
    wb.save(str(book))
    got = ans.read_book(book, sample_model())
    assert got.answers == {"Q-2-MODAL": "б"} and got.problems == []


def test_conflict_reason_is_worded_once(tmp_path):
    """Причина противоречащих дублей в замечании чтения — без
    двойного «не принят» (ревью: «…ни один не принят, оставьте один —
    ответ не принят»)."""
    doc = sample_model()
    book = _book(tmp_path, doc, {"Q-2-MODAL": "б"})
    wb = load_workbook(str(book))
    src = wb[ax.SHEET_TITLES["2"]]
    dup = wb[ax.SHEET_TITLES["1"]]
    dup.append([c.value for c in src[2]])
    dup.cell(row=dup.max_row, column=ax.ANSWER_COL).value = "а"
    wb.save(str(book))
    got = ans.read_book(book, doc)
    assert len(got.problems) == 1
    assert got.problems[0].count("не принят") == 1
    assert "оставьте один" in got.problems[0]
    # Отказ по отпечатку — суффикс «ответ не принят» на месте.
    (tmp_path / "b2").mkdir()
    book2 = _book(tmp_path / "b2", doc, {"Q-2-MODAL": "б"})
    _q(doc)["answer_format"]["options"][1]["text"] = "иначе"
    got = ans.read_book(book2, doc)
    assert got.problems[0].count("не принят") == 1
    assert got.problems[0].endswith("— ответ не принят")


def test_foreign_edition_row_keeps_accepted_answer(tmp_path):
    """Лишняя строка того же вопроса с чужим отпечатком (скопирована
    из старой книги) не снимает принятый ответ строки текущего
    издания: принят «б», чужая строка отклонена с причиной «две
    строки … из другого издания», а не «вопрос переиздан» (книга
    свежая). Итог не зависит от порядка листов."""
    doc = sample_model()
    for foreign_sheet in ("1", "3a"):     # чужая строка ДО и ПОСЛЕ
        sub = tmp_path / foreign_sheet
        sub.mkdir()
        book = _book(sub, doc, {"Q-2-MODAL": "б"})
        wb = load_workbook(str(book))
        src = wb[ax.SHEET_TITLES["2"]]
        dup = wb[ax.SHEET_TITLES[foreign_sheet]]
        dup.append([c.value for c in src[2]])
        dup.cell(row=dup.max_row, column=ax.ANSWER_COL).value = "а"
        dup.cell(row=dup.max_row,
                 column=ax.FINGERPRINT_COL).value = "000000000000"
        wb.save(str(book))
        got = ans.read_book(book, doc)
        assert got.answers == {"Q-2-MODAL": "б"}, foreign_sheet
        assert got.refused == {}
        assert len(got.problems) == 1
        assert "две строки вопроса" in got.problems[0] \
            and "из другого издания" in got.problems[0]
        assert "переиздан" not in got.problems[0]
        assert got.problems[0].startswith(ax.SHEET_TITLES[foreign_sheet])
    # Отрицательный контроль: чужая строка ОДНА (строки текущего
    # издания нет) — прежний отказ «вопрос переиздан».
    (tmp_path / "alone").mkdir()
    book = _book(tmp_path / "alone", doc, {"Q-2-MODAL": "б"})
    wb = load_workbook(str(book))
    wb[ax.SHEET_TITLES["2"]].cell(
        row=2, column=ax.FINGERPRINT_COL).value = "000000000000"
    wb.save(str(book))
    got = ans.read_book(book, doc)
    assert got.answers == {}
    assert "переиздан" in got.refused["Q-2-MODAL"]["reason"]


def test_second_round_book_with_answered_sheet_reads_clean(tmp_path):
    """Второй круг: книга, отрендеренная из модели с принятым ответом,
    несёт штатный лист «Принятые ответы» — он не «лист вне частей»,
    замечаний чтения нет, гейт ответов зелёный (ревью: регрессия
    красного гейта со второго круга)."""
    doc = sample_model()
    ans.apply_answers(doc, {"Q-2-MODAL": "б"})           # круг 1
    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, book)
    wb = load_workbook(str(book))
    assert ax.ANSWERED_TITLE in wb.sheetnames
    assert any("ID вопроса" in str(c.value or "")
               for c in wb[ax.ANSWERED_TITLE][1])
    _fill(book, {"Q-3b-MONTH": "а"})                      # круг 2
    got = ans.read_book(book, doc)
    assert got.problems == [] and got.refused == {}
    assert got.answers == {"Q-3b-MONTH": "а"}
    doc_r2 = copy.deepcopy(doc)               # свод ревизии 2 — для контролей
    summary = ans.apply_answers(doc, got.answers, refused=got.refused)
    summary["read_problems"] = got.problems
    model_path = tmp_path / "model.json"
    summary_path = tmp_path / "answers-summary.json"
    m.dump_model(doc, model_path)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False),
                            encoding="utf-8")
    ax.render_anketa_xlsx(doc, tmp_path / "anketa2.xlsx")
    problems, counters = gate_answers.check(
        model_path, summary_path, book_path=tmp_path / "anketa2.xlsx")
    assert problems == [] and counters["revision"] == 3

    # Отрицательные контроли: чужой лист С колонкой ответа и строками
    # — замечание (живой случай «Переподача»); чужой лист без колонки
    # ответа (пометки со столбцом «ID вопроса») — не ответный.
    wb = load_workbook(str(book))
    ws = wb.create_sheet("Переподача")
    ws.append(list(ax.HEADERS))
    ws.append(["Q-1-REST", "", "", "", "", "а", ""])
    notes = wb.create_sheet("Пометки")
    notes.append(["ID вопроса", "Комментарий"])
    notes.append(["Q-1-REST", "обсудить"])
    wb.save(str(book))
    got = ans.read_book(book, doc_r2)
    assert got.answers == {"Q-3b-MONTH": "а"}
    assert len(got.problems) == 1
    assert got.problems[0].startswith("лист 'Переподача' вне частей")
    assert "Пометки" not in " ".join(got.problems)


def test_second_round_through_panel(client, tmp_path):
    """Панель: круг 1 книгой → «ответы применены»; круг 2 книгой,
    скачанной с /file/анкета-xlsx (несёт «Принятые ответы»), с
    ответом на оставшийся вопрос → снова «ответы применены», без
    замечаний чтения."""
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"

    def _open_answer(model):
        """Открытый вопрос с машинно-однозначным ответом: развилка
        условной строки — ключ ветви; условие расчёта — текст."""
        for r in model["rules"]:
            if r.get("conditional"):
                qid = r["conditional"]["question"]
                return qid, sorted(r["conditional"]["branches"])[0]
        for q in model["questions"]:
            if m.is_open_question(q) \
                    and q.get("answer_format", {}).get("type") == "choice":
                # Развилка без условной строки: ключ принят, строки
                # ждут обработчика — для проверки ЧТЕНИЯ книги этого
                # достаточно (замечаний чтения быть не должно).
                return q["id"], q["answer_format"]["options"][0]["key"]
        raise AssertionError("нет открытого вопроса-развилки")

    def _post(book):
        resp = client.post(f"/api/runs/{run_id}/answers",
                           files={"file": ("anketa.xlsx",
                                           book.read_bytes(), _XL_MIME)})
        assert resp.status_code == 200, resp.text
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            st = client.get(f"/api/runs/{run_id}").json()
            if st["state"] != "применение ответов идёт":
                return resp.json(), st
            time.sleep(0.1)
        raise AssertionError("применение ответов не завершилось")

    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    qid1, key1 = _open_answer(model)
    book1 = tmp_path / "round1.xlsx"
    book1.write_bytes(
        client.get(f"/api/runs/{run_id}/file/анкета-xlsx").content)
    _fill(book1, {qid1: key1})
    got, st = _post(book1)
    assert got["received"] == 1 and got["refused"] == 0
    assert st["state"] == "ответы применены", st
    assert st["answers_summary"]["read_problems"] == 0

    # Круг 2 — книгой, которую панель отдаёт ПОСЛЕ круга 1.
    book2 = tmp_path / "round2.xlsx"
    book2.write_bytes(
        client.get(f"/api/runs/{run_id}/file/анкета-xlsx").content)
    wb = load_workbook(str(book2))
    assert ax.ANSWERED_TITLE in wb.sheetnames
    model2 = client.get(f"/api/runs/{run_id}/file/модель").json()
    assert model2["meta"]["revision"] == 2
    qid2, key2 = _open_answer(model2)
    assert qid2 != qid1
    _fill(book2, {qid2: key2})
    got, st = _post(book2)
    assert got["received"] == 1 and got["revision"] == 3
    assert st["state"] == "ответы применены", st
    assert st["answers_problems"] == []
    assert st["answers_read_problems"] == []
    assert st["answers_summary"]["read_problems"] == 0
    assert st["answers_summary"]["revision"] == 3
    model3 = client.get(f"/api/runs/{run_id}/file/модель").json()
    q2 = next(q for q in model3["questions"] if q["id"] == qid2)
    assert q2["answer"]["key"] == key2 and q2["answer"]["revision"] == 3
    assert q2["answer"]["status"] in ("применён", "ожидает обработчика")


def test_gate_answers_red_on_read_problems(tmp_path):
    """Замечания чтения книги — нарушения гейта ответов, а не строка
    в скачиваемой сводке."""
    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "б"})
    model_path = tmp_path / "model.json"
    summary_path = tmp_path / "answers-summary.json"
    m.dump_model(doc, model_path)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False),
                            encoding="utf-8")
    problems, _ = gate_answers.check(model_path, summary_path)
    assert problems == []
    summary["read_problems"] = ["Q-3a-DAY: в книге два разных ответа "
                                "(«а» и «б») — ни один не принят"]
    summary_path.write_text(json.dumps(summary, ensure_ascii=False),
                            encoding="utf-8")
    problems, _ = gate_answers.check(model_path, summary_path)
    assert len(problems) == 1
    assert problems[0].startswith("чтение книги: Q-3a-DAY")


def test_read_problems_in_run_status_and_red_gate(client, tmp_path):
    """Панель: книга чужой ревизии принимается к разбору (не 400),
    ответ ложится «не разобран», замечание чтения видно в статусе
    прогона и краснит гейт ответов."""
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    rule = next(r for r in model["rules"] if r.get("conditional"))
    qid = rule["conditional"]["question"]
    key = sorted(rule["conditional"]["branches"])[0]
    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(model, book)
    _fill(book, {qid: key})
    wb = load_workbook(str(book))
    ws = wb[ax.SUMMARY_TITLE]
    for i in range(1, ws.max_row + 1):
        if ws.cell(row=i, column=1).value == ax.REVISION_LABEL:
            ws.cell(row=i, column=2).value = 7      # книга чужой ревизии
    wb.save(str(book))
    resp = client.post(f"/api/runs/{run_id}/answers",
                       files={"file": ("anketa.xlsx", book.read_bytes(),
                                       _XL_MIME)})
    assert resp.status_code == 200, resp.text
    # Немедленный ответ: «получено» = принятые + отклонённые — то же
    # число, что в сводке стадии.
    assert resp.json()["received"] == 1
    assert resp.json()["accepted"] == 0 and resp.json()["refused"] == 1
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] != "применение ответов идёт":
            break
        time.sleep(0.1)
    assert st["state"] == "красный гейт ответов", st
    assert st["answers_summary"]["received"] == 1
    assert st["answers_summary"]["applied"] == 0
    assert st["answers_summary"]["unparsed"] == 1
    assert st["answers_summary"]["read_problems"] == 1
    assert st["answers_summary"]["refused"] == 1
    assert any("ревизии 7" in p for p in st["answers_read_problems"])
    assert any(p.startswith("чтение книги:") for p in st["answers_problems"])
    new_model = client.get(f"/api/runs/{run_id}/file/модель").json()
    q = next(x for x in new_model["questions"] if x["id"] == qid)
    assert q["answer"]["status"] == "не разобран"
    new_rule = next(r for r in new_model["rules"] if r["id"] == rule["id"])
    assert "conditional" in new_rule           # развилка не разрешена


# ------------------------------------------------------------ A3.3


def test_caveat_on_fork_is_not_applied_machine_wise():
    """«а, но только для дальних рейсов»: ключ есть, но оговорка —
    ограничение, которое код истолковать не может; развилка и
    блокировка остаются, ответ ждёт обработчика."""
    doc = sample_model()
    assert ans.match_choice("а, но только для дальних рейсов",
                            _q(doc)["answer_format"]["options"]) \
        == ("а", "но только для дальних рейсов")
    kind_before = _rule(doc)["kind"]
    summary = ans.apply_answers(
        doc, {"Q-2-MODAL": "а, но только для дальних рейсов"})
    assert summary["applied"] == 0 and summary["awaiting"] == 1
    assert summary["resolved_rules"] == []
    q = _q(doc)
    assert q["answer"]["status"] == "ожидает обработчика"
    assert q["answer"]["key"] == "а"
    assert q["answer"]["note"] == "но только для дальних рейсов"
    rule = _rule(doc)
    assert rule["kind"] == kind_before and rule.get("conditional")
    assert "Q-2-MODAL" in rule["blockers"]
    assert m.validate_schema(doc) + m.validate_model(doc) == []
    # Открытым вопрос не считается (ответ принят, ждёт внесения), в
    # анкете следующей ревизии он в «Принятых ответах» с оговоркой.
    assert not m.is_open_question(q)
    assert "а (но только для дальних рейсов)" in render.render_anketa(doc)

    # Без оговорки — прежнее машинное применение.
    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "а"})
    assert summary["applied"] == 1 and "conditional" not in _rule(doc)


def test_restated_option_text_is_not_a_caveat():
    """«а) <текст варианта>» — пересказ перечня, не оговорка:
    применяется машинно (живые машинные книги регрессии)."""
    doc = sample_model()
    opt = _q(doc)["answer_format"]["options"][1]
    assert ans.match_choice(f"б) {opt['text']}",
                            _q(doc)["answer_format"]["options"]) \
        == ("б", "")
    assert ans.match_choice(f"б ({opt['text']})",
                            _q(doc)["answer_format"]["options"]) \
        == ("б", "")
    summary = ans.apply_answers(doc, {"Q-2-MODAL": f"б) {opt['text']}"})
    assert summary["applied"] == 1


def test_worker_receives_key_with_caveat_and_completes():
    """Обработчик получает «ключ (оговорка)» и вносит её в строку
    патчем под контрактом; после патча ответ «применён»."""
    doc = sample_model()
    ans.apply_answers(doc, {"Q-2-MODAL": "а, но только для дальних "
                                         "рейсов"})
    q = _q(doc)
    affected = ans._affected_rules(doc, q)
    assert [r["id"] for r in affected] == ["R-S1-П-2-3"]
    patched = copy.deepcopy(affected[0])
    patched["kind"] = "SOFT"
    patched["weight"] = 5
    patched["statement"] += " (только для дальних рейсов)"
    patched.pop("conditional", None)
    patched.pop("blockers", None)
    seen = {}

    def llm(prompt):
        seen["prompt"] = prompt
        return json.dumps({"patches": [
            {"question": "Q-2-MODAL", "change": "оговорка внесена",
             "rules": [patched]}]}, ensure_ascii=False)

    summary = ans.apply_worker_patches(doc, llm)
    assert "ОТВЕТ АВИАКОМПАНИИ: а (но только для дальних рейсов)" \
        in seen["prompt"]
    assert summary["worker_rejected"] == [], summary
    assert summary["worker_applied"] == 1
    assert q["answer"]["status"] == "применён"
    rule = _rule(doc)
    assert "conditional" not in rule and "Q-2-MODAL" not in \
        rule.get("blockers", [])
    assert "только для дальних рейсов" in rule["statement"]
    assert m.validate_schema(doc) + m.validate_model(doc) == []


# Оговорка = ОГРАНИЧЕНИЕ применимости: маркеры в хвосте.
_CAVEATS = (
    "а, но только для дальних рейсов",
    "б, кроме чартеров",
    "а (только зимой)",
    "а — если экипаж согласен",
    "б (при условии согласия экипажа)",
    "а, начиная с 2027 года",
    "а (до 31 декабря)",
    "б — после ввода FRMS",
    "а, в случае эстафеты",
    "а (лишь для КВС)",
    "б (за исключением праздников)",
)
# Хвост без маркеров — пояснение (обоснование выбора).
_EXPLANATIONS = (
    "а (чтение аудитора)",
    "г (баланс величин: код — справочная величина)",
    "б (приоритет у утверждённых отпусков: на конфликтный период "
    "допускается временное снижение нормы резерва, вплоть до 0)",
    "а — считать до получения данных; правило исключается из расчёта "
    "с пометкой",
    "в (справочные данные: на решение не влияют — до появления в базе "
    "шаблонов исполнителя)",
    "б (ежесуточно, не менее указанного числа — допускается больше)",
    "а (ночёвка относится к календарному дню её НАЧАЛА (вечера))",
    "а (нормально: ночёвки учтены)",          # «но» внутри слова — не маркер
    "а (верна формулировка свода)",
)
_OPTS = [{"key": k, "text": f"вариант {k}"} for k in "абвг"]


def test_caveat_is_restriction_by_markers():
    """Классификация хвоста детерминирована и не зависит от набора:
    маркеры ограничения — оговорка; без них — пояснение."""
    for raw in _CAVEATS:
        got = ans.match_choice(raw, _OPTS)
        assert got is not None and got[1], raw
        assert ans._caveat(got[1]) == got[1], raw
    for raw in _EXPLANATIONS:
        got = ans.match_choice(raw, _OPTS)
        assert got is not None and got[1], raw
        assert ans._caveat(got[1]) == "", raw
    assert ans._caveat("") == "" and ans._caveat("()") == ""


def test_explanation_tail_is_applied_with_note():
    """«а (чтение аудитора)»: ключ применяется машинно (развилка
    разрешена), пояснение ложится заметкой ответа — как прежде
    (машинные книги регрессии). Отрицательный контроль — та же
    развилка с оговоркой ждёт обработчика."""
    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "а (чтение аудитора)"})
    assert summary["applied"] == 1 and summary["awaiting"] == 0
    assert summary["resolved_rules"] == ["R-S1-П-2-3"]
    q = _q(doc)
    assert q["answer"]["status"] == "применён" and q["answer"]["key"] == "а"
    assert q["answer"]["note"] == "(чтение аудитора)"
    assert "conditional" not in _rule(doc)
    assert m.validate_schema(doc) + m.validate_model(doc) == []

    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "а (только зимой)"})
    assert summary["applied"] == 0 and summary["awaiting"] == 1
    assert _q(doc)["answer"]["status"] == "ожидает обработчика"
    assert _rule(doc).get("conditional")


def test_llm_note_is_classified_by_the_same_rule():
    """Заметка разборщика (нечёткий ответ) проходит ту же
    классификацию: пояснение — применяется с заметкой, ограничение —
    ждёт обработчика."""
    def llm_for(note):
        return lambda p: json.dumps(
            {"matches": [{"question": "Q-2-MODAL", "key": "б",
                          "note": note}]}, ensure_ascii=False)

    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "считаем жёстким"},
                                llm=llm_for("по решению директора"))
    assert summary["applied"] == 1
    assert _q(doc)["answer"]["note"] == "по решению директора"
    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "считаем жёстким"},
                                llm=llm_for("только для чартеров"))
    assert summary["awaiting"] == 1
    assert _q(doc)["answer"]["status"] == "ожидает обработчика"
