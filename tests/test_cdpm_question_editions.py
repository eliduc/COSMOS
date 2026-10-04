"""A04 scope witnesses: in-memory models and temporary books; no live channel."""
import copy
import hashlib
import json

import pytest
from openpyxl import load_workbook

from cdpm import anketa_xlsx as ax, answers, exchange, model as m, provenance
from cdpm.sample import sample_model


def _doc():
    doc = sample_model()
    q = {
        "id": "Q-3a-edition", "part": "3a", "title": "Длительность",
        "text": "Какое ограничение действует?",
        "answer_format": {"type": "value", "note": "минуты"},
        "rule_refs": [doc["rules"][1]["id"]],
        "source_refs": [{"sid": "S1", "anchor": "П.1.1"}],
    }
    doc["questions"].append(q)
    return doc, q


def _published(q, legacy=False):
    # Independent construction of the agreed snapshot, not the production
    # freezer. Reader/mutation witnesses also run on the unmodified baseline.
    fmt = q["answer_format"]
    ed = {"id": q["id"], "revision": 1,
          "title": m.normalized_question_text(q["title"]),
          "text": m.normalized_question_text(q["text"]),
          "type": fmt["type"], "options": copy.deepcopy(fmt.get("options", [])),
          "note": fmt.get("note", ""), "targets": sorted(q["rule_refs"]),
          "sources": copy.deepcopy(q["source_refs"])}
    encode = lambda x: json.dumps(x, ensure_ascii=False, sort_keys=True,
                                 separators=(",", ":"))
    parts = [ed["id"], "1", ed["title"], ed["text"], ed["type"]]
    parts += [f"{x['key']}={x['text']}" for x in ed["options"]]
    parts += [ed["note"], encode(ed["targets"]), encode(ed["sources"])]
    ed["fingerprint"] = hashlib.sha256("\x1f".join(parts).encode()).hexdigest()[:12]
    if legacy:
        ed["fingerprint_v1"] = ax.question_fingerprint(q)
    q["edition"] = ed
    return ed


def _filled(tmp_path, doc, q):
    path = tmp_path / "answers.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path)
    ws = wb[ax.SHEET_TITLES[q["part"]]]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == q["id"]:
            ws.cell(row, ax.ANSWER_COL).value = "600 минут"
    wb.save(path)
    wb.close()
    return path


def test_reader_uses_full_edition_identity():
    _, q = _doc()
    old = ax.question_fingerprint(q)
    ed = _published(q)
    assert ed["fingerprint"] != old
    assert ax.question_fingerprint(q) == ed["fingerprint"]
    assert provenance.question_edition(q)["fingerprint"] == ed["fingerprint"]


@pytest.mark.parametrize("mutation", [
    lambda q: q["rule_refs"].append("OTHER"),
    lambda q: q["answer_format"].update(note="часы"),
    lambda q: q.update(title="Другой предмет"),
    lambda q: q.update(text="Другой вопрос?"),
    lambda q: q["answer_format"].update(type="text"),
    lambda q: q["source_refs"].clear(),
    lambda q: q["edition"].update(fingerprint="000000000000"),
    lambda q: q["edition"]["targets"].append("OTHER"),
], ids=["scope", "note", "title", "text", "format", "sources", "hash", "snapshot"])
def test_reader_rejects_mutation_after_publication(tmp_path, mutation):
    doc, q = _doc()
    _published(q)
    path = _filled(tmp_path, doc, q)
    mutation(q)
    reading = answers.read_book(path, doc)
    assert q["id"] not in reading.answers
    assert q["id"] in reading.refused
    assert any("издан" in p for p in reading.problems)


def test_technical_source_addition_does_not_reissue(tmp_path):
    doc, q = _doc()
    ed = copy.deepcopy(_published(q))
    path = _filled(tmp_path, doc, q)
    result = exchange._apply_question_patch(doc, {
        "question_id": q["id"],
        "add_source_refs_technical": [{"sid": "S1", "anchor": "П.2.5"}]}, None)
    assert result == []
    assert q["edition"] == ed
    assert answers.read_book(path, doc).answers == {q["id"]: "600 минут"}


@pytest.mark.parametrize("answered", [False, True])
def test_scope_growth_preserves_old_question_and_makes_e2(tmp_path, answered):
    doc, q = _doc()
    _published(q)
    if answered:
        q["answer"] = {"raw": "600 минут", "status": "применён", "revision": 1}
    old = copy.deepcopy(q)
    other = doc["rules"][0]["id"]
    result = exchange._apply_question_patch(doc, {
        "question_id": q["id"], "add_rule_refs": [other]}, None)
    assert result == []
    assert q == old
    new = next(x for x in doc["questions"] if x["id"] == q["id"] + "-E2")
    assert new["edition_of"] == q["id"]
    assert new["edition"]["revision"] == 2
    assert new["edition"]["targets"] == sorted(q["rule_refs"] + [other])
    assert "answer" not in new
    assert ax.question_fingerprint(new) != ax.question_fingerprint(q)
    assert m.validate_schema(doc) == []
    assert m.validate_model(doc) == []
    assert q not in ax.part_questions(doc, q["part"])
    assert new in ax.part_questions(doc, q["part"])
    for rule in doc["rules"]:
        if rule["id"] in new["rule_refs"]:
            assert new["id"] in rule["blockers"]
    if answered:
        row = next(row for _, rows in ax.sheet_plan(doc) for row in rows
                   if row[0] == new["id"])
        assert "применён к изданию 1" in row[4]


def test_old_book_cannot_be_relabelled_as_e2(tmp_path):
    doc, q = _doc()
    _published(q)
    path = _filled(tmp_path, doc, q)
    assert exchange._apply_question_patch(doc, {"question_id": q["id"],
        "add_rule_refs": [doc["rules"][0]["id"]]}, None) == []
    wb = load_workbook(path)
    ws = wb[ax.SHEET_TITLES[q["part"]]]
    new_id = q["id"] + "-E2"
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == q["id"]:
            ws.cell(row, 1).value = new_id
    wb.save(path)
    wb.close()
    result = answers.read_book(path, doc)
    assert new_id not in result.answers
    assert new_id in result.refused


def test_late_blocker_cannot_extend_answer_scope():
    doc, q = _doc()
    _published(q)
    other = doc["rules"][0]
    other.setdefault("blockers", []).append(q["id"])
    before = copy.deepcopy(doc["rules"])
    assert [r["id"] for r in answers._affected_rules(doc, q)] == q["rule_refs"]
    result = answers.apply_answers(doc, {q["id"]: "600 минут"}, None)
    assert result["applied"] == 0
    assert result["refused"]
    assert doc["rules"] == before
    assert q["answer"]["status"] == "не разобран"


def test_provenance_scope_cannot_expand_from_effects():
    doc, q = _doc()
    _published(q)
    other = doc["rules"][0]["id"]
    q["answer"] = {"raw": "600 минут", "status": "применён", "revision": 1,
                   "effects": [{"rule": other, "field": "value", "after": 600}]}
    assert provenance._target_problems(doc, q, other, "", "some_param")
    assert not provenance._target_problems(doc, q, q["rule_refs"][0], "", "some_param")


def test_freeze_is_idempotent_and_legacy_identity_is_preserved():
    from cdpm import question_editions as qe
    doc, q = _doc()
    old = ax.question_fingerprint(q)
    ed = copy.deepcopy(qe.freeze(q, legacy=True))
    assert ax.question_fingerprint(q) == old
    assert provenance.question_edition(q)["fingerprint"] == old
    assert qe.freeze(q) == ed
    assert m.validate_schema(doc) == []
    q["rule_refs"].append(doc["rules"][0]["id"])
    with pytest.raises(qe.EditionRefusal):
        qe.freeze(q)
    assert q["edition"] == ed


def test_previously_printed_legacy_book_survives_explicit_migration(tmp_path):
    from cdpm import question_editions as qe
    doc, q = _doc()
    path = _filled(tmp_path, doc, q)
    old = provenance.question_edition(q)
    qe.freeze(q, legacy=True)
    assert provenance.question_edition(q) == old
    assert answers.read_book(path, doc).answers == {q["id"]: "600 минут"}


@pytest.mark.parametrize("field", ["title", "text", "note", "targets", "sources", "type", "options"])
def test_every_semantic_component_changes_new_edition_fingerprint(field):
    from cdpm import question_editions as qe
    _, q = _doc()
    old = copy.deepcopy(q)
    qe.freeze(old)
    if field in ("title", "text"):
        q[field] += " уточнение"
    elif field == "note":
        q["answer_format"]["note"] = "часы"
    elif field == "targets":
        q["rule_refs"].append("R-new")
    elif field == "sources":
        q["source_refs"].append({"sid": "S1", "anchor": "П.2.5"})
    elif field == "type":
        q["answer_format"]["type"] = "text"
    else:
        q["answer_format"]["options"] = [{"key": "а", "text": "600 минут"}]
    qe.freeze(q)
    assert ax.question_fingerprint(q) != ax.question_fingerprint(old)


def test_chain_is_explicit_and_stale_patch_cannot_fork():
    from cdpm import question_editions as qe
    doc, q = _doc()
    qe.freeze(q)
    second = qe.expand(doc, q, [doc["rules"][0]["id"]])
    before = copy.deepcopy(doc)
    with pytest.raises(qe.EditionRefusal, match="уже переиздан"):
        qe.expand(doc, q, [doc["rules"][2]["id"]])
    assert doc == before
    third = qe.expand(doc, second, [doc["rules"][2]["id"]])
    assert third["id"] == q["id"] + "-E3"
    assert third["edition_of"] == second["id"]
    assert third["edition"]["revision"] == 3


@pytest.mark.parametrize("edition", [None, [], {}, {"revision": True}, {"revision": 0},
                                    {"revision": 1, "options": 4}])
def test_malformed_snapshot_is_named_not_repaired(edition):
    from cdpm import question_editions as qe
    _, q = _doc()
    q["edition"] = edition
    assert qe.problems(q)
    with pytest.raises(qe.EditionRefusal):
        ax.question_fingerprint(q)


def test_snapshot_has_no_shared_mutable_lists():
    from cdpm import question_editions as qe
    _, q = _doc()
    qe.freeze(q)
    ed = copy.deepcopy(q["edition"])
    q["source_refs"][0]["anchor"] = "другой"
    q["rule_refs"].append("другой")
    assert q["edition"] == ed
    assert qe.problems(q)


def test_exchange_trace_and_diff_include_only_real_edition_changes():
    doc, q = _doc()
    _published(q)
    answer = {"remark": "A-1", "verdict": "ПРИНЯТО", "basis": "источник",
              "question_patch": {"question_id": q["id"],
                                 "add_rule_refs": [doc["rules"][0]["id"]]}}
    pack = {"round": 1, "remarks": [{"id": "A-1", "target": {"question": q["id"]},
                                    "defect_class": "потерянная трассировка"}]}
    executed = set()
    new, accepted, rejected, _ = exchange._apply_patches(
        doc, pack, {"round": 1, "answers": [answer]}, 1, executed_out=executed)
    assert not rejected
    assert executed == {"A-1"}
    assert exchange.executed_by_trace(doc, new, [answer]) == {"A-1"}
    assert not exchange.check_diff_only(doc, new, accepted)
    new["rules"][-1]["note"] = "незаявленная правка"
    assert exchange.check_diff_only(doc, new, accepted)


def test_edition_tag_does_not_legalize_unrelated_duplicate():
    from cdpm import question_editions as qe
    doc, q = _doc()
    qe.freeze(q)
    other = copy.deepcopy(q)
    other.pop("edition")
    other["id"] = "Q-3a-unrelated"
    qe.freeze(other)
    doc["questions"].append(other)
    assert any("дубль вопроса" in p for p in m.validate_model(doc))
    other["edition_of"] = q["id"]
    assert m.validate_model(doc)


def test_invalid_growth_is_atomic_and_noop_has_no_trace():
    doc, q = _doc()
    _published(q)
    before = copy.deepcopy(doc)
    assert exchange._apply_question_patch(doc, {"question_id": q["id"],
        "add_rule_refs": ["R-absent"]}, None)
    assert doc == before
    assert exchange._apply_question_patch(doc, {"question_id": q["id"],
        "add_rule_refs": list(q["rule_refs"])}, None) == [exchange.NO_TRACE_REASON]
    assert doc == before
