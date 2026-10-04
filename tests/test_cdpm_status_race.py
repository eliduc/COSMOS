# -*- coding: utf-8 -*-
"""Гонка статуса прогона: фоновая стадия против действий пользователя.

Живой случай Prod 2026-09-04 19:12–19:27, run-9009 (набор
open-bench/polyaris-h0k): пользователь нажал «Прогон целиком», во время
извлечения загрузил книгу данных ID-H0-DATA.xlsx (панель показала
«файлов 1, листов 6, строк 80, данные заморожены») и поставил ей
политику «облако». После завершения цикла статус показал data_files
пусто — «файлов пока нет». Причина: фоновая нить стадии брала снимок
status.json в начале и в конце писала его ЦЕЛИКОМ, затирая ключи
линии данных, записанные приёмником параллельно.

Принцип после правки: статус прогона — набор ключей, у каждого свой
владелец (стадия или приёмник); ЛЮБАЯ запись — слияние по ключам
(read-modify-write под замком нити + файловой блокировкой
status.lock); замена целиком — только рождение прогона и сброс.
Гейт ниже судит и код (ast: вне двух помощников вызовов _write_status
нет), и поведение (четыре стадии в фоне против действий пользователя).
"""
import ast
import inspect
import json
import threading
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from cdpm import webui
from test_cdpm_extract import fake_llm
from test_cdpm_webui import (_allow_cloud, _extract_and_wait, _slow_llm,
                             _upload, _xlsx_bytes)

_XL_MIME = ("application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet")
_DATA_BOOK = "ID-H0-DATA.xlsx"
_TITLE = "H0-К: гонка стадии и загрузки"


@pytest.fixture(autouse=True)
def _s06r2_declared_layout(declared_test_layout):
    """S06-R2 (06.09.2026): облачные стадии панели требуют объявленной
    раскладки ролей — тесты этого файла объявляют явную тестовую
    (conftest.declared_test_layout), молчаливого облака нет."""


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("CDPM_RUNS_DIR", str(tmp_path / "runs"))
    monkeypatch.setenv("CDPM_SETTINGS_FILE",
                       str(tmp_path / "cdpm-settings.json"))
    monkeypatch.setattr(webui, "llm_factory", lambda scope: fake_llm)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: fake_llm)
    return TestClient(webui.app)


def _wait_while(client, run_id, running: str, key: str = "state",
                timeout_s: float = 60.0) -> dict:
    """Ожидание конца стадии: опрос до выхода из состояния «… идёт»."""
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st.get(key) != running:
            return st
        time.sleep(0.05)
    raise AssertionError(f"стадия «{running}» не завершилась за "
                         f"{timeout_s} с")


def _data_book() -> bytes:
    return _xlsx_bytes({"Ростер": [["Табельный", "Дата"],
                                   ["EMP-1", "05.03.2026"],
                                   ["EMP-2", "07.03.2026"]]})


def _user_acts_during_stage(client, run_id: str) -> None:
    """То, что делал пользователь во время стадии: загрузка книги
    данных, политика «облако» ей, раскладка ролей, название прогона.
    Каждое действие обязано пройти (стадия линии правил линию данных
    не занимает)."""
    r = client.post(f"/api/runs/{run_id}/data", files=[
        ("files", (_DATA_BOOK, _data_book(), _XL_MIME))])
    assert r.status_code == 200, r.text
    assert r.json()["state"] == "данные заморожены"
    assert r.json()["counters"] == {"files": 1, "sheets": 1, "rows": 2}
    cfg = client.get(f"/api/runs/{run_id}/files-config").json()["files"]
    assert _DATA_BOOK in cfg, cfg
    cfg[_DATA_BOOK] = {"processing": "cloud", "note": "книга данных"}
    r = client.put(f"/api/runs/{run_id}/files-config",
                   json={"files": cfg})
    assert r.status_code == 200, r.text
    # Путём интерфейса (S06-R2): попап роли шлёт ВСЮ действующую
    # раскладку с правкой одной роли поверх — снимок прогона полный,
    # прочие облачные роли остаются объявленными.
    layout = client.get(f"/api/runs/{run_id}/agents").json()
    r = client.put(f"/api/runs/{run_id}/agents", json={
        **layout,
        "data_compiler": {"channels": {"cloud": {
            "provider": "cloud", "model": "claude-opus-5"}}}})
    assert r.status_code == 200, r.text
    r = client.put(f"/api/runs/{run_id}/title", json={"title": _TITLE})
    assert r.status_code == 200, r.text


def _assert_user_acts_survived(client, run_id: str) -> None:
    """После стадии всё сделанное пользователем на месте: статус,
    ручка линии данных, политика файла, раскладка ролей, название."""
    st = client.get(f"/api/runs/{run_id}").json()
    assert st.get("data_files") == [_DATA_BOOK], \
        f"data_files затёрты стадией: {st}"
    assert st.get("data_state") == "данные заморожены", st
    assert st["data_freeze_counters"]["files"] == 1
    assert st["data_freeze_problems"] == []
    assert st.get("title") == _TITLE, st
    dl = client.get(f"/api/runs/{run_id}/data").json()
    assert dl["files"] == [_DATA_BOOK]
    assert dl["state"] == "данные заморожены"
    cfg = client.get(f"/api/runs/{run_id}/files-config").json()["files"]
    assert cfg[_DATA_BOOK] == {"processing": "cloud",
                               "note": "книга данных"}
    ag = client.get(f"/api/runs/{run_id}/agents").json()
    assert ag["data_compiler"]["channels"]["cloud"]["model"] \
        == "claude-opus-5"
    run_dir = Path(webui.runs_root()) / run_id
    raw = json.loads((run_dir / "status.json").read_text(encoding="utf-8"))
    assert raw["data_files"] == [_DATA_BOOK]
    assert raw["title"] == _TITLE


# ─────────────────────────── стадии линии правил ──────────────────────────

def test_extract_in_background_keeps_parallel_data_upload(client,
                                                          monkeypatch):
    """Живой случай run-9009: извлечение в фоне, во время него —
    загрузка книги данных, политика «облако», раскладка ролей,
    название. После «извлечено» ничего из этого не потеряно."""
    run_id = _upload(client).json()["run"]
    reached, release = threading.Event(), threading.Event()
    slow = _slow_llm(reached, release)
    monkeypatch.setattr(webui, "llm_factory", lambda scope: slow)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: slow)
    r = client.post(f"/api/runs/{run_id}/extract",
                    params={"customer": "Полярис (тест)", "scope": "any"})
    assert r.status_code == 200, r.text
    assert reached.wait(10), "стадия не дошла до модели"
    _user_acts_during_stage(client, run_id)
    release.set()
    st = _wait_while(client, run_id, "извлечение идёт")
    assert st["state"] == "извлечено", st
    assert st["extract_summary"]["rules"] == 6
    assert "progress" not in st
    _assert_user_acts_survived(client, run_id)


def test_full_cycle_in_background_keeps_parallel_data_upload(client,
                                                             monkeypatch):
    """«Прогон целиком» = извлечение + сквозной цикл: цикл в фоне, во
    время него — те же действия пользователя; после финала цикла
    книга данных и её политика на месте."""
    import re as _re
    from test_cdpm_exchange_run import (_make_auditor, _make_worker,
                                        _REMARK_DISTORT)
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    _allow_cloud(client, run_id)
    auditor = _make_auditor([[dict(_REMARK_DISTORT)]] * 12)
    worker = _make_worker({})          # всегда отклоняет — тупик

    def arbiter(prompt):
        ids = _re.findall(r"Тупик \[([^\]]+)\]", prompt)
        return json.dumps({"verdicts": [
            {"deadlock": i, "verdict": "обработчик прав",
             "basis": "по клаузе источника"} for i in ids]},
            ensure_ascii=False)

    reached, release = threading.Event(), threading.Event()

    def route(prompt):
        reached.set()
        release.wait(10)
        if "Тупик [" in prompt:
            return arbiter(prompt)
        if "Клаузы документа" in prompt:
            return auditor(prompt)
        return worker(prompt)

    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: route)
    r = client.post(f"/api/runs/{run_id}/full-cycle",
                    params={"auditor": "cloud", "worker": "cloud",
                            "max_rounds": 4, "max_megacycles": 2})
    assert r.status_code == 200, r.text
    assert reached.wait(10), "цикл не дошёл до модели"
    _user_acts_during_stage(client, run_id)
    release.set()
    st = _wait_while(client, run_id, "сквозной цикл идёт", timeout_s=90)
    assert st["state"] == "сквозной цикл: чисто по Б/К (кругов: 2)", st
    assert st["arbiter_summary"]["by_verdict"].get("обработчик прав") == 1
    _assert_user_acts_survived(client, run_id)


def test_answers_in_background_keeps_parallel_data_upload(client,
                                                          monkeypatch,
                                                          tmp_path):
    """Применение ответов анкеты в фоне против загрузки данных."""
    from openpyxl import load_workbook
    from cdpm import anketa_xlsx as ax
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    model = client.get(f"/api/runs/{run_id}/file/модель").json()
    rule = next(r for r in model["rules"] if r.get("conditional"))
    qid = rule["conditional"]["question"]
    key = sorted(rule["conditional"]["branches"])[0]
    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(model, book)
    wb = load_workbook(str(book))
    for ws in wb.worksheets:
        if ws.title in (ax.SUMMARY_TITLE, ax.ANSWERED_TITLE):
            continue
        for row in ws.iter_rows(min_row=2):
            if row[0].value == qid:
                row[5].value = key
    wb.save(str(book))

    reached, release = threading.Event(), threading.Event()
    orig = webui.answers_mod.apply_answers

    def slow_apply(*a, **k):
        reached.set()
        release.wait(10)
        return orig(*a, **k)
    monkeypatch.setattr(webui.answers_mod, "apply_answers", slow_apply)
    r = client.post(f"/api/runs/{run_id}/answers",
                    files={"file": ("anketa.xlsx", book.read_bytes(),
                                    _XL_MIME)})
    assert r.status_code == 200, r.text
    assert reached.wait(10), "стадия ответов не началась"
    _user_acts_during_stage(client, run_id)
    release.set()
    st = _wait_while(client, run_id, "применение ответов идёт")
    assert st["state"] == "ответы применены", st
    assert st["answers_summary"]["applied"] == 1
    _assert_user_acts_survived(client, run_id)


# ─────────────────────────── стадия линии данных ──────────────────────────

def test_data_run_in_background_keeps_parallel_rules_line_edits(
        client, monkeypatch):
    """Обратное направление: проверки данных в фоне, во время них
    пользователь правит линию правил (название, политика документа,
    признак комплекта). Финал проверок эти ключи не затирает, а свои
    (сводка отчёта) — пишет."""
    import re as _re
    run_id = _upload(client).json()["run"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    r = client.post(f"/api/runs/{run_id}/data", files=[
        ("files", (_DATA_BOOK, _data_book(), _XL_MIME))])
    assert r.status_code == 200, r.text
    reached, release = threading.Event(), threading.Event()

    def compiler(prompt):
        ids = _re.findall(r"^(R-\S+) \|", prompt, _re.M)
        return json.dumps({"checks": [], "skipped": [
            {"rule": rid, "reason": "нет данных для проверки"}
            for rid in ids]}, ensure_ascii=False)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: compiler)
    # Точка блокировки — сама компиляция: на тестовом своде все
    # строки блокированы вопросами либо судятся решением, до модели
    # компилятор не доходит вовсе.
    orig = webui.data_checks_mod.compile_checks

    def slow_compile(*a, **k):
        reached.set()
        release.wait(10)
        return orig(*a, **k)
    monkeypatch.setattr(webui.data_checks_mod, "compile_checks",
                        slow_compile)
    assert client.post(f"/api/runs/{run_id}/data-run"
                       ).status_code == 200
    assert reached.wait(10), "проверки не дошли до компилятора"
    # Линия правил во время проверок данных.
    assert client.put(f"/api/runs/{run_id}/title",
                      json={"title": _TITLE}).status_code == 200
    cfg = client.get(f"/api/runs/{run_id}/files-config").json()["files"]
    cfg["RULES-A.md"] = {"processing": "cloud", "note": "и1"}
    assert client.put(f"/api/runs/{run_id}/files-config",
                      json={"files": cfg}).status_code == 200
    assert client.put(f"/api/runs/{run_id}/kit-scope", json={
        "data_class": "synthetic", "llm_scope": "any",
        "llm_scope_basis": ""}).status_code == 200
    release.set()
    st = _wait_while(client, run_id, "проверки идут", key="data_state")
    assert st["data_state"] == "отчёт готов", st
    assert st["data_summary"]["rows"] == 2
    assert st["title"] == _TITLE, st
    assert st["kit"]["llm_scope"] == "any"
    cfg = client.get(f"/api/runs/{run_id}/files-config").json()["files"]
    assert cfg["RULES-A.md"] == {"processing": "cloud", "note": "и1"}
    # Незаписанный файл — умолчание по допуску комплекта («any» →
    # «cloud», 05.09).
    assert cfg[_DATA_BOOK]["processing"] == "cloud"


# ─────────────────────────── примитив слияния ─────────────────────────────

def test_update_status_merges_under_concurrent_writers(tmp_path):
    """Восемь нитей по 40 обновлений своих ключей + ключ-основа: в
    итоге у каждой нити последнее значение, основа цела, JSON не
    рван (каждая запись — слияние под замком)."""
    run_dir = tmp_path / "run-0001"
    run_dir.mkdir()
    webui._write_status(run_dir, {"state": "заморожено",
                                  "files": ["A.md"]})
    n_threads, n_updates = 8, 40
    errors: list[BaseException] = []

    def hammer(i: int):
        try:
            for k in range(n_updates):
                webui._update_status(run_dir, **{f"k{i}": k})
                if k % 10 == 0:
                    webui._mutate_status(
                        run_dir, lambda st: st.setdefault(
                            "touched", []).append(i))
        except BaseException as e:                     # noqa: BLE001
            errors.append(e)
    threads = [threading.Thread(target=hammer, args=(i,))
               for i in range(n_threads)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(30)
    assert not errors, errors
    st = json.loads((run_dir / "status.json").read_text(encoding="utf-8"))
    assert st["state"] == "заморожено" and st["files"] == ["A.md"]
    for i in range(n_threads):
        assert st[f"k{i}"] == n_updates - 1
    assert sorted(st["touched"]) == sorted(
        [i for i in range(n_threads) for _ in range(4)])
    # None удаляет ключ; замена целиком — отдельный помощник.
    webui._update_status(run_dir, k0=None)
    assert "k0" not in json.loads(
        (run_dir / "status.json").read_text(encoding="utf-8"))
    webui._replace_status(run_dir, {"state": "новый"})
    assert json.loads((run_dir / "status.json").read_text(
        encoding="utf-8")) == {"state": "новый"}
    assert (run_dir / "status.lock").is_file()


def test_status_lock_is_reentrant_and_cross_process_file_lock(tmp_path):
    """Вложенное слияние из-под замка не виснет (RLock + глубина файла);
    файловая блокировка держится, пока идёт слияние: другой процесс
    (здесь — попытка взять status.lock чужим дескриптором) ждёт."""
    run_dir = tmp_path / "run-0002"
    run_dir.mkdir()
    webui._write_status(run_dir, {"state": "новый"})

    def nested(st):
        st["outer"] = 1
        webui._update_status(run_dir, inner=1)   # вложенный вход
    webui._mutate_status(run_dir, nested)
    st = json.loads((run_dir / "status.json").read_text(encoding="utf-8"))
    # Внешнее слияние пишет ПОСЛЕ вложенного тем же снимком — ключ
    # вложенного обязан выжить: снимок перечитывается перед записью.
    assert st["outer"] == 1 and st["inner"] == 1

    held, proceed = threading.Event(), threading.Event()

    def holder(st):
        held.set()
        proceed.wait(5)
    t = threading.Thread(target=webui._mutate_status,
                         args=(run_dir, holder))
    t.start()
    assert held.wait(5)
    fd = webui._open_lock_fd(run_dir)
    try:
        with pytest.raises(TimeoutError):
            webui._lock_fd(fd, timeout_s=0.3)
    finally:
        webui._close_lock_fd(fd)
    proceed.set()
    t.join(5)
    fd = webui._open_lock_fd(run_dir)
    try:
        webui._lock_fd(fd, timeout_s=1.0)   # после отпускания — свободно
        webui._unlock_fd(fd)
    finally:
        webui._close_lock_fd(fd)


def test_status_writes_only_through_merge_helpers():
    """Гейт кода: замена status.json целиком (_write_status) дозволена
    только помощникам слияния/замены; стадии и приёмники пишут
    слиянием по ключам. Судится по ast, не по grep."""
    src = inspect.getsource(webui)
    tree = ast.parse(src)
    offenders: list[str] = []
    for node in tree.body:
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        for sub in ast.walk(node):
            if isinstance(sub, ast.Call) and \
                    isinstance(sub.func, ast.Name) and \
                    sub.func.id == "_write_status":
                offenders.append(f"{node.name}:{sub.lineno}")
    allowed = {"_mutate_status", "_replace_status"}
    bad = [o for o in offenders if o.split(":")[0] not in allowed]
    assert not bad, ("прямая замена статуса вне помощников слияния: "
                     + ", ".join(bad))
    assert {o.split(":")[0] for o in offenders} == allowed


def test_stage_start_is_atomic_against_paused_stage(client, monkeypatch):
    """Стадия на паузе жива (продолжится): запуск другой стадии той же
    линии — 409, а не вторая нить поверх первой (прежде «на паузе: …»
    не считалось занятостью — второй запуск извлечения при паузе
    стартовал параллельную нить)."""
    run_id = _upload(client).json()["run"]
    reached, release = threading.Event(), threading.Event()
    slow = _slow_llm(reached, release)
    monkeypatch.setattr(webui, "llm_factory", lambda scope: slow)
    monkeypatch.setattr(webui, "exchange_llm_factory",
                        lambda scope, requested: slow)
    assert client.post(f"/api/runs/{run_id}/extract",
                       params={"customer": "т", "scope": "any"}
                       ).status_code == 200
    assert reached.wait(10)
    assert client.post(f"/api/runs/{run_id}/control",
                       params={"command": "pause"}).status_code == 200
    release.set()
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] == "на паузе: извлечение":
            break
        time.sleep(0.05)
    assert st["state"] == "на паузе: извлечение", st
    r = client.post(f"/api/runs/{run_id}/extract",
                    params={"customer": "т", "scope": "any"})
    assert r.status_code == 409, r.text
    assert "на паузе" in r.json()["detail"]
    assert client.post(f"/api/runs/{run_id}/control",
                       params={"command": "resume"}).status_code == 200
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        st = client.get(f"/api/runs/{run_id}").json()
        if st["state"] not in ("извлечение идёт", "на паузе: извлечение"):
            break
        time.sleep(0.05)
    assert st["state"] == "извлечено", st


def test_orphan_recovery_covers_paused_states(tmp_path, monkeypatch):
    """Гибель процесса на паузе: «на паузе: <стадия>» обеих линий →
    «остановлено: <стадия>» (иначе занятость паузой держала бы 409
    навсегда)."""
    monkeypatch.setenv("CDPM_RUNS_DIR", str(tmp_path))
    run_dir = tmp_path / "run-0007"
    run_dir.mkdir()
    webui._write_status(run_dir, {"state": "на паузе: аудит-цикл",
                                  "data_state": "на паузе: проверки данных",
                                  "error": "давний отказ"})
    got = webui.recover_orphaned_runs()
    assert len(got) == 2, got
    st = json.loads((run_dir / "status.json").read_text(encoding="utf-8"))
    assert st["state"] == "остановлено: аудит-цикл"
    assert st["data_state"] == "остановлено: проверки данных"
    assert st["error"] == "давний отказ"


# ─────────────────────────── политика файлов ──────────────────────────────

def test_files_config_put_merges_and_keeps_other_entries(client):
    """PUT политики — слияние по файлам: запись, не пришедшая в теле,
    остаётся (интерфейс шлёт то, что видел; если в его снимке файла
    данных не было, политика файла не должна пропасть). Удаляет запись
    только каскад удаления файла."""
    run_id = _upload(client).json()["run"]
    r = client.post(f"/api/runs/{run_id}/data", files=[
        ("files", (_DATA_BOOK, _data_book(), _XL_MIME))])
    assert r.status_code == 200, r.text
    assert client.put(f"/api/runs/{run_id}/files-config", json={
        "files": {_DATA_BOOK: {"processing": "cloud",
                               "note": "книга"}}}).status_code == 200
    # Тело только с документом правил — политика книги цела.
    assert client.put(f"/api/runs/{run_id}/files-config", json={
        "files": {"RULES-A.md": {"processing": "cloud", "note": ""}}}
        ).status_code == 200
    cfg = client.get(f"/api/runs/{run_id}/files-config").json()["files"]
    assert cfg[_DATA_BOOK] == {"processing": "cloud", "note": "книга"}
    assert cfg["RULES-A.md"]["processing"] == "cloud"
    # Незаписанный RULES-B.md — умолчание по допуску комплекта
    # (синтетический → «cloud», 05.09), запись о нём не заводится.
    assert cfg["RULES-B.md"]["processing"] == "cloud"
    run_dir = Path(webui.runs_root()) / run_id
    raw = json.loads((run_dir / "files-config.json")
                     .read_text(encoding="utf-8"))["files"]
    assert set(raw) == {_DATA_BOOK, "RULES-A.md"}
    # Каскад удаления файла снимает запись.
    assert client.delete(f"/api/runs/{run_id}/data/{_DATA_BOOK}"
                         ).status_code == 200
    raw = json.loads((run_dir / "files-config.json")
                     .read_text(encoding="utf-8"))["files"]
    assert set(raw) == {"RULES-A.md"}


# ─────────────────────────── книга анкеты для заполнения ──────────────────

def test_anketa_xlsx_is_downloadable_after_extraction(client):
    """Путь человека к книге анкеты: до извлечения — 404 «не
    построен»; после — скачивание xlsx с именем файла (md-анкета
    по-прежнему отдаётся html-представлением)."""
    from io import BytesIO
    from openpyxl import load_workbook
    from cdpm import anketa_xlsx as ax
    run_id = _upload(client).json()["run"]
    r = client.get(f"/api/runs/{run_id}/file/анкета-xlsx")
    assert r.status_code == 404 and "не построен" in r.json()["detail"]
    st = _extract_and_wait(client, run_id)
    assert st["state"] == "извлечено"
    r = client.get(f"/api/runs/{run_id}/file/анкета-xlsx")
    assert r.status_code == 200, r.text[:200]
    assert r.headers["content-type"].startswith(_XL_MIME)
    cd = r.headers["content-disposition"]
    assert "attachment" in cd and f"anketa-{run_id}.xlsx" in cd, cd
    assert r.content[:2] == b"PK"
    wb = load_workbook(BytesIO(r.content))
    assert ax.SUMMARY_TITLE in wb.sheetnames
    html = client.get(f"/api/runs/{run_id}/file/анкета")
    assert html.status_code == 200
    assert "text/html" in html.headers["content-type"]
    # Прогон, извлечённый до этой правки (книги на диске нет, модель
    # есть): книга достраивается из модели по требованию.
    run_dir = Path(webui.runs_root()) / run_id
    (run_dir / "extract" / "anketa.xlsx").unlink()
    r = client.get(f"/api/runs/{run_id}/file/анкета-xlsx")
    assert r.status_code == 200, r.text[:200]
    assert (run_dir / "extract" / "anketa.xlsx").is_file()
    assert load_workbook(BytesIO(r.content)).sheetnames == wb.sheetnames
