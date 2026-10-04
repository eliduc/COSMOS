"""Actual saved workbook, including Accepted Answers; no services."""
from openpyxl import load_workbook

from cdpm import answers, anketa_xlsx as ax, history_questions as hq, render
from test_cdpm_history_editions import _published, _latest
from test_cdpm_history_1_17_2 import _book


def _closed():
    doc, q = _published()
    answers.apply_answers(doc, {q["id"]: "б"})
    next(r for r in doc["rules"] if r["id"] == "R-H-REST")["kind"] = "OUT"
    hq.ensure_history_questions(doc, _book())
    return doc, _latest(doc)


def _saved_rows(tmp_path, doc):
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [(sheet.title, row) for sheet in wb for row in sheet.iter_rows(values_only=True)]
    finally:
        wb.close()


def test_closed_edition_reason_already_reaches_actual_workbook(tmp_path):
    doc, q = _closed()
    reason = q["answer"]["note"]
    assert reason in render.render_anketa(doc)
    rows = _saved_rows(tmp_path, doc)
    found = [(sheet, row) for sheet, row in rows if row[0] == q["id"]]
    assert len(found) == 1
    assert found[0][0] == ax.ANSWERED_TITLE
    assert any(reason in str(cell) for cell in found[0][1])


def test_closed_edition_links_old_answer_on_both_surfaces(tmp_path):
    doc, q = _closed()
    note = "ответ «б» применён к изданию 1"
    assert note in render.render_anketa(doc)
    row = next(row for _, row in _saved_rows(tmp_path, doc) if row[0] == q["id"])
    assert any(note in str(cell) for cell in row)


def test_summary_explains_editions_are_not_resolved_subjects():
    doc, _ = _closed()
    text = " ".join(str(cell) for row in ax.summary_rows(doc) for cell in row)
    assert "издания вопросов, а не число улаженных предметов" in text


def test_real_book_gate_uses_same_published_trace_and_catches_its_removal(tmp_path):
    from cdpm.gates.gate_anketa_xlsx import check_workbook
    from test_cdpm_history_1_17_2 import H01_JOURNAL
    doc, q = _published()
    answers.apply_answers(doc, {q["id"]: "б"})
    hq.ensure_history_questions(doc, _book(journal_rows=H01_JOURNAL))
    new = _latest(doc)
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    assert check_workbook(doc, path)[0] == []
    wb = load_workbook(path)
    ws = wb[ax.SHEET_TITLES[new["part"]]]
    found = False
    for row in ws.iter_rows(min_row=2):
        if row[0].value == new["id"]:
            assert "применён к изданию 1" in row[4].value
            row[4].value = row[4].value.split("; ответ «")[0]
            found = True
    assert found
    wb.save(path)
    wb.close()
    assert check_workbook(doc, path)[0]
