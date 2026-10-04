# -*- coding: utf-8 -*-
"""Заморозка источников CDPM (П-2): извлечение клауз из docx/xlsx/md/doc,
сохранение смысловой разметки, детерминизм, манифест, гейт.

Ключевые охраняемые уроки: жирный шрифт несёт норму (П.1.9 — потеря
разметки при конвертации стоила аудита), нарезка на уровне клауз со
своим якорем у каждой, вложенные таблицы не теряются, `.doc` без Word —
честный отказ с названной причиной, а не тихий пропуск."""
import json
import subprocess
from pathlib import Path

import pytest

from cdpm import freeze
from cdpm.gates import gate_freeze


# ---------------------------------------------------------------------------
# Изготовленные фикстуры (вымышленное содержимое, У-02)
# ---------------------------------------------------------------------------

def _make_docx(path: Path) -> None:
    import docx
    d = docx.Document()
    d.add_paragraph("1.1. Пейринг не длиннее четырёх дней.")
    p = d.add_paragraph()
    p.add_run("1.9. Отдых на базе ")
    p.add_run("не менее 10 часов").bold = True
    p.add_run(" между сменами.")
    d.add_paragraph("")                       # пустой абзац — пропускается
    d.add_paragraph("Преамбула без номера пункта.")
    d.add_paragraph("1.1. Повтор номера пункта в другом разделе.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Направление"
    t.cell(0, 1).paragraphs[0].add_run("Норма экипажа").bold = True
    t.cell(1, 0).text = "Север"
    t.cell(1, 1).text = "2 пилота"
    nested = t.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "вложенная норма"
    d.save(str(path))


def _make_xlsx(path: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Журнал"
    ws["A1"] = "Замечание"
    ws["B1"] = "Статус"
    ws["A1"].font = openpyxl.styles.Font(b=True)
    ws["A3"] = "Кольцо длиннее нормы"          # строка 2 пустая
    ws["B3"] = "ЗАКРЫТО"
    ws2 = wb.create_sheet("Лист2")
    ws2["A1"] = "Записей нет"
    wb.save(str(path))


_MD_TEXT = """# Дополнение к правилам

2.4. Ночной рейс требует усиленного состава.

Свободная строка без номера.
"""


def _fixture_files(tmp_path: Path) -> list[Path]:
    docx_p = tmp_path / "PRAVILA-OM-A.docx"
    xlsx_p = tmp_path / "ZHURNAL-OP.xlsx"
    md_p = tmp_path / "DOPOLNENIE.md"
    _make_docx(docx_p)
    _make_xlsx(xlsx_p)
    md_p.write_text(_MD_TEXT, encoding="utf-8")
    return [docx_p, xlsx_p, md_p]


# ---------------------------------------------------------------------------
# Извлечение клауз
# ---------------------------------------------------------------------------

def test_docx_clauses_anchors_bold_and_nested_tables(tmp_path):
    p = tmp_path / "a.docx"
    _make_docx(p)
    clauses = freeze.extract_file(p)["clauses"]
    by_anchor = {c["anchor"]: c for c in clauses}
    # Нумерованный пункт становится адресуемым якорем.
    assert "п.1.1" in by_anchor
    # Жирный фрагмент внутри пункта сохранён маркерами (урок П.1.9).
    assert "**не менее 10 часов**" in by_anchor["п.1.9"]["text"]
    # Абзац без номера получает позиционный якорь.
    assert any(a.startswith("абз.") for a in by_anchor)
    # Повтор номера пункта не склеивает две клаузы в один адрес.
    assert "п.1.1 (2)" in by_anchor
    # Таблица извлечена поячеечно, жирная ячейка сохранила маркер.
    assert by_anchor["табл.1, строка 1, кол 2"]["text"] == \
        "**Норма экипажа**"
    # Вложенная таблица не потеряна.
    assert any("вл.табл.1" in a for a in by_anchor)
    # Пустые абзацы не создают клауз.
    assert all(c["text"].strip() for c in clauses)


def test_xlsx_clauses_rows_sheets_and_bold(tmp_path):
    p = tmp_path / "b.xlsx"
    _make_xlsx(p)
    clauses = freeze.extract_file(p)["clauses"]
    by_anchor = {c["anchor"]: c for c in clauses}
    assert by_anchor["лист «Журнал», строка 1"]["text"] == \
        "**Замечание** | Статус"
    # Пустая строка 2 пропущена, номер строки Excel сохранён в якоре.
    assert "лист «Журнал», строка 3" in by_anchor
    assert "лист «Лист2», строка 1" in by_anchor


def test_md_clauses_numbered_and_positional(tmp_path):
    p = tmp_path / "c.md"
    p.write_text(_MD_TEXT, encoding="utf-8")
    clauses = freeze.extract_file(p)["clauses"]
    by_anchor = {c["anchor"]: c for c in clauses}
    assert "п.2.4" in by_anchor
    # Позиционный якорь — номер СТРОКИ файла (пустые строки считаются).
    assert by_anchor["строка 5"]["text"] == "Свободная строка без номера."


def test_extraction_is_deterministic(tmp_path):
    for p in _fixture_files(tmp_path):
        assert freeze.extract_file(p) == freeze.extract_file(p)


def test_unsupported_format_is_an_honest_refusal(tmp_path):
    # .pdf поддержан с 2026-08-26 — непокрытый пример теперь .rtf.
    p = tmp_path / "x.rtf"
    p.write_bytes(b"{\\rtf1")
    with pytest.raises(freeze.FreezeError) as e:
        freeze.extract_file(p)
    assert "не входит в первую очередь" in str(e.value)
    assert ".docx" in str(e.value)      # поддержанные форматы названы


def test_doc_without_word_is_an_honest_refusal(tmp_path, monkeypatch):
    """Симуляция машины без Microsoft Word: отказ называет и файл, и
    причину — тихий пропуск документа запрещён методикой."""
    p = tmp_path / "старый.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0")
    monkeypatch.setattr(freeze.subprocess, "run",
                        lambda *a, **k: subprocess.CompletedProcess(
                            a, returncode=3, stdout="NO-WORD", stderr=""))
    with pytest.raises(freeze.FreezeError) as e:
        freeze.extract_file(p, convert_dir=tmp_path / "conv")
    assert "Microsoft Word" in str(e.value)
    assert "старый.doc" in str(e.value)


# ---------------------------------------------------------------------------
# Комплект заморозки и гейт
# ---------------------------------------------------------------------------

@pytest.fixture()
def kit(tmp_path):
    files = _fixture_files(tmp_path)
    out = tmp_path / "kit"
    manifest = freeze.freeze(files, out)
    return files, out, manifest


def test_freeze_writes_manifest_clauses_and_projections(kit):
    files, out, manifest = kit
    assert [f["sid"] for f in manifest["files"]] == ["S1", "S2", "S3"]
    for f, path in zip(manifest["files"], files):
        assert f["file"] == path.name
        assert f["sha256"] == freeze.sha256_of(path)
        entry = json.loads((out / "clauses" / f"{f['sid']}.json")
                           .read_text(encoding="utf-8"))
        assert f["clause_count"] == len(entry["clauses"]) > 0
        proj = out / "projections" / f"{f['sid']}-{path.stem}.md"
        assert proj.read_text(encoding="utf-8") == \
            freeze.render_projection(entry)
    assert manifest["arbitration"] == freeze.ARBITRATION


def test_freeze_is_deterministic_for_docx_xlsx_md(tmp_path):
    files = _fixture_files(tmp_path)
    m1 = freeze.freeze(files, tmp_path / "k1")
    m2 = freeze.freeze(files, tmp_path / "k2")
    assert m1 == m2
    for sid in ("S1", "S2", "S3"):
        assert (tmp_path / "k1" / "clauses" / f"{sid}.json").read_bytes() \
            == (tmp_path / "k2" / "clauses" / f"{sid}.json").read_bytes()


def test_gate_green_with_originals(kit, capsys):
    files, out, _ = kit
    rc = gate_freeze.run(str(out), [str(p) for p in files])
    outtxt = capsys.readouterr().out
    assert rc == 0, outtxt
    assert "файлов 3" in outtxt
    assert "оригиналов сверено 3" in outtxt
    assert "самотест пройден" in outtxt
    assert "ВЕРДИКТ: комплект заморозки соответствует контракту" in outtxt


def test_gate_without_originals_says_so_loudly(kit, capsys):
    _, out, _ = kit
    rc = gate_freeze.run(str(out))
    outtxt = capsys.readouterr().out
    assert rc == 0
    assert "НЕ сверялись" in outtxt


def test_gate_red_on_tampered_projection(kit, capsys):
    files, out, manifest = kit
    proj = out / "projections" / "S1-PRAVILA-OM-A.md"
    proj.write_text(proj.read_text(encoding="utf-8")
                    .replace("10 часов", "12 часов"), encoding="utf-8")
    rc = gate_freeze.run(str(out), [str(p) for p in files])
    assert rc == 1
    assert "мимо машинной формы" in capsys.readouterr().out


def test_gate_red_on_lost_clause_in_machine_form(kit, capsys):
    files, out, _ = kit
    cl_path = out / "clauses" / "S1.json"
    entry = json.loads(cl_path.read_text(encoding="utf-8"))
    entry["clauses"] = entry["clauses"][:-1]        # потерянная клауза
    cl_path.write_text(json.dumps(entry, ensure_ascii=False, indent=2)
                       + "\n", encoding="utf-8")
    rc = gate_freeze.run(str(out), [str(p) for p in files])
    outtxt = capsys.readouterr().out
    assert rc == 1
    # Ловится дважды: счётчиком манифеста и повторным извлечением.
    assert "в машинной форме" in outtxt
    assert "другой набор клауз" in outtxt or "не совпадает" in outtxt


def test_gate_red_on_changed_original(kit, capsys):
    files, out, _ = kit
    files[2].write_text(_MD_TEXT + "\nНовая строка.\n", encoding="utf-8")
    rc = gate_freeze.run(str(out), [str(p) for p in files])
    assert rc == 1
    assert "изменён после заморозки" in capsys.readouterr().out


def test_gate_red_on_unfrozen_original(kit, tmp_path, capsys):
    files, out, _ = kit
    extra = tmp_path / "ZABYLI.md"
    extra.write_text("1.1. Забытый документ.\n", encoding="utf-8")
    rc = gate_freeze.run(str(out), [str(p) for p in files + [extra]])
    assert rc == 1
    assert "в манифесте заморозки отсутствует" in capsys.readouterr().out


def test_cli_freeze_prints_full_paths_and_counts(tmp_path, capsys):
    files = _fixture_files(tmp_path)
    out = tmp_path / "kit"
    rc = freeze.main([str(p) for p in files] + ["--out-dir", str(out)])
    outtxt = capsys.readouterr().out
    assert rc == 0
    assert str((out / "manifest.json").resolve()) in outtxt
    assert "итого: файлов 3" in outtxt


def test_cli_refusal_names_reason(tmp_path, capsys):
    bad = tmp_path / "x.rtf"
    bad.write_text("{\\rtf1}", encoding="utf-8")
    rc = freeze.main([str(bad), "--out-dir", str(tmp_path / "kit")])
    assert rc == 1
    assert "ОТКАЗ" in capsys.readouterr().out


# ---------------------------------------------------------------------------
# .doc через Microsoft Word (пропускается с названной причиной без Word)
# ---------------------------------------------------------------------------

def _word_present() -> bool:
    if not hasattr(_word_present, "_cached"):
        try:
            _word_present._cached = freeze.word_available()
        except Exception:                                  # noqa: BLE001
            _word_present._cached = False
    return _word_present._cached


@pytest.mark.skipif(not _word_present(),
                    reason="нужен Microsoft Word: конвертация .doc "
                           "выполняется его COM-автоматизацией")
def test_doc_roundtrip_matches_docx_extraction(tmp_path):
    """.docx сохраняется Word-ом как .doc, замораживается — набор клауз
    обязан совпасть с извлечением из исходного .docx (конвертация не
    теряет ни клауз, ни жирной разметки)."""
    src = tmp_path / "PRAVILA.docx"
    _make_docx(src)
    doc_path = tmp_path / "PRAVILA.doc"
    script = tmp_path / "to_doc.ps1"
    script.write_text(
        "param([string]$In,[string]$Out)\n"
        "$ErrorActionPreference='Stop'\n"
        "$w = New-Object -ComObject Word.Application\n"
        "try { $d = $w.Documents.Open((Resolve-Path $In).Path,$false,"
        "$true); $d.SaveAs2([string]$Out, 0); $d.Close($false) } "
        "finally { $w.Quit() }\n", encoding="ascii")
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy",
                    "Bypass", "-File", str(script), "-In", str(src),
                    "-Out", str(doc_path)], check=True, timeout=300,
                   capture_output=True)
    got = freeze.extract_file(doc_path, convert_dir=tmp_path / "conv")
    want = freeze.extract_file(src)["clauses"]
    assert got["converted"] is not None
    assert got["clauses"] == want
