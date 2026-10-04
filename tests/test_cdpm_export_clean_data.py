# -*- coding: utf-8 -*-
"""ЭКС-3: очищенные операционные данные — вычитание адресных Б/К
находок отчёта проверок, отчёт очистки прилагается; прочие находки
остаются перечисленными. Без отчёта проверок endpoint отказывает."""
import io
import json

from cdpm import export_clean_data as ec


def _dataset():
    return {"tables": {"D1": {"Рейсы": {
        "columns": ["Номер", "Дата"],
        "rows": [{"Номер": "PL1", "Дата": "01.01"},
                 {"Номер": "PL1", "Дата": "01.01"},   # дубль (листовая строка 3)
                 {"Номер": "PL2", "Дата": "02.01"}],
    }}}}


def _report():
    return {"customer": "Тест", "findings": [
        {"id": "F-0001", "class": "дубль", "severity": "КРИТИЧЕСКОЕ",
         "target": {"sid": "D1", "sheet": "Рейсы", "row": 3},
         "text": "дубль рейса"},
        {"id": "F-0002", "class": "формат", "severity": "СЕРЬЁЗНОЕ",
         "target": {"sid": "D1", "sheet": "Рейсы", "row": 4},
         "text": "подозрительная дата"},
    ]}


def test_clean_removes_only_addressed_bk_rows():
    tables, cleaning = ec.build_clean(_dataset(), _report())
    rows = tables["D1"]["Рейсы"]["rows"]
    assert len(rows) == 2                # удалена только строка 3
    assert rows[1]["Номер"] == "PL2"          # серьёзная — оставлена
    assert cleaning["rows_removed"] == 1
    assert cleaning["removed"][0]["finding"] == "F-0001"
    assert cleaning["kept_findings"][0]["finding"] == "F-0002"


def test_clean_exports_json_and_xlsx_with_cleaning_sheet():
    body = ec.export_json(_dataset(), _report())
    got = json.loads(body)
    assert got["cleaning"]["rows_removed"] == 1
    raw = ec.export_xlsx(_dataset(), _report())
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    assert "Очистка" in wb.sheetnames
    ws = wb["Рейсы"]
    assert ws.max_row == 3                    # заголовок + 2 строки
    ops = [r[0].value for r in wb["Очистка"].iter_rows(min_row=2)]
    assert ops == ["удалено", "оставлено"]


def test_relational_finding_does_not_remove_rows():
    """Находка сессии COSMOS (мост, 2026-08-22): «минимальный разрыв»
    — свойство ПАРЫ строк; очистка не вычитает строку по реляционной
    находке даже при критической серьёзности — иначе уничтожается
    задача, а не дефект."""
    from cdpm import export_clean_data as cl
    dataset = {"sources": [{"sid": "D", "file": "d.xlsx",
                            "sha256": "0" * 64, "rows_seen": 3}],
               "tables": {"D": {"Рейсы": {
                   "columns": ["Номер"],
                   "rows": [{"Номер": "PL1"}, {"Номер": "PL2"},
                            {"Номер": "PL3"}]}}}}
    report = {
        "customer": "Т",
        "checks": [
            {"id": "C-1", "template": "минимальный-разрыв-часов",
             "status": "прогнана"},
            {"id": "C-2", "template": "уникальность",
             "status": "прогнана"}],
        "findings": [
            {"id": "F-1", "check": "C-1", "severity": "КРИТИЧЕСКОЕ",
             "class": "нарушение правила",
             "target": {"sid": "D", "sheet": "Рейсы", "row": 2},
             "text": "разрыв меньше 0.75 ч"},
            {"id": "F-2", "check": "C-2", "severity": "КРИТИЧЕСКОЕ",
             "class": "нарушение правила",
             "target": {"sid": "D", "sheet": "Рейсы", "row": 4},
             "text": "повтор ключа"}],
    }
    tables, cleaning = cl.build_clean(dataset, report)
    rows = tables["D"]["Рейсы"]["rows"]
    # Реляционная (строка 2) осталась; настоящий дефект
    # (листовая строка 4 = третья запись) вычтен.
    assert [r["Номер"] for r in rows] == ["PL1", "PL2"]
    kept = [e for e in cleaning["kept_findings"]
            if e.get("relational")]
    assert len(kept) == 1 and kept[0]["finding"] == "F-1"
    assert cleaning["rows_removed"] == 1


def test_row_numbering_sheet_convention():
    """Живой Лайт run-0040: находки нумеруют строки ЛИСТОВОЙ
    конвенцией исполнителей (заголовок = 1, первая запись = 2), а
    очистка резала 1-базно и вычитала СОСЕДНЮЮ здоровую строку
    вместо дефектной. Строка 2 находки = ПЕРВАЯ запись."""
    dataset = {"tables": {"D": {"Рейсы": {
        "columns": ["Номер"],
        "rows": [{"Номер": "ДЕФЕКТ"}, {"Номер": "ЗДОРОВАЯ"}]}}}}
    report = {"customer": "Т", "findings": [
        {"id": "F-1", "class": "пропуск данных",
         "severity": "КРИТИЧЕСКОЕ",
         "target": {"sid": "D", "sheet": "Рейсы", "row": 2},
         "text": "пустое обязательное поле"}]}
    tables, cleaning = ec.build_clean(dataset, report)
    rows = tables["D"]["Рейсы"]["rows"]
    assert [r["Номер"] for r in rows] == ["ЗДОРОВАЯ"]
    assert cleaning["rows_removed"] == 1
