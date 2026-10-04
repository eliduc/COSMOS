# -*- coding: utf-8 -*-
"""Приёмник ответов анкеты (интерактивный цикл, ИР-1)."""
import copy
import json

from openpyxl import load_workbook

from cdpm import answers as ans
from cdpm import anketa_xlsx as ax
from cdpm import model as m
from cdpm import render
from cdpm.gates import gate_anketa_xlsx, gate_answers
from cdpm.sample import sample_model


def _book_with(tmp_path, doc, filled: dict[str, str]):
    """Книга анкеты с заполненными ответами: как её вернёт
    авиакомпания."""
    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, book)
    wb = load_workbook(str(book))
    for ws in wb.worksheets:
        if ws.title in (ax.SUMMARY_TITLE, ax.ANSWERED_TITLE):
            continue
        for row in ws.iter_rows(min_row=2):
            qid = row[0].value
            if qid in filled:
                row[5].value = filled[qid]
    wb.save(str(book))
    return book


def test_read_answer_book(tmp_path):
    doc = sample_model()
    book = _book_with(tmp_path, doc, {"Q-2-MODAL": "б",
                                      "Q-3a-DAY": "вариант а"})
    got, problems = ans.read_answer_book(book, doc)
    assert got == {"Q-2-MODAL": "б", "Q-3a-DAY": "вариант а"}
    assert problems == []
    # Ответ на чужой идентификатор — замечание, не смерть чтения.
    wb = load_workbook(str(book))
    ws = wb[ax.SHEET_TITLES["2"]]
    ws.cell(row=1 + ws.max_row, column=1).value = "Q-НЕТ"
    ws.cell(row=ws.max_row, column=6).value = "да"
    wb.save(str(book))
    got, problems = ans.read_answer_book(book, doc)
    assert "Q-2-MODAL" in got
    assert any("Q-НЕТ" in p for p in problems)


def test_match_choice_deterministic():
    options = [{"key": "а", "text": "минимум 10 часов"},
               {"key": "б", "text": "минимум 12 часов"}]
    assert ans.match_choice("а", options) == ("а", "")
    assert ans.match_choice("Б", options) == ("б", "")
    assert ans.match_choice("вариант б", options) == ("б", "")
    assert ans.match_choice("а) как в правилах", options) \
        == ("а", "как в правилах")
    assert ans.match_choice("минимум 12 часов", options) == ("б", "")
    assert ans.match_choice("а, с оговоркой по базе", options) \
        == ("а", "с оговоркой по базе")
    assert ans.match_choice("зависит от сезона", options) is None


def test_resolve_case_orphans():
    """Живой стоп run-0023 на гейте запуска (1.09): записи листа
    оставались CASE после применённого вопроса-развилки к строке-
    носителю; строки-«запросы» целевой висели CASE при применённых
    Q-*-GOAL. Разрешение детерминированно и fail-closed: SOFT не
    наследуется, открытый вопрос не подменяется."""
    doc = {
        "rules": [
            {"id": "R-S2-лист-«Заявки»-строка-2", "kind": "DATA",
             "zone": "ROSTERING", "sources": [
                 {"sid": "S2", "anchor": "лист «Заявки», строка 2"}]},
            {"id": "R-S2-лист-«Заявки»-строка-3", "kind": "CASE",
             "zone": "ROSTERING", "sources": [
                 {"sid": "S2", "anchor": "лист «Заявки», строка 3"}]},
            {"id": "R-S2-лист-«Премии»-строка-2", "kind": "SOFT",
             "zone": "ROSTERING", "weight": 5, "sources": [
                 {"sid": "S2", "anchor": "лист «Премии», строка 2"}]},
            {"id": "R-S2-лист-«Премии»-строка-3", "kind": "CASE",
             "zone": "ROSTERING", "sources": [
                 {"sid": "S2", "anchor": "лист «Премии», строка 3"}]},
            {"id": "R-S1-т2-99-3", "kind": "CASE", "zone": "PAIRING",
             "weight": 3, "sources": [
                 {"sid": "S1",
                  "anchor": "табл.2, строка 2, кол 3"}]},
            {"id": "R-S2-лист-«Отгулы»-строка-3", "kind": "CASE",
             "zone": "ROSTERING", "sources": [
                 {"sid": "S2", "anchor": "лист «Отгулы», строка 3"}]},
        ],
        "questions": [
            {"id": "Q-2-заявки",
             "answer": {"status": "применён", "key": "в"},
             "answer_format": {"type": "choice"},
             "source_refs": [
                 {"sid": "S2", "anchor": "лист «Заявки», строка 2"}],
             "rule_refs": ["R-S2-лист-«Заявки»-строка-2"]},
            {"id": "Q-2-премии",
             "answer": {"status": "применён", "key": "а"},
             "answer_format": {"type": "choice"},
             "source_refs": [
                 {"sid": "S2", "anchor": "лист «Премии», строка 2"}],
             "rule_refs": ["R-S2-лист-«Премии»-строка-2"]},
            {"id": "Q-3a-GOAL", "answer": {"status": "применён"},
             "source_refs": [
                 {"sid": "S1",
                  "anchor": "табл.2, строка 2, кол 3"}],
             "rule_refs": []},
            {"id": "Q-2-отгулы", "rule_refs":
             ["R-S2-лист-«Отгулы»-строка-3"], "source_refs": []},
        ],
    }
    notes = m.resolve_case_orphans(doc)
    by_id = {r["id"]: r for r in doc["rules"]}
    got = by_id["R-S2-лист-«Заявки»-строка-3"]
    assert got["kind"] == "DATA"
    assert "унаследован" in got["note"]
    assert by_id["R-S2-лист-«Премии»-строка-3"]["kind"] == "CASE"
    goal = by_id["R-S1-т2-99-3"]
    assert goal["kind"] == "OUT"
    assert "weight" not in goal
    assert "Q-3a-GOAL" in goal["note"]
    assert by_id["R-S2-лист-«Отгулы»-строка-3"]["kind"] == "CASE"
    assert len(notes) == 2
    assert m.resolve_case_orphans(doc) == []


def test_canon_option_key_strips_decorated_keys():
    """Живой класс run-0019/0022 (тот же, что фантомные номера
    замечаний): LLM-разборщик возвращал ключ формой «а) полный текст
    варианта» — буквальная сверка ставила unparsed, вопрос жил
    вечно. Обрезка законна ТОЛЬКО при детерминированной сверке
    текста с вариантом; гадание запрещено (fail-closed)."""
    options = [
        {"key": "а", "text": "Полные экипажи, сутки по местному "
                             "времени базы"},
        {"key": "б", "text": "местное время базы PLY для всех "
                             "записей"},
    ]
    canon = ans._canon_option_key
    assert canon("а) Полные экипажи, сутки по местному времени базы",
                 options) == "а"
    assert canon("б)  местное время базы PLY для всех записей ",
                 options) == "б"
    assert canon("а)", options) == "а"
    assert canon("Полные экипажи, сутки по местному времени базы",
                 options) == "а"
    assert canon("а) чужой текст, не совпадающий с вариантом",
                 options) is None
    assert canon("в) местное время базы PLY для всех записей",
                 options) is None
    assert canon("зависит от сезона", options) is None
    assert canon(None, options) is None
    same = [{"key": "а", "text": "одинаковый текст"},
            {"key": "б", "text": "одинаковый текст"}]
    assert ans._canon_option_key("одинаковый текст", same) is None


def test_fit_b_answer_denied_goes_to_worker():
    """Живой прогон H0 run-0005 (2.09): ответ «неверно» на B-вопрос
    посадки закрывался статусом без действия — ложное покрытие.
    Отрицание допущения блокирует строку вопросом и отдаёт её
    обработчику; подтверждение «а» — применяется машинно."""
    from cdpm.sample import sample_model
    doc = sample_model()
    r = next(x for x in doc["rules"] if x["kind"] == "HARD")
    r["template_fit"] = {"category": "B", "registry_rev": "rev",
                         "family": "duty_limits",
                         "assumption": "предел 13:00",
                         "params": {"max_duty_duration": {
                             "value": 780, "quote": "до 13:00"}}}
    [qid] = m.ensure_fit_questions(doc)
    ans.apply_answers(doc, {qid: "б) неверно — базовый предел 720, "
                                 "780 — потолок продления"})
    q = next(x for x in doc["questions"] if x["id"] == qid)
    assert q["answer"]["status"] == "ожидает обработчика"
    assert qid in r.get("blockers", [])
    # Подтверждение — применяется машинно, блокировки нет.
    doc2 = sample_model()
    r2 = next(x for x in doc2["rules"] if x["kind"] == "HARD")
    r2["template_fit"] = dict(r["template_fit"])
    [qid2] = m.ensure_fit_questions(doc2)
    ans.apply_answers(doc2, {qid2: "а) допущение верно — подтверждаем"})
    q2 = next(x for x in doc2["questions"] if x["id"] == qid2)
    assert q2["answer"]["status"] == "применён"
    assert qid2 not in r2.get("blockers", [])


def test_worker_patch_may_carry_template_fit():
    """Патч обработчика по B-вопросу переносит посадку: template_fit
    — законное поле патча (registry_rev при приёме ставит код)."""
    assert "template_fit" in ans._PATCH_FIELDS


def test_worker_patch_fit_gets_registry_rev_from_code():
    """Прогон №2 H0 run-0007 (2.09): обработчик по B-вопросу посадки
    пересадил template_fit без registry_rev (промпт велит его не
    писать) — патч отбраковала схема «registry_rev is required»,
    вопрос завис «ожидает обработчика». Отпечаток базы посадке патча
    проставляет код, как в цикле обмена."""
    from cdpm.sample import sample_model
    doc = sample_model()
    r = next(x for x in doc["rules"] if x["kind"] == "HARD")
    r["template_fit"] = {"category": "B", "registry_rev": "rev",
                         "family": "duty_limits",
                         "assumption": "предел 13:00",
                         "params": {"max_duty_duration": {
                             "value": 780, "quote": "до 13:00"}}}
    [qid] = m.ensure_fit_questions(doc)
    ans.apply_answers(doc, {qid: "б) неверно — базовый предел 720"})
    q = next(x for x in doc["questions"] if x["id"] == qid)
    assert q["answer"]["status"] == "ожидает обработчика"
    patched = copy.deepcopy(r)
    patched["template_fit"] = {"category": "A", "family": "duty_limits",
                               "params": {"max_duty_duration": {
                                   "value": 720, "quote": "до 13:00"}}}
    patched.pop("blockers", None)
    summary = ans.apply_worker_patches(
        doc, lambda p: _worker_reply(qid, [patched]))
    assert summary["worker_rejected"] == []
    assert summary["worker_applied"] == 1
    got = next(x for x in doc["rules"] if x["id"] == r["id"])
    assert got["template_fit"]["registry_rev"] == m.registry_fingerprint()
    assert got["template_fit"]["category"] == "A"
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_inapplicable_answer_takes_row_out_of_scope():
    """Запрос одной из линий: закрытый перечень может целиком
    не действовать (условный пример: FRMS-значения, когда FRMS
    не внедрена). Ответ
    каноничной формой применяется: строка-носитель уходит вне
    контура (OUT) со следом, блокировки сняты, хвост ответа — в
    оговорке."""
    doc = sample_model()
    assert ans.match_choice(
        "ни одно из значений не действует: FRMS не внедрена",
        [{"key": "а", "text": "5"}]) == (
            ans.INAPPLICABLE_KEY, "FRMS не внедрена")
    summary = ans.apply_answers(
        doc, {"Q-2-MODAL": "строка неприменима — FRMS не внедрена"})
    assert summary["applied"] == 1
    assert summary["resolved_rules"] == ["R-S1-П-2-3"]
    rule = next(r for r in doc["rules"] if r["id"] == "R-S1-П-2-3")
    assert rule["kind"] == "OUT"
    assert "conditional" not in rule and "blockers" not in rule
    assert "перечень признан неприменимым" in rule["note"]
    q = next(x for x in doc["questions"] if x["id"] == "Q-2-MODAL")
    assert q["answer"]["status"] == "применён"
    assert q["answer"]["key"] == ans.INAPPLICABLE_KEY
    assert q["answer"]["note"] == "FRMS не внедрена"
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_apply_resolves_conditional_and_renders():
    doc = sample_model()
    plain_anketa = render.render_anketa(doc)
    assert "Принятые ответы" not in plain_anketa
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "б"})
    assert summary["applied"] == 1
    assert summary["resolved_rules"] == ["R-S1-П-2-3"]
    assert doc["meta"]["revision"] == 2
    rule = next(r for r in doc["rules"] if r["id"] == "R-S1-П-2-3")
    assert rule["kind"] == "HARD"
    assert "conditional" not in rule and "blockers" not in rule
    q = next(x for x in doc["questions"] if x["id"] == "Q-2-MODAL")
    # Журнал следствий (A3.4): что ответ изменил в строке — вид,
    # снятая развилка, снятые блокировки.
    effects = q["answer"].pop("effects")
    fields = {e["field"] for e in effects if e["rule"] == "R-S1-П-2-3"}
    assert {"conditional", "blockers"} <= fields
    cond_e = next(e for e in effects if e["field"] == "conditional")
    assert cond_e["before"]["question"] == "Q-2-MODAL" \
        and cond_e.get("after_absent") is True
    version = q["answer"]["answer_version"]
    assert version["fingerprint"] == "9aa032336a8e"  # sha256("б")[:12]
    assert version["book_sha256"] is None and version["received_at"]
    assert q["answer"] == {"raw": "б", "status": "применён",
                           "key": "б", "revision": 2, "answer_version": version}
    assert m.validate_schema(doc) + m.validate_model(doc) == []
    anketa = render.render_anketa(doc)
    assert "Принятые ответы" in anketa
    assert "Ревизия анкеты: 2" in anketa
    # Отвеченный вопрос ушёл из частей (остался только в принятых).
    assert "### ВОПРОС Q-2-MODAL" not in anketa
    svod = render.render_svod(doc)
    assert "по Вопросу Q-2-MODAL" not in svod
    assert "Ревизия свода: 2" in svod


def test_apply_soft_branch_binds_objective():
    """Ветвь SOFT без блокировок — автопривязка к вопросу целевой
    функции своей зоны: вес приходит только из ответа."""
    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "а"})
    assert summary["applied"] == 1
    rule = next(r for r in doc["rules"] if r["id"] == "R-S1-П-2-3")
    assert rule["kind"] == "SOFT"
    assert rule["blockers"], "SOFT-строка осталась без вопроса о весе"
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_apply_freeform_awaits_worker():
    """Ответ не-развилки (реестр, значение) записывается и ждёт
    модельного применения (ИР-2); блокировки не трогаются."""
    doc = sample_model()
    blocked_before = [r["id"] for r in doc["rules"]
                      if "Q-3a-OBJ" in r.get("blockers", [])]
    summary = ans.apply_answers(
        doc, {"Q-3a-OBJ": "компоненты: ночёвки 10, кредит 5"})
    assert summary["awaiting"] == 1 and summary["applied"] == 0
    q = next(x for x in doc["questions"] if x["id"] == "Q-3a-OBJ")
    assert q["answer"]["status"] == "ожидает обработчика"
    blocked_after = [r["id"] for r in doc["rules"]
                     if "Q-3a-OBJ" in r.get("blockers", [])]
    assert blocked_after == blocked_before
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_apply_fuzzy_via_llm_and_fail_closed():
    doc = sample_model()

    def llm(prompt):
        assert "Q-2-MODAL" in prompt
        return json.dumps({"matches": [
            {"question": "Q-2-MODAL", "key": "б",
             "note": "жёстко, кроме чартеров"}]}, ensure_ascii=False)

    summary = ans.apply_answers(
        doc, {"Q-2-MODAL": "считаем это жёстким, кроме чартеров"},
        llm=llm)
    # Аудит 04.09, A3.3: ключ с оговоркой («кроме чартеров») машинно
    # не применяется — ждёт обработчика, развилка и блокировка целы.
    assert summary["applied"] == 0 and summary["awaiting"] == 1
    q = next(x for x in doc["questions"] if x["id"] == "Q-2-MODAL")
    assert q["answer"]["status"] == "ожидает обработчика"
    assert q["answer"]["key"] == "б"
    assert q["answer"]["note"] == "жёстко, кроме чартеров"
    rule = next(r for r in doc["rules"] if r["id"] == "R-S1-П-2-3")
    assert rule.get("conditional") and "Q-2-MODAL" in rule["blockers"]
    assert m.validate_schema(doc) + m.validate_model(doc) == []
    # Без оговорки тот же нечёткий ответ применяется машинно.
    doc = sample_model()
    summary = ans.apply_answers(
        doc, {"Q-2-MODAL": "считаем это жёстким"},
        llm=lambda p: json.dumps(
            {"matches": [{"question": "Q-2-MODAL", "key": "б"}]},
            ensure_ascii=False))
    assert summary["applied"] == 1

    # Ключ вне вариантов от модели — честное «не разобран».
    doc = sample_model()
    summary = ans.apply_answers(
        doc, {"Q-2-MODAL": "по-разному"},
        llm=lambda p: json.dumps(
            {"matches": [{"question": "Q-2-MODAL", "key": "щ"}]},
            ensure_ascii=False))
    assert summary["unparsed"] == 1
    q = next(x for x in doc["questions"] if x["id"] == "Q-2-MODAL")
    assert q["answer"]["status"] == "не разобран"
    assert m.validate_schema(doc) + m.validate_model(doc) == []
    # Неразобранный вопрос остаётся в анкете с пометкой.
    anketa = render.render_anketa(doc)
    assert "ВОПРОС Q-2-MODAL" in anketa
    assert "Прежний ответ не разобрался" in anketa

    # Без канала нечёткий ответ тоже честно «не разобран».
    doc = sample_model()
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "по-разному"})
    assert summary["unparsed"] == 1


def test_reanswer_of_answered_question_is_ignored():
    doc = sample_model()
    ans.apply_answers(doc, {"Q-2-MODAL": "б"})
    summary = ans.apply_answers(doc, {"Q-2-MODAL": "а"})
    assert summary["already_answered"] == 1
    assert summary["applied"] == 0
    q = next(x for x in doc["questions"] if x["id"] == "Q-2-MODAL")
    assert q["answer"]["key"] == "б"       # первый ответ не затёрт
    assert doc["meta"]["revision"] == 3    # проход был — ревизия растёт


def test_xlsx_answered_sheet_and_gate(tmp_path):
    doc = sample_model()
    ans.apply_answers(doc, {"Q-2-MODAL": "б"})
    book = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, book)
    wb = load_workbook(str(book))
    assert ax.ANSWERED_TITLE in wb.sheetnames
    # Единственный вопрос части 2 отвечен — лист части честно исчез.
    assert ax.SHEET_TITLES["2"] not in wb.sheetnames
    ws = wb[ax.ANSWERED_TITLE]
    ids = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)
           if row and row[0]]
    assert ids == ["Q-2-MODAL"]
    problems, _ = gate_anketa_xlsx.check_workbook(doc, book)
    assert problems == []


def test_gate_answers_green_and_red(tmp_path):
    from cdpm import answers as ans_mod
    doc = sample_model()
    summary = ans_mod.apply_answers(doc, {"Q-2-MODAL": "б"})
    model_path = tmp_path / "model.json"
    summary_path = tmp_path / "answers-summary.json"
    svod_path = tmp_path / "svod.md"
    anketa_path = tmp_path / "anketa.md"
    book_path = tmp_path / "anketa.xlsx"
    m.dump_model(doc, model_path)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False),
                            encoding="utf-8")
    svod_path.write_text(render.render_svod(doc), encoding="utf-8")
    anketa_path.write_text(render.render_anketa(doc), encoding="utf-8")
    ax.render_anketa_xlsx(doc, book_path)
    assert gate_answers.run(str(model_path), str(summary_path),
                            str(svod_path), str(anketa_path),
                            str(book_path)) == 0
    # Порча свода мимо модели — красный.
    svod_path.write_text(
        svod_path.read_text(encoding="utf-8").replace("HARD", "МЯГК",
                                                      1),
        encoding="utf-8")
    problems, _ = gate_answers.check(model_path, summary_path,
                                     svod_path, anketa_path, book_path)
    assert any("мимо модели" in p for p in problems)
    # Сводка с чужой ревизией — красный.
    svod_path.write_text(render.render_svod(doc), encoding="utf-8")
    summary_path.write_text(json.dumps({**summary, "revision": 9},
                                       ensure_ascii=False),
                            encoding="utf-8")
    problems, _ = gate_answers.check(model_path, summary_path,
                                     svod_path, anketa_path, book_path)
    assert any("не равна meta.revision" in p for p in problems)


def test_answers_checkpoint_called():
    """ПС-1: контрольная точка зовётся перед чанком разборщика."""
    doc = sample_model()
    calls = []
    ans.apply_answers(
        doc, {"Q-2-MODAL": "по-разному"},
        llm=lambda p: json.dumps(
            {"matches": [{"question": "Q-2-MODAL",
                          "unparsed": "ответ не называет вариант"}]},
            ensure_ascii=False),
        control=lambda: calls.append(1))
    assert calls == [1]


def test_cli_roundtrip(tmp_path, monkeypatch):
    doc = sample_model()
    model_path = tmp_path / "model.json"
    m.dump_model(doc, model_path)
    book = _book_with(tmp_path, doc, {"Q-2-MODAL": "б"})
    monkeypatch.setattr(
        ans.llm_mod, "cosmos_llm",
        lambda **kw: lambda p: (_ for _ in ()).throw(
            AssertionError("канал не должен зваться: ответ чёткий")))
    out = tmp_path / "out"
    rc = ans.main([str(model_path), str(book),
                   "--out-dir", str(out)])
    assert rc == 0
    got = m.load_model(out / "model.json")
    assert got["meta"]["revision"] == 2
    import hashlib
    from datetime import datetime
    received = next(q for q in got["questions"] if q["id"] == "Q-2-MODAL")["answer"]
    assert received["answer_version"]["book_sha256"] == hashlib.sha256(book.read_bytes()).hexdigest()
    assert datetime.fromisoformat(received["answer_version"]["received_at"]).utcoffset() is not None
    summary = json.loads((out / "answers-summary.json")
                         .read_text(encoding="utf-8"))
    assert summary["applied"] == 1
    assert summary["book_sha256"]
    assert gate_answers.run(
        str(out / "model.json"),
        str(out / "answers-summary.json"),
        str(out / "svod.md"), str(out / "anketa.md"),
        str(out / "anketa.xlsx")) == 0


def test_old_models_render_unchanged():
    """Обратная совместимость: модель без ответов рендерится без
    ревизий, разделов и листов ответов."""
    doc = sample_model()
    assert "Ревизия" not in render.render_svod(doc)
    anketa = render.render_anketa(doc)
    assert "Ревизия" not in anketa and "Принятые ответы" not in anketa
    assert ax.answered_rows(doc) == []


def test_readiness_states():
    """СГ-1: три состояния линии правил считаются по зонам блокировок
    и открытых вопросов."""
    doc = sample_model()
    got = m.readiness(doc)
    assert got["state"] == "а"
    assert not got["pairing_ready"] and not got["rostering_ready"]
    assert "блокированные строки" in got["pairing_open"]

    # Чистая модель без блокировок и вопросов — состояние (в).
    clean = {"meta": doc["meta"], "sources": doc["sources"],
             "rules": [{"id": "R-1", "zone": "BOTH", "kind": "HARD",
                        "title": "т", "statement": "с",
                        "sources": doc["rules"][0]["sources"]}],
             "questions": []}
    got = m.readiness(clean)
    assert got["state"] == "в"
    assert got["pairing_ready"] and got["rostering_ready"]

    # Открытый вопрос части 3б — пейринг готов, ростеринг нет: (б).
    halfway = copy.deepcopy(clean)
    halfway["questions"] = [{
        "id": "Q-3b-X", "part": "3b", "title": "т", "text": "чч",
        "answer_format": {"type": "value", "note": "часы"},
        "rule_refs": ["R-1"], "source_refs": doc["rules"][0]["sources"]}]
    got = m.readiness(halfway)
    assert got["state"] == "б"
    assert got["pairing_ready"] and not got["rostering_ready"]
    assert got["rostering_open"] == {"открытые вопросы части 3б": 1}
    # Отвеченный вопрос перестаёт держать готовность (если строки не
    # заблокированы).
    halfway["questions"][0]["answer"] = {
        "raw": "10", "status": "ожидает обработчика", "revision": 2}
    halfway["meta"] = {**halfway["meta"], "revision": 2}
    got = m.readiness(halfway)
    assert got["state"] == "в"


def _worker_reply(qid, rules):
    return json.dumps({"patches": [
        {"question": qid, "change": "ответ внесён", "rules": rules}]},
        ensure_ascii=False)


def test_worker_patch_applies_objective_answer():
    """ИР-2: ответ о весах (часть 3а, реестр) вносится обработчиком —
    SOFT-строка получает вес и законно остаётся без блокировки."""
    doc = sample_model()
    ans.apply_answers(doc, {"Q-3a-OBJ": "ночёвки вне базы: вес 10"})
    q = next(x for x in doc["questions"] if x["id"] == "Q-3a-OBJ")
    assert q["answer"]["status"] == "ожидает обработчика"
    affected = ans._affected_rules(doc, q)
    assert affected, "у вопроса о весах нет затронутых строк"
    rule = copy.deepcopy(affected[0])
    rule["note"] = "вес 10 (ответ Q-3a-OBJ)"
    # P-N-SHD 2026-08-30: путь к весу — МАШИННОЕ поле, не заметка;
    # патч без weight валидатор отвергает fail-closed (живой корень
    # прод-демо: «взвешенность» членством в rule_refs пропустила
    # мягкую строку с weight=null в пакет).
    rule["weight"] = 10
    rule.pop("blockers", None)

    def llm(prompt):
        assert "ОТВЕТ АВИАКОМПАНИИ: ночёвки вне базы: вес 10" in prompt
        return _worker_reply("Q-3a-OBJ", [rule])

    summary = ans.apply_worker_patches(doc, llm)
    assert summary["worker_applied"] == 1
    assert summary["worker_rejected"] == []
    assert q["answer"]["status"] == "применён"
    got = next(r for r in doc["rules"] if r["id"] == rule["id"])
    assert got["note"] == "вес 10 (ответ Q-3a-OBJ)"
    assert "Q-3a-OBJ" not in got.get("blockers", [])
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_worker_patch_rejections_fail_closed():
    """Чужая строка, изменённая трассировка, слом контракта и молчание
    канала — откат вопроса с причиной; статус остаётся «ожидает»."""
    doc = sample_model()
    ans.apply_answers(doc, {"Q-3a-OBJ": "вес 10"})
    q = next(x for x in doc["questions"] if x["id"] == "Q-3a-OBJ")
    affected = ans._affected_rules(doc, q)
    rule = copy.deepcopy(affected[0])

    # 1) Патч чужой строки.
    alien = copy.deepcopy(rule)
    alien["id"] = "R-S1-П-1-1"
    summary = ans.apply_worker_patches(
        doc, lambda p: _worker_reply("Q-3a-OBJ", [alien]))
    assert "не затронута вопросом" in \
        summary["worker_rejected"][0]["reasons"][0]
    assert q["answer"]["status"] == "ожидает обработчика"

    # 2) Изменённая трассировка.
    bad_src = copy.deepcopy(rule)
    bad_src["sources"] = [{"sid": "S1", "anchor": "выдумано"}]
    summary = ans.apply_worker_patches(
        doc, lambda p: _worker_reply("Q-3a-OBJ", [bad_src]))
    assert "трассировка" in summary["worker_rejected"][0]["reasons"][0]

    # 3) Слом контракта (тип вне перечня) ловится полной валидацией.
    bad_kind = copy.deepcopy(rule)
    bad_kind["kind"] = "МЯГКОЕ"
    summary = ans.apply_worker_patches(
        doc, lambda p: _worker_reply("Q-3a-OBJ", [bad_kind]))
    assert "не проходит контракт" in \
        summary["worker_rejected"][0]["reasons"][0]

    # 4) Молчание канала.
    summary = ans.apply_worker_patches(doc, lambda p: "не могу")
    assert "не дал патча" in summary["worker_rejected"][0]["reasons"][0]
    assert q["answer"]["status"] == "ожидает обработчика"
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_worker_patch_part1_contradiction_out():
    """Часть 1: действующее требование остаётся, отменённое — kind OUT
    с причиной (удаления строк нет)."""
    doc = sample_model()
    q1 = next(x for x in doc["questions"] if x["part"] == "1")
    key = q1["answer_format"]["options"][0]["key"]
    ans.apply_answers(doc, {q1["id"]: key})
    affected = ans._affected_rules(doc, q1)
    assert len(affected) >= 2
    keep, drop = (copy.deepcopy(affected[0]),
                  copy.deepcopy(affected[1]))
    drop["kind"] = "OUT"
    drop["note"] = f"отменено ответом {q1['id']}: действует {keep['id']}"
    for r in (keep, drop):
        r.pop("blockers", None)
        r.pop("conditional", None)

    summary = ans.apply_worker_patches(
        doc, lambda p: _worker_reply(q1["id"], [keep, drop]))
    assert summary["worker_applied"] == 1, summary["worker_rejected"]
    got = next(r for r in doc["rules"] if r["id"] == drop["id"])
    assert got["kind"] == "OUT"
    assert q1["id"] not in got.get("blockers", [])
    assert m.validate_schema(doc) + m.validate_model(doc) == []


def test_ответы_условий_применяются_машинно_без_обработчика():
    """Живой урок 27-28.08: вопросы условий — носители текста (raw
    читает экспорт условий), а их применение трижды падало на
    мёртвом канале модели, хотя канал не был нужен."""
    from cdpm import answers as am
    doc = {"meta": {"revision": 1},
           "rules": [],
           "questions": [{
               "id": "Q-3a-условия-размах-командировки",
               "part": "3a", "title": "т", "text": "т",
               "answer_format": {"type": "value"}}]}
    summary = am.apply_answers(
        doc, {"Q-3a-условия-размах-командировки":
              "11520 минут (8 суток) с перегонами"}, llm=None)
    a = doc["questions"][0]["answer"]
    assert a["status"] == "применён", a
    assert "1520" in a["raw"]
    assert summary["applied"] == 1 and summary["awaiting"] == 0
