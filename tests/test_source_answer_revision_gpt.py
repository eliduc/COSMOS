"""Gapped answer revisions: exact pre-apply evidence, synthetic/offline only."""
from copy import deepcopy
import io
import json
import socket
import subprocess

import pytest
from openpyxl import load_workbook

from cdpm import answers, anketa_xlsx, source_answer_evidence as ae
from cdpm.sample import sample_model
from test_source_answer_evidence_gpt import (
    answered_run, raw, sha, package_payloads, zip_body, import_sources,
)


PREFIX = "answers/rev-004/"
WITNESS = PREFIX + "source-revision-recovery.json"
SOURCE = PREFIX + "source-model-recovery.json"


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("No provider/process in source-revision fixtures")
    monkeypatch.setattr(socket.socket, "connect", denied)
    monkeypatch.setattr(socket, "create_connection", denied)
    monkeypatch.setattr(subprocess, "Popen", denied)
    monkeypatch.setattr(answers.llm_mod, "cosmos_llm", denied)


@pytest.fixture
def gap(answered_run, tmp_path):
    run, old_doc = answered_run
    source = deepcopy(old_doc)
    source["meta"]["revision"] = 2
    source["questions"][0].pop("answer")
    book = tmp_path / "issued-rev2.xlsx"
    anketa_xlsx.render_anketa_xlsx(source, book)
    wb = load_workbook(book)
    wb[anketa_xlsx.SHEET_TITLES["3b"]].cell(2, anketa_xlsx.ANSWER_COL, "600 минут")
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    body = buffer.getvalue()
    doc = deepcopy(source)
    doc["meta"]["revision"] = 5  # later activity must not infer input revision
    doc["questions"][0]["answer"] = deepcopy(old_doc["questions"][0]["answer"])
    doc["questions"][0]["answer"]["revision"] = 4
    doc["questions"][0]["answer"]["answer_version"]["book_sha256"] = sha(body)
    summary = {"revision": 4, "applied": 1, "worker_applied": 0,
               "refused": [], "worker_rejected": [], "read_problems": [], "book_sha256": sha(body)}
    folder = run / PREFIX
    folder.mkdir()
    (folder / "anketa-otvety.xlsx").write_bytes(body)
    (folder / "answers-summary.json").write_bytes(raw(summary))
    (run / "extract/model.json").write_bytes(raw(doc))
    return run, doc, raw(source)


def build(gap, **kwargs):
    run, doc, source = gap
    return ae.build_recovery_witness(
        run, raw(doc), source, applied_revision=4,
        expected_source_model_sha256=sha(source),
        backup_manifest_sha256="b" * 64,
        backup_member="data/cdpm-runs/run-1/extract/model.json", **kwargs)


def install_fixture(gap):
    files = build(gap)
    assert set(files) == {WITNESS, SOURCE}
    for name, body in files.items():
        with (gap[0] / name).open("xb") as stream:
            stream.write(body)
    return files


def test_receipt_captures_pre_apply_model_revision_not_output_minus_one():
    doc = sample_model()
    doc["meta"]["revision"] = 2
    summary = answers.apply_answers(doc, {"Q-2-MODAL": "б"}, revision=4)
    assert summary["source_revision"] == 2
    assert summary["revision"] == doc["meta"]["revision"] == 4


def test_ordinary_upload_persists_exact_source_revision(tmp_path, monkeypatch):
    from test_answers_revision_allocation_gpt import _receiver, _filled_book, _history
    doc = sample_model()
    doc["meta"]["revision"] = 2
    receiver, run, state = _receiver(tmp_path, doc, monkeypatch, execute_worker=True)
    old = _history(run, 3)
    result = receiver(run.name, "answer.xlsx", _filled_book(tmp_path, doc))
    assert result["revision"] == 4
    summary = json.loads((run / PREFIX / "answers-summary.json").read_bytes())
    assert summary["source_revision"] == 2 and summary["revision"] == 4
    assert state["published"]["meta"]["revision"] == 4
    assert old == {p: p.read_bytes() for p in old}


def test_explicit_gapped_receipt_exports_its_issued_revision(gap):
    run, doc, _ = gap
    receipt = run / PREFIX / "answers-summary.json"
    summary = json.loads(receipt.read_bytes())
    summary["source_revision"] = 2
    receipt.write_bytes(raw(summary))
    files = ae.export_files(run, raw(doc), run.name)
    entry = json.loads(files[ae.EVIDENCE])["answers"][0]
    assert entry["answer"]["source_revision"] == 2
    assert entry["answer"]["revision"] == 4


def test_gap_without_proven_source_stays_refused(gap):
    run, doc, _ = gap
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


def test_posthoc_builder_is_read_only_deterministic_and_export_preserves_history(gap):
    run, doc, source = gap
    before = {p: p.read_bytes() for p in run.rglob("*") if p.is_file()}
    files = build(gap)
    assert files == build(gap)
    assert files[SOURCE] == source
    assert before == {p: p.read_bytes() for p in run.rglob("*") if p.is_file()}
    w = json.loads(files[WITNESS])
    assert w["kind"] == "posthoc-pre-apply-model" and w["source_revision"] == 2
    assert "source_revision" not in json.loads(before[run / PREFIX / "answers-summary.json"])
    install_fixture(gap)
    exported = ae.export_files(run, raw(doc), run.name)
    assert all(exported[name] == body for name, body in files.items())
    assert before == {p: p.read_bytes() for p in before}


@pytest.mark.parametrize("mode", ["explicit", "recovered"])
def test_gapped_package_roundtrip_revalidates_the_same_sources(gap, tmp_path, mode):
    run, doc, _ = gap
    if mode == "recovered":
        install_fixture(gap)
    else:
        p = run / PREFIX / "answers-summary.json"
        p.write_bytes(raw({**json.loads(p.read_bytes()), "source_revision": 2}))
    payloads = package_payloads(run, doc)
    packed = zip_body(payloads, manifest_edit=lambda m: m.update(revision=5))
    dest = tmp_path / "imported/run-1"
    import_sources(packed, dest)
    rows = ae.clauses_from_run(dest)
    assert len(rows) == 1 and "600 минут" in rows[0]["text"]
    assert rows[0]["answer_source"]["answer"]["source_revision"] == 2


@pytest.mark.parametrize("damage", [
    "source-revision", "applied-revision", "run", "book-sha", "summary-sha",
    "source-sha", "foreign-qid", "edition", "old-applied", "duplicate-qid",
    "model-revision", "stale-model", "read-problems",
])
def test_recovery_mutations_refuse_without_falling_back(gap, damage):
    run, doc, _ = gap
    install_fixture(gap)
    w = json.loads((run / WITNESS).read_bytes())
    src = json.loads((run / SOURCE).read_bytes())
    if damage == "source-revision":
        w["source_revision"] = 3
    elif damage == "applied-revision":
        w["revision"] = 5
    elif damage == "run":
        w["source_run"] = "another-run"
    elif damage in {"book-sha", "summary-sha", "source-sha"}:
        key = {"book-sha": "book", "summary-sha": "summary", "source-sha": "source_model"}[damage]
        w[key]["sha256"] = "0" * 64
    elif damage == "foreign-qid":
        w["questions"][0]["id"] = "Q-foreign"
    elif damage == "read-problems":
        p = run / PREFIX / "answers-summary.json"
        p.write_bytes(raw({**json.loads(p.read_bytes()), "read_problems": ["stale book"]}))
        w["summary"].update(sha256=sha(p.read_bytes()), bytes=len(p.read_bytes()))
    else:
        if damage == "edition":
            src["questions"][0]["edition"]["fingerprint"] = "0" * 12
        elif damage == "old-applied":
            src["questions"][0]["answer"] = deepcopy(doc["questions"][0]["answer"])
        elif damage == "duplicate-qid":
            src["questions"].append(deepcopy(src["questions"][0]))
        else:
            src["meta"]["revision"] = 1 if damage == "stale-model" else 3
            if damage == "stale-model":
                w["source_revision"] = 1  # reach real read_book stale guard
        body = raw(src)
        (run / SOURCE).write_bytes(body)
        w["source_model"].update(sha256=sha(body), bytes=len(body))
    (run / WITNESS).write_bytes(raw(w))
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


@pytest.mark.parametrize("name", [WITNESS, SOURCE])
@pytest.mark.parametrize("corrupt", [False, True])
def test_missing_or_corrupt_witness_refuses_export_and_import(gap, tmp_path, name, corrupt):
    run, doc, _ = gap
    install_fixture(gap)
    payloads = package_payloads(run, doc)
    if corrupt:
        (run / name).write_bytes(b"broken")
        payloads[name] = b"broken"
    else:
        (run / name).unlink()
        payloads.pop(name)
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)
    with pytest.raises(ae.EvidenceError):
        import_sources(zip_body(payloads, manifest_edit=lambda m: m.update(revision=5)), tmp_path / "dest/run-1")


@pytest.mark.parametrize("revision", [0, True, "2", 4, 5])
def test_bad_explicit_source_revision_never_uses_recovery(gap, revision):
    run, doc, _ = gap
    install_fixture(gap)
    p = run / PREFIX / "answers-summary.json"
    p.write_bytes(raw({**json.loads(p.read_bytes()), "source_revision": revision}))
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


def test_bytes_reader_has_the_unchanged_stale_and_mutation_guards(gap):
    run, _doc, source = gap
    body = (run / PREFIX / "anketa-otvety.xlsx").read_bytes()
    before = json.loads(source)
    good = answers.read_book(body, deepcopy(before))
    assert not good.problems and good.answers == {"Q-rest-outstation": "600 минут"}
    before["meta"]["revision"] = 3
    bad = answers.read_book(body, before)
    assert bad.refused and not bad.answers


def test_wrong_pinned_source_is_never_attested(gap):
    run, doc, source = gap
    with pytest.raises(ae.EvidenceError, match="pinned"):
        ae.build_recovery_witness(run, raw(doc), source, applied_revision=4,
            expected_source_model_sha256="0" * 64, backup_manifest_sha256="b" * 64,
            backup_member="data/cdpm-runs/run-1/extract/model.json")


@pytest.mark.parametrize("key,value", [("meta", []), ("questions", None), ("questions", [None])])
def test_malformed_pre_model_is_a_named_refusal(gap, key, value):
    run, doc, _ = gap
    install_fixture(gap)
    src = json.loads((run / SOURCE).read_bytes())
    src[key] = value
    body = raw(src)
    (run / SOURCE).write_bytes(body)
    w = json.loads((run / WITNESS).read_bytes())
    w["source_model"].update(sha256=sha(body), bytes=len(body))
    (run / WITNESS).write_bytes(raw(w))
    with pytest.raises(ae.EvidenceError):
        ae.export_files(run, raw(doc), run.name)


def test_post_import_source_change_is_rechecked(gap, tmp_path):
    run, doc, _ = gap
    install_fixture(gap)
    dest = tmp_path / "imported/run-1"
    import_sources(zip_body(package_payloads(run, doc), manifest_edit=lambda m: m.update(revision=5)), dest)
    (dest / SOURCE).write_bytes(b"changed after import")
    with pytest.raises(ae.EvidenceError):
        ae.clauses_from_run(dest)


def test_legacy_consecutive_package_remains_unchanged(answered_run):
    run, doc = answered_run
    files = ae.export_files(run, raw(doc), run.name)
    assert all("recovery" not in name for name in files)
    assert "source_revision" not in json.loads(files[ae.EVIDENCE])["answers"][0]["answer"]
