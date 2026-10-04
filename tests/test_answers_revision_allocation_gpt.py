"""Answer uploads reserve history paths without HTTP, threads or providers."""
import ast
import copy
import hashlib
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from cdpm import answers as ans
from cdpm import anketa_xlsx as ax
from cdpm.sample import sample_model


def _filled_book(tmp_path, doc):
    path = tmp_path / "input.xlsx"
    ax.render_anketa_xlsx(doc, path)
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == "Q-2-MODAL":
                sheet.cell(row[0].row, ax.ANSWER_COL).value = "б"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _receiver(tmp_path, doc, monkeypatch, *, execute_worker=False):
    """Compile the actual handler alone; no webui import/startup or HTTP app."""
    source = Path(__file__).parents[1] / "cdpm" / "webui.py"
    parsed = ast.parse(source.read_text(encoding="utf-8"))
    handler = next(node for node in parsed.body
                   if isinstance(node, ast.FunctionDef)
                   and node.name == "_upload_answers_bytes")
    handler.decorator_list = []
    run = tmp_path / "run-0001"
    (run / "extract").mkdir(parents=True)
    (run / "extract" / "model.json").write_text(json.dumps(doc), encoding="utf-8")
    state = {"journal": [], "read_revisions": [], "status": {}}
    original_read = ans.read_book

    def read_book(path, model):
        state["read_revisions"].append(model["meta"].get("revision", 1))
        return original_read(path, model)

    monkeypatch.setattr(ans, "read_book", read_book)

    class HTTPException(Exception):
        def __init__(self, status_code, detail):
            self.status_code, self.detail = status_code, detail
            super().__init__(detail)

    def forbidden(*args, **kwargs):
        raise AssertionError("No LLM/provider allowed")

    def publish(model, path):
        state["published"] = copy.deepcopy(model)
        path.write_text(json.dumps(model), encoding="utf-8")

    names = ("ensure_conditions_questions", "ensure_goal_questions_sweep",
             "ensure_criterion_question", "ensure_fit_questions",
             "restore_soft_penalty_rules", "resolve_case_orphans",
             "apply_topic_value_answers", "assign_soft_families",
             "derive_all_rule_traits", "contour_carriers_report",
             "validate_schema", "validate_model")
    model_stub = SimpleNamespace(**{name: lambda *args: [] for name in names},
                                 publish_model=publish)
    namespace = dict(
        Path=Path, json=json, HTTPException=HTTPException, answers_mod=ans,
        _run_dir=lambda run_id: run,
        _read_status=lambda path: {"state": "готово"},
        _line_busy=lambda value: False,
        _load_model_or_409=lambda path: copy.deepcopy(doc),
        _strictest_class=lambda path: "local",
        _state_brief=lambda status: status,
        actions_journal=SimpleNamespace(
            note=lambda *args, **kwargs: state["journal"].append(kwargs),
            file_entry=lambda name, data, relative: {
                "path": relative, "sha256": hashlib.sha256(data).hexdigest()}),
        _claim_stage=lambda *args, **kwargs: None,
        _control_path=lambda path: path / "control.json",
        _make_checkpoint=lambda *args: None,
        _channel_for=lambda *args, **kwargs: forbidden,
        _start_stage_thread=lambda path, work: work() if execute_worker else None,
        _update_status=lambda path, **kwargs: state["status"].update(kwargs),
        _stage_failure=lambda stage, error: {"state": stage, "error": str(error)},
        StopRequested=type("StopRequested", (Exception,), {}),
        m_mod=model_stub,
        answers_fp_mod=SimpleNamespace(book_content_sha256=lambda path: "a" * 64),
        family_params_mod=SimpleNamespace(ensure_family_param_questions=lambda model: []),
        render_mod=SimpleNamespace(render_svod=lambda model: "svod",
                                   render_anketa=lambda model: "anketa"),
        anketa_xlsx_mod=SimpleNamespace(render_anketa_xlsx=lambda *args: None),
        gate_answers=SimpleNamespace(check=lambda *args: ([], {})),
    )
    exec(compile(ast.Module(body=[handler], type_ignores=[]), str(source), "exec"), namespace)
    return namespace["_upload_answers_bytes"], run, state


def _history(run, *numbers):
    files = {}
    for revision in numbers:
        directory = run / "answers" / f"rev-{revision:03d}"
        directory.mkdir(parents=True)
        for name in ("anketa-otvety.xlsx", "answers-summary.json"):
            path = directory / name
            path.write_bytes(f"historic-{revision}-{name}".encode())
            files[path] = path.read_bytes()
    return files


def test_reextract_upload_preserves_rev002_rev003(tmp_path, monkeypatch):
    doc = sample_model()
    assert doc["meta"].get("revision", 1) == 1
    receiver, run, state = _receiver(tmp_path, doc, monkeypatch)
    history = _history(run, 2, 3)
    data = _filled_book(tmp_path, doc)
    result = receiver("run-0001", "answers.xlsx", data)
    assert {path: path.read_bytes() for path in history} == history
    assert result["revision"] == 4
    assert (run / "answers/rev-004/anketa-otvety.xlsx").read_bytes() == data
    assert state["read_revisions"] == [1]
    assert state["journal"][0]["files"][0]["path"] == "answers/rev-004/anketa-otvety.xlsx"


@pytest.mark.parametrize("current,occupied,expected", [(1, (), 2), (5, (2, 3), 6)])
def test_healthy_upload_number(tmp_path, monkeypatch, current, occupied, expected):
    doc = sample_model()
    doc["meta"]["revision"] = current
    receiver, run, state = _receiver(tmp_path, doc, monkeypatch)
    history = _history(run, *occupied)
    result = receiver("run-0001", "answers.xlsx", _filled_book(tmp_path, doc))
    assert result["revision"] == expected
    assert {path: path.read_bytes() for path in history} == history
    assert state["read_revisions"] == [current]


def test_actual_answer_revision_matches_reserved_path_model_and_ui(tmp_path, monkeypatch):
    doc = sample_model()
    receiver, run, state = _receiver(tmp_path, doc, monkeypatch, execute_worker=True)
    history = _history(run, 2, 3)
    result = receiver("run-0001", "answers.xlsx", _filled_book(tmp_path, doc))
    assert state["status"].get("error") is None, state["status"]
    assert result["revision"] == state["published"]["meta"]["revision"] == 4
    answer = next(q for q in state["published"]["questions"] if q["id"] == "Q-2-MODAL")["answer"]
    assert answer["revision"] == 4
    assert state["status"]["answers_summary"]["revision"] == 4
    summary = json.loads((run / "answers/rev-004/answers-summary.json").read_text())
    assert summary["revision"] == 4
    assert {path: path.read_bytes() for path in history} == history


def test_old_input_revision_stays_refused_despite_free_output_revision(tmp_path, monkeypatch):
    old = sample_model()
    data = _filled_book(tmp_path, old)
    doc = copy.deepcopy(old)
    doc["meta"]["revision"] = 2
    receiver, run, state = _receiver(tmp_path, doc, monkeypatch, execute_worker=True)
    history = _history(run, 2, 3)
    result = receiver("run-0001", "answers.xlsx", data)
    assert result["accepted"] == 0 and result["refused"] == 1
    assert state["read_revisions"] == [2]
    assert state["status"].get("error") is None, state["status"]
    answer = next(q for q in state["published"]["questions"] if q["id"] == "Q-2-MODAL")["answer"]
    assert answer["status"] == "не разобран"
    assert state["published"]["meta"]["revision"] == 4
    assert {path: path.read_bytes() for path in history} == history


@pytest.mark.parametrize("current,occupied,expected", [
    (1, (), 2), (1, (2, 3), 4), (3, (1, 9), 10), (7, (), 8), (1, (1000,), 1001),
])
def test_allocator_above_model_and_all_occupied_numbers(tmp_path, current, occupied, expected):
    root = tmp_path / "answers"
    root.mkdir()
    for number in occupied:
        # Even an occupied file rather than a valid history directory is kept.
        (root / f"rev-{number:03d}").write_bytes(b"preserve")
    revision, directory = ans.reserve_revision_directory(root, current)
    assert revision == expected and directory.is_dir()
    assert all((root / f"rev-{number:03d}").read_bytes() == b"preserve" for number in occupied)


def test_empty_failed_upload_slot_is_never_reused(tmp_path):
    root = tmp_path / "answers"
    (root / "rev-004").mkdir(parents=True)
    first, first_dir = ans.reserve_revision_directory(root, 1)
    second, second_dir = ans.reserve_revision_directory(root, 1)
    assert (first, second) == (5, 6)
    assert first_dir.is_dir() and second_dir.is_dir()


def test_concurrent_slot_claim_reallocates_without_overwrite(tmp_path, monkeypatch):
    root = tmp_path / "answers"
    (root / "rev-003").mkdir(parents=True)
    original = Path.mkdir
    collided = []

    def racing_mkdir(path, *args, **kwargs):
        if path == root / "rev-004" and not collided:
            collided.append(path)
            original(path)
            (path / "winner.txt").write_bytes(b"other upload")
            raise FileExistsError(path)
        return original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "mkdir", racing_mkdir)
    revision, directory = ans.reserve_revision_directory(root, 1)
    assert revision == 5 and directory.name == "rev-005"
    assert (root / "rev-004/winner.txt").read_bytes() == b"other upload"


def test_book_exclusive_create_preserves_collision_bytes(tmp_path, monkeypatch):
    doc = sample_model()
    receiver, run, state = _receiver(tmp_path, doc, monkeypatch)
    original = ans.reserve_revision_directory

    def collided(*args):
        revision, directory = original(*args)
        (directory / "anketa-otvety.xlsx").write_bytes(b"do not replace")
        return revision, directory

    monkeypatch.setattr(ans, "reserve_revision_directory", collided)
    with pytest.raises(FileExistsError):
        receiver("run-0001", "answers.xlsx", _filled_book(tmp_path, doc))
    assert (run / "answers/rev-002/anketa-otvety.xlsx").read_bytes() == b"do not replace"
    assert not state["journal"]


@pytest.mark.parametrize("revision", [0, 1, True, "4", 2.0])
def test_explicit_output_revision_must_be_monotonic_integer(revision):
    doc = sample_model()
    before = copy.deepcopy(doc)
    with pytest.raises(ans.AnswersError, match="ревизия"):
        ans.apply_answers(doc, {}, revision=revision)
    assert doc == before


def test_legacy_apply_default_still_increments_once():
    doc = sample_model()
    result = ans.apply_answers(doc, {})
    assert result["revision"] == doc["meta"]["revision"] == 2


def test_changed_question_fingerprint_stays_refused(tmp_path, monkeypatch):
    doc = sample_model()
    data = _filled_book(tmp_path, doc)
    question = next(q for q in doc["questions"] if q["id"] == "Q-2-MODAL")
    question["text"] += " Новое издание вопроса."
    receiver, run, state = _receiver(tmp_path, doc, monkeypatch)
    _history(run, 2, 3)
    result = receiver("run-0001", "answers.xlsx", data)
    assert result["accepted"] == 0 and result["refused"] == 1
    assert result["revision"] == 4 and state["read_revisions"] == [1]
