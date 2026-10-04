"""Explain actual scope growth on both delivered questionnaire surfaces."""
import copy

import pytest
from openpyxl import load_workbook

from cdpm import anketa_xlsx as ax, model as m, question_editions as qe, render
from test_cdpm_conditions_editions import ROOT, initial, rule, topics


def shown(doc, q, surface, tmp_path):
    if surface == "md":
        text = render.render_anketa(doc)
        assert q["id"] in text
        return text
    path = tmp_path / "anketa.xlsx"
    ax.render_anketa_xlsx(doc, path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [row for ws in wb for row in ws.iter_rows(values_only=True)
                if row and row[0] == q["id"]]
        assert len(rows) == 1
        return "\n".join(str(cell) for cell in rows[0] if cell is not None)
    finally:
        wb.close()


@pytest.mark.parametrize("surface", ["md", "xlsx"])
@pytest.mark.parametrize("writer", ["ensure_conditions_questions", "ensure_pairing_clarification_questions"])
@pytest.mark.parametrize("change", ["target_new_sources", "target_same_sources", "source_only", "no_op"])
def test_scope_presentation_names_actual_change(topics, monkeypatch, tmp_path, surface, writer, change):
    monkeypatch.setattr(m, "_clarification_topics", lambda customer: topics)
    ensure = getattr(m, writer)
    doc, q = initial(2)
    qe.freeze(q)
    q["answer"] = {"raw": "10 часов", "status": "применён", "revision": 1}
    before = copy.deepcopy(q)
    if change.startswith("target_"):
        added = rule(2)
        if change == "target_same_sources":
            added["sources"] = copy.deepcopy(doc["rules"][0]["sources"])
        doc["rules"].append(added)
    elif change == "source_only":
        doc["rules"][0]["sources"].append({"sid": "S1", "anchor": "п.новый"})
    assert ensure(doc) == ([] if change == "no_op" else [ROOT + "-E2"])
    current = qe.current_for_root(doc, ROOT)
    snapshot = copy.deepcopy(doc)
    text = shown(doc, current, surface, tmp_path)
    assert doc == snapshot and q == before
    if change.startswith("target_"):
        assert "Добавлены строки в область вопроса: R-2" in text
        assert "Ответ на это издание относится ко всем 3 строкам" in text
    else:
        assert "Добавлены строки в область вопроса" not in text
    assert ("Добавлены нормативные источники" in text) == (change in ("target_new_sources", "source_only"))


@pytest.mark.parametrize("surface", ["md", "xlsx"])
def test_scope_presentation_names_only_this_editions_growth(topics, tmp_path, surface):
    doc, q = initial()
    qe.freeze(q)
    doc["rules"].append(rule(1))
    assert m.ensure_conditions_questions(doc) == [ROOT + "-E2"]
    previous = copy.deepcopy(doc["questions"])
    doc["rules"].append(rule(2))
    assert m.ensure_conditions_questions(doc) == [ROOT + "-E3"]
    current = qe.current_for_root(doc, ROOT)
    text = shown(doc, current, surface, tmp_path)
    assert "Добавлены строки в область вопроса: R-2" in text
    assert "Добавлены строки в область вопроса: R-1" not in text
    assert doc["questions"][:-1] == previous
