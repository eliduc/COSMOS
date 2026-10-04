# -*- coding: utf-8 -*-
"""Книга данных набора ИДАК H0 (ультимативная задача владельца 03.09,
пункт 2): правила H0 на ней выполнены и СВЯЗЫВАЮТ — у каждого порога
случай в пределах часа, книга замерзает и проходит гейт заморозки,
листы и колонки — по контракту данных пакета (COSMOS)."""
import importlib.util
import json
import sys
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent.parent
KIT = HERE / "open-bench" / "polyaris-h0"


def _builder():
    spec = importlib.util.spec_from_file_location(
        "build_data_h0", KIT / "build_data_h0.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules["build_data_h0"] = mod
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def bd():
    return _builder()


def test_self_check_clean_and_rules_bind(bd):
    problems, facts = bd.self_check()
    assert problems == [], problems
    # Каждое жёсткое правило с порогом имеет случай у самого порога:
    # запас не больше часа, иначе правило книгу не связывает (Д8).
    for rid, gap in facts["slack_minutes"].items():
        assert 0 <= gap <= 60, (rid, gap)
    assert facts["roles"] == ["Второй пилот", "КВС"]        # crew_size = 2
    assert facts["weekday_of_first_day"] == "Friday"        # надбавка VN-D91
    assert facts["crew_by_base"] == {"PLY": 12, "SVG": 6, "GRT": 6}


def test_reserve_is_proved_by_witness_assignment(bd):
    """Резерв доказывается СВИДЕТЕЛЬСТВОМ (находка решателя 03.09: −4hard
    «Reserve crews»): смены раздаются базам, и в худший день свободных
    пилотов не меньше порога — ночёвка занимает экипаж два дня."""
    witness, problems = bd.reserve_witness()
    assert problems == [], problems
    assert witness["need"] == {"PLY": 4, "SVG": 2, "GRT": 2}
    for base, need in witness["need"].items():
        assert witness["free_min"][base] >= need, (base, witness["free_min"])
    # Ночёвка занимает два календарных дня — это видно в плане смен.
    plan = bd.duty_plan()
    night = [x for x in plan if x["duty"] == "ночёвка"]
    assert night and all(len(x["days"]) == 2 for x in night)
    # Прежний состав 8/4/4 свидетельства не даёт (был −4hard у решателя).
    old = bd.CREW_AT
    try:
        bd.CREW_AT = (("PLY", 8), ("SVG", 4), ("GRT", 4))
        _w, probs = bd.reserve_witness()
        assert probs, "на прежнем составе свидетельство обязано не строиться"
    finally:
        bd.CREW_AT = old


def test_допуск_по_типу_выполним_на_всём_плановом_периоде(bd):
    """VN-070 (расширение H0 03.09): тип каждого рейса — из парка
    набора, и у КАЖДОГО пилота допуск на этот тип действует до конца
    планового периода. Иначе правило допуска делает книгу
    невыполнимой ещё до решателя."""
    problems, facts = bd.self_check()
    assert problems == [], problems
    assert facts["flight_types"] == [bd.AIRCRAFT]
    assert facts["crew_with_valid_rating"] == len(bd.crews())
    # Просроченный допуск обязан ломать самопроверку.
    orig = bd.qualifications
    try:
        bd.qualifications = lambda crew: [
            dict(r, **({"Действует до": "2026-01-01"} if i == 0 else {}))
            for i, r in enumerate(orig(crew))]
        probs, _f = bd.self_check()
        assert any("VN-070" in p for p in probs), probs
    finally:
        bd.qualifications = orig


def test_tables_match_package_data_contract(bd):
    t = bd.tables()
    assert set(t) == {"Рейсы", "Экипажи", "Отпуска", "Допуски", "Свойства",
                      "Аэропорты"}
    assert t["Рейсы"]["columns"][:6] == ["Номер рейса", "Дата", "Вылет",
                                         "Время вылета", "Прилёт", "Время прилёта"]
    assert {"Табельный номер", "База", "Должность"} <= set(t["Экипажи"]["columns"])
    assert t["Отпуска"]["columns"] == ["Табельный номер", "С", "По"]
    keys = {r["ключ"] for r in t["Свойства"]["rows"]}
    assert keys == {"planSince", "planUntil"}
    assert {r["Код"] for r in t["Аэропорты"]["rows"]} == set(bd.BASES)
    # Коды баз и тип судна встречаются в данных (сверка кодов приёмника).
    assert {r["Вылет"] for r in t["Рейсы"]["rows"]} | \
           {r["Прилёт"] for r in t["Рейсы"]["rows"]} == set(bd.BASES)
    assert {r["Тип ВС"] for r in t["Рейсы"]["rows"]} == {bd.AIRCRAFT}


def test_book_freezes_and_passes_freeze_gate(bd, tmp_path):
    from cdpm import data_freeze as df
    from cdpm.gates import gate_data_freeze as gate
    src = KIT / "data" / "ID-H0-DATA.xlsx"
    if not src.is_file():
        bd.write_xlsx(src)
    out = tmp_path / "data-freeze"
    counters = df.freeze_data([src], out)
    assert counters["files"] == 1 and counters["sheets"] == 6
    problems, gate_counters = gate.check(out, src.parent)
    assert problems == [], problems
    assert gate_counters["rows"] == counters["rows"]
    dataset = json.loads((out / "dataset.json").read_text(encoding="utf-8"))
    assert dataset["tables"]


def test_written_book_matches_generator(bd):
    """Файл в наборе — тот, что порождает генератор (детерминированность:
    книга не правится руками)."""
    import openpyxl
    src = KIT / "data" / "ID-H0-DATA.xlsx"
    assert src.is_file(), "книги данных нет — запустите build_data_h0.py --write"
    wb = openpyxl.load_workbook(str(src), read_only=True)
    got = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)]
           for ws in wb.worksheets}
    want = bd.tables()
    assert set(got) == set(want)
    for name, sheet in want.items():
        rows = got[name]
        assert rows[0] == sheet["columns"], name
        assert len(rows) - 1 == len(sheet["rows"]), name
