# -*- coding: utf-8 -*-
"""Выгрузки решений текущего пакета постановки в известных форматах
(слово владельца 06.09.2026: «пейринг и ростеринг выгружать также в
разных известных форматах, в том числе Meridian»).

Чистые функции: модуль не читает диск и не знает «сейчас». Артефакты
выдачи (delivery.json, builder-catalog.json, builder-payload.json,
solution.json, solved.json, manifest.json и байты model.json пакета)
собирает `webapp/package_runs.export_sources`, сюда они приходят уже
разобранными (`PairingSources` / `RosterSources`). Выгрузка полностью
восстановима из артефактов, поэтому не хранится: ручка строит файл по
требованию, а `?attempt=<id>` даёт выгрузку любого снимка попытки.

Форматы: пейринг — CSV, XLSX, JSON, Meridian (спецификация COS-Meridian
1.17, форма гейта линии BA); ростер — CSV, XLSX, JSON, iCal (RFC 5545),
Meridian (ПРОЕКТ формы `type: "roster"`, спецификацией не подтверждён —
пометка в описании файла, заголовке ответа и интерфейсе).

Принципы:
  * сертификация: рабочие виды (CSV/XLSX/JSON/iCal) несертифицированного
    решения выгружаются С ПОМЕТКОЙ (тело, имя файла `-uncertified`,
    заголовок X-COSMOS-Certified: false); Meridian — вход системы учёта
    авиакомпании — отказ словами, кроме явного `allow_uncertified`;
    `certified: null` читается как «не установлено» = не сертифицировано;
  * шкала времени: решение пишется в шкале НАБОРА, которому отвечает.
    Умолчание — метки как есть (ровно так пишет набор для Meridian
    `cdpm/export_meridian.py`: legs[].start/end — метки листа, legId по
    дате листа); перевод в UTC — только явным параметром `utc=1`, и тогда
    переводятся участки, границы смен и planSince/planUntil, а legId
    строится по переведённой дате. Шкала — из контракта пакета
    (`manifest.time`: time_scale + base_time_zone IANA, раздел 12у
    руководства), не из догадок; не объявлена — как есть с пометкой;
    объявлена `shifted` (ненулевой сдвиг CDPM) — перевод не определён
    контрактом, отказ от перевода назван этим словом, не «не объявлена»;
  * `created` — метка журнала кнопки (момент выдачи, не «сейчас»):
    журнал пишет местное время машины без пояса — пояс не выдумывается,
    об этом сказано в описании файла и заголовке X-COSMOS-Created-Scale;
    появится `at_utc` либо `at` со смещением — выгрузка перейдёт на UTC
    сама; журнала нет — отказ словами;
  * идентификаторы: datasetId/presetId — явный параметр человека →
    объявление источника в манифесте пакета (`source.meridian`) →
    отпечаток пакета / sha256 байт model.json с оговоркой в описании;
  * алфавит роли участка — параметр (`gate`: Р/П — форма гейта BA и
    данных заказчика; `source`: константы писателя `_payload_to_meridian`
    и импортёра `_meridian_to_payload` — Р/«D/H»), записан в описании;
    третьего словаря в коде нет — берутся существующие константы;
  * персональные данные: табельный номер — достаточная личность для всех
    форматов; «Фамилия И.О.» и «Должность» листа «Экипажи» не выгружаются;
  * детерминированность: сортировка по ключам, `created` из выдачи,
    фиксированные даты, система создания и права записей zip у XLSX
    (одни байты и на Windows, и на Linux), DTSTAMP = created у iCal (в
    форме UTC по RFC 5545, шкала штампа объявлена словами) — два вызова
    дают одни байты;
  * отказы словами: `ExportNotFound` (404 — нет пакета/выдачи/попытки/
    сотрудника/формата), `ExportRefused` (409 — выдача-остановка, формат
    не выражает решение, Meridian несертифицированного без флага),
    `ExportBadRequest` (400 — кривой параметр).
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from typing import Any, Callable

FORM_VERSION = "cosmos-export/1"

FORMATS = {"pairing": ("csv", "xlsx", "json", "meridian"),
           "roster": ("csv", "xlsx", "json", "ics", "meridian")}
MEDIA = {"csv": "text/csv; charset=utf-8",
         "xlsx": ("application/vnd.openxmlformats-officedocument"
                  ".spreadsheetml.sheet"),
         "json": "application/json",
         "ics": "text/calendar; charset=utf-8",
         "meridian": "application/json"}
EXT = {"csv": "csv", "xlsx": "xlsx", "json": "json", "ics": "ics",
       "meridian": "json"}

#: Пути ростера: каталог выдачи по имени пути.
ROSTER_PATHS = {"cold": "button", "on-pairing": "button-on-pairing"}

#: Алфавит роли участка в наряде Meridian: «gate» — форма гейта линии BA
#: и данных заказчика (Р = работа, П = пассажиром); «source» — константы
#: существующих писателя и импортёра COSMOS (Р / «D/H»). Третьей копии
#: словаря здесь нет: «source» читается из модулей, у которых он живёт.
GATE_ROLES = ("Р", "П")


def role_alphabets() -> dict[str, tuple[str, str]]:
    try:
        from _payload_to_meridian import OPERATING_DUTY_ROLE as _work
        from _meridian_to_payload import DEADHEAD_DUTY_ROLE as _dh
        source = (str(_work), str(_dh))
    except Exception:                                      # noqa: BLE001
        source = ("Р", "D/H")
    return {"gate": GATE_ROLES, "source": source}


#: Деление FD/CC формата Meridian над каноном ролей COSMOS
#: (webapp.task_package.ROLE_ORDER = Captain, FirstOfficer,
#: FlightAttendant; стенд паритета следит, что объединение равно канону).
#: Роль вне словаря — отказ словами; объявление `division` в манифесте
#: пакета первично (просьба к CDPM — открытый вопрос владельцу).
DIVISION_ROLES = {"FD": frozenset({"Captain", "FirstOfficer"}),
                  "CC": frozenset({"FlightAttendant"})}

#: Защита от формул в CSV — то же правило, что `webapp/results._csv_safe`
#: (копия пяти строк: vmp не тянет webapp).
_CSV_DANGEROUS = ("=", "+", "-", "@", "\t", "\r")

_TRUE_WORDS = {"да", "true", "1", "yes", "y", "истина"}
_ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}$")
_SLOT_RE = re.compile(r"^(.*):slot:\d+$")
_LEGID_RE = re.compile(r"^(.+)\$(\d{2})\.(\d{2})\.(\d{4})\$(\d+)$")
_CARRIER_RE = re.compile(r"^([A-Za-zА-Яа-я0-9]{2})")
_ASCII_NAME_RE = re.compile(r"^[A-Za-z0-9_.-]{1,40}$")


# ────────────────────────────────────────────────────────── отказы словами

class ExportError(RuntimeError):
    """Выгрузка невозможна — с названной причиной, а не молча."""
    http_status = 409


class ExportNotFound(ExportError):
    """Нет того, что просили: пакета, выдачи, снимка попытки, сотрудника,
    формата (404)."""
    http_status = 404


class ExportRefused(ExportError):
    """Есть, но формат этого не выражает либо правило запрещает (409)."""
    http_status = 409


class ExportBadRequest(ExportError):
    """Параметр запроса вне словаря (400)."""
    http_status = 400


# ─────────────────────────────────────────────────────────────── источники

@dataclass(frozen=True)
class TimeScale:
    """Шкала меток пакета: kind ∈ base-local | utc | shifted | unknown;
    zone — IANA пояс базы (у base-local и shifted, когда объявлен);
    offset_min — объявленный сдвиг. `shifted` — слово контракта пакета
    (CDPM пишет его при ненулевом base_tz_offset_min): шкала объявлена,
    но смысл сдвига контрактом не закреплён (раздел 12у) — перевод в UTC
    не определён, и об этом говорится словами, а не «не объявлена»."""
    kind: str
    zone: str | None
    offset_min: int = 0

    @property
    def convertible(self) -> bool:
        return self.kind == "utc" or (self.kind == "base-local"
                                      and bool(self.zone))

    def words(self) -> str:
        if self.kind == "base-local":
            return ("base-local, пояс %s по контракту пакета"
                    % (self.zone or "не объявлен"))
        if self.kind == "utc":
            return "UTC по контракту пакета"
        if self.kind == "shifted":
            return ("shifted, сдвиг %d мин по контракту пакета — перевод в "
                    "UTC не определён" % self.offset_min)
        return "не объявлена пакетом"

    def header(self) -> str:
        """Значение заголовка X-COSMOS-Time-Scale для меток «как есть»:
        `base-local[;zone=<IANA>]` | `utc` | `shifted;offset=<мин>
        [;zone=<IANA>]` | `unknown`."""
        if self.kind == "shifted":
            return "shifted;offset=%d" % self.offset_min + \
                ((";zone=" + self.zone) if self.zone else "")
        return self.kind + ((";zone=" + self.zone) if self.zone else "")


@dataclass(frozen=True)
class PairingSources:
    delivery: dict
    catalog: dict
    payload: dict
    manifest: dict
    rules: dict | None
    certificate: dict | None
    model_sha256: str
    package_sha: str
    attempt: str | None = None


@dataclass(frozen=True)
class RosterSources:
    delivery: dict
    solution: dict
    solved: dict | None
    payload: dict
    manifest: dict
    certificate: dict | None
    model_sha256: str
    package_sha: str
    attempt: str | None = None
    path: str | None = None
    pairing_catalog: dict | None = None
    pairing_catalog_note: str | None = None


@dataclass(frozen=True)
class Options:
    """Параметры выгрузки (все — из запроса, умолчания честные)."""
    utc: bool = False
    roles: str | None = None
    dataset_id: str | None = None
    preset_id: str | None = None
    allow_uncertified: bool = False
    employee: str | None = None


# ─────────────────────────────────────────────────────────── время и метки

def _dt(value: Any) -> datetime:
    s = str(value)
    try:
        d = datetime.fromisoformat(s)
    except ValueError as exc:
        raise ExportRefused("метка времени не ISO 8601: %r" % s) from exc
    if d.tzinfo is not None:
        d = d.astimezone(timezone.utc).replace(tzinfo=None)
    return d.replace(microsecond=0)


def _iso(d: datetime) -> str:
    return d.strftime("%Y-%m-%dT%H:%M:%S")


def _norm(value: Any) -> str:
    return _iso(_dt(value))


def _minutes(t0: datetime, t1: datetime) -> int:
    return int(round((t1 - t0).total_seconds() / 60))


def time_scale(manifest: dict | None) -> TimeScale:
    """Шкала меток — из контракта пакета (manifest.time, раздел 12у):
    `time_scale` base-local + `base_time_zone` (IANA, проверяется через
    zoneinfo) либо utc; `shifted` (ненулевой `base_tz_offset_min` у CDPM)
    — объявлена, но непереводима; не объявлена — unknown."""
    t = (manifest or {}).get("time")
    if not isinstance(t, dict):
        return TimeScale("unknown", None, 0)
    kind = str(t.get("time_scale") or "").strip().lower()
    try:
        off = int(t.get("base_tz_offset_min") or 0)
    except (TypeError, ValueError):
        off = 0
    if kind == "utc":
        return TimeScale("utc", None, off)
    if kind in ("base-local", "shifted"):
        zone = str(t.get("base_time_zone") or "").strip() or None
        if zone:
            try:
                from zoneinfo import ZoneInfo
                ZoneInfo(zone)
            except Exception:                              # noqa: BLE001
                zone = None
        return TimeScale(kind, zone, off)
    return TimeScale("unknown", None, off)


def to_utc(label: Any, scale: TimeScale) -> datetime:
    """Перевод метки по объявленной шкале; unknown/непереводимая — как
    есть (вызывающий обязан сказать об этом словами)."""
    d = _dt(label)
    if scale.kind == "base-local" and scale.zone:
        from zoneinfo import ZoneInfo
        return (d.replace(tzinfo=ZoneInfo(scale.zone))
                .astimezone(timezone.utc).replace(tzinfo=None))
    return d


def _stamp(label: Any, scale: TimeScale, utc: bool) -> str:
    return _iso(to_utc(label, scale) if utc else _dt(label))


def scale_note(scale: TimeScale, utc: bool) -> str:
    """Слова о шкале меток файла — для описания Meridian и паспорта."""
    if utc:
        if scale.convertible:
            return ("шкала меток: UTC — переведено из %s по параметру "
                    "utc=1; legId построены по переведённой дате"
                    % scale.words())
        if scale.kind == "shifted":
            return ("шкала меток: shifted (сдвиг %d мин) — перевод в UTC "
                    "запрошен (utc=1), но не выполнен: смысл сдвига "
                    "контрактом пакета не закреплён — записано без "
                    "перевода" % scale.offset_min)
        return ("шкала меток: перевод в UTC запрошен (utc=1), но шкала "
                "пакетом не объявлена — записано без перевода")
    return ("шкала меток: %s; перевод в UTC не выполнен — как в наборе "
            "CDPM для Meridian" % scale.words())


def created_stamp(delivery: dict) -> tuple[str, str]:
    """Момент выдачи — последняя запись журнала кнопки: (метка, шкала).
    Шкала «utc», когда журнал несёт `at_utc` либо `at` со смещением;
    иначе «machine-local» (местное время машины нажатия, пояс не
    записан — vmp/run/button.py пишет time.strftime без %z)."""
    journal = (delivery or {}).get("journal")
    if not isinstance(journal, list) or not journal:
        raise ExportRefused("в выдаче нет журнала — момент выдачи не "
                            "установлен, выгрузка без штампа невозможна")
    last = journal[-1] if isinstance(journal[-1], dict) else {}
    at_utc = last.get("at_utc")
    if at_utc:
        return _norm(at_utc), "utc"
    at = last.get("at")
    if not at:
        raise ExportRefused("последняя запись журнала выдачи без момента "
                            "(at) — момент выдачи не установлен")
    try:
        parsed = datetime.fromisoformat(str(at))
    except ValueError as exc:
        raise ExportRefused("момент выдачи в журнале не ISO 8601: %r"
                            % at) from exc
    if parsed.tzinfo is not None:
        return _iso(parsed.astimezone(timezone.utc)
                    .replace(tzinfo=None)), "utc"
    return _iso(parsed.replace(microsecond=0)), "machine-local"


def created_note(scale_word: str) -> str:
    if scale_word == "utc":
        return "created: момент выдачи по журналу кнопки, UTC"
    return ("created: метка журнала кнопки, местное время машины, пояс "
            "не записан")


# ─────────────────────────────────────────────────────────── сертификация

def certification(delivery: dict) -> dict:
    """{certified, established, verdict, certificate_verdict, findings,
    result}: `certified: null`/отсутствие — не установлено = не
    сертифицировано."""
    d = delivery or {}
    raw = d.get("certified")
    return {"certified": raw is True,
            "established": isinstance(raw, bool),
            "verdict": d.get("verdict"),
            "certificate_verdict": d.get("certificate_verdict"),
            "findings": d.get("findings"),
            "result": d.get("result")}


def require_solution(delivery: dict, task: str) -> None:
    """Выдача-остановка или без решения — отказ словами (409)."""
    d = delivery or {}
    verdict = str(d.get("verdict") or "")
    if verdict.startswith("остановка"):
        raise ExportRefused("выдача %s — %s: решения нет, выгружать нечего"
                            % (task, verdict))
    sched = d.get("schedule")
    if not isinstance(sched, dict) or not sched.get("solver_run_id"):
        raise ExportRefused("выдача %s не несёт решения (нет schedule."
                            "solver_run_id, вердикт «%s») — выгружать "
                            "нечего" % (task, verdict or "нет"))


def _uncertified_gate(cert: dict, opts: Options, task: str) -> None:
    if not cert["certified"] and not opts.allow_uncertified:
        raise ExportRefused(
            "решение %s не сертифицировано (%s) — Meridian принимает только "
            "сертифицированное; выгрузка с пометкой «не сертифицировано» — "
            "явно параметром allow_uncertified=1"
            % (task, "сертификация не установлена" if not cert["established"]
               else "вердикт сертификата «%s»" % (cert.get("certificate_verdict")
                                                     or cert.get("verdict"))))


def uncertified_words(cert: dict) -> str:
    if cert["certified"]:
        return ""
    return ("НЕ СЕРТИФИЦИРОВАНО: %s"
            % ("сертификация не установлена" if not cert["established"]
               else "вердикт сертификата «%s»"
               % (cert.get("certificate_verdict") or cert.get("verdict"))))


# ───────────────────────────────────────────────────────── идентификаторы

def _source_meridian(manifest: dict | None) -> dict:
    src = (manifest or {}).get("source")
    if isinstance(src, dict):
        m = src.get("meridian")
        if isinstance(m, dict):
            return m
        if str(src.get("kind") or "").lower() == "meridian":
            return {k: v for k, v in src.items() if k != "kind"}
    return {}


def default_roles(manifest: dict | None) -> str:
    """Умолчание алфавита ролей: пакет из источника Meridian → «source»,
    иначе «gate»."""
    return "source" if _source_meridian(manifest) else "gate"


def ids(sources, opts: Options) -> dict:
    """datasetId / presetId с оговорками: явный параметр → объявление
    источника в манифесте → отпечаток пакета / sha256 model.json."""
    src = _source_meridian(sources.manifest)
    notes = []
    if opts.dataset_id:
        dataset = str(opts.dataset_id)[:150]
        notes.append("datasetId задан явно параметром запроса")
    elif src.get("datasetId"):
        dataset = str(src["datasetId"])[:150]
        notes.append("datasetId — объявление источника Meridian в "
                     "манифесте пакета")
    else:
        dataset = sources.package_sha
        notes.append("datasetId набора Meridian пакетом не объявлен — "
                     "записан отпечаток пакета")
    if opts.preset_id:
        preset = str(opts.preset_id)[:150]
        notes.append("presetId задан явно параметром запроса")
    elif src.get("presetId"):
        preset = str(src["presetId"])[:150]
        notes.append("presetId — объявление источника Meridian в "
                     "манифесте пакета")
    else:
        preset = sources.model_sha256
        notes.append("presetId — sha256 байт model.json пакета "
                     "(модель правил ВМП)")
    return {"datasetId": dataset, "presetId": preset, "notes": notes}


def division_of(seat_roles, manifest: dict | None) -> str:
    """FD/CC: объявление манифеста первично, иначе по ролям мест над
    каноном; роль вне словаря — отказ словами."""
    declared = str((manifest or {}).get("division") or "").upper()
    if declared in ("FD", "CC"):
        return declared
    roles = {str(r) for r in (seat_roles or []) if r}
    if not roles:
        raise ExportRefused("division не определим: у пакета нет ролей мест "
                            "и манифест не объявляет division")
    for code, known in DIVISION_ROLES.items():
        if roles <= known:
            return code
    raise ExportRefused("division не определим по ролям %s (словарь FD: %s; "
                        "CC: %s) — нужно объявление division в манифесте "
                        "пакета" % (sorted(roles),
                                    ", ".join(sorted(DIVISION_ROLES["FD"])),
                                    ", ".join(sorted(DIVISION_ROLES["CC"]))))


def _carrier(flight: dict) -> str:
    """Перевозчик участка: поле `carrier` рейса, иначе две первые буквы
    номера — то же правило, что у набора CDPM для Meridian
    (`cdpm/export_meridian._carrier_of`); иного источника нет (`_coverage`
    пакета перевозчиков обязательств не несёт), и отказ называет ровно
    то, что проверено."""
    if flight.get("carrier"):
        return str(flight["carrier"]).upper()
    m = _CARRIER_RE.match(str(flight.get("flightNumber") or ""))
    if not m:
        raise ExportRefused("перевозчик не определим для рейса %r: ни поля "
                            "carrier, ни двухзначного кода в номере"
                            % flight.get("flightNumber"))
    return m.group(1).upper()


# ────────────────────────────────────────────────────────────── участки

def leg_index(flights: list[dict], scale: TimeScale, utc: bool) -> dict[str, dict]:
    """Участки Meridian по всем рейсам данных (порядок — по вылету):
    legId = sourceLegId источника (если отличен от fn), иначе
    `<номер>$<дд.мм.гггг>$<n>` по дате start в ШКАЛЕ ФАЙЛА (n —
    порядковый среди того же номера и даты)."""
    out: dict[str, dict] = {}
    seen: dict[tuple[str, str], int] = {}
    ordered = sorted(flights, key=lambda f: (str(f.get("departureUTCDateTime")),
                                             str(f.get("flightNumber"))))
    for f in ordered:
        fn = str(f.get("flightNumber") or "")
        if not fn:
            raise ExportRefused("рейс данных без flightNumber")
        start = _stamp(f["departureUTCDateTime"], scale, utc)
        end = _stamp(f["arrivalUTCDateTime"], scale, utc)
        num = fn.split("|")[0]
        src = f.get("sourceLegId")
        if not src:
            fl_id = f.get("flightLegId")
            if fl_id and str(fl_id) != fn:
                src = fl_id
        if src and str(src) != fn:
            leg_id = str(src)
            m = _LEGID_RE.match(leg_id)
            leg_no = int(m.group(5)) if m else 1
        else:
            day = datetime.fromisoformat(start).strftime("%d.%m.%Y")
            n = seen.get((num, day), 0) + 1
            seen[(num, day)] = n
            leg_id, leg_no = "%s$%s$%d" % (num, day, n), n
        out[fn] = {
            "legId": leg_id, "flightNumber": num, "legNumber": leg_no,
            "start": start, "end": end,
            "departureStation": str(f.get("departureAirport") or ""),
            "arrivalStation": str(f.get("arrivalAirport") or ""),
            "carrier": _carrier(f),
            "fleet": str(f.get("aircraftType") or f.get("fleet") or ""),
            "properties": str(f.get("properties") or "REGULAR"),
            "serviceType": str(f.get("serviceType") or "J"),
            "deadhead": bool(f.get("transportOnly")),
            "dutyCrewTypes": {}}
    return out


def _leg_notes(flights: list[dict]) -> list[str]:
    notes = []
    if not any(f.get("properties") for f in flights):
        notes.append("свойства рейса (properties) пакетом не объявлены — "
                     "REGULAR по умолчанию формата")
    if not any(f.get("serviceType") for f in flights):
        notes.append("код обслуживания (serviceType) пакетом не объявлен — "
                     "J по умолчанию формата")
    return notes


# ───────────────────────────────────────────────── пейринг: модель выгрузки

def _label12(catalog: dict, package_sha: str) -> str:
    return (str(catalog.get("label") or "")[-12:] or package_sha[:12])


def _rest_floors(rules: dict | None, catalog: dict) -> tuple[int, int] | None:
    """Полы отдыха дома / вне базы для восстановления смен: source-rules
    (min_rest_* либо первый ярус layover_tiers) → config_sig каталога;
    нет — None (число ниоткуда не берётся)."""
    r = rules or {}
    home, away = r.get("min_rest_home_min"), r.get("min_rest_outstation_min")
    tiers = r.get("layover_tiers")
    if (home is None or away is None) and isinstance(tiers, list) and tiers \
            and isinstance(tiers[0], (list, tuple)) and len(tiers[0]) >= 3:
        home = tiers[0][1] if home is None else home
        away = tiers[0][2] if away is None else away
    sig = catalog.get("config_sig") if isinstance(catalog.get("config_sig"),
                                                  dict) else {}
    if home is None:
        home = sig.get("rest_base")
    if away is None:
        away = sig.get("rest_off")
    if home is None or away is None:
        return None
    return int(home), int(away)


def _chain_periods(unit: dict, fns: list[str], active: set[str],
                   fl: dict[str, dict], floors: tuple[int, int] | None,
                   bases: set[str]) -> tuple[list[tuple[datetime, datetime,
                                                        list[str]]], str | None]:
    """Смены цепочки: (1) периоды единой модели цены построителя
    (price.accounting_scope.periods), если они покрывают все участки;
    (2) границы duty_fns построителя с перегонами у ближайшей смены
    (правило писателя _payload_to_meridian); (3) _duty_periods.build_periods
    по полам отдыха. Своей нарезки смен здесь нет."""
    def _key(fn: str) -> datetime:
        return _dt(fl[fn]["departureUTCDateTime"])

    price = unit.get("price") if isinstance(unit.get("price"), dict) else {}
    scope = price.get("accounting_scope") if isinstance(
        price.get("accounting_scope"), dict) else {}
    per = scope.get("periods")
    if isinstance(per, list) and per and all(isinstance(p, dict) for p in per):
        union = {str(x) for p in per for x in (p.get("fns") or [])}
        if union == set(fns) and all(p.get("start") and p.get("end")
                                     for p in per):
            out = [(_dt(p["start"]), _dt(p["end"]),
                    sorted((str(x) for x in p["fns"]), key=_key))
                   for p in per]
            out.sort(key=lambda t: t[0])
            return out, None
    groups = [[str(x) for x in g] for g in (unit.get("duty_fns") or []) if g]
    if groups:
        if {x for g in groups for x in g} != {x for x in fns if x in active}:
            groups = []
    if groups:
        groups = [sorted(g, key=_key) for g in groups]
        groups.sort(key=lambda g: _key(g[0]))
        ride = sorted((x for x in fns if x not in active), key=_key)
        for x in ride:
            t = _key(x)
            best = min(range(len(groups)),
                       key=lambda k: abs((t - _key(groups[k][0])).total_seconds()))
            groups[best].append(x)
            groups[best].sort(key=_key)
        out = [(_key(g[0]), max(_dt(fl[x]["arrivalUTCDateTime"]) for x in g), g)
               for g in groups]
        return out, ("границы смен — по вылету и прилёту участков по "
                     "duty_fns построителя (подготовка/разбор пакетом не "
                     "объявлены)")
    if floors is None:
        raise ExportRefused("смены цепочки восстановить нельзя: каталог не "
                            "несёт периодов и границ смен, а полы отдыха "
                            "(source-rules) не объявлены")
    import _duty_periods as dp
    acts = [dp.Activity(x, _key(x), _dt(fl[x]["arrivalUTCDateTime"]),
                        fl[x]["departureAirport"], fl[x]["arrivalAirport"],
                        x in active, x) for x in fns]
    home, away = floors
    pers = dp.build_periods(acts, lambda p, st: home if st in bases else away,
                            0, 0)
    out = [(p.start, p.end, [a.payload for a in p.activities]) for p in pers]
    return out, ("периоды восстановлены единой моделью _duty_periods "
                 "(подготовка/разбор 0 — пакетом не объявлены)")


def _demand(payload: dict) -> dict[str, dict[str, int]]:
    """Спрос мест участка по ролям — из flightAssignments данных."""
    out: dict[str, dict[str, int]] = {}
    for a in payload.get("flightAssignments") or []:
        if not isinstance(a, dict) or not a.get("flight") or not a.get("requiredSkill"):
            continue
        d = out.setdefault(str(a["flight"]), {})
        d[str(a["requiredSkill"])] = d.get(str(a["requiredSkill"]), 0) + 1
    return out


def _expand(counts: dict[str, int], order: list[str]) -> list[str]:
    roles = sorted(counts, key=lambda r: (order.index(r) if r in order
                                          else len(order), r))
    return [r for r in roles for _ in range(counts[r])]


def pairing_chains(src: PairingSources) -> list[dict]:
    """Единая модель выгрузки пейринга: цепочки с развёрнутыми участками
    (аэропорты/времена/тип ВС), сменами, местами, ценой. Все форматы
    строятся из неё. Метки — как в артефактах (шкала пакета)."""
    payload, catalog = src.payload, src.catalog
    fl = {str(f.get("flightNumber")): f for f in payload.get("flights") or []}
    if not fl:
        raise ExportRefused("данные пакета не несут рейсов — выгружать нечего")
    units = catalog.get("selected")
    if not isinstance(units, list) or not units:
        raise ExportRefused("каталог цепочек пуст — решения пейринга нет")
    label12 = _label12(catalog, src.package_sha)
    floors = _rest_floors(src.rules, catalog)
    bases = {str(e.get("homeAirport")) for e in payload.get("employees") or []
             if isinstance(e, dict) and e.get("homeAirport")}
    demand = _demand(payload)
    role_order = [str(r) for r in payload.get("_seatRoles") or []]
    chains: list[dict] = []
    for idx, unit in enumerate(units, start=1):
        if not isinstance(unit, dict):
            raise ExportRefused("запись цепочки %d каталога не объект" % idx)
        fns = [str(x) for x in (unit.get("fns") or unit.get("active_fns") or [])]
        if not fns:
            raise ExportRefused("цепочка %d каталога без единого участка" % idx)
        missing = [x for x in fns if x not in fl]
        if missing:
            raise ExportRefused("рейс цепочки %d не найден в данных пакета: %s"
                                % (idx, ", ".join(missing[:5])))
        active = {str(x) for x in (unit.get("active_fns") or fns)}
        periods, note = _chain_periods(unit, fns, active, fl, floors, bases)
        chain_id = "COSMOS-%s-%04d" % (label12, idx)
        notes = [note] if note else []
        legs, duties = [], []
        leg_no = 0
        for n, (start, end, members) in enumerate(periods, start=1):
            work_min = _minutes(start, end)
            flight_min = 0
            stations: list[str] = []
            need: dict[str, int] = {}
            for x in members:
                f = fl[x]
                leg_no += 1
                dep, arr = _dt(f["departureUTCDateTime"]), _dt(f["arrivalUTCDateTime"])
                working = x in active
                if working:
                    flight_min += _minutes(dep, arr)
                    for role, cnt in demand.get(x, {}).items():
                        need[role] = max(need.get(role, 0), cnt)
                for st in (f["departureAirport"], f["arrivalAirport"]):
                    if st not in stations:
                        stations.append(st)
                legs.append({"fn": x, "flight_no": x.split("|")[0],
                             "leg_no": leg_no, "duty_no": n,
                             "date": _iso(dep)[:10],
                             "dep": f["departureAirport"], "dep_time": _iso(dep),
                             "arr": f["arrivalAirport"], "arr_time": _iso(arr),
                             "aircraft": str(f.get("aircraftType") or ""),
                             "activity": "work" if working else "deadhead"})
            duties.append({"no": n, "start": _iso(start), "end": _iso(end),
                           "legs": len(members), "flight_min": flight_min,
                           "work_min": work_min, "first_dep": fl[members[0]]["departureAirport"],
                           "last_arr": fl[members[-1]]["arrivalAirport"],
                           "stations": stations, "demand": need,
                           "_start": start, "_end": end})
        rest_after = int(round(float(unit.get("minRestAfterMin") or 0)))
        for a, b in zip(duties, duties[1:]):
            a["rest_after_min"] = _minutes(a["_end"], b["_start"])
        duties[-1]["rest_after_min"] = rest_after
        for d in duties:
            d.pop("_start"); d.pop("_end")
        seats = [str(s) for s in (unit.get("seats") or [])]
        if not seats:
            need_all: dict[str, int] = {}
            for d in duties:
                for role, cnt in d["demand"].items():
                    need_all[role] = max(need_all.get(role, 0), cnt)
            seats = _expand(need_all, role_order)
            notes.append("места цепочки каталогом не объявлены — состав по "
                         "спросу участков (поролевые максимумы)")
        seat_counts: dict[str, int] = {}
        for s in seats:
            seat_counts[s] = seat_counts.get(s, 0) + 1
        for d in duties:
            for role, cnt in d["demand"].items():
                if seat_counts.get(role, 0) < cnt:
                    raise ExportRefused(
                        "места цепочки %s не покрывают спрос смены %d: роль %s "
                        "— мест %d, участок требует %d"
                        % (chain_id, d["no"], role, seat_counts.get(role, 0), cnt))
            d["extended"] = len(seats) > sum(d["demand"].values()) if d["demand"] else False
        price = unit.get("price") if isinstance(unit.get("price"), dict) else {}
        chains.append({
            "id": chain_id, "no": idx, "k": idx - 1,
            "base": fl[legs[0]["fn"]]["departureAirport"] if legs else "",
            "start": duties[0]["start"], "end": duties[-1]["end"],
            "legs": legs, "duties": duties,
            "seats": seats,
            "seat_codes": [[str(c) for c in (grp or [])]
                           for grp in (unit.get("seatCodes") or [])],
            "pinned": bool(unit.get("pinned")),
            "source_unit": str(unit.get("sourceUnit") or ""),
            "kind": str(unit.get("kind") or ""),
            "rest_after_min": rest_after,
            "price": {"cost_known": price.get("known_subtotal",
                                              unit.get("known_subtotal",
                                                       unit.get("cost"))),
                      "full_cost": price.get("full_cost", unit.get("full_cost")),
                      "certainty": price.get("cost_certainty",
                                             unit.get("cost_certainty")),
                      "nights": price.get("nights"), "crew": price.get("crew")},
            "notes": notes})
    return chains


def _coverage_numbers(catalog: dict, chains: list[dict],
                      payload: dict) -> dict:
    cov = catalog.get("coverage") if isinstance(catalog.get("coverage"), dict) else {}
    flights = payload.get("flights") or []
    obligations = cov.get("legs_total")
    if obligations is None:
        obligations = sum(1 for f in flights if not f.get("transportOnly"))
    obligations = int(obligations)
    uncovered = int(cov.get("uncovered") or 0)
    covered = cov.get("covered")
    covered = int(covered) if covered is not None else obligations - uncovered
    if obligations <= 0:
        raise ExportRefused("в пакете нет рейсов-обязательств (legs_total = 0) "
                            "— выгружать нечего")
    return {"legs_total": len(flights), "obligations": obligations,
            "covered": covered, "uncovered": uncovered,
            "pct": round(100.0 * covered / obligations, 2),
            "verdict": cov.get("verdict"), "complete": cov.get("complete"),
            "uncovered_legs": list(catalog.get("proven_impossible_legs") or []),
            "proofs": list(catalog.get("proven_impossible_proofs") or [])}


# ───────────────────────────────────────────────────────────── паспорт

def _passport(sources, cert: dict, task: str, scale: TimeScale,
              created: str, created_scale: str, extra: dict | None = None) -> dict:
    man = sources.manifest or {}
    d = sources.delivery or {}
    kit = man.get("kit") if isinstance(man.get("kit"), dict) else {}
    out = {"form": FORM_VERSION, "task": task,
           "package": {"sha256": sources.package_sha,
                       "package_fingerprint": (d.get("trace") or {}).get("package_fingerprint"),
                       "spec_version": man.get("spec_version"),
                       "source_run": man.get("source_run"),
                       "data_class": kit.get("data_class"),
                       "model_sha256": sources.model_sha256},
           "run": {"solver_run_id": (d.get("schedule") or {}).get("solver_run_id"),
                   "attempt": sources.attempt,
                   "roster_path": getattr(sources, "path", None)},
           "certification": dict(cert),
           "time": {"scale": scale.kind, "zone": scale.zone,
                    "offset_min": scale.offset_min,
                    "labels": "как в артефактах (шкала пакета)"},
           "created": created, "created_scale": created_scale,
           "personal_data": "только табельный номер; «Фамилия И.О.» и "
                            "«Должность» не выгружаются"}
    if extra:
        out.update(extra)
    return out


def _passport_rows(passport: dict) -> list[list]:
    rows: list[list] = []

    def walk(prefix: str, obj):
        if isinstance(obj, dict):
            for k in sorted(obj):
                walk((prefix + "." if prefix else "") + str(k), obj[k])
        elif isinstance(obj, list):
            rows.append([prefix, ", ".join(str(x) for x in obj)])
        else:
            rows.append([prefix, "" if obj is None else obj])
    walk("", passport)
    return rows


def _cert_summary(certificate: dict | None) -> dict | None:
    if not isinstance(certificate, dict):
        return None
    c = certificate.get("certificate")
    return c if isinstance(c, dict) else None


# ────────────────────────────────────────────────────────── писатели: общее

def _csv_safe(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    s = str(value)
    if s and s[0] in _CSV_DANGEROUS:
        return "'" + s
    return s


def _csv_text(header: list[str], rows: list[list]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(header)
    for r in rows:
        w.writerow([_csv_safe(v) for v in r])
    return buf.getvalue()


#: Записи zip-контейнера XLSX: система создания и права записи
#: фиксированы (Unix, rw-------), а не унаследованы от площадки —
#: стандартная библиотека ставит create_system 0 на Windows и 3 на
#: прочих, и без фиксации один и тот же файл с Local и с legion
#: отличался бы на байт в каждой записи центрального каталога.
_ZIP_CREATE_SYSTEM = 3
_ZIP_EXTERNAL_ATTR = 0o600 << 16


def _stable_zip(data: bytes, created: datetime) -> bytes:
    """Переписывание zip-контейнера с фиксированными датами, системой
    создания и правами записей — байты XLSX не зависят ни от момента
    вызова, ни от площадки."""
    stamp = (max(created.year, 1980), created.month, created.day,
             created.hour, created.minute, created.second)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=stamp)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = _ZIP_CREATE_SYSTEM
            zi.external_attr = _ZIP_EXTERNAL_ATTR
            zout.writestr(zi, zin.read(info.filename))
    return out.getvalue()


def _cell(v):
    if v is None or isinstance(v, (bool, int, float)):
        return v
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False, sort_keys=True)
    return str(v)


def _xlsx(sheets: list[tuple[str, list[str], list[list]]],
          created: datetime) -> bytes:
    """Книга: заголовки жирные, заморозка первой строки, автофильтр,
    ширина по содержимому, без формул и макросов; детерминированные
    байты (свойства книги и даты записей zip = created)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(list(header))
        for c in ws[1]:
            c.font = Font(bold=True)
        widths = [len(str(h)) for h in header]
        for r in rows:
            vals = [_cell(v) for v in r]
            ws.append(vals)
            row_cells = ws[ws.max_row]
            for i, (cell, v) in enumerate(zip(row_cells, vals)):
                if isinstance(v, str):
                    # Строка, похожая на формулу, остаётся строкой.
                    cell.data_type = "s"
                if i < len(widths):
                    widths[i] = max(widths[i], min(60, len(str(v)) if v is not None else 0))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(60, w + 2)
    wb.properties.creator = "COSMOS"
    wb.properties.lastModifiedBy = "COSMOS"
    wb.properties.created = created
    wb.properties.modified = created
    # Писатель вызывается напрямую: `Workbook.save` → `save_workbook`
    # перебивает `properties.modified` текущим моментом (openpyxl 3.1), и
    # два вызова по разные стороны секунды давали разные docProps/core.xml —
    # обещание «одни байты» держалось только удачей.
    from openpyxl.writer.excel import ExcelWriter
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(wb, archive).write_data()
    return _stable_zip(buf.getvalue(), created)


def _json_bytes(doc: dict, *, sort_keys: bool) -> bytes:
    return (json.dumps(doc, ensure_ascii=False, indent=1, sort_keys=sort_keys)
            + "\n").encode("utf-8")


def _delivery_view(delivery: dict, keys: tuple[str, ...]) -> dict:
    return {k: delivery.get(k) for k in keys if k in delivery}


PAIRING_DELIVERY_KEYS = ("verdict", "certified", "certificate_verdict", "result",
                         "mode", "objective", "mode_note", "coverage_verdict",
                         "chains", "coverage", "schedule", "trace",
                         "catalog_sha256", "findings", "proven_impossible_legs",
                         "assumptions")
ROSTER_DELIVERY_KEYS = ("verdict", "certified", "certificate_verdict", "result",
                        "mode", "objective", "mode_note", "schedule", "trace",
                        "pairing_run_id", "pairing_chains",
                        "pairing_catalog_sha256", "pairing_verdict", "solver_hard",
                        "findings", "counted_share", "window_coverage",
                        "cost_certainty", "deadheads", "pairing_selection")


# ─────────────────────────────────────────────────── пейринг: CSV/XLSX/JSON

PAIRING_CSV_HEADER = [
    "chain_id", "chain_no", "base", "duty_no", "leg_no", "flight_no",
    "flight_key", "date", "dep_airport", "dep_time", "arr_airport", "arr_time",
    "aircraft_type", "activity", "seats", "seat_codes", "pinned", "source_unit",
    "chain_cost_known", "chain_cost_full", "cost_certainty", "min_rest_after_min"]

PAIRING_XLSX_LEGS_HEADER = [
    "Цепочка", "№ цепочки", "База", "Смена", "№ участка", "Рейс", "Ключ рейса",
    "Дата", "Аэропорт вылета", "Время вылета", "Аэропорт прилёта",
    "Время прилёта", "Тип ВС", "Вид (работа/перегон)", "Места", "Коды мест",
    "Закреплена", "Источник", "Цена известная", "Цена полная",
    "Определённость цены", "Отдых после, мин"]


def _pairing_leg_rows(chains: list[dict]) -> list[list]:
    rows = []
    for ch in chains:
        for lg in ch["legs"]:
            rows.append([ch["id"], ch["no"], ch["base"], lg["duty_no"],
                         lg["leg_no"], lg["flight_no"], lg["fn"], lg["date"],
                         lg["dep"], lg["dep_time"], lg["arr"], lg["arr_time"],
                         lg["aircraft"], lg["activity"], "+".join(ch["seats"]),
                         "; ".join("+".join(g) for g in ch["seat_codes"]),
                         ch["pinned"], ch["source_unit"],
                         ch["price"]["cost_known"], ch["price"]["full_cost"],
                         ch["price"]["certainty"], ch["rest_after_min"]])
    return rows


def pairing_csv(src: PairingSources) -> str:
    """Одна строка на (цепочка, участок); UTF-8 без BOM (BOM добавляет
    ручка), CRLF, защита от формул."""
    return _csv_text(PAIRING_CSV_HEADER, _pairing_leg_rows(pairing_chains(src)))


def pairing_xlsx(src: PairingSources) -> bytes:
    chains = pairing_chains(src)
    cert = certification(src.delivery)
    scale = time_scale(src.manifest)
    created, cscale = created_stamp(src.delivery)
    passport = _passport(src, cert, "pairing", scale, created, cscale)
    d, cat = src.delivery, src.catalog
    cov = _coverage_numbers(cat, chains, src.payload)
    chain_rows = [[c["id"], c["no"], c["base"], c["start"], c["end"],
                   len(c["duties"]), len(c["legs"]),
                   sum(1 for lg in c["legs"] if lg["activity"] == "deadhead"),
                   "+".join(c["seats"]), "; ".join("+".join(g) for g in c["seat_codes"]),
                   c["pinned"], c["source_unit"], c["price"]["cost_known"],
                   c["price"]["full_cost"], c["price"]["certainty"],
                   c["price"]["nights"], c["rest_after_min"]] for c in chains]
    duty_rows = [[c["id"], du["no"], du["start"], du["end"], du["first_dep"],
                  du["last_arr"], du["legs"], du["flight_min"], du["work_min"],
                  du["rest_after_min"], ", ".join(du["stations"])]
                 for c in chains for du in c["duties"]]
    cov_rows = [["Рейсов в данных", cov["legs_total"]],
                ["Рейсов-обязательств", cov["obligations"]],
                ["Покрыто", cov["covered"]], ["Не покрыто", cov["uncovered"]],
                ["Вердикт покрытия", cov["verdict"]],
                ["Покрытие полное", cov["complete"]]]
    for leg in cov["uncovered_legs"]:
        cov_rows.append(["Не покрыт (доказано)", leg])
    for p in cov["proofs"]:
        cov_rows.append(["Доказательство", json.dumps(p, ensure_ascii=False,
                                                       sort_keys=True)])
    res = d.get("result") if isinstance(d.get("result"), dict) else {}
    kpi = [["Вердикт", d.get("verdict")], ["Сертифицировано", cert["certified"]],
           ["Сертификация установлена", cert["established"]],
           ["Вердикт сертификата", d.get("certificate_verdict")],
           ["Режим", d.get("mode")], ["Критерий", d.get("objective")],
           ["Вердикт покрытия", d.get("coverage_verdict")],
           ["Допустимость (result.feasibility)", res.get("feasibility")],
           ["Стоимость (result.cost)", res.get("cost")],
           ["Оптимальность (result.optimality)", res.get("optimality")],
           ["ip_status", cat.get("ip_status")], ["ip_objective", cat.get("ip_objective")],
           ["Область ip_objective", cat.get("ip_objective_scope")],
           ["known_subtotal", cat.get("known_subtotal")],
           ["full_cost", cat.get("full_cost")],
           ["Определённость цены каталога",
            json.dumps(cat.get("price_certainty"), ensure_ascii=False,
                       sort_keys=True) if cat.get("price_certainty") is not None else ""],
           ["catalog_sha256", d.get("catalog_sha256")],
           ["solver_run_id", (d.get("schedule") or {}).get("solver_run_id")],
           ["sha пакета", src.package_sha],
           ["Версия спецификации", (src.manifest or {}).get("spec_version")],
           ["created", created], ["Шкала created", cscale],
           ["Шкала меток", scale.words()]]
    sheets = [("Паспорт", ["Поле", "Значение"], _passport_rows(passport)),
              ("Цепочки", ["Цепочка", "№", "База", "Начало", "Конец", "Смен",
                           "Участков", "Перегонов", "Места", "Коды мест",
                           "Закреплена", "Источник", "Цена известная",
                           "Цена полная", "Определённость цены", "Ночей",
                           "Отдых после, мин"], chain_rows),
              ("Участки", PAIRING_XLSX_LEGS_HEADER, _pairing_leg_rows(chains)),
              ("Смены", ["Цепочка", "№ смены", "Начало", "Конец", "Первый вылет",
                         "Последний прилёт", "Участков", "Полётное время, мин",
                         "Рабочее время, мин", "Отдых после, мин", "Станции"],
               duty_rows),
              ("Покрытие", ["Показатель", "Значение"], cov_rows),
              ("Показатели", ["Показатель", "Значение"], kpi)]
    return _xlsx(sheets, _dt(created))


def pairing_json(src: PairingSources) -> bytes:
    chains = pairing_chains(src)
    cert = certification(src.delivery)
    scale = time_scale(src.manifest)
    created, cscale = created_stamp(src.delivery)
    cat = src.catalog
    doc = {"kind": "cosmos-pairing-export/1",
           "passport": _passport(src, cert, "pairing", scale, created, cscale),
           "delivery": _delivery_view(src.delivery, PAIRING_DELIVERY_KEYS),
           "catalog_head": {k: cat.get(k) for k in
                            ("ip_status", "ip_objective", "ip_objective_scope",
                             "known_subtotal", "full_cost", "price_certainty",
                             "config_sig", "sit_window", "label")},
           "coverage": _coverage_numbers(cat, chains, src.payload),
           "chains": chains,
           "uncovered": list(cat.get("proven_impossible_legs") or []),
           "certificate": _cert_summary(src.certificate)}
    return _json_bytes(doc, sort_keys=True)


# ─────────────────────────────────────────────────────── пейринг: Meridian

def _metric(mtype: str, code: str, value, descr: str) -> dict:
    return {"metricType": mtype, "metricCode": code,
            "metricValue": str(value), "metricDescription": descr}


def _roles_of(opts: Options, manifest: dict | None) -> tuple[str, tuple[str, str]]:
    name = (opts.roles or default_roles(manifest)).strip().lower()
    alph = role_alphabets()
    if name not in alph:
        raise ExportBadRequest("алфавит ролей %r вне словаря: %s"
                               % (name, ", ".join(sorted(alph))))
    return name, alph[name]


def _plan_bounds(legs: list[dict], manifest: dict | None,
                 scale: TimeScale, utc: bool) -> tuple[str, str, str]:
    """planSince/planUntil — те же, что у набора Meridian: объявление
    источника в манифесте, иначе min(start)/max(end) участков (правило
    cdpm/export_meridian.build_dataset)."""
    src = _source_meridian(manifest)
    if src.get("planSince") and src.get("planUntil"):
        return (_stamp(src["planSince"], scale, utc),
                _stamp(src["planUntil"], scale, utc),
                "planSince/planUntil — объявление источника Meridian в манифесте")
    return (min(l["start"] for l in legs), max(l["end"] for l in legs),
            "planSince/planUntil — границы участков набора (правило набора "
            "CDPM для Meridian)")


def meridian_pairing_doc(src: PairingSources, opts: Options | None = None) -> dict:
    """Решение пейринга формата Meridian 1.17 (форма гейта линии BA)."""
    opts = opts or Options()
    cert = certification(src.delivery)
    require_solution(src.delivery, "пейринга")
    _uncertified_gate(cert, opts, "пейринга")
    scale = time_scale(src.manifest)
    utc = bool(opts.utc)
    role_name, (r_work, r_dh) = _roles_of(opts, src.manifest)
    chains = pairing_chains(src)
    flights = src.payload.get("flights") or []
    legs_by_fn = leg_index(flights, scale, utc)
    created, cscale = created_stamp(src.delivery)
    idn = ids(src, opts)
    notes: list[str] = list(idn["notes"])
    notes.append(scale_note(scale, utc))
    notes.append(created_note(cscale))
    notes.append("алфавит ролей участка: %s/%s (%s)"
                 % (r_work, r_dh, "гейт BA" if role_name == "gate"
                    else "писатель и импортёр COSMOS"))
    notes.extend(_leg_notes(flights))
    pairings, duties = [], []
    for ch in chains:
        notes.extend("%s: %s" % (ch["id"], n) for n in ch["notes"])
        legs_by_duty: dict[int, list[dict]] = {}
        for lg in ch["legs"]:
            legs_by_duty.setdefault(lg["duty_no"], []).append(lg)
        recs = []
        for du in ch["duties"]:
            did = "%s$%d" % (ch["id"], du["no"])
            if len(did) > 50:
                raise ExportRefused("dutyId длиннее 50 знаков: %s" % did)
            members = legs_by_duty.get(du["no"], [])
            for lg in members:
                legs_by_fn[lg["fn"]]["dutyCrewTypes"][did] = (
                    r_work if lg["activity"] == "work" else r_dh)
            recs.append({
                "dutyId": did,
                "start": _stamp(du["start"], scale, utc),
                "end": _stamp(du["end"], scale, utc),
                "restOverride": int(du["rest_after_min"]),
                "assignmentType": "FLIGHT",
                "flightTimeCredit": int(du["flight_min"]),
                "workTimeCredit": int(du["work_min"]),
                "split": False,
                "pairingId": ch["id"],
                "numberOfDuty": int(du["no"]),
                "firstDepartureStation": du["first_dep"],
                "lastArrivalStation": du["last_arr"],
                "stationsVisited": list(du["stations"]),
                "complements": [{"employeeId": "", "complementRank": r,
                                 "requirements": [], "locked": False}
                                for r in ch["seats"]]})
        duties.extend(recs)
        pairings.append({
            "pairingId": ch["id"],
            "start": recs[0]["start"], "end": recs[-1]["end"],
            "restOverride": int(ch["rest_after_min"]),
            "requirements": {},
            "locked": bool(ch["pinned"]),
            "flightMinutesCredit": sum(d["flightTimeCredit"] for d in recs),
            "workMinutesCredit": sum(d["workTimeCredit"] for d in recs),
            "primaryPairing": True})
    legs = [legs_by_fn[k] for k in sorted(legs_by_fn, key=lambda k: (
        legs_by_fn[k]["start"], legs_by_fn[k]["legId"]))]
    plan_since, plan_until, plan_note = _plan_bounds(legs, src.manifest, scale, utc)
    notes.append(plan_note)
    cov = _coverage_numbers(src.catalog, chains, src.payload)
    d = src.delivery
    metrics = [
        _metric("INTEGER", "PAIRINGS_COUNT", len(pairings), "Количество пейрингов"),
        _metric("INTEGER", "DUTIES_COUNT", len(duties), "Количество полётных смен"),
        _metric("INTEGER", "LEGS_TOTAL", len(legs), "Участков в файле (все рейсы данных)"),
        _metric("INTEGER", "LEGS_OBLIGATIONS", cov["obligations"],
                "Рейсов-обязательств покрытия (без транспортных)"),
        _metric("INTEGER", "LEGS_COVERED", cov["covered"], "Покрыто рейсов-обязательств"),
        _metric("INTEGER", "LEGS_UNCOVERED", cov["uncovered"], "Не покрыто рейсов-обязательств"),
        _metric("PERCENT", "COVERAGE_PCT", cov["pct"], "Доля покрытия обязательств, процентов"),
        _metric("INTEGER", "DUTIES_EXTENDED_CREW",
                sum(1 for ch in chains for du in ch["duties"] if du.get("extended")),
                "Смен с составом сверх спроса участков (усиленный состав)"),
        _metric("DURATION", "FLIGHT_MINUTES", sum(p["flightMinutesCredit"] for p in pairings),
                "Суммарное полётное время, минут"),
        _metric("DURATION", "WORK_MINUTES", sum(p["workMinutesCredit"] for p in pairings),
                "Суммарное рабочее время, минут"),
        _metric("OTHER", "CERTIFIED", 1 if cert["certified"] else 0,
                "Решение сертифицировано (1) / нет (0)"),
        _metric("OTHER", "COST_CERTAINTY", d.get("cost_certainty")
                or (d.get("result") or {}).get("cost") or "", "Определённость цены решения"),
        _metric("OTHER", "SOLVER_RUN_ID", (d.get("schedule") or {}).get("solver_run_id") or "",
                "Прогон решателя COSMOS"),
        _metric("OTHER", "ROLE_ALPHABET", "%s/%s" % (r_work, r_dh),
                "Алфавит роли участка: работа/пассажиром"),
        _metric("OTHER", "TIME_SCALE", "utc" if (utc and scale.convertible)
                else scale.kind, "Шкала меток файла")]
    words = uncertified_words(cert)
    description = ("COSMOS: решение пейринга по пакету постановки %s (прогон %s); %s%s"
                   % (src.package_sha[:12],
                      (d.get("schedule") or {}).get("solver_run_id"),
                      "; ".join(_dedupe(notes)),
                      ("; " + words) if words else ""))
    return {"datasetId": idn["datasetId"], "presetId": idn["presetId"],
            "format": 1, "type": "pairing",
            "division": division_of(src.payload.get("_seatRoles") or
                                    sorted({s for ch in chains for s in ch["seats"]}),
                                    src.manifest),
            "planSince": plan_since, "planUntil": plan_until,
            "created": created, "description": description,
            "legs": legs, "pairings": pairings, "duties": duties,
            "groundActivities": [], "metrics": metrics}


def _dedupe(items: list[str]) -> list[str]:
    seen, out = set(), []
    for x in items:
        if x not in seen:
            seen.add(x)
            out.append(x)
    return out


def pairing_meridian(src: PairingSources, opts: Options | None = None) -> bytes:
    return _json_bytes(meridian_pairing_doc(src, opts), sort_keys=False)


def meridian_gate_inputs(src: PairingSources, doc: dict) -> tuple[dict, dict]:
    """Вход гейта линии BA: каталог с coverage.covered и данные с
    flightLegId := legId файла (та же привязка, что у документа)."""
    cat = json.loads(json.dumps(src.catalog))
    cov = cat.get("coverage") if isinstance(cat.get("coverage"), dict) else {}
    cov = dict(cov)
    if "covered" not in cov:
        cov["covered"] = int(cov.get("legs_total") or 0) - int(cov.get("uncovered") or 0)
    cat["coverage"] = cov
    ids_by_num = {}
    for l in doc.get("legs") or []:
        ids_by_num.setdefault((l["flightNumber"], l["start"]), l["legId"])
    pay = json.loads(json.dumps(src.payload))
    scale = time_scale(src.manifest)
    utc = any(m.get("metricCode") == "TIME_SCALE" and m.get("metricValue") == "utc"
              and scale.kind != "utc" for m in doc.get("metrics") or [])
    for f in pay.get("flights") or []:
        fn = str(f["flightNumber"])
        key = (fn.split("|")[0], _stamp(f["departureUTCDateTime"], scale, utc))
        f["flightLegId"] = ids_by_num.get(key, f.get("flightLegId") or fn)
    return cat, pay


# ─────────────────────────────────────────────────────── ростер: модель

def _truthy(v) -> bool:
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    return str(v).strip().lower() in _TRUE_WORDS


def _sheet_rows(solution: dict, sheet_marks: tuple[str, ...]) -> list[dict]:
    """Строки листа справочника решения (reference_tables[*][лист])."""
    out = []
    rt = solution.get("reference_tables")
    if not isinstance(rt, dict):
        return out
    for tabs in rt.values():
        if not isinstance(tabs, dict):
            continue
        for name, t in tabs.items():
            if not any(m in str(name).lower() for m in sheet_marks):
                continue
            rows = t.get("rows") if isinstance(t, dict) else t
            if isinstance(rows, list):
                out.extend(r for r in rows if isinstance(r, dict))
    return out


def _pick(row: dict, marks: tuple[str, ...]):
    for k, v in row.items():
        low = str(k).lower()
        if any(m in low for m in marks):
            return v
    return None


def roster_rows(src: RosterSources) -> dict:
    """Нормализованные строки ростера: назначения (deadhead → bool,
    date всегда, flight_no), штат (с людьми без рейсов), смены,
    цепочки, горизонт. Метки — как в артефактах."""
    sol = src.solution
    raw = sol.get("assignments")
    if not isinstance(raw, list):
        raise ExportRefused("решение ростера не несёт назначений (assignments)")
    fl = {str(f.get("flightNumber")): f for f in (src.payload.get("flights") or [])
          if isinstance(f, dict)}
    rows = []
    for a in raw:
        if not isinstance(a, dict) or not a.get("employee") or not a.get("flight"):
            continue
        key = str(a["flight"])
        dep = a.get("dep_time") or (fl.get(key) or {}).get("departureUTCDateTime")
        arr = a.get("arr_time") or (fl.get(key) or {}).get("arrivalUTCDateTime")
        if not dep or not arr:
            raise ExportRefused("назначение %s/%s без времён вылета/прилёта"
                                % (a.get("employee"), key))
        dep, arr = _norm(dep), _norm(arr)
        rows.append({
            "employee": str(a["employee"]),
            "base": str(a.get("base") or ""),
            "date": str(a.get("date") or dep[:10])[:10],
            "flight_no": key.split("|")[0], "flight_key": key,
            "dep_airport": str(a.get("dep_airport") or (fl.get(key) or {}).get("departureAirport") or ""),
            "dep_time": dep,
            "arr_airport": str(a.get("arr_airport") or (fl.get(key) or {}).get("arrivalAirport") or ""),
            "arr_time": arr,
            "role": str(a.get("role") or ""),
            "deadhead": _truthy(a.get("deadhead")),
            "duty_id": str(a.get("duty_id") or "%s/%s" % (a["employee"], dep[:10])),
            "pairing_id": str(a["pairing_id"]) if a.get("pairing_id") else "",
            "aircraft_type": str((fl.get(key) or {}).get("aircraftType") or "")})
    rows.sort(key=lambda r: (r["employee"], r["dep_time"], r["flight_key"]))
    # Штат: окно решения ∪ лист «Экипажи» ∪ назначения (люди без рейсов
    # входят); только табельный номер и база.
    emp: dict[str, dict] = {}
    for w in sol.get("window_employees") or []:
        if isinstance(w, dict) and w.get("employee"):
            emp.setdefault(str(w["employee"]), {"id": str(w["employee"]),
                                                 "base": str(w.get("base") or "")})
    for r in _sheet_rows(sol, ("экипаж", "employee", "crew")):
        eid = _pick(r, ("табельн", "employeeid", "employee_id"))
        if eid is None or str(eid).strip() == "":
            continue
        eid = str(eid).strip()
        base = _pick(r, ("база", "base"))
        rec = emp.setdefault(eid, {"id": eid, "base": ""})
        if base and not rec["base"]:
            rec["base"] = str(base).strip()
    for r in rows:
        rec = emp.setdefault(r["employee"], {"id": r["employee"], "base": ""})
        if r["base"] and not rec["base"]:
            rec["base"] = r["base"]
    for rec in emp.values():
        mine = [r for r in rows if r["employee"] == rec["id"]]
        rec["legs"] = len(mine)
        rec["deadheads"] = sum(1 for r in mine if r["deadhead"])
        rec["flight_min"] = sum(_minutes(_dt(r["dep_time"]), _dt(r["arr_time"]))
                                for r in mine if not r["deadhead"])
        rec["days_with_flights"] = len({r["date"] for r in mine})
    employees = [emp[k] for k in sorted(emp)]
    duties = []
    for d in sol.get("duties") or []:
        if not isinstance(d, dict) or not d.get("id"):
            continue
        mine = [r for r in rows if r["duty_id"] == str(d["id"])]
        duties.append({"id": str(d["id"]), "employee": str(d.get("employee") or ""),
                       "start": _norm(d["start"]) if d.get("start") else "",
                       "end": _norm(d["end"]) if d.get("end") else "",
                       "base": str(d.get("base") or ""),
                       "rest_at": str(d.get("last_arr_airport") or ""),
                       "legs": len(mine),
                       "deadheads": sum(1 for r in mine if r["deadhead"])})
    duties.sort(key=lambda d: (d["employee"], d["start"], d["id"]))
    chains = []
    for p in sol.get("pairings") or []:
        if not isinstance(p, dict) or not p.get("id"):
            continue
        mine = [r for r in rows if r["pairing_id"] == str(p["id"])
                and r["employee"] == str(p.get("employee") or "")]
        chains.append({"id": str(p["id"]), "employee": str(p.get("employee") or ""),
                       "role": mine[0]["role"] if mine else "",
                       "start": _norm(p["start"]) if p.get("start") else "",
                       "end": _norm(p["end"]) if p.get("end") else "",
                       "legs": len(mine)})
    chains.sort(key=lambda c: (c["id"], c["employee"]))
    wh = sol.get("window_horizon") if isinstance(sol.get("window_horizon"), dict) else {}
    h_start = wh.get("start_inclusive") or (min(r["dep_time"] for r in rows) if rows else None)
    h_end = wh.get("end_exclusive") or (max(r["arr_time"] for r in rows) if rows else None)
    horizon_days: list[str] = []
    if h_start and h_end:
        d0, d1 = _dt(h_start).date(), _dt(h_end)
        last = d1.date() if d1.time() != d1.min.time() else d1.date() - timedelta(days=1)
        cur = d0
        while cur <= last:
            horizon_days.append(cur.isoformat())
            cur += timedelta(days=1)
    return {"assignments": rows, "employees": employees, "duties": duties,
            "chains": chains, "horizon_days": horizon_days,
            "horizon": {"start_inclusive": _norm(h_start) if h_start else None,
                        "end_exclusive": _norm(h_end) if h_end else None}}


ROSTER_CSV_HEADER = ["employee", "base", "date", "flight_no", "flight_key",
                     "dep_airport", "dep_time", "arr_airport", "arr_time", "role",
                     "activity", "duty_id", "pairing_id", "aircraft_type"]

#: Заголовки листа «Назначения» — как у входа сертификатора.
ROSTER_XLSX_ASSIGN_HEADER = ["Табельный номер", "Рейс", "Дата", "Аэропорт вылета",
                             "Аэропорт прилёта", "Время вылета", "Время прилёта",
                             "Роль", "База приписки", "Смена", "Перегон (да/нет)",
                             "Цепочка", "Тип ВС"]


def _roster_csv_rows(rows: dict) -> list[list]:
    return [[r["employee"], r["base"], r["date"], r["flight_no"], r["flight_key"],
             r["dep_airport"], r["dep_time"], r["arr_airport"], r["arr_time"],
             r["role"], "deadhead" if r["deadhead"] else "work", r["duty_id"],
             r["pairing_id"], r["aircraft_type"]] for r in rows["assignments"]]


def roster_csv(src: RosterSources) -> str:
    return _csv_text(ROSTER_CSV_HEADER, _roster_csv_rows(roster_rows(src)))


def _board_rows(rows: dict) -> tuple[list[str], list[list]]:
    days = rows["horizon_days"]
    header = ["Табельный номер", "База", "Роли"] + [
        "%s.%s" % (d[8:10], d[5:7]) for d in days]
    out = []
    for e in rows["employees"]:
        mine = [r for r in rows["assignments"] if r["employee"] == e["id"]]
        roles = sorted({r["role"] for r in mine if r["role"]})
        line = [e["id"], e["base"], ", ".join(roles)]
        for d in days:
            cell = [("*" if r["deadhead"] else "") + r["flight_no"]
                    for r in mine if r["date"] == d]
            line.append(" ".join(cell))
        out.append(line)
    return header, out


def roster_xlsx(src: RosterSources) -> bytes:
    rows = roster_rows(src)
    cert = certification(src.delivery)
    scale = time_scale(src.manifest)
    created, cscale = created_stamp(src.delivery)
    passport = _passport(src, cert, "roster", scale, created, cscale)
    d = src.delivery
    assign_rows = [[r["employee"], r["flight_no"], r["date"], r["dep_airport"],
                    r["arr_airport"], r["dep_time"], r["arr_time"], r["role"],
                    r["base"], r["duty_id"], "да" if r["deadhead"] else "нет",
                    r["pairing_id"], r["aircraft_type"]] for r in rows["assignments"]]
    duty_rows = [[du["employee"], du["id"], du["start"], du["end"], du["base"],
                  du["rest_at"], du["legs"], du["deadheads"]] for du in rows["duties"]]
    chain_rows = [[c["id"], c["employee"], c["role"], c["start"], c["end"], c["legs"]]
                  for c in rows["chains"]]
    board_header, board = _board_rows(rows)
    staff = [[e["id"], e["base"], e["legs"], e["flight_min"], e["deadheads"],
              e["days_with_flights"]] for e in rows["employees"]]
    res = d.get("result") if isinstance(d.get("result"), dict) else {}
    kpi = [["Вердикт", d.get("verdict")], ["Сертифицировано", cert["certified"]],
           ["Сертификация установлена", cert["established"]],
           ["Вердикт сертификата", d.get("certificate_verdict")],
           ["Допустимость (result.feasibility)", res.get("feasibility")],
           ["Стоимость (result.cost)", res.get("cost")],
           ["Оптимальность (result.optimality)", res.get("optimality")],
           ["score", (d.get("schedule") or {}).get("score")],
           ["solver_hard", d.get("solver_hard")], ["findings", d.get("findings")],
           ["counted_share", json.dumps(d.get("counted_share"), ensure_ascii=False,
                                        sort_keys=True) if d.get("counted_share") is not None else ""],
           ["window_coverage", json.dumps({k: v for k, v in (d.get("window_coverage") or {}).items()
                                           if k != "checks"}, ensure_ascii=False, sort_keys=True)
            if d.get("window_coverage") else ""],
           ["pairing_run_id", d.get("pairing_run_id")],
           ["solver_run_id", (d.get("schedule") or {}).get("solver_run_id")],
           ["Путь ростера", src.path], ["sha пакета", src.package_sha],
           ["Версия спецификации", (src.manifest or {}).get("spec_version")],
           ["created", created], ["Шкала created", cscale],
           ["Шкала меток", scale.words()]]
    sheets = [("Паспорт", ["Поле", "Значение"], _passport_rows(passport)),
              ("Назначения", ROSTER_XLSX_ASSIGN_HEADER, assign_rows),
              ("Смены", ["Табельный номер", "Идентификатор", "Начало", "Конец",
                         "База приписки", "Место отдыха", "Участков", "Перегонов"],
               duty_rows),
              ("Цепочки", ["Идентификатор", "Табельный номер", "Роль", "Начало",
                           "Конец", "Участков"], chain_rows),
              ("Доска", board_header, board),
              ("Штат", ["Табельный номер", "База", "Назначений",
                        "Полётное время, мин", "Перегонов", "Дней с рейсами"], staff),
              ("Показатели", ["Показатель", "Значение"], kpi)]
    return _xlsx(sheets, _dt(created))


def roster_json(src: RosterSources) -> bytes:
    rows = roster_rows(src)
    cert = certification(src.delivery)
    scale = time_scale(src.manifest)
    created, cscale = created_stamp(src.delivery)
    doc = {"kind": "cosmos-roster-export/1",
           "passport": _passport(src, cert, "roster", scale, created, cscale),
           "delivery": _delivery_view(src.delivery, ROSTER_DELIVERY_KEYS),
           "window_horizon": rows["horizon"],
           "time": {"scale": scale.kind, "zone": scale.zone},
           "assignments": rows["assignments"], "duties": rows["duties"],
           "pairings": rows["chains"], "employees": rows["employees"],
           "certificate": _cert_summary(src.certificate)}
    return _json_bytes(doc, sort_keys=True)


# ─────────────────────────────────────────────────────────── ростер: iCal

def _ics_escape(s: str) -> str:
    return (str(s).replace("\\", "\\\\").replace(";", "\\;")
            .replace(",", "\\,").replace("\r\n", "\\n").replace("\n", "\\n"))


def _ics_fold(line: str) -> list[str]:
    """Свёртка строки по 75 октетов без разрыва многобайтовых символов."""
    raw = line.encode("utf-8")
    if len(raw) <= 75:
        return [line]
    out, pos, limit = [], 0, 75
    while pos < len(raw):
        cut = min(len(raw), pos + limit)
        while cut < len(raw) and cut > pos and (raw[cut] & 0xC0) == 0x80:
            cut -= 1
        chunk = raw[pos:cut].decode("utf-8")
        out.append(chunk if not out else " " + chunk)
        pos = cut
        limit = 74
    return out


def _ics_dt(label: str, scale: TimeScale, utc: bool) -> tuple[str, str]:
    """(параметры, значение) для DTSTART/DTEND: utc → …Z; base-local с
    поясом → TZID; неизвестная шкала → плавающее время без Z."""
    if utc and scale.convertible or scale.kind == "utc":
        return "", _iso(to_utc(label, scale)).replace("-", "").replace(":", "") + "Z"
    val = _iso(_dt(label)).replace("-", "").replace(":", "")
    if scale.kind == "base-local" and scale.zone:
        return ";TZID=" + scale.zone, val
    return "", val


def _vtimezone(zone: str, days: list[str], times: list[str]) -> list[str]:
    """VTIMEZONE по правилам zoneinfo на горизонте файла: компонент на
    каждый наблюдаемый переход смещения."""
    from zoneinfo import ZoneInfo
    z = ZoneInfo(zone)
    stamps = sorted({_dt(t) for t in times} | {_dt(d + "T00:00:00") for d in days})
    if not stamps:
        return []
    lo, hi = stamps[0] - timedelta(days=1), stamps[-1] + timedelta(days=1)

    def off(d: datetime) -> timedelta:
        return d.replace(tzinfo=z).utcoffset() or timedelta(0)

    def fmt(td: timedelta) -> str:
        s = int(td.total_seconds())
        sign = "+" if s >= 0 else "-"
        s = abs(s)
        return "%s%02d%02d" % (sign, s // 3600, (s % 3600) // 60)

    comps: list[tuple[str, datetime, timedelta, timedelta]] = []
    cur, prev = lo, off(lo)
    comps.append(("STANDARD", lo, prev, prev))
    step = timedelta(hours=1)
    while cur <= hi:
        nxt = cur + step
        o = off(nxt)
        if o != prev:
            dst = (nxt.replace(tzinfo=z).dst() or timedelta(0)) != timedelta(0)
            comps.append(("DAYLIGHT" if dst else "STANDARD", nxt, prev, o))
            prev = o
        cur = nxt
    lines = ["BEGIN:VTIMEZONE", "TZID:" + zone]
    for kind, start, o_from, o_to in comps:
        lines += ["BEGIN:" + kind,
                  "DTSTART:" + _iso(start).replace("-", "").replace(":", ""),
                  "TZOFFSETFROM:" + fmt(o_from), "TZOFFSETTO:" + fmt(o_to),
                  "END:" + kind]
    lines.append("END:VTIMEZONE")
    return lines


def roster_ics(src: RosterSources, opts: Options | None = None) -> str:
    """Календарь ростера (RFC 5545): VEVENT на назначение; `employee` —
    календарь одного сотрудника (неизвестный табельный — отказ 404)."""
    opts = opts or Options()
    rows = roster_rows(src)
    cert = certification(src.delivery)
    scale = time_scale(src.manifest)
    created, cscale = created_stamp(src.delivery)
    run_id = str((src.delivery.get("schedule") or {}).get("solver_run_id") or "")
    items = rows["assignments"]
    if opts.employee is not None:
        known = {e["id"] for e in rows["employees"]}
        if opts.employee not in known:
            raise ExportNotFound("сотрудник %r в решении ростера не найден; "
                                 "табельные: %s" % (opts.employee,
                                                    ", ".join(sorted(known)[:20])))
        items = [r for r in items if r["employee"] == opts.employee]
    utc = bool(opts.utc)
    if utc and scale.convertible or scale.kind == "utc":
        tscale = "utc" + ("" if scale.kind == "utc" else
                          ";converted-from=%s" % scale.zone)
    elif scale.kind == "base-local" and scale.zone:
        tscale = "base-local;zone=%s" % scale.zone
    elif scale.kind == "shifted":
        tscale = (scale.header() + ";note=перевод в UTC контрактом пакета не "
                  "определён — метки записаны как плавающее время без перевода")
    else:
        tscale = "unknown;note=метки записаны как плавающее время без перевода"
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0",
             "PRODID:-//COSMOS//export_solutions 1//RU", "CALSCALE:GREGORIAN",
             "METHOD:PUBLISH",
             "X-WR-CALNAME:" + _ics_escape("COSMOS ростер %s%s" % (
                 src.package_sha[:12], (" " + opts.employee) if opts.employee else "")),
             "X-COSMOS-PACKAGE:" + src.package_sha,
             "X-COSMOS-SOLVER-RUN:" + run_id,
             "X-COSMOS-CERTIFIED:" + ("true" if cert["certified"] else "false"),
             "X-COSMOS-TIME-SCALE:" + tscale,
             "X-COSMOS-CREATED-SCALE:" + cscale,
             "X-COSMOS-FORM:" + FORM_VERSION]
    # DTSTAMP — по RFC 5545 §3.8.7.2 строго в форме UTC («…Z»): это метка
    # версии объекта, не время события. При шкале журнала machine-local
    # значение — метка журнала кнопки без перевода (пояс не выдумывается),
    # и об этом сказано словами: X-COSMOS-CREATED-SCALE и заметка ниже.
    # Плавающий DTSTAMP стандарт не допускает — строгие импортёры отвергали
    # бы файл целиком.
    if cscale != "utc":
        lines.append("X-COSMOS-DTSTAMP-NOTE:" + _ics_escape(
            "DTSTAMP записан в форме UTC по RFC 5545; значение — метка журнала "
            "кнопки (местное время машины, пояс не записан) без перевода"))
    if scale.kind == "base-local" and scale.zone and not (utc and scale.convertible):
        lines += _vtimezone(scale.zone, rows["horizon_days"],
                            [r["dep_time"] for r in items] + [r["arr_time"] for r in items])
    stamp = _iso(_dt(created)).replace("-", "").replace(":", "") + "Z"
    for r in items:
        uid = hashlib.sha1("|".join([src.package_sha, run_id, r["employee"],
                                     r["flight_key"]]).encode("utf-8")).hexdigest()[:24]
        p0, v0 = _ics_dt(r["dep_time"], scale, utc)
        p1, v1 = _ics_dt(r["arr_time"], scale, utc)
        kind = "deadhead" if r["deadhead"] else "work"
        summary = "%s%s %s→%s (%s)" % ("DH " if r["deadhead"] else "", r["flight_no"],
                                       r["dep_airport"], r["arr_airport"], r["role"])
        descr = ("сотрудник %s; смена %s; цепочка %s; роль %s; база %s; перегон %s; "
                 "сертифицировано %s" % (r["employee"], r["duty_id"],
                                         r["pairing_id"] or "—", r["role"], r["base"],
                                         "да" if r["deadhead"] else "нет",
                                         "да" if cert["certified"] else "нет"))
        lines += ["BEGIN:VEVENT", "UID:%s@cosmos" % uid, "DTSTAMP:" + stamp,
                  "DTSTART%s:%s" % (p0, v0), "DTEND%s:%s" % (p1, v1),
                  "SUMMARY:" + _ics_escape(summary),
                  "LOCATION:" + _ics_escape(r["dep_airport"]),
                  "DESCRIPTION:" + _ics_escape(descr),
                  "CATEGORIES:" + _ics_escape(r["employee"]) + "," + kind,
                  "END:VEVENT"]
    lines.append("END:VCALENDAR")
    folded: list[str] = []
    for ln in lines:
        folded.extend(_ics_fold(ln))
    return "\r\n".join(folded) + "\r\n"


# ────────────────────────────────────────────────── ростер: Meridian (проект)

ROSTER_FORM_NOTE = ("ФОРМА РОСТЕРА — ПРОЕКТ: не подтверждена спецификацией, "
                    "на сверку владельцу")


def chain_key(pairing_id: str) -> str:
    """Ключ цепочки по метке слоя цепочек решателя: «каталог:<k>:slot:<j>»
    (PackagedRuns.kt: j — indexInFlight, не место каталога) → «каталог:<k>»;
    иная форма — сама метка."""
    m = _SLOT_RE.match(str(pairing_id))
    return m.group(1) if m else str(pairing_id)


def _catalog_index(key: str) -> int | None:
    m = re.match(r"^каталог:(\d+)$", key)
    return int(m.group(1)) if m else None


def _threshold(solution: dict, name: str) -> int | None:
    th = solution.get("thresholds") if isinstance(solution.get("thresholds"), dict) else {}
    for rec in th.get("applied") or []:
        if isinstance(rec, dict) and str(rec.get("threshold")) == name:
            try:
                return int(rec.get("value"))
            except (TypeError, ValueError):
                return None
    return None


def meridian_roster_doc(src: RosterSources, opts: Options | None = None) -> dict:
    """ПРОЕКТ формы `type: "roster"`: структура пейринга, у мест заполнены
    табельные. Источник истины — назначения ростера; каталог пейринга —
    только идентификаторы цепочек и minRestAfterMin."""
    opts = opts or Options()
    cert = certification(src.delivery)
    require_solution(src.delivery, "ростеринга")
    _uncertified_gate(cert, opts, "ростеринга")
    scale = time_scale(src.manifest)
    utc = bool(opts.utc)
    role_name, (r_work, r_dh) = _roles_of(opts, src.manifest)
    rows = roster_rows(src)
    if not rows["assignments"]:
        raise ExportRefused("решение ростера без единого назначения — выгружать нечего")
    if not any(r["pairing_id"] for r in rows["assignments"]):
        raise ExportRefused("сквозной ростер без слоя цепочек (pairing_id у "
                            "назначений нет) — в форму Meridian не выражается; "
                            "выгрузите CSV/XLSX/JSON/iCal либо выполните "
                            "ростеринг на утверждённом пейринге")
    flights = src.payload.get("flights") or []
    legs_by_fn = leg_index(flights, scale, utc)
    for r in rows["assignments"]:
        if r["flight_key"] not in legs_by_fn:
            raise ExportRefused("рейс назначения %s не найден в данных пакета"
                                % r["flight_key"])
    created, cscale = created_stamp(src.delivery)
    idn = ids(src, opts)
    notes: list[str] = [ROSTER_FORM_NOTE] + list(idn["notes"])
    notes.append(scale_note(scale, utc))
    notes.append(created_note(cscale))
    notes.append("алфавит ролей участка: %s/%s (%s)"
                 % (r_work, r_dh, "гейт BA" if role_name == "gate"
                    else "писатель и импортёр COSMOS"))
    notes.extend(_leg_notes(flights))
    cat = src.pairing_catalog
    if src.pairing_catalog_note:
        notes.append(src.pairing_catalog_note)
    label12 = _label12(cat, src.package_sha) if cat else src.package_sha[:12]
    units = (cat.get("selected") or []) if cat else []
    # Закрепления и незанятые места — из решения решателя (solved.json).
    pinned: set[tuple[str, str]] = set()
    unfilled: dict[str, list[tuple[str, str, str]]] = {}   # chain → (slot, flight, role)
    for fa in (src.solved or {}).get("flightAssignments") or []:
        if not isinstance(fa, dict) or not fa.get("flight"):
            continue
        ck = chain_key(str(fa.get("chainId") or ""))
        if fa.get("employee"):
            if fa.get("pinned"):
                pinned.add((str(fa["employee"]), str(fa["flight"])))
        elif ck:
            unfilled.setdefault(ck, []).append(
                (str(fa.get("chainId")), str(fa["flight"]), str(fa.get("requiredSkill") or "")))
    # Цепочки: группа мест одной метки; проверки — одно множество участков
    # и одно разбиение на смены (как множества множеств участков; duty_id
    # — только ярлык, у людей он персональный).
    by_chain: dict[str, dict[str, list[dict]]] = {}
    for r in rows["assignments"]:
        if not r["pairing_id"]:
            raise ExportRefused("назначение %s/%s без метки цепочки при слое "
                                "цепочек у остальных — ростер смешанный, в "
                                "форму Meridian не выражается"
                                % (r["employee"], r["flight_key"]))
        by_chain.setdefault(chain_key(r["pairing_id"]), {}).setdefault(
            r["employee"], []).append(r)
    home_rest = _threshold(src.solution, "min_rest_home_min")
    away_rest = _threshold(src.solution, "min_rest_outstation_min")
    # Участки каждой смены решателя у сотрудника — по ВСЕМ его
    # назначениям (смена решателя может охватывать несколько цепочек).
    emp_duty_legs: dict[tuple[str, str], set[str]] = {}
    for r in rows["assignments"]:
        emp_duty_legs.setdefault((r["employee"], r["duty_id"]), set()).add(r["flight_key"])
    pairings, duties = [], []
    seats_total = seats_filled = seats_unfilled = 0
    diverged = 0
    numbered = sorted(by_chain, key=lambda k: (_catalog_index(k) is None,
                                               _catalog_index(k) or 0, k))
    # Номера цепочек: индексная метка «каталог:<k>» — k+1 (как у выгрузки
    # пейринга того же каталога); метка иной формы — с первого номера,
    # свободного и от каталога, и от индексных меток. Порядковый номер в
    # списке не годится: при разрежённых индексах («каталог:0», «каталог:2»,
    # «прочее») он совпадал бы с каталожным — дубль pairingId и dutyId.
    free_no = max([len(units)] + [_catalog_index(k) + 1 for k in numbered
                                  if _catalog_index(k) is not None]) + 1
    extra_no: dict[str, int] = {}
    for ck in numbered:
        if _catalog_index(ck) is None:
            extra_no[ck] = free_no
            free_no += 1
    if extra_no:
        notes.append("цепочек с меткой вне формы «каталог:<k>»: %d — пронумерованы "
                     "после каталожных, с номера %04d"
                     % (len(extra_no), min(extra_no.values())))
    for ck in numbered:
        members = by_chain[ck]
        leg_sets = {e: frozenset(r["flight_key"] for r in rs) for e, rs in members.items()}
        parts = {e: frozenset(frozenset(r["flight_key"] for r in rs if r["duty_id"] == d)
                              for d in {r["duty_id"] for r in rs})
                 for e, rs in members.items()}
        if len(set(leg_sets.values())) != 1:
            raise ExportRefused("цепочка %s разделена между сотрудниками по "
                                "участкам (%s) — в форму Meridian не выражается"
                                % (ck, "; ".join("%s: %d уч." % (e, len(s))
                                                  for e, s in sorted(leg_sets.items()))))
        if len(set(parts.values())) != 1:
            raise ExportRefused("смены цепочки %s у сотрудников разбиты по-разному "
                                "(%s) — в форму Meridian не выражается"
                                % (ck, "; ".join("%s: %d смен" % (e, len(p))
                                                  for e, p in sorted(parts.items()))))
        k = _catalog_index(ck)
        unit = units[k] if (k is not None and 0 <= k < len(units)
                            and isinstance(units[k], dict)) else None
        pid = "COSMOS-%s-%04d" % (label12, (k + 1) if k is not None else extra_no[ck])
        if unit is not None:
            if {str(x) for x in (unit.get("fns") or [])} != set(next(iter(leg_sets.values()))):
                diverged += 1
                notes.append("%s: множество участков цепочки в ростере отличается "
                             "от каталога пейринга — истина по назначениям" % pid)
        first_emp = sorted(members)[0]
        rs0 = members[first_emp]
        duty_groups = sorted(next(iter(parts.values())),
                             key=lambda s: min(r["dep_time"] for r in rs0 if r["flight_key"] in s))
        seat_order = [str(s) for s in (unit.get("seats") or [])] if unit else []
        recs = []
        for n, group in enumerate(duty_groups, start=1):
            did = "%s$%d" % (pid, n)
            legs_in = sorted((r for r in rs0 if r["flight_key"] in group),
                             key=lambda r: (r["dep_time"], r["flight_key"]))
            start = min(r["dep_time"] for r in legs_in)
            end = max(r["arr_time"] for r in legs_in)
            # Границы смены — по записям смен решения, когда запись смены
            # сотрудника лежит целиком в этой группе участков; смена
            # решателя, охватывающая несколько цепочек (две короткие
            # цепочки в один день), границы чужой цепочки не растягивает.
            for e, rs in members.items():
                dids = {r["duty_id"] for r in rs if r["flight_key"] in group}
                for du in rows["duties"]:
                    if (du["employee"] == e and du["id"] in dids and du["start"]
                            and du["end"]
                            and emp_duty_legs.get((e, du["id"]), set()) <= set(group)):
                        start, end = min(start, du["start"]), max(end, du["end"])
            flight_min = 0
            stations: list[str] = []
            for r in legs_in:
                dh_all = all(x["deadhead"] for e, rs in members.items()
                             for x in rs if x["flight_key"] == r["flight_key"])
                legs_by_fn[r["flight_key"]]["dutyCrewTypes"][did] = r_dh if dh_all else r_work
                if not dh_all:
                    flight_min += _minutes(_dt(r["dep_time"]), _dt(r["arr_time"]))
                for st in (r["dep_airport"], r["arr_airport"]):
                    if st not in stations:
                        stations.append(st)
            comps = []
            for e in members:
                mine = [r for r in members[e] if r["flight_key"] in group]
                roles = sorted({r["role"] for r in mine if r["role"]})
                rank = roles[0] if roles else ""
                if not rank:
                    raise ExportRefused("у сотрудника %s на смене %s нет роли — "
                                        "complementRank не определим" % (e, did))
                locked = any((e, r["flight_key"]) in pinned for r in mine)
                comps.append({"employeeId": e, "complementRank": rank,
                              "requirements": [], "locked": locked})
            seen_slots: set[str] = set()
            for slot, flight, role in unfilled.get(ck, []):
                if flight in group and slot not in seen_slots:
                    seen_slots.add(slot)
                    comps.append({"employeeId": "", "complementRank": role,
                                  "requirements": [], "locked": False})
            comps.sort(key=lambda c: (seat_order.index(c["complementRank"])
                                      if c["complementRank"] in seat_order else len(seat_order),
                                      c["complementRank"], c["employeeId"] == "", c["employeeId"]))
            seats_total += len(comps)
            seats_filled += sum(1 for c in comps if c["employeeId"])
            seats_unfilled += sum(1 for c in comps if not c["employeeId"])
            recs.append({"dutyId": did, "start": _stamp(start, scale, utc),
                         "end": _stamp(end, scale, utc), "restOverride": 0,
                         "assignmentType": "FLIGHT",
                         "flightTimeCredit": int(flight_min),
                         "workTimeCredit": _minutes(_dt(start), _dt(end)),
                         "split": False, "pairingId": pid, "numberOfDuty": n,
                         "firstDepartureStation": legs_in[0]["dep_airport"],
                         "lastArrivalStation": legs_in[-1]["arr_airport"],
                         "stationsVisited": stations, "complements": comps,
                         "_s": _dt(start), "_e": _dt(end)})
        for a, b in zip(recs, recs[1:]):
            a["restOverride"] = _minutes(a["_e"], b["_s"])
        if unit is not None and unit.get("minRestAfterMin") is not None:
            rest_after = int(round(float(unit["minRestAfterMin"])))
        else:
            last_st = recs[-1]["lastArrivalStation"]
            bases = {members[e][0]["base"] for e in members if members[e][0]["base"]}
            rest_after = (home_rest if last_st in bases else away_rest) or 0
            notes.append("%s: отдых после цепочки — порог решателя по станции "
                         "окончания (%s), каталога пейринга нет"
                         % (pid, "дома" if last_st in bases else "вне базы"))
        # Пол отдыха после цепочки — переопределение только там, где
        # решение его действительно держит: сотрудник, продолжающий в
        # той же смене решателя следующей цепочкой (разрыв меньше пола —
        # стыковка, не отдых), получает 0 = «без переопределения», иначе
        # файл утверждал бы, что решение нарушает собственный отдых.
        chain_end = recs[-1]["_e"]
        for e, rs in members.items():
            nxt = [_dt(r["dep_time"]) for r in rows["assignments"]
                   if r["employee"] == e and _dt(r["dep_time"]) >= chain_end
                   and r["flight_key"] not in leg_sets[e]]
            if nxt and _minutes(chain_end, min(nxt)) < rest_after:
                notes.append("%s: сотрудник %s продолжает следующей цепочкой в "
                             "той же смене решателя — переопределение отдыха "
                             "после цепочки не пишется (0)" % (pid, e))
                rest_after = 0
        recs[-1]["restOverride"] = rest_after
        for rcd in recs:
            rcd.pop("_s"); rcd.pop("_e")
        duties.extend(recs)
        pairings.append({"pairingId": pid, "start": recs[0]["start"],
                         "end": recs[-1]["end"], "restOverride": rest_after,
                         "requirements": {},
                         "locked": bool(unit.get("pinned")) if unit else False,
                         "flightMinutesCredit": sum(d["flightTimeCredit"] for d in recs),
                         "workMinutesCredit": sum(d["workTimeCredit"] for d in recs),
                         "primaryPairing": True})
    legs = [legs_by_fn[k] for k in sorted(legs_by_fn, key=lambda k: (
        legs_by_fn[k]["start"], legs_by_fn[k]["legId"]))]
    plan_since, plan_until, plan_note = _plan_bounds(legs, src.manifest, scale, utc)
    notes.append(plan_note)
    obligations = [l for l in legs if not l["deadhead"]]
    covered = sum(1 for l in obligations if r_work in l["dutyCrewTypes"].values())
    if not obligations:
        raise ExportRefused("в пакете нет рейсов-обязательств — выгружать нечего")
    used = {r["employee"] for r in rows["assignments"]}
    per_emp = {e["id"]: e["flight_min"] for e in rows["employees"]}
    d = src.delivery
    metrics = [
        _metric("INTEGER", "PAIRINGS_COUNT", len(pairings), "Количество пейрингов"),
        _metric("INTEGER", "DUTIES_COUNT", len(duties), "Количество полётных смен"),
        _metric("INTEGER", "LEGS_TOTAL", len(legs), "Участков в файле (все рейсы данных)"),
        _metric("INTEGER", "LEGS_OBLIGATIONS", len(obligations), "Рейсов-обязательств покрытия"),
        _metric("INTEGER", "LEGS_COVERED", covered, "Рейсов-обязательств с рабочим назначением"),
        _metric("INTEGER", "LEGS_UNCOVERED", len(obligations) - covered, "Рейсов-обязательств без рабочего назначения"),
        _metric("PERCENT", "COVERAGE_PCT", round(100.0 * covered / len(obligations), 2),
                "Доля покрытия обязательств, процентов"),
        _metric("INTEGER", "SEATS_TOTAL", seats_total, "Мест в сменах всего"),
        _metric("INTEGER", "SEATS_FILLED", seats_filled, "Мест с назначенным сотрудником"),
        _metric("INTEGER", "SEATS_UNFILLED", seats_unfilled, "Мест без сотрудника"),
        _metric("INTEGER", "EMPLOYEES_USED", len(used), "Сотрудников с назначениями"),
        _metric("DURATION", "FLIGHT_MINUTES_ASSIGNED", sum(p["flightMinutesCredit"] for p in pairings),
                "Назначенное полётное время, минут"),
        _metric("DURATION", "WORK_MINUTES_ASSIGNED", sum(p["workMinutesCredit"] for p in pairings),
                "Назначенное рабочее время, минут"),
        _metric("DURATION", "MAX_EMPLOYEE_FLIGHT_MINUTES", max(per_emp.values()) if per_emp else 0,
                "Наибольшее полётное время одного сотрудника, минут"),
        _metric("OTHER", "CERTIFIED", 1 if cert["certified"] else 0,
                "Решение сертифицировано (1) / нет (0)"),
        _metric("OTHER", "FORM_STATUS", "draft-unconfirmed",
                "Форма ростера Meridian — проект, спецификацией не подтверждена"),
        _metric("OTHER", "CHAINS_DIVERGED_FROM_CATALOG", diverged,
                "Цепочек, чьё множество участков в ростере отличается от каталога"),
        _metric("OTHER", "SOLVER_RUN_ID", (d.get("schedule") or {}).get("solver_run_id") or "",
                "Прогон решателя COSMOS"),
        _metric("OTHER", "ROLE_ALPHABET", "%s/%s" % (r_work, r_dh),
                "Алфавит роли участка: работа/пассажиром"),
        _metric("OTHER", "TIME_SCALE", "utc" if (utc and scale.convertible) else scale.kind,
                "Шкала меток файла")]
    words = uncertified_words(cert)
    description = ("%s; COSMOS: решение ростеринга по пакету постановки %s (прогон %s, путь %s); %s%s"
                   % (ROSTER_FORM_NOTE, src.package_sha[:12],
                      (d.get("schedule") or {}).get("solver_run_id"), src.path or "?",
                      "; ".join(_dedupe(notes[1:])), ("; " + words) if words else ""))
    all_roles = sorted({r["role"] for r in rows["assignments"] if r["role"]}
                       | {c["complementRank"] for du in duties for c in du["complements"]})
    return {"datasetId": idn["datasetId"], "presetId": idn["presetId"],
            "format": 1, "type": "roster",
            "division": division_of(all_roles, src.manifest),
            "planSince": plan_since, "planUntil": plan_until,
            "created": created, "description": description,
            "legs": legs, "pairings": pairings, "duties": duties,
            "groundActivities": [], "metrics": metrics}


def roster_meridian(src: RosterSources, opts: Options | None = None) -> bytes:
    return _json_bytes(meridian_roster_doc(src, opts), sort_keys=False)


# ───────────────────────────────────────────────────── зеркало гейта BA

def meridian_check(doc: dict, kind: str, *, catalog: dict | None = None,
                   payload: dict | None = None, rows: dict | None = None,
                   roles: tuple[str, str] = GATE_ROLES) -> list[str]:
    """Проверки П1…П7 гейта линии BA (ba_meridian_gate.py) пункт в пункт
    для pairing (П6/П7 — при переданных каталоге и данных гейта) и Р1…Р4
    для проекта формы roster. Пустой список = PASS; каждая находка —
    словами с номером пункта. Гейт BA судит только type=pairing (на
    type≠pairing он падает assert-ом, не находкой) — ростер судится
    только здесь."""
    fails: list[str] = []
    r_work, r_dh = roles

    def bad(where: str, msg: str) -> None:
        fails.append("[%s] %s" % (where, msg))

    def dt(s):
        return datetime.fromisoformat(str(s))

    root_req = ["datasetId", "presetId", "format", "type", "division", "planSince",
                "planUntil", "created", "legs", "pairings", "duties",
                "groundActivities", "metrics"]
    for k in root_req:
        if k not in doc:
            bad("П1", "нет обязательного поля %s" % k)
    if not (isinstance(doc.get("format"), int) and doc["format"] > 0):
        bad("П1", "format не положительное целое")
    if doc.get("type") != kind:
        bad("П1", "type %r ≠ %r" % (doc.get("type"), kind))
    if doc.get("division") not in ("FD", "CC"):
        bad("П1", "division %r вне FD/CC" % doc.get("division"))
    for k in ("planSince", "planUntil", "created"):
        if not _ISO_RE.match(str(doc.get(k, ""))):
            bad("П1", "%s: не ISO 8601 без зоны — %r" % (k, doc.get(k)))
    for k in ("datasetId", "presetId"):
        if len(str(doc.get(k, ""))) > 150:
            bad("П1", "%s длиннее 150" % k)
    leg_req = ["legId", "flightNumber", "legNumber", "start", "end",
               "departureStation", "arrivalStation", "carrier", "fleet",
               "properties", "deadhead", "dutyCrewTypes"]
    seen_ids: dict[str, int] = {}
    for l in doc.get("legs") or []:
        for k in leg_req:
            if k not in l:
                bad("П2", "%s: нет поля %s" % (l.get("legId"), k))
        if not _ISO_RE.match(str(l.get("start"))) or not _ISO_RE.match(str(l.get("end"))):
            bad("П2", "%s: время не по формату" % l.get("legId"))
            continue
        if dt(l["end"]) <= dt(l["start"]):
            bad("П2", "%s: прилёт не позже вылета" % l.get("legId"))
        if not re.fullmatch(r"\w{2}", str(l.get("carrier"))):
            bad("П2", "%s: перевозчик не по маске — %s" % (l.get("legId"), l.get("carrier")))
        if l.get("properties") not in ("REGULAR", "CHARTER", "TRANSFER"):
            bad("П2", "%s: properties вне списка" % l.get("legId"))
        if not isinstance(l.get("deadhead"), bool):
            bad("П2", "%s: deadhead не логический" % l.get("legId"))
        if not isinstance(l.get("legNumber"), int) or l["legNumber"] <= 0:
            bad("П2", "%s: legNumber не положительный" % l.get("legId"))
        seen_ids[str(l.get("legId"))] = seen_ids.get(str(l.get("legId")), 0) + 1
    dup = [k for k, v in seen_ids.items() if v > 1]
    if dup:
        bad("П2", "дубли legId: %s" % dup[:3])
    duty_req = ["dutyId", "start", "end", "flightTimeCredit", "workTimeCredit",
                "stationsVisited", "split", "pairingId", "numberOfDuty",
                "firstDepartureStation", "lastArrivalStation", "complements"]
    duty_by_id: dict[str, dict] = {}
    for d in doc.get("duties") or []:
        for k in duty_req:
            if k not in d:
                bad("П3", "%s: нет поля %s" % (d.get("dutyId"), k))
        if len(str(d.get("dutyId"))) > 50:
            bad("П3", "%s: длина свыше 50" % d.get("dutyId"))
        try:
            if dt(d["end"]) <= dt(d["start"]):
                bad("П3", "%s: конец не позже начала" % d.get("dutyId"))
        except (KeyError, ValueError):
            bad("П3", "%s: границы смены не разобраны" % d.get("dutyId"))
        if not isinstance(d.get("split"), bool):
            bad("П3", "%s: split не логический" % d.get("dutyId"))
        for c in d.get("complements") or []:
            if "complementRank" not in c or not str(c.get("complementRank")):
                bad("П3", "%s: место без ранга" % d.get("dutyId"))
            if kind == "pairing" and c.get("employeeId"):
                bad("П3", "%s: пейринг несёт назначенного сотрудника %s (правило 0080)"
                    % (d.get("dutyId"), c["employeeId"]))
        duty_by_id[str(d.get("dutyId"))] = d
    pair_req = ["pairingId", "start", "end", "restOverride", "flightMinutesCredit",
                "workMinutesCredit", "primaryPairing"]
    pair_by_id: dict[str, dict] = {}
    for p in doc.get("pairings") or []:
        for k in pair_req:
            if k not in p:
                bad("П4", "%s: нет поля %s" % (p.get("pairingId"), k))
        if len(str(p.get("pairingId"))) > 150:
            bad("П4", "%s: длина свыше 150" % p.get("pairingId"))
        pair_by_id[str(p.get("pairingId"))] = p
    duties_of: dict[str, list[dict]] = {}
    for d in doc.get("duties") or []:
        if str(d.get("pairingId")) not in pair_by_id:
            bad("П4", "%s: ссылка на несуществующий пейринг" % d.get("dutyId"))
        duties_of.setdefault(str(d.get("pairingId")), []).append(d)
    for pid, ds in duties_of.items():
        if pid not in pair_by_id:
            continue
        ds.sort(key=lambda x: x.get("numberOfDuty") or 0)
        if [x.get("numberOfDuty") for x in ds] != list(range(1, len(ds) + 1)):
            bad("П4", "%s: номера смен не сплошные" % pid)
        p = pair_by_id[pid]
        if p.get("start") != ds[0].get("start") or p.get("end") != ds[-1].get("end"):
            bad("П4", "%s: границы пейринга не совпали со сменами" % pid)
        try:
            if p.get("flightMinutesCredit") != sum(x["flightTimeCredit"] for x in ds):
                bad("П4", "%s: налёт пейринга не равен сумме смен" % pid)
            if p.get("workMinutesCredit") != sum(x["workTimeCredit"] for x in ds):
                bad("П4", "%s: рабочее время не равно сумме смен" % pid)
        except (KeyError, TypeError):
            bad("П4", "%s: кредиты смен не числа" % pid)
    orphan = set(pair_by_id) - set(duties_of)
    if orphan:
        bad("П4", "пейринги без смен: %s" % sorted(orphan)[:3])
    for l in doc.get("legs") or []:
        for did, role in (l.get("dutyCrewTypes") or {}).items():
            if did not in duty_by_id:
                bad("П5", "%s: роль ссылается на несуществующую смену %s" % (l.get("legId"), did))
                continue
            if role not in (r_work, r_dh):
                bad("П5", "%s: роль %r вне алфавита %s/%s" % (l.get("legId"), role, r_work, r_dh))
            d = duty_by_id[did]
            try:
                if not (dt(d["start"]) <= dt(l["start"]) and dt(l["end"]) <= dt(d["end"])):
                    bad("П5", "%s: участок вне окна смены %s" % (l.get("legId"), did))
            except (KeyError, ValueError):
                pass
    mcodes = {}
    for m in doc.get("metrics") or []:
        for k in ("metricType", "metricCode", "metricValue", "metricDescription"):
            if not str(m.get(k, "")):
                bad("П7", "метрика без поля %s" % k)
        if m.get("metricType") not in ("INTEGER", "PERCENT", "NUMBER", "DURATION", "OTHER"):
            bad("П7", "%s: тип вне списка" % m.get("metricCode"))
        if not isinstance(m.get("metricValue"), str):
            bad("П7", "%s: metricValue не строка" % m.get("metricCode"))
        mcodes[str(m.get("metricCode"))] = m
    if kind == "pairing" and catalog is not None and payload is not None:
        work_legs = {l["legId"] for l in doc.get("legs") or []
                     if r_work in (l.get("dutyCrewTypes") or {}).values()}
        dh_legs = {l["legId"] for l in doc.get("legs") or []
                   if r_dh in (l.get("dutyCrewTypes") or {}).values()}
        leg_id_of = {f["flightNumber"]: f["flightLegId"] for f in payload["flights"]}
        cat_work = {leg_id_of[fn] for ch in catalog["selected"] for fn in ch["active_fns"]}
        cat_dh = {leg_id_of[fn] for ch in catalog["selected"] for fn in ch["fns"]
                  if fn not in set(ch["active_fns"])}
        if work_legs != cat_work:
            bad("П6", "рабочие участки разошлись: только в решении %d, только в каталоге %d"
                % (len(work_legs - cat_work), len(cat_work - work_legs)))
        if dh_legs != cat_dh:
            bad("П6", "перегоны разошлись: %d расхождений" % len(dh_legs ^ cat_dh))
        if len(pair_by_id) != len(catalog["selected"]):
            bad("П6", "цепочек %d против %d" % (len(pair_by_id), len(catalog["selected"])))
        transport_ids = {f["flightLegId"] for f in payload["flights"] if f.get("transportOnly")}
        uncovered_sol = [l["legId"] for l in doc.get("legs") or []
                         if r_work not in (l.get("dutyCrewTypes") or {}).values()
                         and l["legId"] not in transport_ids]
        if len(uncovered_sol) != catalog["coverage"]["uncovered"]:
            bad("П6", "непокрытых %d против %d в каталоге"
                % (len(uncovered_sol), catalog["coverage"]["uncovered"]))
        if mcodes.get("LEGS_COVERED", {}).get("metricValue") != str(catalog["coverage"]["covered"]):
            bad("П7", "метрика покрытия не равна каталогу")
    if kind == "roster":
        known = {e["id"] for e in (rows or {}).get("employees") or []} if rows else None
        restored: set[tuple[str, str, str, bool]] = set()
        for d in doc.get("duties") or []:
            did = str(d.get("dutyId"))
            comps = d.get("complements") or []
            for c in comps:
                if not str(c.get("complementRank") or ""):
                    bad("Р1", "%s: место без ранга" % did)
                eid = c.get("employeeId")
                if eid is None or not isinstance(eid, str):
                    bad("Р1", "%s: employeeId не строка" % did)
                elif eid and known is not None and eid not in known:
                    bad("Р1", "%s: сотрудник %s вне штата решения" % (did, eid))
            for l in doc.get("legs") or []:
                role = (l.get("dutyCrewTypes") or {}).get(did)
                if role is None:
                    continue
                for c in comps:
                    if c.get("employeeId"):
                        restored.add((c["employeeId"], l["legId"], c["complementRank"],
                                      role == r_dh))
        if rows is not None:
            legs_by_key = {}
            for l in doc.get("legs") or []:
                legs_by_key.setdefault((l["flightNumber"], l["start"]), l["legId"])
            expected: set[tuple[str, str, str, bool]] = set()
            for r in rows.get("assignments") or []:
                lid = None
                for (num, start), v in legs_by_key.items():
                    if num == r["flight_no"] and start[:16] == r["dep_time"][:16]:
                        lid = v
                        break
                if lid is None:
                    lid = [l["legId"] for l in doc.get("legs") or []
                           if l["flightNumber"] == r["flight_no"] and
                           l["departureStation"] == r["dep_airport"]]
                    lid = lid[0] if len(lid) == 1 else r["flight_key"]
                expected.add((r["employee"], lid, r["role"], r["deadhead"]))
            if restored != expected:
                bad("Р2", "назначения из файла не совпали с решением: только в файле %d, "
                    "только в решении %d" % (len(restored - expected), len(expected - restored)))
        per_emp: dict[str, list[dict]] = {}
        for d in doc.get("duties") or []:
            for c in d.get("complements") or []:
                if c.get("employeeId"):
                    per_emp.setdefault(c["employeeId"], []).append(d)
        for e, ds in per_emp.items():
            ds.sort(key=lambda x: x["start"])
            for a, b in zip(ds, ds[1:]):
                if dt(b["start"]) < dt(a["end"]):
                    bad("Р3", "%s: смены %s и %s пересекаются" % (e, a["dutyId"], b["dutyId"]))
                elif a.get("restOverride") and _minutes(dt(a["end"]), dt(b["start"])) < int(a["restOverride"]):
                    bad("Р3", "%s: разрыв между %s и %s меньше restOverride" % (e, a["dutyId"], b["dutyId"]))
        try:
            if (int(mcodes["SEATS_FILLED"]["metricValue"]) + int(mcodes["SEATS_UNFILLED"]["metricValue"])
                    != int(mcodes["SEATS_TOTAL"]["metricValue"])):
                bad("Р4", "SEATS_FILLED + SEATS_UNFILLED ≠ SEATS_TOTAL")
        except (KeyError, ValueError):
            bad("Р4", "метрики мест отсутствуют")
    return fails


# ───────────────────────────────────────────────────────── имена и выгрузка

def _ascii_tag(value: str) -> str:
    v = str(value)
    if _ASCII_NAME_RE.match(v):
        return v
    return hashlib.sha1(v.encode("utf-8")).hexdigest()[:8]


def file_name(solution: str, fmt: str, package_sha: str, attempt: str | None,
              certified: bool, employee: str | None = None) -> tuple[str, str]:
    """(имя ASCII, имя UTF-8): cosmos-<solution>-<sha12>[-<attempt>]
    [-<employee>][-uncertified][-meridian].<ext>; для Content-Disposition
    ASCII-имя обязательно, UTF-8 — через filename* (RFC 5987), когда
    табельный вне ASCII."""
    parts = ["cosmos", solution, package_sha[:12]]
    parts_u = list(parts)
    if attempt:
        parts.append(_ascii_tag(attempt)); parts_u.append(str(attempt))
    if employee:
        parts.append(_ascii_tag(employee)); parts_u.append(str(employee))
    if not certified:
        parts.append("uncertified"); parts_u.append("uncertified")
    if fmt == "meridian":
        parts.append("meridian"); parts_u.append("meridian")
    return ("-".join(parts) + "." + EXT[fmt], "-".join(parts_u) + "." + EXT[fmt])


def content_disposition(name_ascii: str, name_utf8: str) -> str:
    from urllib.parse import quote
    out = 'attachment; filename="%s"' % name_ascii
    if name_utf8 != name_ascii:
        out += "; filename*=UTF-8''%s" % quote(name_utf8, safe="")
    return out


def export(sources, fmt: str, opts: Options | None = None
           ) -> tuple[bytes, str, str, str, dict]:
    """Единая точка: (байты, имя ASCII, имя UTF-8, media type, meta).
    CSV — с BOM (для табличных редакторов), iCal — как есть."""
    opts = opts or Options()
    kind = "pairing" if isinstance(sources, PairingSources) else "roster"
    if fmt not in FORMATS[kind]:
        raise ExportNotFound("формат %r для %s неизвестен; перечень: %s"
                             % (fmt, kind, ", ".join(FORMATS[kind])))
    cert = certification(sources.delivery)
    task = "пейринга" if kind == "pairing" else "ростеринга"
    require_solution(sources.delivery, task)
    scale = time_scale(sources.manifest)
    created, cscale = created_stamp(sources.delivery)
    meta = {"certified": cert["certified"], "established": cert["established"],
            "package_sha": sources.package_sha,
            "solver_run_id": (sources.delivery.get("schedule") or {}).get("solver_run_id"),
            "time_scale": ("utc" + ((";zone=" + scale.zone) if scale.zone else "")
                           if (opts.utc and scale.convertible) else scale.header()),
            "created": created, "created_scale": cscale,
            "form_status": None, "employee": opts.employee}
    if kind == "pairing":
        if fmt == "csv":
            data = ("\ufeff" + pairing_csv(sources)).encode("utf-8")
        elif fmt == "xlsx":
            data = pairing_xlsx(sources)
        elif fmt == "json":
            data = pairing_json(sources)
        else:
            data = pairing_meridian(sources, opts)
    else:
        if fmt == "csv":
            data = ("\ufeff" + roster_csv(sources)).encode("utf-8")
        elif fmt == "xlsx":
            data = roster_xlsx(sources)
        elif fmt == "json":
            data = roster_json(sources)
        elif fmt == "ics":
            data = roster_ics(sources, opts).encode("utf-8")
        else:
            data = roster_meridian(sources, opts)
            meta["form_status"] = "draft"
    name_a, name_u = file_name(kind, fmt, sources.package_sha, sources.attempt,
                               cert["certified"],
                               opts.employee if fmt == "ics" else None)
    return data, name_a, name_u, MEDIA[fmt], meta


def availability(sources, kind: str) -> dict:
    """Указатель доступности форматов для панели: по каждому формату ok
    либо причина словами сервера; Meridian строится пробно (дёшево).
    Базовая проверка — та же, что у export(): решение есть И штамп выдачи
    установлен (журнал кнопки); иначе ни один формат не строится, и
    указатель говорит причину, а не рисует живую ссылку на 409."""
    cert = certification(sources.delivery)
    out = {"present": True, "certified": cert["certified"],
           "established": cert["established"], "verdict": cert["verdict"],
           "attempt": sources.attempt, "formats": {}}
    try:
        require_solution(sources.delivery,
                         "пейринга" if kind == "pairing" else "ростеринга")
        created_stamp(sources.delivery)
        base_ok, base_reason = True, None
    except ExportError as e:
        base_ok, base_reason = False, str(e)
    for fmt in FORMATS[kind]:
        rec: dict = {"ok": base_ok}
        if not base_ok:
            rec["reason"] = base_reason
        elif fmt == "meridian":
            try:
                if kind == "pairing":
                    meridian_pairing_doc(sources, Options(allow_uncertified=True))
                else:
                    meridian_roster_doc(sources, Options(allow_uncertified=True))
                if not cert["certified"]:
                    rec.update(ok=False, needs_confirm=True,
                               reason="решение не сертифицировано — Meridian только "
                                      "с явным подтверждением (allow_uncertified=1)")
                if kind == "roster":
                    rec["form_status"] = "draft"
            except ExportError as e:
                rec.update(ok=False, reason=str(e))
        out["formats"][fmt] = rec
    if kind == "roster":
        try:
            out["employees"] = [e["id"] for e in roster_rows(sources)["employees"]]
        except ExportError:
            out["employees"] = []
    return out
